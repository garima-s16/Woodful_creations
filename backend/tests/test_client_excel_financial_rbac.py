"""Test for a real gap found this turn - clients.xlsx included a
company-wide financial summary (total order value, total outstanding)
and per-client financial columns, downloadable by any authenticated
role."""
import io
from openpyxl import load_workbook
from app.core.security import hash_password
from app.models.user import User


def _login(client, test_user):
    resp = client.post("/api/auth/login", json={"identifier": "test@example.com", "password": "TestPass123!"})
    assert resp.status_code == 200


def test_master_client_export_includes_financial_columns(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/reports/clients.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["Clients"]
    header_row = [c.value for c in ws[5]]
    assert "Total Order Value" in header_row
    assert "Outstanding" in header_row


def test_employee_client_export_excludes_financial_columns(client, test_user, db_session):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "Client Excel RBAC Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    user = User(
        username="clientexcelrbacuser", email="clientexcelrbacuser@example.com", full_name="Client Excel RBAC User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "clientexcelrbacuser@example.com", "password": "EmpPass1!"})

    resp = client.get("/api/reports/clients.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["Clients"]
    header_row = [c.value for c in ws[5]]
    assert "Total Order Value" not in header_row
    assert "Total Paid" not in header_row
    assert "Outstanding" not in header_row
    # Non-financial contact/status columns remain useful and present.
    assert "Client Name" in header_row
    assert "Phone" in header_row
    assert "Status" in header_row
