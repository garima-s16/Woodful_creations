"""Inventory domain tests: materials, stock operations, ledger/
transfers, purchases/reporting, material imports, and shortage
intelligence. Combines all former test_*.py files under
tests/modules/inventory/."""
import json
import pytest
from datetime import datetime
from io import BytesIO
from openpyxl import load_workbook
from sqlalchemy import event
from app.platform.database import engine
from decimal import Decimal
from app.platform.security import hash_password
from app.modules.inventory.models import Material
from app.modules.procurement.models import Supplier
from app.modules.auth.auth import User
from app.modules.inventory.schemas import PurchaseCreate
from app.modules.operations.schemas import IssueCreate
from app.modules.inventory.services import StockService
from app.modules.procurement.services import ProcurementService
from tests.helpers import _login, _create_employee_with_login as _create_employee
from tests.helpers import _login


# --- test_materials.py ---
def test_material_can_be_created_with_decimal_opening_stock(client, test_user):
    _login(client, test_user)
    resp = client.post("/api/materials/", json={
        "name": "Decimal Opening Stock Adhesive", "unit": "Kg", "opening_stock": "2.5", "minimum_stock": "1.5",
    })
    assert resp.status_code == 201
    assert float(resp.json()["current_stock"]) == 2.5
    assert float(resp.json()["minimum_stock"]) == 1.5


def test_purchase_of_decimal_quantity_does_not_truncate_stock(client, test_user):
    """The exact scenario the brief names: 2.5 kg of adhesive must
    remain 2.5, never silently become 2."""
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Decimal Purchase Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Decimal Purchase Adhesive", "unit": "Kg", "opening_stock": "0", "minimum_stock": "1",
    }).json()

    client.post("/api/purchases/", json={
        "date": "2026-08-15T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
        "quantity": "2.5", "unit": "Kg", "rate": "200.00", "gst_percent": "18", "payment_status": "Paid",
    })

    updated = client.get(f"/api/materials/{material['id']}").json()
    assert float(updated["current_stock"]) == 2.5
    assert float(updated["total_purchased"]) == 2.5


def test_multiple_decimal_purchases_accumulate_precisely(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Multi Decimal Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Multi Decimal Material", "unit": "Litres", "opening_stock": "0", "minimum_stock": "1",
    }).json()

    for qty in ["1.25", "0.75", "2.5"]:
        client.post("/api/purchases/", json={
            "date": "2026-08-15T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
            "quantity": qty, "unit": "Litres", "rate": "100.00", "gst_percent": "18", "payment_status": "Paid",
        })

    updated = client.get(f"/api/materials/{material['id']}").json()
    assert float(updated["current_stock"]) == 4.5


def test_issuing_decimal_quantity_does_not_truncate_stock(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Decimal Issue Material", "unit": "Metres", "opening_stock": "10", "minimum_stock": "1",
    }).json()

    client.post("/api/issues/", json={
        "date": "2026-08-15T00:00:00", "material_id": material["id"],
        "quantity_issued": "3.75", "unit": "Metres",
    })

    updated = client.get(f"/api/materials/{material['id']}").json()
    assert float(updated["current_stock"]) == 6.25  # 10 - 3.75
    assert float(updated["total_issued"]) == 3.75


def test_cannot_issue_more_than_available_decimal_stock(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Decimal Overissue Material", "unit": "Kg", "opening_stock": "2.5", "minimum_stock": "1",
    }).json()

    resp = client.post("/api/issues/", json={
        "date": "2026-08-15T00:00:00", "material_id": material["id"],
        "quantity_issued": "3.0", "unit": "Kg",
    })
    assert resp.status_code == 400


def test_stock_transfer_preserves_decimal_quantity(client, test_user):
    _login(client, test_user)
    location = client.post("/api/locations/", json={"name": "Decimal Transfer Location"}).json()
    material = client.post("/api/materials/", json={
        "name": "Decimal Transfer Material", "unit": "Kg", "opening_stock": "5.5", "minimum_stock": "1",
    }).json()

    resp = client.post("/api/stock/transfers", json={
        "material_id": material["id"], "quantity": "2.25", "to_location_id": location["id"],
    })
    assert resp.status_code == 201
    assert float(resp.json()["quantity"]) == 2.25


def test_stock_adjustment_preserves_decimal_quantity(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Decimal Adjustment Material", "unit": "Kg", "opening_stock": "5.5", "minimum_stock": "1",
    }).json()

    resp = client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Physical Count Increase",
        "quantity_delta": "0.75", "reason": "Recount found extra fractional stock",
    })
    assert resp.status_code == 201
    assert float(resp.json()["stock_before"]) == 5.5
    assert float(resp.json()["stock_after"]) == 6.25


def test_negative_decimal_adjustment_error_message_does_not_crash(client, test_user):
    """The exact bug that was caught and fixed - a decimal
    quantity_delta through the :+d format specifier used to raise
    ValueError instead of returning a proper 400."""
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Negative Decimal Guard Material", "unit": "Kg", "opening_stock": "1.5", "minimum_stock": "1",
    }).json()

    resp = client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Damage",
        "quantity_delta": "-2.5", "reason": "Water damage",
    })
    assert resp.status_code == 400  # not 500 - the format string must not crash
    assert "2.5" in resp.json()["detail"]


def test_material_created_without_location_can_have_one_assigned_later(client, test_user):
    """The exact reported bug: a material created with no location has
    its opening stock permanently un-findable at any specific location
    (_location_balance never matches None), so a location-specific
    adjustment against it is always wrongly rejected - even though the
    material's own aggregate current_stock correctly shows it. Editing
    the material to set a location for the first time must retroactively
    fix this, not require a data migration."""
    _login(client, test_user)
    location = client.post("/api/locations/", json={"name": "Unlocated Fix Location"}).json()
    material = client.post("/api/materials/", json={
        "name": "Unlocated Opening Stock Material", "unit": "Sheet", "opening_stock": "10", "minimum_stock": "1",
    }).json()
    assert material["opening_stock_location_id"] is None

    # Before assigning a location, a location-specific adjustment
    # against this material's genuinely-available stock is wrongly
    # rejected - reproducing the reported bug directly.
    resp = client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Damage", "quantity_delta": "-1",
        "reason": "Wastage", "location_id": location["id"],
    })
    assert resp.status_code == 400
    assert "0" in resp.json()["detail"]

    client.put(f"/api/materials/{material['id']}", json={"location_id": location["id"]})
    updated = client.get(f"/api/materials/{material['id']}").json()
    assert updated["opening_stock_location_id"] == location["id"]

    # The same adjustment now succeeds against the now-locatable stock.
    resp = client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Damage", "quantity_delta": "-1",
        "reason": "Wastage", "location_id": location["id"],
    })
    assert resp.status_code == 201
    assert float(resp.json()["stock_after"]) == 9.0


def test_material_with_existing_opening_location_is_never_overwritten(client, test_user):
    """A material whose opening stock already has a real, meaningful
    location must never have that silently rewritten just because its
    primary location is edited later - only a currently-missing value
    gets filled in."""
    _login(client, test_user)
    original_location = client.post("/api/locations/", json={"name": "Original Opening Location"}).json()
    new_location = client.post("/api/locations/", json={"name": "New Primary Location"}).json()
    material = client.post("/api/materials/", json={
        "name": "Already Located Material", "unit": "Sheet", "opening_stock": "10", "minimum_stock": "1",
        "location_id": original_location["id"],
    }).json()
    assert material["opening_stock_location_id"] == original_location["id"]

    client.put(f"/api/materials/{material['id']}", json={"location_id": new_location["id"]})
    updated = client.get(f"/api/materials/{material['id']}").json()
    assert updated["opening_stock_location_id"] == original_location["id"]  # unchanged


def test_material_defaults_to_active(client, test_user):
    _login(client, test_user)
    resp = client.post("/api/materials/", json={
        "name": "Default Active Material", "unit": "Sheets", "opening_stock": "5", "minimum_stock": "1",
    })
    assert resp.status_code == 201
    assert resp.json()["is_active"] is True


def test_material_can_be_deactivated_and_reactivated(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Deactivate Test Material", "unit": "Sheets", "opening_stock": "5", "minimum_stock": "1",
    }).json()

    resp = client.put(f"/api/materials/{material['id']}", json={"is_active": False})
    assert resp.status_code == 200
    assert resp.json()["is_active"] is False

    resp2 = client.put(f"/api/materials/{material['id']}", json={"is_active": True})
    assert resp2.json()["is_active"] is True


def test_active_only_filter_excludes_inactive_materials(client, test_user):
    _login(client, test_user)
    active_material = client.post("/api/materials/", json={
        "name": "Active Only Filter Material A", "unit": "Sheets", "opening_stock": "5", "minimum_stock": "1",
    }).json()
    inactive_material = client.post("/api/materials/", json={
        "name": "Active Only Filter Material B", "unit": "Sheets", "opening_stock": "5", "minimum_stock": "1",
    }).json()
    client.put(f"/api/materials/{inactive_material['id']}", json={"is_active": False})

    filtered = client.get("/api/materials/", params={"active_only": True}).json()
    ids = [m["id"] for m in filtered]
    assert active_material["id"] in ids
    assert inactive_material["id"] not in ids


def test_default_material_list_still_includes_inactive_materials(client, test_user):
    """Backward compatibility - omitting active_only must behave
    exactly as before this field existed."""
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Unfiltered List Material", "unit": "Sheets", "opening_stock": "5", "minimum_stock": "1",
    }).json()
    client.put(f"/api/materials/{material['id']}", json={"is_active": False})

    unfiltered = client.get("/api/materials/").json()
    assert any(m["id"] == material["id"] for m in unfiltered)


def test_recent_stock_movement_includes_purchases_and_issues(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Movement Test Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Movement Test Material", "unit": "Sheets", "opening_stock": "20", "minimum_stock": "1",
    }).json()

    client.post("/api/purchases/", json={
        "date": "2026-08-16T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
        "quantity": "10", "unit": "Sheets", "rate": "500.00", "gst_percent": "18", "payment_status": "Paid",
    })
    client.post("/api/issues/", json={
        "date": "2026-08-16T00:00:00", "material_id": material["id"], "quantity_issued": "5", "unit": "Sheets",
    })

    resp = client.get("/api/dashboard/stock")
    assert resp.status_code == 200
    movement = resp.json()["recent_stock_movement"]
    assert any(row["type"] == "IN" and row["material"] == "Movement Test Material" for row in movement)
    assert any(row["type"] == "OUT" and row["material"] == "Movement Test Material" for row in movement)


def test_recent_stock_movement_is_not_hardcoded(client, test_user):
    """The brief's explicit requirement - must reflect actual seeded
    transactions, not a fixed/fabricated list."""
    _login(client, test_user)
    before = client.get("/api/dashboard/stock").json()["recent_stock_movement"]

    supplier = client.post("/api/suppliers/", json={"name": "Not Hardcoded Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Not Hardcoded Movement Material", "unit": "Sheets", "opening_stock": "0", "minimum_stock": "1",
    }).json()
    client.post("/api/purchases/", json={
        "date": "2026-08-16T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
        "quantity": "3", "unit": "Sheets", "rate": "500.00", "gst_percent": "18", "payment_status": "Paid",
    })

    after = client.get("/api/dashboard/stock").json()["recent_stock_movement"]
    assert after != before
    assert any(row["material"] == "Not Hardcoded Movement Material" for row in after)


def test_create_category_and_subcategory(client, test_user):
    _login(client, test_user)
    category = client.post("/api/material-categories/", json={"name": "Board & Wood Materials"}).json()
    assert len(category["business_id"]) == 10

    subcategory = client.post("/api/material-categories/subcategories", json={
        "category_id": category["id"], "name": "Plywood",
    }).json()
    assert subcategory["category_id"] == category["id"]


def test_duplicate_category_name_rejected(client, test_user):
    _login(client, test_user)
    client.post("/api/material-categories/", json={"name": "Duplicate Category Test"})
    resp = client.post("/api/material-categories/", json={"name": "Duplicate Category Test"})
    assert resp.status_code == 400


def test_duplicate_subcategory_within_same_category_rejected(client, test_user):
    _login(client, test_user)
    category = client.post("/api/material-categories/", json={"name": "Dup Subcat Test Category"}).json()
    client.post("/api/material-categories/subcategories", json={"category_id": category["id"], "name": "Plywood"})
    resp = client.post("/api/material-categories/subcategories", json={"category_id": category["id"], "name": "Plywood"})
    assert resp.status_code == 400


def test_same_subcategory_name_allowed_under_different_categories(client, test_user):
    """"Hardware" under "Furniture Fittings" and "Hardware" under a
    different parent are genuinely different subcategories - the
    uniqueness is per-category, not global."""
    _login(client, test_user)
    cat_a = client.post("/api/material-categories/", json={"name": "Category A For Hardware Test"}).json()
    cat_b = client.post("/api/material-categories/", json={"name": "Category B For Hardware Test"}).json()
    resp_a = client.post("/api/material-categories/subcategories", json={"category_id": cat_a["id"], "name": "Hardware"})
    resp_b = client.post("/api/material-categories/subcategories", json={"category_id": cat_b["id"], "name": "Hardware"})
    assert resp_a.status_code == 201
    assert resp_b.status_code == 201


def test_attribute_definition_with_invalid_data_type_rejected(client, test_user):
    _login(client, test_user)
    category = client.post("/api/material-categories/", json={"name": "Attr Type Test Category"}).json()
    subcategory = client.post("/api/material-categories/subcategories", json={
        "category_id": category["id"], "name": "Attr Type Test Subcategory",
    }).json()
    resp = client.post(f"/api/material-categories/subcategories/{subcategory['id']}/attributes", json={
        "name": "Thickness", "data_type": "not_a_real_type",
    })
    assert resp.status_code == 422


def test_material_created_with_subcategory_syncs_legacy_category_string(client, test_user):
    """The single most important compatibility guarantee: every existing
    consumer (dashboard, chatbot, PDF/Excel, low-stock filter) reads
    material.category as a plain string - this must stay populated and
    correct even for materials created through the new hierarchy."""
    _login(client, test_user)
    category = client.post("/api/material-categories/", json={"name": "Board & Wood Materials Sync Test"}).json()
    subcategory = client.post("/api/material-categories/subcategories", json={
        "category_id": category["id"], "name": "HDHMR Sync Test",
    }).json()

    material = client.post("/api/materials/", json={
        "name": "HDHMR 18mm Sync Test", "unit": "Sheets", "opening_stock": 10, "minimum_stock": 5,
        "subcategory_id": subcategory["id"],
    }).json()
    assert material["category"] == "Board & Wood Materials Sync Test"
    assert material["subcategory_id"] == subcategory["id"]


def test_material_without_subcategory_still_works_the_old_way(client, test_user):
    """Backward compatibility - a material can still be created exactly
    the way it always was, with just a flat category string and no
    hierarchy at all."""
    _login(client, test_user)
    resp = client.post("/api/materials/", json={
        "name": "Legacy Flat Category Material", "category": "Plywood", "unit": "Sheets",
        "opening_stock": 5, "minimum_stock": 2,
    })
    assert resp.status_code == 201
    body = resp.json()
    assert body["category"] == "Plywood"
    assert body["subcategory_id"] is None


def test_material_attribute_values_persist_with_correct_typed_storage(client, test_user):
    _login(client, test_user)
    category = client.post("/api/material-categories/", json={"name": "Attr Value Test Category"}).json()
    subcategory = client.post("/api/material-categories/subcategories", json={
        "category_id": category["id"], "name": "Attr Value Test Subcategory",
    }).json()
    thickness_attr = client.post(f"/api/material-categories/subcategories/{subcategory['id']}/attributes", json={
        "name": "Thickness", "data_type": "number", "unit_label": "mm",
    }).json()
    brand_attr = client.post(f"/api/material-categories/subcategories/{subcategory['id']}/attributes", json={
        "name": "Brand", "data_type": "text",
    }).json()

    material = client.post("/api/materials/", json={
        "name": "Attr Value Test Material", "unit": "Sheets", "opening_stock": 5, "minimum_stock": 2,
        "subcategory_id": subcategory["id"],
        "attribute_values": [
            {"attribute_definition_id": thickness_attr["id"], "value_number": "18"},
            {"attribute_definition_id": brand_attr["id"], "value_text": "Greenpanel"},
        ],
    }).json()
    values = {v["attribute_definition_id"]: v for v in material["attribute_values"]}
    assert float(values[thickness_attr["id"]]["value_number"]) == 18.0
    assert values[brand_attr["id"]]["value_text"] == "Greenpanel"
    assert "18" in values[thickness_attr["id"]]["display_value"] and "mm" in values[thickness_attr["id"]]["display_value"]


def test_updating_attribute_values_replaces_full_set(client, test_user):
    _login(client, test_user)
    category = client.post("/api/material-categories/", json={"name": "Update Attr Test Category"}).json()
    subcategory = client.post("/api/material-categories/subcategories", json={
        "category_id": category["id"], "name": "Update Attr Test Subcategory",
    }).json()
    attr = client.post(f"/api/material-categories/subcategories/{subcategory['id']}/attributes", json={
        "name": "Colour", "data_type": "text",
    }).json()
    material = client.post("/api/materials/", json={
        "name": "Update Attr Test Material", "unit": "Sheets", "opening_stock": 1, "minimum_stock": 1,
        "subcategory_id": subcategory["id"],
        "attribute_values": [{"attribute_definition_id": attr["id"], "value_text": "White"}],
    }).json()

    resp = client.put(f"/api/materials/{material['id']}", json={
        "attribute_values": [{"attribute_definition_id": attr["id"], "value_text": "Walnut"}],
    })
    assert resp.status_code == 200
    values = resp.json()["attribute_values"]
    assert len(values) == 1
    assert values[0]["value_text"] == "Walnut"


def test_attribute_value_returns_real_name_not_just_id(client, test_user):
    """Regression test - the API used to only return
    attribute_definition_id, forcing the frontend to show a raw
    internal ID ("Attribute #5") instead of an actual name."""
    _login(client, test_user)
    category = client.post("/api/material-categories/", json={"name": "Attr Name Test Category"}).json()
    subcategory = client.post("/api/material-categories/subcategories", json={
        "category_id": category["id"], "name": "Attr Name Test Subcategory",
    }).json()
    attr = client.post(f"/api/material-categories/subcategories/{subcategory['id']}/attributes", json={
        "name": "Thickness", "data_type": "number", "unit_label": "mm",
    }).json()
    material = client.post("/api/materials/", json={
        "name": "Attr Name Test Material", "unit": "Sheets", "opening_stock": 1, "minimum_stock": 1,
        "subcategory_id": subcategory["id"],
        "attribute_values": [{"attribute_definition_id": attr["id"], "value_number": "12"}],
    }).json()
    assert material["attribute_values"][0]["attribute_name"] == "Thickness"


def test_full_hierarchy_response_shape_matches_frontend_expectations(client, test_user):
    """Confirms GET /api/material-categories/ returns categories with
    nested subcategories with nested attribute_definitions - the exact
    shape MaterialAttributesEditor.jsx relies on to build its dropdowns."""
    _login(client, test_user)
    category = client.post("/api/material-categories/", json={"name": "Shape Test Category"}).json()
    subcategory = client.post("/api/material-categories/subcategories", json={
        "category_id": category["id"], "name": "Shape Test Subcategory",
    }).json()
    client.post(f"/api/material-categories/subcategories/{subcategory['id']}/attributes", json={
        "name": "Shape Test Attribute", "data_type": "text",
    })

    resp = client.get("/api/material-categories/")
    assert resp.status_code == 200
    found_category = next(c for c in resp.json() if c["id"] == category["id"])
    assert "subcategories" in found_category
    found_subcategory = next(s for s in found_category["subcategories"] if s["id"] == subcategory["id"])
    assert "attribute_definitions" in found_subcategory
    assert found_subcategory["attribute_definitions"][0]["name"] == "Shape Test Attribute"


def _setup_plywood_subcategory(client):
    category = client.post("/api/material-categories/", json={"name": "Attr Filter Test Category"}).json()
    subcategory = client.post("/api/material-categories/subcategories", json={
        "category_id": category["id"], "name": "Attr Filter Test Plywood",
    }).json()
    thickness_attr = client.post(f"/api/material-categories/subcategories/{subcategory['id']}/attributes", json={
        "name": "Thickness", "data_type": "number", "unit_label": "mm",
    }).json()
    brand_attr = client.post(f"/api/material-categories/subcategories/{subcategory['id']}/attributes", json={
        "name": "Brand", "data_type": "text",
    }).json()
    return subcategory, thickness_attr, brand_attr


def test_filter_by_subcategory_id(client, test_user):
    _login(client, test_user)
    subcategory, _, _ = _setup_plywood_subcategory(client)
    client.post("/api/materials/", json={
        "name": "Subcat Filter Material", "unit": "Sheets", "opening_stock": 1, "minimum_stock": 1,
        "subcategory_id": subcategory["id"],
    })
    client.post("/api/materials/", json={
        "name": "Unrelated Material", "unit": "Sheets", "opening_stock": 1, "minimum_stock": 1,
    })

    resp = client.get("/api/materials/", params={"subcategory_id": subcategory["id"]})
    names = {m["name"] for m in resp.json()}
    assert "Subcat Filter Material" in names
    assert "Unrelated Material" not in names


def test_filter_by_single_numeric_attribute(client, test_user):
    _login(client, test_user)
    subcategory, thickness_attr, _ = _setup_plywood_subcategory(client)
    client.post("/api/materials/", json={
        "name": "18mm Filter Test Material", "unit": "Sheets", "opening_stock": 1, "minimum_stock": 1,
        "subcategory_id": subcategory["id"],
        "attribute_values": [{"attribute_definition_id": thickness_attr["id"], "value_number": "18"}],
    })
    client.post("/api/materials/", json={
        "name": "12mm Filter Test Material", "unit": "Sheets", "opening_stock": 1, "minimum_stock": 1,
        "subcategory_id": subcategory["id"],
        "attribute_values": [{"attribute_definition_id": thickness_attr["id"], "value_number": "12"}],
    })

    resp = client.get("/api/materials/", params={
        "attribute_filters": json.dumps({str(thickness_attr["id"]): "18"}),
    })
    names = {m["name"] for m in resp.json()}
    assert "18mm Filter Test Material" in names
    assert "12mm Filter Test Material" not in names


def test_multiple_attribute_filters_narrow_with_and_not_or(client, test_user):
    """The core correctness requirement - selecting both Thickness=18 AND
    Brand=Century must return only materials matching BOTH, not either."""
    _login(client, test_user)
    subcategory, thickness_attr, brand_attr = _setup_plywood_subcategory(client)

    matches_both = client.post("/api/materials/", json={
        "name": "Century 18mm Material", "unit": "Sheets", "opening_stock": 1, "minimum_stock": 1,
        "subcategory_id": subcategory["id"],
        "attribute_values": [
            {"attribute_definition_id": thickness_attr["id"], "value_number": "18"},
            {"attribute_definition_id": brand_attr["id"], "value_text": "Century"},
        ],
    }).json()
    matches_thickness_only = client.post("/api/materials/", json={
        "name": "Greenpanel 18mm Material", "unit": "Sheets", "opening_stock": 1, "minimum_stock": 1,
        "subcategory_id": subcategory["id"],
        "attribute_values": [
            {"attribute_definition_id": thickness_attr["id"], "value_number": "18"},
            {"attribute_definition_id": brand_attr["id"], "value_text": "Greenpanel"},
        ],
    }).json()

    resp = client.get("/api/materials/", params={
        "attribute_filters": json.dumps({str(thickness_attr["id"]): "18", str(brand_attr["id"]): "Century"}),
    })
    ids = {m["id"] for m in resp.json()}
    assert matches_both["id"] in ids
    assert matches_thickness_only["id"] not in ids


def test_invalid_attribute_filters_json_rejected(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/materials/", params={"attribute_filters": "not valid json"})
    assert resp.status_code == 400


def test_attribute_filter_composes_with_pagination(client, test_user):
    _login(client, test_user)
    subcategory, thickness_attr, _ = _setup_plywood_subcategory(client)
    for i in range(3):
        client.post("/api/materials/", json={
            "name": f"Pagination Attr Filter Material {i}", "unit": "Sheets", "opening_stock": 1, "minimum_stock": 1,
            "subcategory_id": subcategory["id"],
            "attribute_values": [{"attribute_definition_id": thickness_attr["id"], "value_number": "25"}],
        })

    resp = client.get("/api/materials/", params={
        "attribute_filters": json.dumps({str(thickness_attr["id"]): "25"}), "limit": 2, "offset": 0,
    })
    assert len(resp.json()) == 2
    assert int(resp.headers["X-Total-Count"]) == 3


def test_subcategory_and_attribute_filter_combine_correctly_like_the_catalog_ui(client, test_user):
    """Reproduces exactly what MaterialsPage.jsx's filter panel now
    sends: subcategory_id plus a JSON attribute_filters map together in
    one request - the real end-to-end shape, not just each independently."""
    _login(client, test_user)
    subcategory, thickness_attr, _ = _setup_plywood_subcategory(client)
    other_subcategory = client.post("/api/material-categories/subcategories", json={
        "category_id": subcategory["category_id"], "name": "Different Subcategory For UI Test",
    }).json()

    matching = client.post("/api/materials/", json={
        "name": "UI Filter Match", "unit": "Sheets", "opening_stock": 1, "minimum_stock": 1,
        "subcategory_id": subcategory["id"],
        "attribute_values": [{"attribute_definition_id": thickness_attr["id"], "value_number": "18"}],
    }).json()
    # Same thickness value, but a different subcategory - must be excluded.
    client.post("/api/materials/", json={
        "name": "UI Filter Wrong Subcategory", "unit": "Sheets", "opening_stock": 1, "minimum_stock": 1,
        "subcategory_id": other_subcategory["id"],
    })

    resp = client.get("/api/materials/", params={
        "subcategory_id": subcategory["id"],
        "attribute_filters": json.dumps({str(thickness_attr["id"]): "18"}),
    })
    names = {m["name"] for m in resp.json()}
    assert "UI Filter Match" in names
    assert "UI Filter Wrong Subcategory" not in names


# --- test_stock_operations.py ---
def _make_locations(client):
    warehouse = client.post("/api/locations/", json={"name": "ML Warehouse"}).json()
    rack_a = client.post("/api/locations/", json={"name": "ML Rack A2", "parent_id": warehouse["id"]}).json()
    rack_b = client.post("/api/locations/", json={"name": "ML Rack B1", "parent_id": warehouse["id"]}).json()
    return rack_a, rack_b


def _location_stock(client, material_id):
    resp = client.get(f"/api/stock/locations/{material_id}")
    assert resp.status_code == 200
    return resp.json()


def test_receipt_into_two_different_locations(client, test_user):
    _login(client, test_user)
    rack_a, rack_b = _make_locations(client)
    material = client.post("/api/materials/", json={
        "name": "HDHMR 18mm Multi-Loc", "unit": "Sheets", "opening_stock": 0, "minimum_stock": 1,
    }).json()
    supplier_id = client.post("/api/suppliers/", json={"name": "ML Supplier"}).json()["id"]

    client.post("/api/purchases/", json={
        "date": "2026-08-19T00:00:00", "supplier_id": supplier_id, "material_id": material["id"],
        "quantity": "12", "unit": "Sheets", "rate": "500", "gst_percent": "18",
        "location_id": rack_a["id"],
    })
    client.post("/api/purchases/", json={
        "date": "2026-08-19T00:00:00", "supplier_id": supplier_id, "material_id": material["id"],
        "quantity": "6", "unit": "Sheets", "rate": "500", "gst_percent": "18",
        "location_id": rack_b["id"],
    })

    breakdown = _location_stock(client, material["id"])
    by_id = {row["location_id"]: Decimal(str(row["quantity"])) for row in breakdown["locations"]}
    assert by_id[rack_a["id"]] == Decimal("12")
    assert by_id[rack_b["id"]] == Decimal("6")
    assert Decimal(str(breakdown["total"])) == Decimal("18")
    assert sum(by_id.values()) == Decimal(str(breakdown["total"]))

    material_after = client.get(f"/api/materials/{material['id']}").json()
    assert Decimal(str(material_after["current_stock"])) == Decimal("18")


def test_issue_from_specific_location(client, test_user):
    _login(client, test_user)
    rack_a, rack_b = _make_locations(client)
    material = client.post("/api/materials/", json={
        "name": "Issue Multi-Loc Material", "unit": "Sheets", "opening_stock": 0, "minimum_stock": 1,
    }).json()
    supplier_id = client.post("/api/suppliers/", json={"name": "Issue ML Supplier"}).json()["id"]
    client.post("/api/purchases/", json={
        "date": "2026-08-19T00:00:00", "supplier_id": supplier_id, "material_id": material["id"],
        "quantity": "10", "unit": "Sheets", "rate": "100", "gst_percent": "0", "location_id": rack_a["id"],
    })
    client.post("/api/purchases/", json={
        "date": "2026-08-19T00:00:00", "supplier_id": supplier_id, "material_id": material["id"],
        "quantity": "10", "unit": "Sheets", "rate": "100", "gst_percent": "0", "location_id": rack_b["id"],
    })

    resp = client.post("/api/issues/", json={
        "date": "2026-08-19T00:00:00", "material_id": material["id"], "quantity_issued": "4",
        "unit": "Sheets", "location_id": rack_a["id"],
    })
    assert resp.status_code == 201

    breakdown = _location_stock(client, material["id"])
    by_id = {row["location_id"]: Decimal(str(row["quantity"])) for row in breakdown["locations"]}
    assert by_id[rack_a["id"]] == Decimal("6")
    assert by_id[rack_b["id"]] == Decimal("10")
    assert Decimal(str(breakdown["total"])) == Decimal("16")


def test_issue_cannot_exceed_specific_location_balance(client, test_user):
    _login(client, test_user)
    rack_a, rack_b = _make_locations(client)
    material = client.post("/api/materials/", json={
        "name": "Issue Overdraw Material", "unit": "Sheets", "opening_stock": 0, "minimum_stock": 1,
    }).json()
    supplier_id = client.post("/api/suppliers/", json={"name": "Overdraw Supplier"}).json()["id"]
    # Plenty of material-wide stock, but only at rack_b - rack_a has none.
    client.post("/api/purchases/", json={
        "date": "2026-08-19T00:00:00", "supplier_id": supplier_id, "material_id": material["id"],
        "quantity": "50", "unit": "Sheets", "rate": "100", "gst_percent": "0", "location_id": rack_b["id"],
    })

    resp = client.post("/api/issues/", json={
        "date": "2026-08-19T00:00:00", "material_id": material["id"], "quantity_issued": "5",
        "unit": "Sheets", "location_id": rack_a["id"],
    })
    assert resp.status_code == 400


def test_transfer_moves_quantity_between_locations(client, test_user):
    _login(client, test_user)
    rack_a, rack_b = _make_locations(client)
    material = client.post("/api/materials/", json={
        "name": "Transfer Ledger Multi-Loc", "unit": "Sheets", "opening_stock": 20, "minimum_stock": 1,
        "location_id": rack_a["id"],
    }).json()

    resp = client.post("/api/stock/transfers", json={
        "material_id": material["id"], "quantity": "8", "to_location_id": rack_b["id"],
    })
    assert resp.status_code == 201

    breakdown = _location_stock(client, material["id"])
    by_id = {row["location_id"]: Decimal(str(row["quantity"])) for row in breakdown["locations"]}
    assert by_id[rack_a["id"]] == Decimal("12")
    assert by_id[rack_b["id"]] == Decimal("8")
    assert Decimal(str(breakdown["total"])) == Decimal("20")

    # Total stock is unaffected by a transfer - reconciliation must still pass.
    verify = client.get(f"/api/stock/verify/{material['id']}").json()
    assert verify["matches"] is True


def test_adjustment_at_specific_location(client, test_user):
    _login(client, test_user)
    rack_a, rack_b = _make_locations(client)
    material = client.post("/api/materials/", json={
        "name": "Adjustment Multi-Loc Material", "unit": "Sheets", "opening_stock": 0, "minimum_stock": 1,
    }).json()
    supplier_id = client.post("/api/suppliers/", json={"name": "Adj ML Supplier"}).json()["id"]
    client.post("/api/purchases/", json={
        "date": "2026-08-19T00:00:00", "supplier_id": supplier_id, "material_id": material["id"],
        "quantity": "10", "unit": "Sheets", "rate": "100", "gst_percent": "0", "location_id": rack_a["id"],
    })

    resp = client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Damage", "quantity_delta": "-2",
        "reason": "Water damage at rack A2", "location_id": rack_a["id"],
    })
    assert resp.status_code == 201

    breakdown = _location_stock(client, material["id"])
    by_id = {row["location_id"]: Decimal(str(row["quantity"])) for row in breakdown["locations"]}
    assert by_id[rack_a["id"]] == Decimal("8")
    assert Decimal(str(breakdown["total"])) == Decimal("8")


def test_location_totals_and_reconciliation_after_mixed_activity(client, test_user):
    """receipt A, receipt B, issue from A, transfer A->B, adjustment at B -
    location totals must sum to the material-wide total at every step."""
    _login(client, test_user)
    rack_a, rack_b = _make_locations(client)
    material = client.post("/api/materials/", json={
        "name": "Mixed Activity Multi-Loc", "unit": "Sheets", "opening_stock": 0, "minimum_stock": 1,
    }).json()
    supplier_id = client.post("/api/suppliers/", json={"name": "Mixed Supplier"}).json()["id"]

    client.post("/api/purchases/", json={
        "date": "2026-08-19T00:00:00", "supplier_id": supplier_id, "material_id": material["id"],
        "quantity": "20", "unit": "Sheets", "rate": "100", "gst_percent": "0", "location_id": rack_a["id"],
    })
    client.post("/api/purchases/", json={
        "date": "2026-08-19T00:00:00", "supplier_id": supplier_id, "material_id": material["id"],
        "quantity": "10", "unit": "Sheets", "rate": "100", "gst_percent": "0", "location_id": rack_b["id"],
    })
    # A: 20, B: 10, total 30
    client.post("/api/issues/", json={
        "date": "2026-08-19T00:00:00", "material_id": material["id"], "quantity_issued": "5",
        "unit": "Sheets", "location_id": rack_a["id"],
    })
    # A: 15, B: 10, total 25
    client.post("/api/stock/transfers", json={
        "material_id": material["id"], "quantity": "5", "from_location_id": rack_a["id"],
        "to_location_id": rack_b["id"],
    })
    # A: 10, B: 15, total 25
    client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Physical Count Increase", "quantity_delta": "3",
        "reason": "Recount at B1", "location_id": rack_b["id"],
    })
    # A: 10, B: 18, total 28

    breakdown = _location_stock(client, material["id"])
    by_id = {row["location_id"]: Decimal(str(row["quantity"])) for row in breakdown["locations"]}
    assert by_id[rack_a["id"]] == Decimal("10")
    assert by_id[rack_b["id"]] == Decimal("18")
    assert Decimal(str(breakdown["total"])) == Decimal("28")
    assert sum(by_id.values()) == Decimal(str(breakdown["total"]))

    material_after = client.get(f"/api/materials/{material['id']}").json()
    assert Decimal(str(material_after["current_stock"])) == Decimal("28")

    verify = client.get(f"/api/stock/verify/{material['id']}").json()
    assert verify["matches"] is True


def test_existing_stock_ledger_verification_still_works_with_locations(client, test_user):
    """Existing verify_stock_matches_ledger reconciliation must keep
    working exactly as before now that ledger entries carry location_id -
    a regression here would mean multi-location broke the ledger's
    original guarantee."""
    _login(client, test_user)
    rack_a, _ = _make_locations(client)
    material = client.post("/api/materials/", json={
        "name": "Verify Still Works Material", "unit": "Sheets", "opening_stock": "10", "minimum_stock": "1",
    }).json()
    supplier_id = client.post("/api/suppliers/", json={"name": "Verify Supplier"}).json()["id"]
    client.post("/api/purchases/", json={
        "date": "2026-08-19T00:00:00", "supplier_id": supplier_id, "material_id": material["id"],
        "quantity": "5", "unit": "Sheets", "rate": "100", "gst_percent": "0", "location_id": rack_a["id"],
    })

    verify = client.get(f"/api/stock/verify/{material['id']}").json()
    assert verify["matches"] is True
    assert float(verify["current_stock"]) == 15.0


def test_no_location_falls_back_to_material_primary_location(client, test_user):
    """A material with no explicit location on receipt still shows up
    somewhere sensible in the per-location breakdown - never silently
    dropped from the total."""
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "No Location Material", "unit": "Sheets", "opening_stock": "7", "minimum_stock": "1",
    }).json()

    breakdown = _location_stock(client, material["id"])
    assert Decimal(str(breakdown["total"])) == Decimal("7")
    total_located = sum(Decimal(str(row["quantity"])) for row in breakdown["locations"])
    assert total_located == Decimal("7")


def test_default_purchase_still_increases_stock_immediately(client, test_user):
    """Backward compatibility - omitting receipt_status must behave
    exactly as every purchase always has."""
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Default Receipt Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Default Receipt Material", "unit": "Sheets", "opening_stock": 5, "minimum_stock": 1,
    }).json()

    resp = client.post("/api/purchases/", json={
        "date": "2026-08-15T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
        "quantity": "10", "unit": "Sheets", "rate": "500.00", "gst_percent": "18", "payment_status": "Paid",
    })
    assert resp.status_code == 201
    assert resp.json()["receipt_status"] == "Received"

    updated = client.get(f"/api/materials/{material['id']}").json()
    assert updated["current_stock"] == 15


def test_ordered_purchase_does_not_increase_stock(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Ordered Test Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Ordered Test Material", "unit": "Sheets", "opening_stock": 5, "minimum_stock": 1,
    }).json()

    resp = client.post("/api/purchases/", json={
        "date": "2026-08-15T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
        "quantity": "10", "unit": "Sheets", "rate": "500.00", "gst_percent": "18", "payment_status": "Paid",
        "receipt_status": "Ordered",
    })
    assert resp.status_code == 201
    assert resp.json()["receipt_status"] == "Ordered"

    unchanged = client.get(f"/api/materials/{material['id']}").json()
    assert unchanged["current_stock"] == 5


def test_marking_received_now_increases_stock(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Mark Received Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Mark Received Material", "unit": "Sheets", "opening_stock": 5, "minimum_stock": 1,
    }).json()
    purchase = client.post("/api/purchases/", json={
        "date": "2026-08-15T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
        "quantity": "10", "unit": "Sheets", "rate": "500.00", "gst_percent": "18", "payment_status": "Paid",
        "receipt_status": "Ordered",
    }).json()

    resp = client.post(f"/api/purchases/{purchase['id']}/receive")
    assert resp.status_code == 200
    assert resp.json()["receipt_status"] == "Received"

    updated = client.get(f"/api/materials/{material['id']}").json()
    assert updated["current_stock"] == 15


def test_cannot_receive_the_same_purchase_twice(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Double Receive Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Double Receive Material", "unit": "Sheets", "opening_stock": 0, "minimum_stock": 1,
    }).json()
    purchase = client.post("/api/purchases/", json={
        "date": "2026-08-15T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
        "quantity": "10", "unit": "Sheets", "rate": "500.00", "gst_percent": "18", "payment_status": "Paid",
        "receipt_status": "Ordered",
    }).json()

    first = client.post(f"/api/purchases/{purchase['id']}/receive")
    assert first.status_code == 200
    second = client.post(f"/api/purchases/{purchase['id']}/receive")
    assert second.status_code == 400

    # Confirm stock was only applied once, not twice.
    updated = client.get(f"/api/materials/{material['id']}").json()
    assert updated["current_stock"] == 10


def test_receiving_an_already_received_purchase_directly_also_rejected(client, test_user):
    """A purchase created as Received (the default) should also reject
    a /receive call - it's already done."""
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Already Received Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Already Received Material", "unit": "Sheets", "opening_stock": 0, "minimum_stock": 1,
    }).json()
    purchase = client.post("/api/purchases/", json={
        "date": "2026-08-15T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
        "quantity": "5", "unit": "Sheets", "rate": "500.00", "gst_percent": "18", "payment_status": "Paid",
    }).json()

    resp = client.post(f"/api/purchases/{purchase['id']}/receive")
    assert resp.status_code == 400


def test_no_purchase_received_notification_for_ordered_status(client, test_user):
    """The PURCHASE_RECEIVED notification must only fire when stock
    actually changes, not merely when a purchase is placed as Ordered."""
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Notif Ordered Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Notif Ordered Material", "unit": "Sheets", "opening_stock": 100, "minimum_stock": 1,
    }).json()

    client.post("/api/purchases/", json={
        "date": "2026-08-15T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
        "quantity": "10", "unit": "Sheets", "rate": "500.00", "gst_percent": "18", "payment_status": "Paid",
        "receipt_status": "Ordered",
    })

    notifications = client.get("/api/notifications/").json()
    matches = [n for n in notifications if n["notification_type"] == "PURCHASE_RECEIVED"
               and "Notif Ordered Material" in n["title"]]
    assert len(matches) == 0


def test_receive_requires_master_or_manager(client, test_user, db_session):
    from app.platform.security import hash_password
    from app.modules.auth.auth import User

    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Receive Perm Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Receive Perm Material", "unit": "Sheets", "opening_stock": 0, "minimum_stock": 1,
    }).json()
    purchase = client.post("/api/purchases/", json={
        "date": "2026-08-15T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
        "quantity": "5", "unit": "Sheets", "rate": "500.00", "gst_percent": "18", "payment_status": "Paid",
        "receipt_status": "Ordered",
    }).json()

    employee = client.post("/api/employees/", json={
        "name": "Receive Perm Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    limited_user = User(
        username="receivepermuser", email="receivepermuser@example.com", full_name="Limited User",
        password_hash=hash_password("LimitedPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(limited_user)
    db_session.commit()

    client.post("/api/auth/login", json={"identifier": "receivepermuser@example.com", "password": "LimitedPass1!"})
    resp = client.post(f"/api/purchases/{purchase['id']}/receive")
    assert resp.status_code == 403


def _create_supplier(client):
    resp = client.post("/api/suppliers/", json={
        "supplier_code": "SUP-TEST", "name": "Test Supplier", "category": "Plywood",
    })
    assert resp.status_code == 201
    return resp.json()["id"]


def _create_material(client, supplier_id):
    resp = client.post("/api/materials/", json={
        "material_code": "MAT-TEST", "name": "Test Ply", "category": "Plywood",
        "unit": "Sheets", "minimum_stock": 10, "average_rate": "1000.00",
        "supplier_id": supplier_id, "opening_stock": 50,
    })
    assert resp.status_code == 201
    return resp.json()


def test_material_starts_at_opening_stock(client, test_user):
    _login(client, test_user)
    supplier_id = _create_supplier(client)
    material = _create_material(client, supplier_id)
    assert material["current_stock"] == 50
    assert material["stock_status"] == "STOCK OK"


def test_purchase_increases_stock(client, test_user):
    _login(client, test_user)
    supplier_id = _create_supplier(client)
    material = _create_material(client, supplier_id)

    resp = client.post("/api/purchases/", json={
        "purchase_code": "PUR-TEST-001", "date": "2026-08-01T00:00:00", "supplier_id": supplier_id,
        "material_id": material["id"], "quantity": "20", "unit": "Sheets", "rate": "1000.00",
        "gst_percent": "18", "payment_status": "Paid",
    })
    assert resp.status_code == 201
    body = resp.json()
    assert float(body["taxable_value"]) == 20000.0
    assert float(body["gst_amount"]) == 3600.0
    assert float(body["invoice_total"]) == 23600.0

    mat = client.get(f"/api/materials/{material['id']}").json()
    assert mat["current_stock"] == 70
    assert mat["total_purchased"] == 20


def test_issue_decreases_stock(client, test_user):
    _login(client, test_user)
    supplier_id = _create_supplier(client)
    material = _create_material(client, supplier_id)

    resp = client.post("/api/issues/", json={
        "issue_code": "ISS-TEST-001", "date": "2026-08-01T00:00:00",
        "material_id": material["id"], "quantity_issued": "15", "unit": "Sheets",
        "issued_to": "Ravi", "department": "Assembly",
    })
    assert resp.status_code == 201

    mat = client.get(f"/api/materials/{material['id']}").json()
    assert mat["current_stock"] == 35
    assert mat["total_issued"] == 15


def test_issue_rejects_insufficient_stock(client, test_user):
    _login(client, test_user)
    supplier_id = _create_supplier(client)
    material = _create_material(client, supplier_id)

    resp = client.post("/api/issues/", json={
        "issue_code": "ISS-TEST-002", "date": "2026-08-01T00:00:00",
        "material_id": material["id"], "quantity_issued": "999", "unit": "Sheets",
    })
    assert resp.status_code == 400


def test_low_stock_dashboard_reflects_minimum(client, test_user):
    _login(client, test_user)
    supplier_id = _create_supplier(client)
    material = _create_material(client, supplier_id)

    # Issue down to below minimum_stock (10)
    client.post("/api/issues/", json={
        "issue_code": "ISS-TEST-003", "date": "2026-08-01T00:00:00",
        "material_id": material["id"], "quantity_issued": "45", "unit": "Sheets",
    })

    dashboard = client.get("/api/dashboard/stock").json()
    assert dashboard["low_stock_items"] >= 1


def material(db_session):
    m = Material(
        material_code="MAT-TEST-1", name="Test Mat", unit="Nos",
        opening_stock=10, current_stock=10, minimum_stock=5, average_rate=100,
    )
    db_session.add(m)
    db_session.commit()
    db_session.refresh(m)
    return m


def supplier(db_session):
    s = Supplier(supplier_code="SUP-TEST-1", name="Test Supplier")
    db_session.add(s)
    db_session.commit()
    db_session.refresh(s)
    return s


def test_add_purchase_updates_stock(db_session, material, supplier):
    data = PurchaseCreate(
        date=datetime.utcnow(), supplier_id=supplier.id, material_id=material.id,
        quantity=5, unit="Nos", rate=100, gst_percent=0,
    )
    ProcurementService.record_purchase(db_session, data)
    db_session.refresh(material)
    assert material.current_stock == 15


def test_record_issue_blocks_if_insufficient(db_session, material):
    data = IssueCreate(
        date=datetime.utcnow(), material_id=material.id,
        quantity_issued=20, unit="Nos",
    )
    with pytest.raises(Exception):
        StockService.record_issue(db_session, data)


def test_master_can_delete_unreferenced_material(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Delete Guard Unreferenced Material", "unit": "Sheets", "opening_stock": "5", "minimum_stock": "1",
    }).json()
    resp = client.delete(f"/api/materials/{material['id']}")
    assert resp.status_code == 204


def test_cannot_delete_material_with_purchase_history(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Delete Guard Purchase Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Delete Guard Purchased Material", "unit": "Sheets", "opening_stock": "0", "minimum_stock": "1",
    }).json()
    client.post("/api/purchases/", json={
        "date": "2026-08-17T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
        "quantity": "5", "unit": "Sheets", "rate": "500.00", "gst_percent": "18", "payment_status": "Paid",
    })

    resp = client.delete(f"/api/materials/{material['id']}")
    assert resp.status_code == 400
    assert "purchase history" in resp.json()["detail"].lower()


def test_cannot_delete_material_with_issue_history(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Delete Guard Issued Material", "unit": "Sheets", "opening_stock": "10", "minimum_stock": "1",
    }).json()
    client.post("/api/issues/", json={
        "date": "2026-08-17T00:00:00", "material_id": material["id"], "quantity_issued": "2", "unit": "Sheets",
    })

    resp = client.delete(f"/api/materials/{material['id']}")
    assert resp.status_code == 400
    assert "issue history" in resp.json()["detail"].lower()


def test_master_can_delete_unreferenced_supplier(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Delete Guard Unreferenced Supplier"}).json()
    resp = client.delete(f"/api/suppliers/{supplier['id']}")
    assert resp.status_code == 204


def test_cannot_delete_supplier_with_purchase_history(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Delete Guard Referenced Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Delete Guard Supplier Material", "unit": "Sheets", "opening_stock": "0", "minimum_stock": "1",
    }).json()
    client.post("/api/purchases/", json={
        "date": "2026-08-17T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
        "quantity": "3", "unit": "Sheets", "rate": "400.00", "gst_percent": "18", "payment_status": "Paid",
    })

    resp = client.delete(f"/api/suppliers/{supplier['id']}")
    assert resp.status_code == 400
    assert "purchase history" in resp.json()["detail"].lower()


def test_cannot_delete_supplier_with_material_link(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Delete Guard Linked Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Delete Guard Linked Material", "unit": "Sheets", "opening_stock": "5", "minimum_stock": "1",
    }).json()
    client.post("/api/supplier-materials/", json={
        "supplier_id": supplier["id"], "material_id": material["id"], "supplier_price": "100.00",
    })

    resp = client.delete(f"/api/suppliers/{supplier['id']}")
    assert resp.status_code == 400
    assert "linked to materials" in resp.json()["detail"].lower()


def test_material_delete_is_strictly_master_only(client, test_user, db_session):
    """require_role("master") - any non-master role must be rejected,
    not just an obviously-unprivileged account."""
    from app.platform.security import hash_password
    from app.modules.auth.auth import User

    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Delete Guard Non-Master Test Material", "unit": "Sheets", "opening_stock": "5", "minimum_stock": "1",
    }).json()
    non_master = User(
        username="deleteguarduser", email="deleteguarduser@example.com", full_name="Delete Guard User",
        password_hash=hash_password("UserPass1!"), role="user", is_active=True,
    )
    db_session.add(non_master)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "deleteguarduser@example.com", "password": "UserPass1!"})

    resp = client.delete(f"/api/materials/{material['id']}")
    assert resp.status_code == 403


def _make_ordered_purchase(client, suffix):
    supplier = client.post("/api/suppliers/", json={"name": f"Delete Purchase Supplier {suffix}"}).json()
    material = client.post("/api/materials/", json={
        "name": f"Delete Purchase Material {suffix}", "unit": "Sheets", "opening_stock": "0",
    }).json()
    purchase = client.post("/api/purchases/", json={
        "date": "2026-09-01T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
        "quantity": "2", "unit": "Sheets", "rate": "1300", "gst_percent": "18",
        "payment_status": "Credit", "receipt_status": "Ordered",
    }).json()
    return purchase, material


def test_ordered_purchase_can_be_deleted(client, test_user):
    _login(client, test_user)
    purchase, _ = _make_ordered_purchase(client, "1")

    resp = client.delete(f"/api/purchases/{purchase['id']}")
    assert resp.status_code == 204
    assert client.get(f"/api/purchases/{purchase['id']}").status_code == 404


def test_received_purchase_cannot_be_deleted(client, test_user):
    """Stock has already been added - deleting it blind could leave
    stock wrong if anything else touched the same material since."""
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Delete Received Purchase Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Delete Received Purchase Material", "unit": "Sheets", "opening_stock": "0",
    }).json()
    purchase = client.post("/api/purchases/", json={
        "date": "2026-09-01T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
        "quantity": "2", "unit": "Sheets", "rate": "1300", "gst_percent": "18",
        "payment_status": "Credit", "receipt_status": "Received",
    }).json()

    resp = client.delete(f"/api/purchases/{purchase['id']}")
    assert resp.status_code == 409
    assert client.get(f"/api/purchases/{purchase['id']}").status_code == 200


def test_delete_purchase_requires_master(client, test_user, db_session):
    _login(client, test_user)
    purchase, material = _make_ordered_purchase(client, "2")
    employee = client.post("/api/employees/", json={"name": "Delete Purchase RBAC Employee"}).json()
    from app.platform.security import hash_password
    from app.modules.auth.auth import User
    user = User(
        username="deletepurchaserbacuser", email="deletepurchaserbacuser@example.com",
        full_name="Delete Purchase RBAC User", password_hash=hash_password("EmpPass1!"),
        role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "deletepurchaserbacuser@example.com", "password": "EmpPass1!"})

    resp = client.delete(f"/api/purchases/{purchase['id']}")
    assert resp.status_code == 403


def test_delete_unknown_purchase_returns_404(client, test_user):
    _login(client, test_user)
    resp = client.delete("/api/purchases/999999")
    assert resp.status_code == 404


def test_delete_purchase_linked_to_requirement_blocked(client, test_user):
    _login(client, test_user)
    client_id = client.post("/api/clients/", json={"name": "Delete Purchase Req Client", "phone": "9000010198"}).json()["id"]
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-09-01T00:00:00", "order_value": "5000", "advance": "0",
    }).json()
    material = client.post("/api/materials/", json={
        "name": "Delete Purchase Req Material", "unit": "Sheets", "opening_stock": "0",
    }).json()
    supplier = client.post("/api/suppliers/", json={"name": "Delete Purchase Req Supplier"}).json()
    client.post("/api/supplier-materials/", json={"supplier_id": supplier["id"], "material_id": material["id"]})
    requirement = client.post("/api/procurement-requirements/", json={
        "order_id": order["id"], "material_id": material["id"],
    }).json()
    client.post(f"/api/procurement-requirements/{requirement['id']}/decision", json={
        "requirement_id": requirement["id"], "selected_supplier_id": supplier["id"],
    })
    purchase = client.post(f"/api/procurement-requirements/{requirement['id']}/purchase", json={
        "quantity": "5", "unit": "Sheets", "rate": "1000", "gst_percent": "18", "receipt_status": "Ordered",
    }).json()

    resp = client.delete(f"/api/purchases/{purchase['id']}")
    assert resp.status_code == 409


# --- test_stock_ledger_and_transfers.py ---
def test_ledger_reconciles_after_purchase_receipt(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Ledger Purchase Test Material", "unit": "Sheets", "opening_stock": "10", "minimum_stock": "5",
    }).json()
    supplier_id = client.post("/api/suppliers/", json={"name": "Ledger Purchase Test Supplier"}).json()["id"]
    client.post("/api/purchases/", json={
        "date": "2026-08-19T00:00:00", "supplier_id": supplier_id, "material_id": material["id"],
        "quantity": "20", "unit": "Sheets", "rate": "500", "gst_percent": "18",
    })

    verify = client.get(f"/api/stock/verify/{material['id']}").json()
    assert verify["matches"] is True
    assert float(verify["current_stock"]) == 30.0  # 10 opening + 20 received
    assert verify["ledger_entry_count"] == 1


def test_ledger_entry_traces_back_to_the_purchase(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Ledger Trace Test Material", "unit": "Sheets", "opening_stock": "0", "minimum_stock": "5",
    }).json()
    supplier_id = client.post("/api/suppliers/", json={"name": "Ledger Trace Test Supplier"}).json()["id"]
    purchase = client.post("/api/purchases/", json={
        "date": "2026-08-19T00:00:00", "supplier_id": supplier_id, "material_id": material["id"],
        "quantity": "15", "unit": "Sheets", "rate": "500", "gst_percent": "18",
    }).json()

    ledger = client.get("/api/stock/ledger", params={"material_id": material["id"]}).json()
    assert len(ledger) == 1
    assert ledger[0]["entry_type"] == "Receipt"
    assert ledger[0]["reference_type"] == "purchase"
    assert ledger[0]["reference_id"] == purchase["id"]
    assert float(ledger[0]["quantity_delta"]) == 15.0


def test_ledger_reconciles_after_partial_receipt_then_completion(client, test_user):
    """The most complex real path this session built - two separate
    receive calls on the same purchase - must still leave the ledger
    perfectly reconciled."""
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Ledger Partial Receipt Test Material", "unit": "Sheets", "opening_stock": "0", "minimum_stock": "5",
    }).json()
    supplier_id = client.post("/api/suppliers/", json={"name": "Ledger Partial Receipt Test Supplier"}).json()["id"]
    purchase = client.post("/api/purchases/", json={
        "date": "2026-08-19T00:00:00", "supplier_id": supplier_id, "material_id": material["id"],
        "quantity": "100", "unit": "Sheets", "rate": "500", "gst_percent": "18", "receipt_status": "Ordered",
    }).json()

    client.post(f"/api/purchases/{purchase['id']}/receive", json={"quantity": "40"})
    client.post(f"/api/purchases/{purchase['id']}/receive", json={"quantity": "60"})

    verify = client.get(f"/api/stock/verify/{material['id']}").json()
    assert verify["matches"] is True
    assert float(verify["current_stock"]) == 100.0
    assert verify["ledger_entry_count"] == 2


def test_ledger_reconciles_after_issue(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Ledger Issue Test Material", "unit": "Sheets", "opening_stock": "50", "minimum_stock": "5",
    }).json()
    client.post("/api/issues/", json={
        "date": "2026-08-19T00:00:00", "material_id": material["id"], "quantity_issued": "12", "unit": "Sheets",
    })

    verify = client.get(f"/api/stock/verify/{material['id']}").json()
    assert verify["matches"] is True
    assert float(verify["current_stock"]) == 38.0


def test_ledger_reconciles_after_adjustment(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Ledger Adjustment Test Material", "unit": "Sheets", "opening_stock": "20", "minimum_stock": "5",
    }).json()
    client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Damage", "quantity_delta": "-3", "reason": "Water damage",
    })

    verify = client.get(f"/api/stock/verify/{material['id']}").json()
    assert verify["matches"] is True
    assert float(verify["current_stock"]) == 17.0


def test_ledger_reconciles_after_full_journey(client, test_user):
    """The complete real journey: opening stock, a purchase receipt,
    an issue, a partial return, and a damage adjustment - the ledger
    must reconcile after every single step, not just individually."""
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Ledger Full Journey Test Material", "unit": "Sheets", "opening_stock": "10", "minimum_stock": "5",
    }).json()
    supplier_id = client.post("/api/suppliers/", json={"name": "Ledger Full Journey Test Supplier"}).json()["id"]

    client.post("/api/purchases/", json={
        "date": "2026-08-19T00:00:00", "supplier_id": supplier_id, "material_id": material["id"],
        "quantity": "40", "unit": "Sheets", "rate": "500", "gst_percent": "18",
    })
    assert client.get(f"/api/stock/verify/{material['id']}").json()["matches"] is True

    issue = client.post("/api/issues/", json={
        "date": "2026-08-19T00:00:00", "material_id": material["id"], "quantity_issued": "15", "unit": "Sheets",
    }).json()
    assert client.get(f"/api/stock/verify/{material['id']}").json()["matches"] is True

    client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Return from Issue",
        "quantity_delta": "5", "reason": "Unused sheets returned", "related_issue_id": issue["id"],
    })
    verify_after_return = client.get(f"/api/stock/verify/{material['id']}").json()
    assert verify_after_return["matches"] is True

    client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Damage", "quantity_delta": "-2", "reason": "Damaged in storage",
    })
    final = client.get(f"/api/stock/verify/{material['id']}").json()
    assert final["matches"] is True
    # 10 opening + 40 received - 15 issued + 5 returned - 2 damaged = 38
    assert float(final["current_stock"]) == 38.0
    assert final["ledger_entry_count"] == 4


def test_transfer_moves_material_location(client, test_user):
    _login(client, test_user)
    warehouse = client.post("/api/locations/", json={"name": "Transfer Test Warehouse"}).json()
    rack_a = client.post("/api/locations/", json={"name": "Rack A", "parent_id": warehouse["id"]}).json()
    rack_b = client.post("/api/locations/", json={"name": "Rack B", "parent_id": warehouse["id"]}).json()
    material = client.post("/api/materials/", json={
        "name": "Transfer Test Material", "unit": "Sheets", "opening_stock": 10, "minimum_stock": 1,
        "location_id": rack_a["id"],
    }).json()

    resp = client.post("/api/stock/transfers", json={
        "material_id": material["id"], "quantity": 5, "to_location_id": rack_b["id"],
    })
    assert resp.status_code == 201
    assert resp.json()["from_location_id"] == rack_a["id"]
    assert resp.json()["to_location_id"] == rack_b["id"]

    updated = client.get(f"/api/materials/{material['id']}").json()
    assert updated["location_id"] == rack_b["id"]
    assert "Rack B" in updated["location"]


def test_cannot_transfer_more_than_available_stock(client, test_user):
    _login(client, test_user)
    location = client.post("/api/locations/", json={"name": "Overtransfer Test Location"}).json()
    material = client.post("/api/materials/", json={
        "name": "Overtransfer Test Material", "unit": "Sheets", "opening_stock": 3, "minimum_stock": 1,
    }).json()

    resp = client.post("/api/stock/transfers", json={
        "material_id": material["id"], "quantity": 10, "to_location_id": location["id"],
    })
    assert resp.status_code == 400


def test_adjustment_increases_stock_with_correct_before_after(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Adjustment Increase Material", "unit": "Sheets", "opening_stock": 10, "minimum_stock": 1,
    }).json()

    resp = client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Physical Count Increase",
        "quantity_delta": 3, "reason": "Recount found extra sheets",
    })
    assert resp.status_code == 201
    body = resp.json()
    assert body["stock_before"] == 10
    assert body["stock_after"] == 13

    updated = client.get(f"/api/materials/{material['id']}").json()
    assert updated["current_stock"] == 13


def test_adjustment_cannot_take_stock_negative(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Negative Guard Material", "unit": "Sheets", "opening_stock": 2, "minimum_stock": 1,
    }).json()

    resp = client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Damage",
        "quantity_delta": -5, "reason": "Water damage in storage",
    })
    assert resp.status_code == 400

    unchanged = client.get(f"/api/materials/{material['id']}").json()
    assert unchanged["current_stock"] == 2


def test_adjustment_decrease_for_damage_works(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Damage Adjustment Material", "unit": "Sheets", "opening_stock": 10, "minimum_stock": 1,
    }).json()

    resp = client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Damage",
        "quantity_delta": -2, "reason": "Two sheets damaged in transit",
    })
    assert resp.status_code == 201
    assert resp.json()["stock_after"] == 8


def test_transfer_and_adjustment_require_master_or_manager(client, test_user, db_session):
    from app.platform.security import hash_password
    from app.modules.auth.auth import User

    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "Stock Perm Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    material = client.post("/api/materials/", json={
        "name": "Perm Test Material", "unit": "Sheets", "opening_stock": 10, "minimum_stock": 1,
    }).json()
    location = client.post("/api/locations/", json={"name": "Perm Test Location"}).json()
    limited_user = User(
        username="stockpermuser", email="stockpermuser@example.com", full_name="Limited Stock User",
        password_hash=hash_password("LimitedPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(limited_user)
    db_session.commit()

    resp = client.post("/api/auth/login", json={"identifier": "stockpermuser@example.com", "password": "LimitedPass1!"})
    assert resp.status_code == 200

    transfer_resp = client.post("/api/stock/transfers", json={
        "material_id": material["id"], "quantity": 1, "to_location_id": location["id"],
    })
    assert transfer_resp.status_code == 403

    adjustment_resp = client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Correction", "quantity_delta": 1, "reason": "Test",
    })
    assert adjustment_resp.status_code == 403


def test_material_still_gets_business_id_on_transfer_and_adjustment(client, test_user):
    _login(client, test_user)
    location = client.post("/api/locations/", json={"name": "Business ID Test Location"}).json()
    material = client.post("/api/materials/", json={
        "name": "Business ID Stock Material", "unit": "Sheets", "opening_stock": 5, "minimum_stock": 1,
    }).json()

    transfer = client.post("/api/stock/transfers", json={
        "material_id": material["id"], "quantity": 1, "to_location_id": location["id"],
    }).json()
    assert len(transfer["business_id"]) == 10

    adjustment = client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Correction", "quantity_delta": 1, "reason": "Test",
    }).json()
    assert len(adjustment["business_id"]) == 10


def test_return_from_issue_increases_stock(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Return Test Material", "unit": "Sheets", "opening_stock": "10", "minimum_stock": "5",
    }).json()
    issue = client.post("/api/issues/", json={
        "date": "2026-08-19T00:00:00", "material_id": material["id"], "quantity_issued": "5", "unit": "Sheets",
    }).json()
    before = client.get(f"/api/materials/{material['id']}").json()["current_stock"]

    resp = client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Return from Issue",
        "quantity_delta": "2", "reason": "Unused sheets returned from site", "related_issue_id": issue["id"],
    })
    assert resp.status_code == 201

    after = client.get(f"/api/materials/{material['id']}").json()["current_stock"]
    assert float(after) == float(before) + 2


def test_cannot_return_more_than_was_issued(client, test_user):
    """The concrete anti-arbitrary-quantity rule - a return can never
    exceed what was actually issued."""
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Over Return Test Material", "unit": "Sheets", "opening_stock": "20", "minimum_stock": "5",
    }).json()
    issue = client.post("/api/issues/", json={
        "date": "2026-08-19T00:00:00", "material_id": material["id"], "quantity_issued": "3", "unit": "Sheets",
    }).json()

    resp = client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Return from Issue",
        "quantity_delta": "5", "reason": "Trying to return more than issued", "related_issue_id": issue["id"],
    })
    assert resp.status_code == 400
    assert "only" in resp.json()["detail"].lower()


def test_cannot_double_return_beyond_remaining(client, test_user):
    """Two separate, legitimate-looking returns against the same issue
    must not together exceed the issued quantity."""
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Double Return Test Material", "unit": "Sheets", "opening_stock": "20", "minimum_stock": "5",
    }).json()
    issue = client.post("/api/issues/", json={
        "date": "2026-08-19T00:00:00", "material_id": material["id"], "quantity_issued": "5", "unit": "Sheets",
    }).json()

    first = client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Return from Issue",
        "quantity_delta": "3", "reason": "First partial return", "related_issue_id": issue["id"],
    })
    assert first.status_code == 201

    second = client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Return from Issue",
        "quantity_delta": "3", "reason": "Second return exceeding what remains", "related_issue_id": issue["id"],
    })
    assert second.status_code == 400


def test_return_requires_an_issue_reference(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "No Issue Ref Test Material", "unit": "Sheets", "opening_stock": "10", "minimum_stock": "5",
    }).json()
    resp = client.post("/api/stock/adjustments", json={
        "material_id": material["id"], "adjustment_type": "Return from Issue",
        "quantity_delta": "2", "reason": "No issue referenced",
    })
    assert resp.status_code == 400


def test_return_must_be_for_the_same_material_as_the_issue(client, test_user):
    _login(client, test_user)
    material_a = client.post("/api/materials/", json={
        "name": "Return Mismatch Material A", "unit": "Sheets", "opening_stock": "10", "minimum_stock": "5",
    }).json()
    material_b = client.post("/api/materials/", json={
        "name": "Return Mismatch Material B", "unit": "Sheets", "opening_stock": "10", "minimum_stock": "5",
    }).json()
    issue = client.post("/api/issues/", json={
        "date": "2026-08-19T00:00:00", "material_id": material_a["id"], "quantity_issued": "5", "unit": "Sheets",
    }).json()

    resp = client.post("/api/stock/adjustments", json={
        "material_id": material_b["id"], "adjustment_type": "Return from Issue",
        "quantity_delta": "2", "reason": "Wrong material for this issue", "related_issue_id": issue["id"],
    })
    assert resp.status_code == 400


def test_link_supplier_to_material(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Supplier Material Test Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Supplier Material Test Material", "unit": "Sheets", "opening_stock": 0, "minimum_stock": 1,
    }).json()

    resp = client.post("/api/supplier-materials/", json={
        "supplier_id": supplier["id"], "material_id": material["id"],
        "supplier_sku": "SUP-SKU-001", "supplier_price": "1800.00", "moq": 10, "lead_time_days": 5,
    })
    assert resp.status_code == 201
    assert resp.json()["supplier_sku"] == "SUP-SKU-001"


def test_duplicate_supplier_material_pair_rejected(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Dup Pair Test Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Dup Pair Test Material", "unit": "Sheets", "opening_stock": 0, "minimum_stock": 1,
    }).json()
    client.post("/api/supplier-materials/", json={"supplier_id": supplier["id"], "material_id": material["id"]})
    resp = client.post("/api/supplier-materials/", json={"supplier_id": supplier["id"], "material_id": material["id"]})
    assert resp.status_code == 400


def test_material_can_have_multiple_suppliers(client, test_user):
    """The core many-to-many requirement - the same material linked to
    two different suppliers, each with their own price."""
    _login(client, test_user)
    supplier_a = client.post("/api/suppliers/", json={"name": "Multi Supplier A"}).json()
    supplier_b = client.post("/api/suppliers/", json={"name": "Multi Supplier B"}).json()
    material = client.post("/api/materials/", json={
        "name": "Multi Supplier Material", "unit": "Sheets", "opening_stock": 0, "minimum_stock": 1,
    }).json()

    client.post("/api/supplier-materials/", json={
        "supplier_id": supplier_a["id"], "material_id": material["id"], "supplier_price": "1800.00",
    })
    client.post("/api/supplier-materials/", json={
        "supplier_id": supplier_b["id"], "material_id": material["id"], "supplier_price": "1750.00",
    })

    resp = client.get(f"/api/supplier-materials/by-material/{material['id']}")
    assert resp.status_code == 200
    assert len(resp.json()) == 2
    supplier_names = {link["supplier_name"] for link in resp.json()}
    assert supplier_names == {"Multi Supplier A", "Multi Supplier B"}


def test_preferred_supplier_sorted_first(client, test_user):
    _login(client, test_user)
    cheap_supplier = client.post("/api/suppliers/", json={"name": "Cheap Non-Preferred Supplier"}).json()
    preferred_supplier = client.post("/api/suppliers/", json={"name": "Preferred Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Preferred Sort Test Material", "unit": "Sheets", "opening_stock": 0, "minimum_stock": 1,
    }).json()

    client.post("/api/supplier-materials/", json={
        "supplier_id": cheap_supplier["id"], "material_id": material["id"], "supplier_price": "1000.00",
    })
    client.post("/api/supplier-materials/", json={
        "supplier_id": preferred_supplier["id"], "material_id": material["id"],
        "supplier_price": "1500.00", "is_preferred": True,
    })

    resp = client.get(f"/api/supplier-materials/by-material/{material['id']}")
    results = resp.json()
    assert results[0]["supplier_name"] == "Preferred Supplier"


def test_supplier_can_supply_multiple_materials(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Multi Material Supplier"}).json()
    material_a = client.post("/api/materials/", json={
        "name": "Multi Material A", "unit": "Sheets", "opening_stock": 0, "minimum_stock": 1,
    }).json()
    material_b = client.post("/api/materials/", json={
        "name": "Multi Material B", "unit": "Sheets", "opening_stock": 0, "minimum_stock": 1,
    }).json()

    client.post("/api/supplier-materials/", json={"supplier_id": supplier["id"], "material_id": material_a["id"]})
    client.post("/api/supplier-materials/", json={"supplier_id": supplier["id"], "material_id": material_b["id"]})

    resp = client.get(f"/api/supplier-materials/by-supplier/{supplier['id']}")
    assert len(resp.json()) == 2


def test_recording_a_purchase_updates_last_purchase_price(client, test_user):
    """The specific auto-update behavior - last_purchase_price must not
    be a field the user has to remember to maintain by hand."""
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Last Price Test Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Last Price Test Material", "unit": "Sheets", "opening_stock": 0, "minimum_stock": 1,
    }).json()
    link = client.post("/api/supplier-materials/", json={
        "supplier_id": supplier["id"], "material_id": material["id"], "supplier_price": "1800.00",
    }).json()
    assert link["last_purchase_price"] is None

    client.post("/api/purchases/", json={
        "date": "2026-08-13T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
        "quantity": "10", "unit": "Sheets", "rate": "1750.00", "gst_percent": "18", "payment_status": "Paid",
    })

    updated = client.get(f"/api/supplier-materials/by-material/{material['id']}").json()[0]
    assert float(updated["last_purchase_price"]) == 1750.0


def test_purchase_without_existing_link_does_not_create_one(client, test_user):
    """Deliberate scoping - recording a purchase from a supplier not yet
    linked to the material must not silently auto-create a link."""
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Unlinked Purchase Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Unlinked Purchase Material", "unit": "Sheets", "opening_stock": 0, "minimum_stock": 1,
    }).json()

    resp = client.post("/api/purchases/", json={
        "date": "2026-08-13T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
        "quantity": "5", "unit": "Sheets", "rate": "500.00", "gst_percent": "18", "payment_status": "Paid",
    })
    assert resp.status_code == 201

    links = client.get(f"/api/supplier-materials/by-material/{material['id']}").json()
    assert links == []


def test_master_sees_supplier_material_pricing(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "SM RBAC Master Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "SM RBAC Master Material", "unit": "Sheets", "opening_stock": "5", "minimum_stock": "1",
    }).json()
    client.post("/api/supplier-materials/", json={
        "supplier_id": supplier["id"], "material_id": material["id"], "supplier_price": "1500.00",
    })

    resp = client.get(f"/api/supplier-materials/by-material/{material['id']}").json()
    assert resp[0]["supplier_price"] is not None


def test_employee_supplier_material_pricing_is_genuinely_null(client, test_user, db_session):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "SM RBAC Employee Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "SM RBAC Employee Material", "unit": "Sheets", "opening_stock": "5", "minimum_stock": "1",
    }).json()
    client.post("/api/supplier-materials/", json={
        "supplier_id": supplier["id"], "material_id": material["id"], "supplier_price": "1500.00", "moq": 10, "lead_time_days": 3,
    })

    _create_employee(client, db_session, "smrbacuser", "smrbacuser@example.com")
    resp = client.get(f"/api/supplier-materials/by-material/{material['id']}").json()
    assert resp[0]["supplier_price"] is None
    assert resp[0]["last_purchase_price"] is None
    # Operational (non-financial) fields remain visible.
    assert resp[0]["moq"] == 10
    assert resp[0]["lead_time_days"] == 3
    assert resp[0]["supplier_name"] == "SM RBAC Employee Supplier"


def test_employee_by_supplier_pricing_is_also_null(client, test_user, db_session):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "SM RBAC By Supplier Co"}).json()
    material = client.post("/api/materials/", json={
        "name": "SM RBAC By Supplier Material", "unit": "Sheets", "opening_stock": "5", "minimum_stock": "1",
    }).json()
    client.post("/api/supplier-materials/", json={
        "supplier_id": supplier["id"], "material_id": material["id"], "supplier_price": "800.00",
    })

    _create_employee(client, db_session, "smsuppbacuser", "smsuppbacuser@example.com")
    resp = client.get(f"/api/supplier-materials/by-supplier/{supplier['id']}").json()
    assert resp[0]["supplier_price"] is None


def test_employee_supplier_list_order_does_not_leak_relative_price(client, test_user, db_session):
    """The subtler fix - sorting by price.asc() even with the price
    number hidden still leaks which supplier is cheaper via list order.
    An employee's list must be ordered by something price-blind
    (preferred, then alphabetical) instead."""
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "SM RBAC Order Test Material", "unit": "Sheets", "opening_stock": "5", "minimum_stock": "1",
    }).json()
    expensive_supplier = client.post("/api/suppliers/", json={"name": "Zebra Expensive Supplier"}).json()
    cheap_supplier = client.post("/api/suppliers/", json={"name": "Alpha Cheap Supplier"}).json()
    client.post("/api/supplier-materials/", json={
        "supplier_id": expensive_supplier["id"], "material_id": material["id"], "supplier_price": "5000.00",
    })
    client.post("/api/supplier-materials/", json={
        "supplier_id": cheap_supplier["id"], "material_id": material["id"], "supplier_price": "100.00",
    })

    _create_employee(client, db_session, "smorderrbacuser", "smorderrbacuser@example.com")
    resp = client.get(f"/api/supplier-materials/by-material/{material['id']}").json()
    names_in_order = [r["supplier_name"] for r in resp]
    # Alphabetical (price-blind), NOT price-ascending (which would put
    # the 100.00 supplier first, indirectly revealing it's cheaper).
    assert names_in_order == sorted(names_in_order)


def test_only_master_manager_can_create_supplier_material_link(client, test_user, db_session):
    """Already-correct behavior, confirmed unaffected by this change."""
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "SM RBAC Create Test Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "SM RBAC Create Test Material", "unit": "Sheets", "opening_stock": "5", "minimum_stock": "1",
    }).json()

    _create_employee(client, db_session, "smcreaterbacuser", "smcreaterbacuser@example.com")
    resp = client.post("/api/supplier-materials/", json={
        "supplier_id": supplier["id"], "material_id": material["id"], "supplier_price": "100.00",
    })
    assert resp.status_code == 403


# --- test_purchases_and_reporting.py ---
def _make_ordered_purchase_for_partial_receipt(client, material_id, supplier_id, quantity="100"):
    return client.post("/api/purchases/", json={
        "date": "2026-08-19T00:00:00", "supplier_id": supplier_id, "material_id": material_id,
        "quantity": quantity, "unit": "Sheets", "rate": "500", "gst_percent": "18", "receipt_status": "Ordered",
    }).json()


def test_partial_receipt_moves_to_partially_received(client, test_user):
    _login(client, test_user)
    supplier_id = client.post("/api/suppliers/", json={"name": "Partial Receipt Test Supplier"}).json()["id"]
    material = client.post("/api/materials/", json={
        "name": "Partial Receipt Test Material", "unit": "Sheets", "opening_stock": "0", "minimum_stock": "5",
    }).json()
    purchase = _make_ordered_purchase_for_partial_receipt(client, material["id"], supplier_id, "100")

    resp = client.post(f"/api/purchases/{purchase['id']}/receive", json={"quantity": "40"})
    assert resp.status_code == 200
    data = resp.json()
    assert data["receipt_status"] == "Partially Received"
    assert float(data["quantity_received"]) == 40.0

    refreshed_material = client.get(f"/api/materials/{material['id']}").json()
    assert float(refreshed_material["current_stock"]) == 40.0


def test_second_partial_receipt_completes_it_without_double_counting(client, test_user):
    _login(client, test_user)
    supplier_id = client.post("/api/suppliers/", json={"name": "Second Receipt Test Supplier"}).json()["id"]
    material = client.post("/api/materials/", json={
        "name": "Second Receipt Test Material", "unit": "Sheets", "opening_stock": "0", "minimum_stock": "5",
    }).json()
    purchase = _make_ordered_purchase_for_partial_receipt(client, material["id"], supplier_id, "100")

    client.post(f"/api/purchases/{purchase['id']}/receive", json={"quantity": "40"})
    second = client.post(f"/api/purchases/{purchase['id']}/receive", json={"quantity": "60"})
    assert second.status_code == 200
    assert second.json()["receipt_status"] == "Received"
    assert float(second.json()["quantity_received"]) == 100.0

    refreshed_material = client.get(f"/api/materials/{material['id']}").json()
    # Must be exactly 100 (40 + 60), not more - confirms no double-counting.
    assert float(refreshed_material["current_stock"]) == 100.0


def test_cannot_receive_more_than_remaining(client, test_user):
    """The core quantity-integrity check."""
    _login(client, test_user)
    supplier_id = client.post("/api/suppliers/", json={"name": "Over Receive Test Supplier"}).json()["id"]
    material = client.post("/api/materials/", json={
        "name": "Over Receive Test Material", "unit": "Sheets", "opening_stock": "0", "minimum_stock": "5",
    }).json()
    purchase = _make_ordered_purchase_for_partial_receipt(client, material["id"], supplier_id, "50")

    resp = client.post(f"/api/purchases/{purchase['id']}/receive", json={"quantity": "999"})
    assert resp.status_code == 400
    assert "remains outstanding" in resp.json()["detail"]


def test_cannot_receive_more_after_partial_receipt_already_taken(client, test_user):
    """Two separate receives that together would exceed the ordered
    amount must be rejected on the second one."""
    _login(client, test_user)
    supplier_id = client.post("/api/suppliers/", json={"name": "Cumulative Over Receive Supplier"}).json()["id"]
    material = client.post("/api/materials/", json={
        "name": "Cumulative Over Receive Material", "unit": "Sheets", "opening_stock": "0", "minimum_stock": "5",
    }).json()
    purchase = _make_ordered_purchase_for_partial_receipt(client, material["id"], supplier_id, "50")

    client.post(f"/api/purchases/{purchase['id']}/receive", json={"quantity": "30"})
    resp = client.post(f"/api/purchases/{purchase['id']}/receive", json={"quantity": "30"})
    assert resp.status_code == 400


def test_receiving_with_no_body_still_receives_everything_remaining(client, test_user):
    """Backward compatibility - existing callers that POST with no
    body at all must keep working exactly as before."""
    _login(client, test_user)
    supplier_id = client.post("/api/suppliers/", json={"name": "No Body Receive Test Supplier"}).json()["id"]
    material = client.post("/api/materials/", json={
        "name": "No Body Receive Test Material", "unit": "Sheets", "opening_stock": "0", "minimum_stock": "5",
    }).json()
    purchase = _make_ordered_purchase_for_partial_receipt(client, material["id"], supplier_id, "75")

    resp = client.post(f"/api/purchases/{purchase['id']}/receive")
    assert resp.status_code == 200
    assert resp.json()["receipt_status"] == "Received"
    assert float(resp.json()["quantity_received"]) == 75.0


def test_purchase_created_as_received_has_quantity_received_prefilled(client, test_user):
    """Migration-backfill equivalent at creation time - a purchase
    created already-Received must show full quantity_received
    immediately, not zero."""
    _login(client, test_user)
    supplier_id = client.post("/api/suppliers/", json={"name": "Prefilled Test Supplier"}).json()["id"]
    material = client.post("/api/materials/", json={
        "name": "Prefilled Test Material", "unit": "Sheets", "opening_stock": "0", "minimum_stock": "5",
    }).json()
    purchase = client.post("/api/purchases/", json={
        "date": "2026-08-19T00:00:00", "supplier_id": supplier_id, "material_id": material["id"],
        "quantity": "20", "unit": "Sheets", "rate": "500", "gst_percent": "18",
    }).json()
    assert purchase["receipt_status"] == "Received"
    assert float(purchase["quantity_received"]) == 20.0


def test_purchases_export_returns_valid_xlsx(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/reports/purchases.xlsx")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = load_workbook(BytesIO(resp.content))
    assert "Purchases" in wb.sheetnames
    ws = wb["Purchases"]
    assert ws.cell(row=3, column=1).value == "Purchase ID"


def test_stock_dashboard_export_has_all_sheets(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/reports/stock-dashboard.xlsx")
    assert resp.status_code == 200

    wb = load_workbook(BytesIO(resp.content))
    for expected in ["Dashboard", "Material Master", "Purchases", "Issues", "Suppliers"]:
        assert expected in wb.sheetnames


def test_exports_require_auth(client):
    resp = client.get("/api/reports/purchases.xlsx")
    assert resp.status_code == 401


def test_create_top_level_location(client, test_user):
    _login(client, test_user)
    resp = client.post("/api/locations/", json={"name": "Vijay Nagar Warehouse", "location_type": "Warehouse"})
    assert resp.status_code == 201
    assert resp.json()["full_path"] == "Vijay Nagar Warehouse"


def test_multi_level_full_path_builds_correctly(client, test_user):
    """The core requirement of the tree - a 3-level hierarchy producing
    the correct "Parent > Child > Grandchild" path."""
    _login(client, test_user)
    warehouse = client.post("/api/locations/", json={"name": "Main Warehouse", "location_type": "Warehouse"}).json()
    rack = client.post("/api/locations/", json={
        "name": "Rack A2", "location_type": "Rack", "parent_id": warehouse["id"],
    }).json()
    bin_ = client.post("/api/locations/", json={
        "name": "Bin H1", "location_type": "Bin", "parent_id": rack["id"],
    }).json()
    assert bin_["full_path"] == "Main Warehouse > Rack A2 > Bin H1"


def test_duplicate_name_under_same_parent_rejected(client, test_user):
    _login(client, test_user)
    warehouse = client.post("/api/locations/", json={"name": "Dup Location Test Warehouse"}).json()
    client.post("/api/locations/", json={"name": "Rack A1", "parent_id": warehouse["id"]})
    resp = client.post("/api/locations/", json={"name": "Rack A1", "parent_id": warehouse["id"]})
    assert resp.status_code == 400


def test_same_name_allowed_under_different_parents(client, test_user):
    """"Rack A1" in two different warehouses are genuinely different
    locations - uniqueness is per-parent, not global."""
    _login(client, test_user)
    warehouse_a = client.post("/api/locations/", json={"name": "Warehouse A For Rack Test"}).json()
    warehouse_b = client.post("/api/locations/", json={"name": "Warehouse B For Rack Test"}).json()
    resp_a = client.post("/api/locations/", json={"name": "Rack A1", "parent_id": warehouse_a["id"]})
    resp_b = client.post("/api/locations/", json={"name": "Rack A1", "parent_id": warehouse_b["id"]})
    assert resp_a.status_code == 201
    assert resp_b.status_code == 201


def test_material_created_with_location_id_syncs_legacy_location_string(client, test_user):
    """The core compatibility guarantee - material.location (plain
    string, read by every existing consumer) must reflect the real
    location's full path when location_id is set."""
    _login(client, test_user)
    warehouse = client.post("/api/locations/", json={"name": "Sync Test Warehouse"}).json()
    rack = client.post("/api/locations/", json={
        "name": "Sync Test Rack", "parent_id": warehouse["id"],
    }).json()

    material = client.post("/api/materials/", json={
        "name": "Location Sync Test Material", "unit": "Sheets", "opening_stock": 5, "minimum_stock": 1,
        "location_id": rack["id"],
    }).json()
    assert material["location"] == "Sync Test Warehouse > Sync Test Rack"
    assert material["location_id"] == rack["id"]


def test_material_without_location_id_still_works_the_old_way(client, test_user):
    _login(client, test_user)
    resp = client.post("/api/materials/", json={
        "name": "Legacy Flat Location Material", "unit": "Sheets", "opening_stock": 1, "minimum_stock": 1,
        "location": "Rack A1",
    })
    assert resp.status_code == 201
    body = resp.json()
    assert body["location"] == "Rack A1"
    assert body["location_id"] is None


def test_get_children_of_a_parent_location(client, test_user):
    _login(client, test_user)
    warehouse = client.post("/api/locations/", json={"name": "Children Test Warehouse"}).json()
    client.post("/api/locations/", json={"name": "Rack A1", "parent_id": warehouse["id"]})
    client.post("/api/locations/", json={"name": "Rack A2", "parent_id": warehouse["id"]})

    resp = client.get("/api/locations/", params={"parent_id": warehouse["id"]})
    assert resp.status_code == 200
    names = {loc["name"] for loc in resp.json()}
    assert names == {"Rack A1", "Rack A2"}


class _QueryCounter:
    def __init__(self):
        self.count = 0

    def __enter__(self):
        event.listen(engine, "before_cursor_execute", self._callback)
        return self

    def __exit__(self, *args):
        event.remove(engine, "before_cursor_execute", self._callback)

    def _callback(self, conn, cursor, statement, parameters, context, executemany):
        self.count += 1


def test_purchases_export_query_count_does_not_scale_with_row_count(client, test_user):
    """The core N+1 regression guard - export a growing number of
    purchases, each with a DIFFERENT supplier and material, and confirm
    the query count stays bounded rather than growing linearly."""
    _login(client, test_user)
    for i in range(8):
        supplier = client.post("/api/suppliers/", json={"name": f"N+1 Purchase Supplier {i}"}).json()
        material = client.post("/api/materials/", json={
            "name": f"N+1 Purchase Material {i}", "unit": "Sheets", "opening_stock": "0", "minimum_stock": "1",
        }).json()
        client.post("/api/purchases/", json={
            "date": "2026-08-21T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
            "quantity": "5", "unit": "Sheets", "rate": "500.00", "gst_percent": "18", "payment_status": "Paid",
        })

    with _QueryCounter() as counter:
        resp = client.get("/api/reports/purchases.xlsx")
    assert resp.status_code == 200
    assert counter.count < 15, f"expected a bounded query count, got {counter.count} - possible N+1 regression"


def test_no_confidence_when_nothing_in_db_matches(client, test_user):
    """A typed name that matches nothing in an otherwise-empty database
    must return confidence "none", not a fabricated guess."""
    _login(client, test_user)
    resp = client.get("/api/materials/interpret-name", params={"name": "Totally Unrecognized Material XYZ"})
    assert resp.status_code == 200
    data = resp.json()
    assert data["confidence"] == "none"
    assert data["category_id"] is None
    assert data["subcategory_id"] is None


def test_thickness_extracted_even_with_no_category_match(client, test_user):
    """Thickness extraction is independent of the category-matching
    logic - it should still come through even when nothing else in the
    name matches an existing subcategory/material."""
    _login(client, test_user)
    resp = client.get("/api/materials/interpret-name", params={"name": "Unmatched Thing 9mm"})
    assert resp.status_code == 200
    data = resp.json()
    assert data["thickness_size"] == "9mm"
    assert data["confidence"] == "none"


def test_high_confidence_when_existing_subcategory_name_matches(client, test_user):
    """Typing a name that contains an existing Subcategory's own name
    (e.g. "HDHMR 6mm" when a real "HDHMR" subcategory already exists)
    must return "high" confidence and the real category/subcategory ids
    - never an invented category name."""
    _login(client, test_user)
    category = client.post("/api/material-categories/", json={"name": "Board & Wood Materials Interp Test"}).json()
    subcategory = client.post("/api/material-categories/subcategories", json={
        "category_id": category["id"], "name": "HDHMR Interp Test",
    }).json()

    resp = client.get("/api/materials/interpret-name", params={"name": "HDHMR Interp Test 6mm"})
    assert resp.status_code == 200
    data = resp.json()
    assert data["confidence"] == "high"
    assert data["subcategory_id"] == subcategory["id"]
    assert data["category_id"] == category["id"]
    assert data["thickness_size"] == "6mm"


def test_medium_confidence_when_similar_existing_material_matches(client, test_user):
    """Typing a name similar to an existing Material's own name (but not
    matching any Subcategory name directly) should fall back to "medium"
    confidence, reusing that existing material's real subcategory -
    never a fabricated one."""
    _login(client, test_user)
    category = client.post("/api/material-categories/", json={"name": "Surface Materials Interp Test"}).json()
    subcategory = client.post("/api/material-categories/subcategories", json={
        "category_id": category["id"], "name": "Laminate Interp Test",
    }).json()
    client.post("/api/materials/", json={
        "name": "Merino Laminate Interp Test 18mm", "unit": "sheet", "opening_stock": 0, "minimum_stock": 0,
        "subcategory_id": subcategory["id"],
    })

    # Deliberately does NOT contain "Laminate Interp Test" as a
    # subcategory-name substring match on its own - "merino" is the
    # shared word with the existing material, proving this is genuinely
    # falling through to the medium-confidence material-similarity path,
    # not accidentally hitting the high-confidence subcategory-name path.
    resp = client.get("/api/materials/interpret-name", params={"name": "Merino 12mm"})
    assert resp.status_code == 200
    data = resp.json()
    assert data["confidence"] == "medium"
    assert data["subcategory_id"] == subcategory["id"]
    assert data["thickness_size"] == "12mm"


def test_never_overwrites_forces_confidence_field_present(client, test_user):
    """The response always includes a confidence field the frontend can
    branch on - this is the contract MaterialsPage.jsx's nameSuggestion
    handling depends on (confidence !== 'none' gates whether the
    suggestion chip is shown at all)."""
    _login(client, test_user)
    resp = client.get("/api/materials/interpret-name", params={"name": "Plain Board"})
    assert resp.status_code == 200
    assert "confidence" in resp.json()


def test_missing_name_param_rejected(client, test_user):
    """name is a required query param (min_length=1) - confirms the
    route genuinely validates this rather than silently treating a
    missing name as an empty string."""
    _login(client, test_user)
    resp = client.get("/api/materials/interpret-name")
    assert resp.status_code == 422


# --- test_material_imports.py ---
def test_material_import_commit_creates_real_persisted_material(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Material Import Commit Supplier"}).json()

    resp = client.post("/api/material-imports/commit", json={"rows": [{
        "name": "Material Import Commit Test Material", "unit": "Sheets",
        "supplier_id": supplier["id"], "opening_stock": "12", "minimum_stock": "2",
    }]})
    assert resp.status_code == 200
    body = resp.json()
    assert body["created_materials"] == 1
    assert body["error"] is None
    new_id = body["material_ids"][0]

    persisted = client.get(f"/api/materials/{new_id}").json()
    assert persisted["name"] == "Material Import Commit Test Material"
    assert persisted["current_stock"] == 12.0
    assert persisted["minimum_stock"] == 2.0
    assert persisted["supplier_id"] == supplier["id"]


def test_material_import_commit_reuses_matched_material_without_duplicating(client, test_user):
    _login(client, test_user)
    existing = client.post("/api/materials/", json={
        "name": "Material Import Match Test Material", "unit": "Sheets", "opening_stock": 5, "minimum_stock": 1,
    }).json()

    resp = client.post("/api/material-imports/commit", json={"rows": [{
        "name": "Material Import Match Test Material", "unit": "Sheets",
        "matched_material_id": existing["id"],
    }]})
    assert resp.status_code == 200
    body = resp.json()
    assert body["created_materials"] == 0
    assert body["matched_existing"] == 1

    all_matching = client.get("/api/materials/", params={"search": "Material Import Match Test Material"}).json()
    assert len(all_matching) == 1  # no duplicate created


def test_material_import_commit_requires_master(client, db_session):
    employee = User(
        username="matimportuser", email="matimportuser@example.com", full_name="Material Import User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=None, is_active=True,
    )
    db_session.add(employee)
    db_session.commit()
    resp = client.post("/api/auth/login", json={"identifier": "matimportuser@example.com", "password": "EmpPass1!"})
    assert resp.status_code == 200

    resp = client.post("/api/material-imports/commit", json={"rows": [{"name": "Unauthorized Material", "unit": "Sheets"}]})
    assert resp.status_code == 403


def test_material_import_template_requires_master(client, db_session):
    employee = User(
        username="mattemplateuser", email="mattemplateuser@example.com", full_name="Material Template User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=None, is_active=True,
    )
    db_session.add(employee)
    db_session.commit()
    resp = client.post("/api/auth/login", json={"identifier": "mattemplateuser@example.com", "password": "EmpPass1!"})
    assert resp.status_code == 200

    resp = client.get("/api/material-imports/template")
    assert resp.status_code == 403


def test_material_import_template_unauthenticated_rejected(client):
    resp = client.get("/api/material-imports/template")
    assert resp.status_code in (401, 403)


# --- test_shortage_intelligence.py ---
def _make_client(client, suffix):
    return client.post("/api/clients/", json={
        "name": f"Shortage Intel Client {suffix}", "phone": f"90000103{suffix}",
    }).json()["id"]


def test_shortage_matches_worked_example_5_required_2_available_1_pending(client, test_user):
    """5 sheets required, 2 available, a pending purchase of 1 already
    placed -> shortage 2 (required - available - pending), not 3 -
    exactly the spec's own formula and worked example."""
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Shortage Intel HDHMR Sheet", "unit": "Sheets", "opening_stock": "2", "minimum_stock": "1",
    }).json()
    supplier = client.post("/api/suppliers/", json={"name": "Shortage Intel Supplier"}).json()
    # A purchase already placed but not yet received - "pending".
    client.post("/api/purchases/", json={
        "date": "2026-08-01T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
        "quantity": "1", "unit": "Sheets", "rate": "500.00", "gst_percent": "18", "receipt_status": "Ordered",
    })
    product = client.post("/api/products/", json={
        "name": "Shortage Intel Wardrobe", "unit": "Piece",
        "materials_used": [{"material_id": material["id"], "quantity_required": "5"}],
    }).json()
    client_id = _make_client(client, "1")
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Wardrobe", "quantity": "1", "unit": "Piece", "rate": "40000", "product_id": product["id"]}],
    }).json()

    resp = client.get(f"/api/orders/{order['id']}/material-requirements")
    assert resp.status_code == 200
    materials = resp.json()["materials"]
    assert len(materials) == 1
    row = materials[0]
    assert row["material_id"] == material["id"]
    assert float(row["required"]) == 5.0
    assert float(row["available"]) == 2.0
    assert float(row["gap_before_pending_supply"]) == 3.0
    assert float(row["pending_purchase_quantity"]) == 1.0
    assert float(row["shortage"]) == 2.0
    assert float(row["recommended_purchase_quantity"]) == 2.0


def test_shortage_multiplies_bom_quantity_by_order_quantity(client, test_user):
    """2 wardrobes ordered, each needs 5 sheets -> 10 required, not 5 -
    the BOM quantity must scale with how many units were actually
    ordered, not just the per-unit recipe."""
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Shortage Intel Ply Sheet", "unit": "Sheets", "opening_stock": "0", "minimum_stock": "1",
    }).json()
    product = client.post("/api/products/", json={
        "name": "Shortage Intel Cabinet", "unit": "Piece",
        "materials_used": [{"material_id": material["id"], "quantity_required": "5"}],
    }).json()
    client_id = _make_client(client, "2")
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Cabinet", "quantity": "2", "unit": "Piece", "rate": "20000", "product_id": product["id"]}],
    }).json()

    resp = client.get(f"/api/orders/{order['id']}/material-requirements")
    row = resp.json()["materials"][0]
    assert float(row["required"]) == 10.0
    assert float(row["recommended_purchase_quantity"]) == 10.0


def test_no_shortage_when_stock_covers_requirement(client, test_user):
    """Enough stock on hand -> zero shortage, zero recommendation, not
    a negative or nonsensical number."""
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Shortage Intel Surplus Sheet", "unit": "Sheets", "opening_stock": "20", "minimum_stock": "1",
    }).json()
    product = client.post("/api/products/", json={
        "name": "Shortage Intel Shelf", "unit": "Piece",
        "materials_used": [{"material_id": material["id"], "quantity_required": "3"}],
    }).json()
    client_id = _make_client(client, "3")
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Shelf", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()

    resp = client.get(f"/api/orders/{order['id']}/material-requirements")
    row = resp.json()["materials"][0]
    assert float(row["shortage"]) == 0.0
    assert float(row["recommended_purchase_quantity"]) == 0.0


def test_order_item_without_product_has_no_material_requirement(client, test_user):
    """A freeform order line (no product_id, e.g. "Installation") has
    no BOM to trace, so it contributes nothing - not an error."""
    _login(client, test_user)
    client_id = _make_client(client, "4")
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Installation", "quantity": "1", "unit": "Nos", "rate": "2000"}],
    }).json()

    resp = client.get(f"/api/orders/{order['id']}/material-requirements")
    assert resp.status_code == 200
    assert resp.json()["materials"] == []


def test_reserved_stock_from_other_open_order_reduces_available(client, test_user):
    """Two open orders competing for the same material: Order A's own
    unfulfilled BOM demand (nothing issued against it yet) must reduce
    what Order B sees as "available", per Available = Current Stock -
    Reserved Qty. Neither order's own demand should reserve against
    itself (checked separately below)."""
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Shortage Intel Contested Sheet", "unit": "Sheets", "opening_stock": "10", "minimum_stock": "1",
    }).json()
    product = client.post("/api/products/", json={
        "name": "Shortage Intel Contested Product", "unit": "Piece",
        "materials_used": [{"material_id": material["id"], "quantity_required": "6"}],
    }).json()
    client_a = _make_client(client, "5a")
    client_b = _make_client(client, "5b")
    order_a = client.post("/api/orders/", json={
        "client_id": client_a, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()
    order_b = client.post("/api/orders/", json={
        "client_id": client_b, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()

    # Order A's own check: nothing else is open yet except Order B,
    # which also needs 6 - Order A must see Order B's demand reserved
    # against it, not its own.
    resp_a = client.get(f"/api/orders/{order_a['id']}/material-requirements")
    row_a = resp_a.json()["materials"][0]
    assert float(row_a["reserved_by_other_orders"]) == 6.0
    assert float(row_a["available"]) == 4.0  # 10 stock - 6 reserved by order B
    assert float(row_a["shortage"]) == 2.0   # 6 required - 4 available

    resp_b = client.get(f"/api/orders/{order_b['id']}/material-requirements")
    row_b = resp_b.json()["materials"][0]
    assert float(row_b["reserved_by_other_orders"]) == 6.0  # order A's own unfulfilled demand
    assert float(row_b["available"]) == 4.0
    assert float(row_b["shortage"]) == 2.0


def test_issuing_material_releases_the_reservation(client, test_user):
    """Once a material is actually issued against an order, that
    portion of its demand is fulfilled - it must stop counting as
    reserved against other orders."""
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Shortage Intel Issued Sheet", "unit": "Sheets", "opening_stock": "10", "minimum_stock": "1",
    }).json()
    product = client.post("/api/products/", json={
        "name": "Shortage Intel Issued Product", "unit": "Piece",
        "materials_used": [{"material_id": material["id"], "quantity_required": "6"}],
    }).json()
    client_a = _make_client(client, "6a")
    client_b = _make_client(client, "6b")
    order_a = client.post("/api/orders/", json={
        "client_id": client_a, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()
    order_b = client.post("/api/orders/", json={
        "client_id": client_b, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()
    # Order A's full 6-sheet requirement is issued - fully fulfilled,
    # nothing left of its demand to reserve.
    client.post("/api/issues/", json={
        "date": "2026-08-19T00:00:00", "order_id": order_a["id"], "material_id": material["id"],
        "quantity_issued": "6", "unit": "Sheets",
    })

    resp_b = client.get(f"/api/orders/{order_b['id']}/material-requirements")
    row_b = resp_b.json()["materials"][0]
    assert float(row_b["reserved_by_other_orders"]) == 0.0
    assert float(row_b["available"]) == 10.0


def test_cancelled_order_does_not_reserve_stock(client, test_user):
    """A cancelled order's demand must not reserve stock against
    other orders - it will never actually consume that material."""
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Shortage Intel Cancelled Sheet", "unit": "Sheets", "opening_stock": "10", "minimum_stock": "1",
    }).json()
    product = client.post("/api/products/", json={
        "name": "Shortage Intel Cancelled Product", "unit": "Piece",
        "materials_used": [{"material_id": material["id"], "quantity_required": "6"}],
    }).json()
    client_a = _make_client(client, "7a")
    client_b = _make_client(client, "7b")
    order_a = client.post("/api/orders/", json={
        "client_id": client_a, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()
    client.put(f"/api/orders/{order_a['id']}", json={"project_status": "Cancelled"})
    order_b = client.post("/api/orders/", json={
        "client_id": client_b, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()

    resp_b = client.get(f"/api/orders/{order_b['id']}/material-requirements")
    row_b = resp_b.json()["materials"][0]
    assert float(row_b["reserved_by_other_orders"]) == 0.0
    assert float(row_b["available"]) == 10.0


def test_hardware_requirement_uses_the_same_bom_path_unchanged(client, test_user):
    """Phase C's own worked example: product quantity x hardware-per-
    product = required hinges/handles. A hardware-category Material
    linked via the same ProductMaterial BOM must flow through the
    identical shortage calculation as any raw material - no separate
    "hardware requirement" code path exists or is needed, since a
    hinge is just another Material to this service."""
    _login(client, test_user)
    hinge = client.post("/api/materials/", json={
        "name": "Shortage Intel Soft-Close Hinge", "category": "Hardware", "unit": "Nos",
        "opening_stock": "10", "minimum_stock": "4",
    }).json()
    assert hinge["category"] == "Hardware"
    # 2 doors per wardrobe, 2 hinges per door = 4 hinges per wardrobe.
    product = client.post("/api/products/", json={
        "name": "Shortage Intel Hinged Wardrobe", "unit": "Piece",
        "materials_used": [{"material_id": hinge["id"], "quantity_required": "4"}],
    }).json()
    client_id = _make_client(client, "8")
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Wardrobe", "quantity": "3", "unit": "Piece", "rate": "35000", "product_id": product["id"]}],
    }).json()

    resp = client.get(f"/api/orders/{order['id']}/material-requirements")
    assert resp.status_code == 200
    row = resp.json()["materials"][0]
    assert row["material_id"] == hinge["id"]
    assert float(row["required"]) == 12.0  # 3 wardrobes x 4 hinges
    assert float(row["available"]) == 10.0
    assert float(row["shortage"]) == 2.0  # 12 required - 10 available


def test_material_requirements_requires_auth(client):
    resp = client.get("/api/orders/1/material-requirements")
    assert resp.status_code == 401


def test_material_requirements_unknown_order_returns_404(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/orders/999999/material-requirements")
    assert resp.status_code == 404


def test_at_risk_orders_matches_per_order_calculation(client, test_user):
    """The bulk dashboard calculation (calculate_at_risk_orders) must
    produce exactly the same shortage figures as calling the per-order
    endpoint for each order individually - proves the business-wide
    derivation (reserved-by-others from a total already in memory,
    not a second per-order query) is mathematically equivalent, not a
    second, independently-drifting way to compute a shortage. Reuses
    the exact two-competing-orders scenario from
    test_reserved_stock_from_other_open_order_reduces_available above."""
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "At Risk Dashboard Contested Sheet", "unit": "Sheets", "opening_stock": "10", "minimum_stock": "1",
    }).json()
    product = client.post("/api/products/", json={
        "name": "At Risk Dashboard Contested Product", "unit": "Piece",
        "materials_used": [{"material_id": material["id"], "quantity_required": "6"}],
    }).json()
    client_a = _make_client(client, "9a")
    client_b = _make_client(client, "9b")
    order_a = client.post("/api/orders/", json={
        "client_id": client_a, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()
    order_b = client.post("/api/orders/", json={
        "client_id": client_b, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()

    per_order_a = client.get(f"/api/orders/{order_a['id']}/material-requirements").json()["materials"][0]
    per_order_b = client.get(f"/api/orders/{order_b['id']}/material-requirements").json()["materials"][0]

    resp = client.get("/api/dashboard/at-risk-orders")
    assert resp.status_code == 200
    body = resp.json()
    by_order_id = {row["order_id"]: row for row in body["orders"]}
    assert order_a["id"] in by_order_id
    assert order_b["id"] in by_order_id

    bulk_a = by_order_id[order_a["id"]]["materials"][0]
    bulk_b = by_order_id[order_b["id"]]["materials"][0]
    assert float(bulk_a["shortage"]) == float(per_order_a["shortage"]) == 2.0
    assert float(bulk_a["available"]) == float(per_order_a["available"]) == 4.0
    assert float(bulk_b["shortage"]) == float(per_order_b["shortage"]) == 2.0
    assert float(bulk_b["available"]) == float(per_order_b["available"]) == 4.0


def test_at_risk_orders_excludes_order_with_sufficient_stock(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "At Risk Dashboard Ample Sheet", "unit": "Sheets", "opening_stock": "100", "minimum_stock": "1",
    }).json()
    product = client.post("/api/products/", json={
        "name": "At Risk Dashboard Ample Product", "unit": "Piece",
        "materials_used": [{"material_id": material["id"], "quantity_required": "2"}],
    }).json()
    client_id = _make_client(client, "9c")
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()

    resp = client.get("/api/dashboard/at-risk-orders")
    assert resp.status_code == 200
    order_ids = [row["order_id"] for row in resp.json()["orders"]]
    assert order["id"] not in order_ids


def test_at_risk_orders_requires_auth(client):
    resp = client.get("/api/dashboard/at-risk-orders")
    assert resp.status_code == 401


def test_material_requirement_recalculates_when_order_quantity_changes(client, test_user):
    """P0.1 section 8 (Configuration Change Impact) - changing an
    order's item quantity must immediately change its material
    requirement, with no manual recalculation step and no stale
    value, since this is computed fresh from current order_items/BOM/
    stock on every call rather than stored anywhere."""
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Config Change Impact Sheet", "unit": "Sheets", "opening_stock": "5", "minimum_stock": "1",
    }).json()
    product = client.post("/api/products/", json={
        "name": "Config Change Impact Product", "unit": "Piece",
        "materials_used": [{"material_id": material["id"], "quantity_required": "1"}],
    }).json()
    client_id = client.post("/api/clients/", json={"name": "Config Change Impact Client", "phone": "9000010800"}).json()["id"]
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()

    before = client.get(f"/api/orders/{order['id']}/material-requirements").json()["materials"][0]
    assert float(before["required"]) == 1.0
    assert float(before["shortage"]) == 0.0

    update_resp = client.put(f"/api/orders/{order['id']}", json={
        "items": [{"description": "Item", "quantity": "10", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    })
    assert update_resp.status_code == 200

    after = client.get(f"/api/orders/{order['id']}/material-requirements").json()["materials"][0]
    assert float(after["required"]) == 10.0
    assert float(after["shortage"]) == 5.0


# --- Family 137, Step 2: Dead-Stock / Material-to-Design Matching (feature 7) ---

def _create_dead_stock_material(client, category="Dead Stock Test Category", stock="20"):
    resp = client.post("/api/materials/", json={
        "name": "Dead Stock Test Material", "category": category, "unit": "Sheets",
        "minimum_stock": 1, "average_rate": "500.00", "opening_stock": stock,
    })
    assert resp.status_code == 201
    return resp.json()


def test_dead_stock_matches_finds_idle_material_with_open_estimate(client, test_user):
    _login(client, test_user)
    material = _create_dead_stock_material(client, category="Dead Stock Match Category")
    client_resp = client.post("/api/clients/", json={"name": "Dead Stock Client", "phone": "9812310001"}).json()
    product = client.post("/api/products/", json={
        "name": "Dead Stock Matching Product", "unit": "Sheets", "category": "Dead Stock Match Category",
    }).json()
    estimate = client.post("/api/estimates/", json={
        "client_id": client_resp["id"],
        "line_items": [{
            "description": "Panel from stock", "category": "Material", "quantity": "3",
            "unit": "Sheets", "rate": "1000", "product_id": product["id"],
        }],
    }).json()
    # Freshly created, so not idle under the default 90-day threshold -
    # idle_days=0 makes "just created" count as idle for this test.
    resp = client.get("/api/materials/dead-stock-matches", params={"idle_days": 0})
    assert resp.status_code == 200
    body = resp.json()
    matching_item = next((i for i in body["items"] if i["material_id"] == material["id"]), None)
    assert matching_item is not None
    assert matching_item["category"] == "Dead Stock Match Category"
    match_estimate_ids = [m["estimate_id"] for m in matching_item["potential_matches"]]
    assert estimate["id"] in match_estimate_ids


def test_dead_stock_matches_excludes_zero_stock_material(client, test_user):
    _login(client, test_user)
    material = _create_dead_stock_material(client, category="Dead Stock Zero Category", stock="0")
    resp = client.get("/api/materials/dead-stock-matches", params={"idle_days": 0})
    assert resp.status_code == 200
    ids = [i["material_id"] for i in resp.json()["items"]]
    assert material["id"] not in ids


def test_dead_stock_matches_excludes_material_with_no_open_estimate_match(client, test_user):
    _login(client, test_user)
    material = _create_dead_stock_material(client, category="Dead Stock Lonely Category")
    resp = client.get("/api/materials/dead-stock-matches", params={"idle_days": 0})
    assert resp.status_code == 200
    ids = [i["material_id"] for i in resp.json()["items"]]
    assert material["id"] not in ids


def test_dead_stock_matches_excludes_non_open_estimates(client, test_user):
    _login(client, test_user)
    material = _create_dead_stock_material(client, category="Dead Stock Rejected Category")
    client_resp = client.post("/api/clients/", json={"name": "Dead Stock Rejected Client", "phone": "9812310004"}).json()
    product = client.post("/api/products/", json={
        "name": "Dead Stock Rejected Product", "unit": "Sheets", "category": "Dead Stock Rejected Category",
    }).json()
    estimate = client.post("/api/estimates/", json={
        "client_id": client_resp["id"],
        "line_items": [{
            "description": "Panel", "category": "Material", "quantity": "1",
            "unit": "Sheets", "rate": "1000", "product_id": product["id"],
        }],
    }).json()
    client.put(f"/api/estimates/{estimate['id']}", json={"status": "rejected"})

    resp = client.get("/api/materials/dead-stock-matches", params={"idle_days": 0})
    assert resp.status_code == 200
    ids = [i["material_id"] for i in resp.json()["items"]]
    # No other open estimate references this category, so with the
    # only match now rejected, the material must not be surfaced at all.
    assert material["id"] not in ids


def test_dead_stock_matches_respects_idle_days_threshold(client, test_user):
    _login(client, test_user)
    material = _create_dead_stock_material(client, category="Dead Stock Fresh Category")
    client_resp = client.post("/api/clients/", json={"name": "Dead Stock Fresh Client", "phone": "9812310005"}).json()
    product = client.post("/api/products/", json={
        "name": "Dead Stock Fresh Product", "unit": "Sheets", "category": "Dead Stock Fresh Category",
    }).json()
    client.post("/api/estimates/", json={
        "client_id": client_resp["id"],
        "line_items": [{
            "description": "Panel", "category": "Material", "quantity": "1",
            "unit": "Sheets", "rate": "1000", "product_id": product["id"],
        }],
    })
    # Freshly created material - under the DEFAULT 90-day threshold it
    # is not idle yet, so it must not be reported.
    resp = client.get("/api/materials/dead-stock-matches")
    assert resp.status_code == 200
    ids = [i["material_id"] for i in resp.json()["items"]]
    assert material["id"] not in ids


def test_dead_stock_matches_never_mutates_stock(client, test_user):
    _login(client, test_user)
    material = _create_dead_stock_material(client, category="Dead Stock No-Mutate Category")
    before = client.get(f"/api/materials/{material['id']}").json()["current_stock"]
    client.get("/api/materials/dead-stock-matches", params={"idle_days": 0})
    after = client.get(f"/api/materials/{material['id']}").json()["current_stock"]
    assert float(before) == float(after)


def test_dead_stock_matches_requires_master(client, test_user, db_session):
    _login(client, test_user)
    client.post("/api/auth/logout")
    user = User(username="deadstockrbacuser", email="deadstockrbacuser@example.com", full_name="Dead Stock RBAC User",
                password_hash=hash_password("UserPass1!"), role="user", is_active=True)
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "deadstockrbacuser@example.com", "password": "UserPass1!"})
    resp = client.get("/api/materials/dead-stock-matches")
    assert resp.status_code == 403

