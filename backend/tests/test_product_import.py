"""Family 21 - reusable bulk import infrastructure, exercised through
its first real consumer (Product Master import): Download Template ->
Upload -> Validate/Preview -> Commit. IDs are never accepted from the
uploaded file - product_code/business_id are always generated at
commit time, matching every other creation path in this app."""
import io

import openpyxl

from app.utils.product_import import PRODUCT_IMPORT_COLUMNS


def _login(client, test_user):
    resp = client.post("/api/auth/login", json={"identifier": "test@example.com", "password": "TestPass123!"})
    assert resp.status_code == 200


def _build_workbook(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(PRODUCT_IMPORT_COLUMNS)
    for row in rows:
        ws.append([row.get(col, "") for col in PRODUCT_IMPORT_COLUMNS])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()


def test_download_template_is_a_real_xlsx(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/product-imports/template")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.worksheets[0]
    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=10, values_only=True):
        if row and "Product Name" in row:
            header_row = row
            break
    assert header_row is not None


def test_preview_valid_row_with_existing_subcategory(client, test_user):
    _login(client, test_user)
    category = client.post("/api/product-categories/", json={"name": "Import Test Category"}).json()
    client.post("/api/product-categories/subcategories", json={
        "category_id": category["id"], "name": "Import Test Subcategory",
    })
    file_bytes = _build_workbook([{
        "Product Name": "Imported Wardrobe", "SKU": "IMP-001", "Type": "standard",
        "Category": "Import Test Category", "Subcategory": "Import Test Subcategory",
        "Cost Price": 30000, "Selling Price": 45000, "Tax %": 18,
    }])
    resp = client.post(
        "/api/product-imports/preview",
        files={"file": ("import.xlsx", file_bytes,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["total_rows"] == 1
    assert body["error_rows"] == 0
    assert body["rows"][0]["subcategory_id"] is not None


def test_preview_flags_unknown_subcategory_without_writing_anything(client, test_user, db_session):
    _login(client, test_user)
    file_bytes = _build_workbook([{
        "Product Name": "Bad Subcategory Product", "Type": "standard",
        "Category": "Nonexistent Category", "Subcategory": "Nonexistent Subcategory",
        "Cost Price": 1000, "Selling Price": 2000,
    }])
    resp = client.post(
        "/api/product-imports/preview",
        files={"file": ("import.xlsx", file_bytes,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["error_rows"] == 1
    assert any("Nonexistent Subcategory" in e for e in body["rows"][0]["errors"])

    from app.models.product import Product
    assert db_session.query(Product).filter(Product.name == "Bad Subcategory Product").first() is None


def test_commit_creates_products_with_server_generated_ids(client, test_user):
    _login(client, test_user)
    resp = client.post("/api/product-imports/commit", json={"rows": [
        {"name": "Committed Import Product", "product_type": "standard", "unit": "Piece",
         "cost_price": "5000", "selling_price": "8000"},
    ]})
    assert resp.status_code == 200
    body = resp.json()
    assert body["created_products"] == 1
    assert body["error"] is None

    product = client.get(f"/api/products/{body['product_ids'][0]}").json()
    assert product["product_code"].startswith("PROD-")
    assert product["business_id"] is not None


def test_non_master_cannot_commit_import(client, test_user, db_session):
    from app.core.security import hash_password
    from app.models.user import User

    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "Import RBAC Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    user = User(username="importrbacuser", email="importrbacuser@example.com", full_name="Import RBAC Employee",
                password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True)
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "importrbacuser@example.com", "password": "EmpPass1!"})

    resp = client.post("/api/product-imports/commit", json={"rows": [
        {"name": "Should Not Be Created", "unit": "Piece"},
    ]})
    assert resp.status_code == 403
