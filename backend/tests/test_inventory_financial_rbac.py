"""Tests for this turn's access control brief - Master sees everything,
Employee sees stock quantities/status but never financial data (rate,
stock value, purchase cost/records)."""
from app.core.security import hash_password
from app.models.user import User


def _login(client, test_user):
    resp = client.post("/api/auth/login", json={"identifier": "test@example.com", "password": "TestPass123!"})
    assert resp.status_code == 200


def _create_employee(client, db_session, username, email):
    employee = client.post("/api/employees/", json={
        "name": username, "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    user = User(
        username=username, email=email, full_name=username,
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    resp = client.post("/api/auth/login", json={"identifier": email, "password": "EmpPass1!"})
    assert resp.status_code == 200
    return user


def test_master_sees_material_financial_fields(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Master Visible Material", "unit": "Sheets", "opening_stock": "10",
        "minimum_stock": "1", "average_rate": "500.00",
    }).json()

    resp = client.get(f"/api/materials/{material['id']}").json()
    assert resp["average_rate"] is not None
    assert resp["stock_value"] is not None


def test_employee_material_financial_fields_are_genuinely_null(client, test_user, db_session):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Employee Hidden Rate Material", "unit": "Sheets", "opening_stock": "10",
        "minimum_stock": "1", "average_rate": "500.00",
    }).json()

    _create_employee(client, db_session, "matrbacuser", "matrbacuser@example.com")
    resp = client.get(f"/api/materials/{material['id']}").json()
    assert resp["average_rate"] is None
    assert resp["stock_value"] is None
    # Non-financial fields remain visible.
    assert resp["current_stock"] is not None
    assert resp["stock_status"] is not None


def test_employee_material_list_also_redacts_financial_fields(client, test_user, db_session):
    _login(client, test_user)
    client.post("/api/materials/", json={
        "name": "List Redaction Material", "unit": "Sheets", "opening_stock": "10",
        "minimum_stock": "1", "average_rate": "300.00",
    })

    _create_employee(client, db_session, "listrbacuser", "listrbacuser@example.com")
    materials = client.get("/api/materials/").json()
    match = next(m for m in materials if m["name"] == "List Redaction Material")
    assert match["average_rate"] is None
    assert match["stock_value"] is None


def test_employee_cannot_view_purchases_at_all(client, test_user, db_session):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Purchase RBAC Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Purchase RBAC Material", "unit": "Sheets", "opening_stock": "0", "minimum_stock": "1",
    }).json()
    purchase = client.post("/api/purchases/", json={
        "date": "2026-08-16T00:00:00", "supplier_id": supplier["id"], "material_id": material["id"],
        "quantity": "5", "unit": "Sheets", "rate": "500.00", "gst_percent": "18", "payment_status": "Paid",
    }).json()

    _create_employee(client, db_session, "purchaserbacuser", "purchaserbacuser@example.com")
    list_resp = client.get("/api/purchases/")
    assert list_resp.status_code == 403
    get_resp = client.get(f"/api/purchases/{purchase['id']}")
    assert get_resp.status_code == 403


def test_master_can_view_purchases(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/purchases/")
    assert resp.status_code == 200


def test_employee_dashboard_stock_value_is_null(client, test_user, db_session):
    _login(client, test_user)
    client.post("/api/materials/", json={
        "name": "Dashboard RBAC Material", "unit": "Sheets", "opening_stock": "10",
        "minimum_stock": "1", "average_rate": "200.00",
    })

    _create_employee(client, db_session, "dashrbacuser", "dashrbacuser@example.com")
    resp = client.get("/api/dashboard/stock").json()
    assert resp["total_stock_value"] is None
    assert resp["purchase_value"] is None
    # Non-financial figures remain real.
    assert resp["low_stock_items"] is not None
    assert resp["out_of_stock_items"] is not None
    assert resp["recent_stock_movement"] is not None


def test_master_dashboard_stock_value_is_real(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/dashboard/stock").json()
    assert resp["total_stock_value"] is not None


def test_employee_category_summary_stock_value_is_null(client, test_user, db_session):
    _login(client, test_user)
    client.post("/api/materials/", json={
        "name": "Category Summary RBAC Material", "unit": "Sheets", "opening_stock": "10",
        "minimum_stock": "1", "average_rate": "200.00", "category": "TestCategoryRBAC",
    })

    _create_employee(client, db_session, "catrbacuser", "catrbacuser@example.com")
    resp = client.get("/api/dashboard/stock").json()
    match = next((c for c in resp["category_summary"] if c["category"] == "TestCategoryRBAC"), None)
    assert match is not None
    assert match["stock_value"] is None


def test_employee_excel_export_excludes_financial_sheets_and_columns(client, test_user, db_session):
    _login(client, test_user)
    _create_employee(client, db_session, "excelrbacuser", "excelrbacuser@example.com")

    resp = client.get("/api/reports/stock-dashboard.xlsx")
    assert resp.status_code == 200

    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp.content))
    sheet_names = wb.sheetnames
    assert "Dashboard" not in sheet_names
    assert "Purchases" not in sheet_names
    assert "Material Master" in sheet_names
    assert "Issues" in sheet_names

    material_sheet = wb["Material Master"]
    header_row = [c.value for c in material_sheet[4]]
    assert "Average Rate" not in header_row
    assert "Stock Value" not in header_row


def test_master_excel_export_includes_all_sheets(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/reports/stock-dashboard.xlsx")
    assert resp.status_code == 200

    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp.content))
    assert set(["Dashboard", "Material Master", "Purchases", "Issues", "Suppliers"]).issubset(set(wb.sheetnames))


def test_chatbot_employee_denied_inventory_value_question(client, test_user, db_session):
    _login(client, test_user)
    _create_employee(client, db_session, "chatvaluerbacuser", "chatvaluerbacuser@example.com")

    resp = client.post("/api/chat/", json={"message": "What is the inventory value?"})
    assert "don't have access" in resp.json()["response"].lower()


def test_chatbot_employee_denied_purchase_cost_question(client, test_user, db_session):
    _login(client, test_user)
    _create_employee(client, db_session, "chatcostrbacuser", "chatcostrbacuser@example.com")

    resp = client.post("/api/chat/", json={"message": "What was the purchase cost?"})
    assert "don't have access" in resp.json()["response"].lower()


def test_chatbot_employee_denied_spend_question(client, test_user, db_session):
    _login(client, test_user)
    _create_employee(client, db_session, "chatspendrbacuser", "chatspendrbacuser@example.com")

    resp = client.post("/api/chat/", json={"message": "How much did we spend on plywood?"})
    assert "don't have access" in resp.json()["response"].lower()


def test_chatbot_employee_still_gets_real_stock_status_answers(client, test_user, db_session):
    """The three explicitly-permitted questions from the brief must
    still work normally for an employee."""
    _login(client, test_user)
    client.post("/api/materials/", json={
        "name": "Chat Employee Allowed Material", "unit": "Sheets", "opening_stock": "23", "minimum_stock": "20",
    })
    _create_employee(client, db_session, "chatallowedrbacuser", "chatallowedrbacuser@example.com")

    resp1 = client.post("/api/chat/", json={"message": "How much Chat Employee Allowed Material do we have?"})
    assert "23" in resp1.json()["response"]

    resp2 = client.post("/api/chat/", json={"message": "Which materials are low in stock?"})
    assert resp2.status_code == 200
    assert "don't have access" not in resp2.json()["response"].lower()

    resp3 = client.post("/api/chat/", json={"message": "Which materials are out of stock?"})
    assert resp3.status_code == 200
    assert "don't have access" not in resp3.json()["response"].lower()


def test_chatbot_employee_denied_recent_purchases(client, test_user, db_session):
    """Matches the Purchases API restriction - an employee should not
    learn what was purchased recently via chat if they can't view it
    on the Purchases page either."""
    _login(client, test_user)
    _create_employee(client, db_session, "chatpurchaserbacuser", "chatpurchaserbacuser@example.com")

    resp = client.post("/api/chat/", json={"message": "What materials were purchased recently?"})
    assert "master accounts only" in resp.json()["response"].lower()


def test_chatbot_master_still_gets_real_inventory_value(client, test_user):
    _login(client, test_user)
    resp = client.post("/api/chat/", json={"message": "What is the inventory value?"})
    assert "don't have access" not in resp.json()["response"].lower()
    assert "Rs" in resp.json()["response"]
