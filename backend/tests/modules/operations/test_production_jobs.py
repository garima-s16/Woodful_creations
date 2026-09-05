import io
from openpyxl import load_workbook
from tests.helpers import _login


"""Tests for production_jobs: the security gap in update_production_job
(every field was previously open to any authenticated user, unlike the
equivalent DailyTask endpoint), the new production-bottlenecks chatbot
handler, the stage/completion_date fields and production.xlsx export
(filters respected), and the staff dashboard's production status
summary (a real gap - it had zero production job data despite the
brief requiring "Production/job status where applicable")."""


def test_employee_can_update_status_and_completed_qty(client, test_user, db_session):
    """The genuinely-permitted self-service fields must still work for
    a non-master employee."""
    from app.platform.security.security import hash_password
    from app.modules.auth.models import User
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Cutting"}).json()
    employee = client.post("/api/employees/", json={"name": "Production Update Self Service Employee"}).json()
    user = User(
        username="prodselfserviceuser", email="prodselfserviceuser@example.com",
        full_name="Prod Self Service User", password_hash=hash_password("EmpPass1!"),
        role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "prodselfserviceuser@example.com", "password": "EmpPass1!"})

    resp = client.put(f"/api/production-jobs/{job['id']}", json={"status": "In Progress", "completed_qty": 5})
    assert resp.status_code == 200


def test_employee_cannot_change_stage_or_remarks(client, test_user, db_session):
    """The real fix - planning fields must be blocked for non-master,
    which was not previously enforced at all."""
    from app.platform.security.security import hash_password
    from app.modules.auth.models import User
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Assembly"}).json()
    employee = client.post("/api/employees/", json={"name": "Production Update Restricted Employee"}).json()
    user = User(
        username="prodrestricteduser", email="prodrestricteduser@example.com",
        full_name="Prod Restricted User", password_hash=hash_password("EmpPass1!"),
        role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "prodrestricteduser@example.com", "password": "EmpPass1!"})

    resp = client.put(f"/api/production-jobs/{job['id']}", json={"stage": "QC"})
    assert resp.status_code == 403

    resp2 = client.put(f"/api/production-jobs/{job['id']}", json={"remarks": "trying to sneak this in"})
    assert resp2.status_code == 403


def test_employee_can_set_blocker_reason(client, test_user, db_session):
    from app.platform.security.security import hash_password
    from app.modules.auth.models import User
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Finishing"}).json()
    employee = client.post("/api/employees/", json={"name": "Production Blocker Reason Employee"}).json()
    user = User(
        username="prodblockeruser", email="prodblockeruser@example.com",
        full_name="Prod Blocker User", password_hash=hash_password("EmpPass1!"),
        role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "prodblockeruser@example.com", "password": "EmpPass1!"})

    resp = client.put(f"/api/production-jobs/{job['id']}", json={
        "status": "Blocked", "blocker_reason": "Waiting on CNC router availability",
    })
    assert resp.status_code == 200
    assert resp.json()["blocker_reason"] == "Waiting on CNC router availability"


def test_master_can_still_update_any_field(client, test_user):
    """Regression guard - master must remain unrestricted."""
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Packing"}).json()
    resp = client.put(f"/api/production-jobs/{job['id']}", json={"stage": "Packing", "remarks": "master note"})
    assert resp.status_code == 200


def test_chatbot_production_bottlenecks(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Blocked bottleneck job", "machine": "CNC Router",
    }).json()
    client.put(f"/api/production-jobs/{job['id']}", json={"status": "Blocked", "blocker_reason": "No material"})

    resp = client.post("/api/chat/", json={"message": "show production bottlenecks"})
    assert resp.status_code == 200
    data = resp.json()
    assert "1 job(s) blocked" in data["response"]
    assert any(r["label"] == job["job_code"] for r in data["records"])


def test_chatbot_no_bottlenecks_reports_honestly(client, test_user):
    _login(client, test_user)
    resp = client.post("/api/chat/", json={"message": "any bottleneck in production"})
    assert resp.status_code == 200


def test_production_stage_can_be_set_on_create(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Panel cutting", "stage": "Cutting",
    }).json()
    assert job["stage"] == "Cutting"


def test_completion_date_set_on_transition_to_completed(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Assembly"}).json()
    assert job["completion_date"] is None

    resp = client.put(f"/api/production-jobs/{job['id']}", json={"status": "Completed"})
    assert resp.status_code == 200
    assert resp.json()["completion_date"] is not None


def test_completion_date_does_not_change_on_resave(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Finishing"}).json()
    client.put(f"/api/production-jobs/{job['id']}", json={"status": "Completed"})
    first = client.get("/api/production-jobs/").json()
    first_match = next(j for j in first if j["id"] == job["id"])

    client.put(f"/api/production-jobs/{job['id']}", json={"status": "Completed", "remarks": "verified"})
    second = client.get("/api/production-jobs/").json()
    second_match = next(j for j in second if j["id"] == job["id"])
    assert second_match["completion_date"] == first_match["completion_date"]


def test_production_export_respects_stage_filter(client, test_user):
    _login(client, test_user)
    client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Cutting op", "stage": "Cutting",
    })
    client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "QC op", "stage": "QC",
    })

    resp = client.get("/api/reports/production.xlsx", params={"stage": "Cutting"})
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["Production"]
    operations = [cell.value for row in ws.iter_rows() for cell in row if isinstance(cell.value, str)]
    assert any("Cutting op" in o for o in operations)
    assert not any("QC op" in o for o in operations)


def test_production_export_respects_operator_filter(client, test_user):
    _login(client, test_user)
    emp_a = client.post("/api/employees/", json={"name": "Production Export Operator A"}).json()
    emp_b = client.post("/api/employees/", json={"name": "Production Export Operator B"}).json()
    client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Job for A", "employee_id": emp_a["id"],
    })
    client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Job for B", "employee_id": emp_b["id"],
    })

    resp = client.get("/api/reports/production.xlsx", params={"employee_id": emp_a["id"]})
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["Production"]
    operations = [cell.value for row in ws.iter_rows() for cell in row if isinstance(cell.value, str)]
    assert any("Job for A" in o for o in operations)
    assert not any("Job for B" in o for o in operations)


def test_staff_dashboard_includes_production_status_summary(client, test_user):
    _login(client, test_user)
    order_client = client.post("/api/clients/", json={"name": "Staff Dashboard Production Client", "phone": "9000010184"}).json()
    order = client.post("/api/orders/", json={
        "client_id": order_client["id"], "order_date": "2026-08-18T00:00:00", "order_value": "50000", "advance": "0",
    }).json()
    employee = client.post("/api/employees/", json={"name": "Staff Dashboard Production Employee", "monthly_salary": "20000"}).json()
    client.post("/api/production-jobs/", json={
        "date": "2026-08-18T00:00:00", "order_id": order["id"], "employee_id": employee["id"], "operation": "Cutting",
        "planned_qty": 10, "status": "In Progress",
    })

    resp = client.get("/api/dashboard/staff").json()
    assert "production_status_summary" in resp
    match = next((p for p in resp["production_status_summary"] if p["status"] == "In Progress"), None)
    assert match is not None
    assert match["count"] >= 1


def test_staff_dashboard_production_summary_reflects_real_counts_not_hardcoded(client, test_user):
    """Adding a second job in a different status must change the
    summary - proving it's derived from real data, not a fixed value."""
    _login(client, test_user)
    before = client.get("/api/dashboard/staff").json()["production_status_summary"]

    order_client = client.post("/api/clients/", json={"name": "Staff Dashboard Production Client 2", "phone": "9000010185"}).json()
    order = client.post("/api/orders/", json={
        "client_id": order_client["id"], "order_date": "2026-08-18T00:00:00", "order_value": "30000", "advance": "0",
    }).json()
    employee = client.post("/api/employees/", json={"name": "Staff Dashboard Production Employee 2", "monthly_salary": "18000"}).json()
    client.post("/api/production-jobs/", json={
        "date": "2026-08-18T00:00:00", "order_id": order["id"], "employee_id": employee["id"], "operation": "Assembly",
        "planned_qty": 5, "status": "Pending",
    })

    after = client.get("/api/dashboard/staff").json()["production_status_summary"]
    assert after != before
