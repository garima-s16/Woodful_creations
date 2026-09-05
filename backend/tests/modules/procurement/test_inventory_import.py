"""Tests for the Purchase Excel import: template download, preview
matching (existing material/supplier), commit behavior, and
authorization - plus the import's tolerance for real-world header/
value variations (see HEADER_ALIASES/normalize_number/normalize_date/
normalize_match_key in app/utils/purchase_import.py)."""
import io
from openpyxl import Workbook
from tests.helpers import _login


def _build_upload_fixed_headers(rows):
    """Builds a minimal .xlsx with the exact fixed header row the
    parser looks for, plus whatever data rows the test needs."""
    wb = Workbook()
    ws = wb.active
    headers = ["Material", "Specification", "Quantity", "Unit", "Supplier", "Rate", "GST %", "Invoice Number", "Invoice Date", "Remarks"]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_template_downloads_and_has_correct_headers(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/purchase-imports/template")
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    header_row = [c.value for c in ws[5]]  # row 5: logo row, title, subtitle, blank, headers
    assert "Material" in header_row and "Supplier" in header_row and "GST %" in header_row


def test_preview_matches_existing_material_and_supplier(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Import Test Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Import Test Material", "unit": "Sheets", "opening_stock": 5, "minimum_stock": 1,
    }).json()

    upload = _build_upload_fixed_headers([{
        "Material": "Import Test Material", "Specification": "x", "Quantity": 5, "Unit": "Sheets",
        "Supplier": "Import Test Supplier", "Rate": 500, "GST %": 18,
        "Invoice Number": "INV-1", "Invoice Date": "2026-08-01", "Remarks": None,
    }])
    resp = client.post("/api/purchase-imports/preview", files={"file": ("test.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert resp.status_code == 200
    body = resp.json()
    assert body["matched_rows"] == 1
    assert body["error_rows"] == 0
    assert body["rows"][0]["matched_material_id"] == material["id"]
    assert body["rows"][0]["matched_supplier_id"] == supplier["id"]


def test_preview_flags_unmatched_material_as_new(client, test_user):
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Unmatched Material Test Supplier"})
    upload = _build_upload_fixed_headers([{
        "Material": "Totally New Material XYZ", "Specification": None, "Quantity": 3, "Unit": "Sheets",
        "Supplier": "Unmatched Material Test Supplier", "Rate": 200, "GST %": 18,
        "Invoice Number": None, "Invoice Date": None, "Remarks": None,
    }])
    resp = client.post("/api/purchase-imports/preview", files={"file": ("test.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    body = resp.json()
    assert body["new_material_rows"] == 1
    assert body["rows"][0]["is_new_material"] is True
    assert body["rows"][0]["errors"] == []


def test_preview_flags_unmatched_supplier_as_error(client, test_user):
    """Unlike materials, suppliers are never auto-created on import -
    an unmatched supplier is a real error, not a "new" row."""
    _login(client, test_user)
    client.post("/api/materials/", json={"name": "Supplier Error Test Material", "unit": "Sheets", "opening_stock": 1, "minimum_stock": 1})
    upload = _build_upload_fixed_headers([{
        "Material": "Supplier Error Test Material", "Specification": None, "Quantity": 3, "Unit": "Sheets",
        "Supplier": "Nonexistent Supplier XYZ", "Rate": 200, "GST %": 18,
        "Invoice Number": None, "Invoice Date": None, "Remarks": None,
    }])
    resp = client.post("/api/purchase-imports/preview", files={"file": ("test.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    body = resp.json()
    assert body["error_rows"] == 1
    assert any("supplier" in e.lower() for e in body["rows"][0]["errors"])


def test_commit_creates_real_purchase_for_matched_row(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Commit Test Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Commit Test Material", "unit": "Sheets", "opening_stock": 5, "minimum_stock": 1,
    }).json()

    resp = client.post("/api/purchase-imports/commit", json={"rows": [{
        "material_name": "Commit Test Material", "matched_material_id": material["id"],
        "matched_supplier_id": supplier["id"], "quantity": "10", "unit": "Sheets",
        "rate": "500.00", "gst_percent": "18", "create_new_material": False,
    }]})
    assert resp.status_code == 200
    body = resp.json()
    assert body["created_purchases"] == 1
    assert body["created_materials"] == 0
    assert body["error"] is None

    updated_material = client.get(f"/api/materials/{material['id']}").json()
    assert updated_material["current_stock"] == 15  # 5 opening + 10 imported


def test_commit_creates_new_material_when_requested(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "New Material Commit Supplier"}).json()

    resp = client.post("/api/purchase-imports/commit", json={"rows": [{
        "material_name": "Brand New Imported Material", "specification": "test spec",
        "matched_material_id": None, "matched_supplier_id": supplier["id"],
        "quantity": "5", "unit": "Sheets", "rate": "300.00", "gst_percent": "18",
        "create_new_material": True,
    }]})
    assert resp.status_code == 200
    body = resp.json()
    assert body["created_materials"] == 1
    assert body["created_purchases"] == 1

    materials = client.get("/api/materials/", params={"search": "Brand New Imported Material"}).json()
    assert len(materials) == 1
    assert materials[0]["current_stock"] == 5


def test_commit_stops_and_reports_partial_success_on_bad_row(client, test_user):
    """The honest, actually-implemented behavior: not atomic - rows
    before the failure are genuinely committed, and the response says so."""
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Partial Commit Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Partial Commit Material", "unit": "Sheets", "opening_stock": 5, "minimum_stock": 1,
    }).json()

    resp = client.post("/api/purchase-imports/commit", json={"rows": [
        {
            "material_name": "Partial Commit Material", "matched_material_id": material["id"],
            "matched_supplier_id": supplier["id"], "quantity": "2", "unit": "Sheets",
            "rate": "500.00", "gst_percent": "18", "create_new_material": False,
        },
        {
            "material_name": "No Material Row", "matched_material_id": None,
            "matched_supplier_id": supplier["id"], "quantity": "2", "unit": "Sheets",
            "rate": "500.00", "gst_percent": "18", "create_new_material": False,  # no material - must fail
        },
    ]})
    assert resp.status_code == 200
    body = resp.json()
    assert body["created_purchases"] == 1  # first row succeeded
    assert body["error"] is not None
    assert "row 2" in body["error"].lower()


def test_import_requires_master_or_manager_role(client, test_user, db_session):
    from app.platform.security.security import hash_password
    from app.modules.auth.models import User

    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "Import Perm Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    limited_user = User(
        username="importpermuser", email="importpermuser@example.com", full_name="Limited Import User",
        password_hash=hash_password("LimitedPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(limited_user)
    db_session.commit()

    resp = client.post("/api/auth/login", json={"identifier": "importpermuser@example.com", "password": "LimitedPass1!"})
    assert resp.status_code == 200

    upload = _build_upload_fixed_headers([])
    preview_resp = client.post("/api/purchase-imports/preview", files={"file": ("test.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert preview_resp.status_code == 403


def _build_upload(headers, rows):
    """Like _build_upload_fixed_headers above, but takes the
    header row explicitly so a test can use whatever real-world header
    variation it's checking, instead of always the exact template text."""
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _preview(client, headers, rows):
    upload = _build_upload(headers, rows)
    return client.post(
        "/api/purchase-imports/preview",
        files={"file": ("test.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


CANONICAL_HEADERS = ["Material", "Specification", "Quantity", "Unit", "Supplier",
                     "Rate", "GST %", "Invoice Number", "Invoice Date", "Remarks"]


def test_header_capitalization_variation(client, test_user):
    """1. ALL CAPS / different-case headers still resolve to the same columns."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Norm Header Case Supplier"})
    headers = ["MATERIAL", "SPECIFICATION", "QUANTITY", "UNIT", "SUPPLIER", "RATE", "GST %", "INVOICE NUMBER", "INVOICE DATE", "REMARKS"]
    resp = _preview(client, headers, [["Totally New Header Case Material", "x", 3, "Sheets", "Norm Header Case Supplier", 200, 18, "INV-1", "2026-08-01", ""]])
    assert resp.status_code == 200
    body = resp.json()
    assert body["total_rows"] == 1
    assert body["rows"][0]["errors"] == []


def test_header_whitespace_variation(client, test_user):
    """2. Leading/trailing whitespace around a header cell is ignored."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Norm Header Whitespace Supplier"})
    headers = [" Material ", "Specification", "Quantity", "Unit", " Supplier", "Rate", "GST %", "Invoice Number", "Invoice Date", "Remarks"]
    resp = _preview(client, headers, [["Whitespace Header Material", "x", 2, "Sheets", "Norm Header Whitespace Supplier", 100, 18, "INV-2", "2026-08-01", ""]])
    assert resp.status_code == 200
    assert resp.json()["total_rows"] == 1


def test_material_name_header_alias(client, test_user):
    """3. "Material Name" is accepted as an alias for "Material"."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Material Alias Supplier"})
    headers = ["Material Name", "Specification", "Quantity", "Unit", "Supplier", "Rate", "GST %", "Invoice Number", "Invoice Date", "Remarks"]
    resp = _preview(client, headers, [["Material Alias Header Material", "x", 1, "Sheets", "Material Alias Supplier", 100, 18, "INV-3", "2026-08-01", ""]])
    assert resp.status_code == 200
    body = resp.json()
    assert body["total_rows"] == 1
    assert body["rows"][0]["material_name"] == "Material Alias Header Material"


def test_supplier_name_header_alias(client, test_user):
    """4. "Supplier Name" is accepted as an alias for "Supplier"."""
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Supplier Alias Header Supplier"}).json()
    headers = ["Material", "Specification", "Quantity", "Unit", "Supplier Name", "Rate", "GST %", "Invoice Number", "Invoice Date", "Remarks"]
    resp = _preview(client, headers, [["Supplier Alias Header Material", "x", 1, "Sheets", "Supplier Alias Header Supplier", 100, 18, "INV-4", "2026-08-01", ""]])
    assert resp.status_code == 200
    body = resp.json()
    assert body["rows"][0]["matched_supplier_id"] == supplier["id"]


def test_gst_percent_header_variation(client, test_user):
    """5. "GST%" (no space) is accepted as an alias for "GST %"."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "GST Header Variation Supplier"})
    headers = ["Material", "Specification", "Quantity", "Unit", "Supplier", "Rate", "GST%", "Invoice Number", "Invoice Date", "Remarks"]
    resp = _preview(client, headers, [["GST Header Variation Material", "x", 1, "Sheets", "GST Header Variation Supplier", 100, 18, "INV-5", "2026-08-01", ""]])
    assert resp.status_code == 200
    assert resp.json()["total_rows"] == 1


def test_invoice_number_header_variation(client, test_user):
    """6. "Invoice No" is accepted as an alias for "Invoice Number"."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Invoice No Header Supplier"})
    headers = ["Material", "Specification", "Quantity", "Unit", "Supplier", "Rate", "GST %", "Invoice No", "Invoice Date", "Remarks"]
    resp = _preview(client, headers, [["Invoice No Header Material", "x", 1, "Sheets", "Invoice No Header Supplier", 100, 18, "INV-6", "2026-08-01", ""]])
    assert resp.status_code == 200
    assert resp.json()["total_rows"] == 1


def test_invoice_date_header_variation(client, test_user):
    """7. "Date" is accepted as an alias for "Invoice Date"."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Invoice Date Header Supplier"})
    headers = ["Material", "Specification", "Quantity", "Unit", "Supplier", "Rate", "GST %", "Invoice Number", "Date", "Remarks"]
    resp = _preview(client, headers, [["Invoice Date Header Material", "x", 1, "Sheets", "Invoice Date Header Supplier", 100, 18, "INV-7", "2026-08-01", ""]])
    assert resp.status_code == 200
    assert resp.json()["total_rows"] == 1


def test_comma_formatted_quantity(client, test_user):
    """8. "1,250" is parsed as the number 1250."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Comma Qty Supplier"})
    resp = _preview(client, CANONICAL_HEADERS, [["Comma Qty Material", "x", "1,250", "Sheets", "Comma Qty Supplier", 100, 18, "INV-8", "2026-08-01", ""]])
    assert resp.status_code == 200
    body = resp.json()
    assert body["rows"][0]["errors"] == []
    assert float(body["rows"][0]["quantity"]) == 1250


def test_currency_formatted_rate(client, test_user):
    """9. "₹1,250" is parsed as the number 1250 for Rate."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Currency Rate Supplier"})
    resp = _preview(client, CANONICAL_HEADERS, [["Currency Rate Material", "x", 2, "Sheets", "Currency Rate Supplier", "₹1,250", 18, "INV-9", "2026-08-01", ""]])
    assert resp.status_code == 200
    body = resp.json()
    assert body["rows"][0]["errors"] == []
    assert float(body["rows"][0]["rate"]) == 1250


def test_common_date_format_variation(client, test_user):
    """10. A day-first "DD-MM-YYYY" date string is accepted (in addition
    to the existing YYYY-MM-DD and native Excel-date support)."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Date Format Supplier"})
    resp = _preview(client, CANONICAL_HEADERS, [["Date Format Material", "x", 1, "Sheets", "Date Format Supplier", 100, 18, "INV-10", "01-08-2026", ""]])
    assert resp.status_code == 200
    body = resp.json()
    assert body["rows"][0]["errors"] == []
    assert body["rows"][0]["invoice_date"].startswith("2026-08-01")


def test_unknown_material_still_flagged(client, test_user):
    """11. Normalization never invents a match - a genuinely unknown
    material is still flagged (as "new", not silently matched)."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Unknown Material Norm Supplier"})
    resp = _preview(client, CANONICAL_HEADERS, [["Genuinely Unknown Norm Material XYZ", "x", 1, "Sheets", "Unknown Material Norm Supplier", 100, 18, "INV-11", "2026-08-01", ""]])
    assert resp.status_code == 200
    body = resp.json()
    assert body["new_material_rows"] == 1
    assert body["rows"][0]["is_new_material"] is True


def test_unknown_supplier_still_flagged(client, test_user):
    """12. A genuinely unknown supplier is still a hard error, not
    silently matched or auto-created."""
    _login(client, test_user)
    client.post("/api/materials/", json={"name": "Unknown Supplier Norm Material", "unit": "Sheets", "opening_stock": 1, "minimum_stock": 1})
    resp = _preview(client, CANONICAL_HEADERS, [["Unknown Supplier Norm Material", "x", 1, "Sheets", "Genuinely Unknown Norm Supplier XYZ", 100, 18, "INV-12", "2026-08-01", ""]])
    assert resp.status_code == 200
    body = resp.json()
    assert body["error_rows"] == 1
    assert any("supplier" in e.lower() for e in body["rows"][0]["errors"])


def test_ambiguous_header_is_not_guessed(client, test_user):
    """13. Two columns that both claim "Material" (the canonical header
    AND an alias) makes the header row unresolvable - the importer
    refuses rather than guessing which one is real."""
    _login(client, test_user)
    headers = ["Material", "Material Name", "Quantity", "Unit", "Supplier", "Rate", "GST %", "Invoice Number", "Invoice Date", "Remarks"]
    resp = _preview(client, headers, [["A", "B", 1, "Sheets", "X", 100, 18, "INV-13", "2026-08-01", ""]])
    assert resp.status_code == 400
    assert "header" in resp.json()["detail"].lower()


def test_invalid_quantity_and_date_still_rejected(client, test_user):
    """14. Normalization tolerates *formatting*, not invalid data -
    a non-numeric quantity and an impossible/unrecognized date are
    still reported as errors, never silently coerced."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Invalid Data Norm Supplier"})
    resp = _preview(client, CANONICAL_HEADERS, [["Invalid Data Norm Material", "x", "not-a-number", "Sheets", "Invalid Data Norm Supplier", 100, 18, "INV-14", "31-02-2026", ""]])
    assert resp.status_code == 200
    body = resp.json()
    assert body["error_rows"] == 1
    errors = " ".join(body["rows"][0]["errors"]).lower()
    assert "quantity" in errors
    assert "date" in errors
