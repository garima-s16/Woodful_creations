"""Tests for Family 3 (Sales) - genuine gaps found in estimate-to-
order conversion (no status check existed at all, meaning a rejected
estimate could become an order, and re-converting an already-
converted estimate would silently orphan the first order's link),
plus the new estimates.xlsx export and pending-estimates chatbot
handler."""
import io
from openpyxl import load_workbook


def _login(client, test_user):
    resp = client.post("/api/auth/login", json={"identifier": "test@example.com", "password": "TestPass123!"})
    assert resp.status_code == 200


def test_rejected_estimate_cannot_be_converted_to_order(client, test_user):
    _login(client, test_user)
    client_id = client.post("/api/clients/", json={"name": "Conversion Test Rejected Client", "phone": "9000010171"}).json()["id"]
    estimate = client.post("/api/estimates/", json={"client_id": client_id, "material_cost": "10000"}).json()
    client.put(f"/api/estimates/{estimate['id']}", json={"status": "rejected"})

    resp = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00", "from_estimate_id": estimate["id"],
    })
    assert resp.status_code == 400
    assert "rejected" in resp.json()["detail"].lower()


def test_approved_estimate_converts_successfully(client, test_user):
    _login(client, test_user)
    client_id = client.post("/api/clients/", json={"name": "Conversion Test Approved Client", "phone": "9000010172"}).json()["id"]
    estimate = client.post("/api/estimates/", json={"client_id": client_id, "material_cost": "20000"}).json()
    client.put(f"/api/estimates/{estimate['id']}", json={"status": "approved"})

    resp = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00", "from_estimate_id": estimate["id"],
    })
    assert resp.status_code == 200


def test_estimate_cannot_be_converted_twice(client, test_user):
    _login(client, test_user)
    client_id = client.post("/api/clients/", json={"name": "Conversion Test Double Client", "phone": "9000010173"}).json()["id"]
    estimate = client.post("/api/estimates/", json={"client_id": client_id, "material_cost": "15000"}).json()
    client.put(f"/api/estimates/{estimate['id']}", json={"status": "approved"})
    first = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00", "from_estimate_id": estimate["id"],
    })
    assert first.status_code == 200

    second = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00", "from_estimate_id": estimate["id"],
    })
    assert second.status_code == 400
    assert "already" in second.json()["detail"].lower()


def test_estimates_export_requires_master(client, test_user, db_session):
    from app.core.security import hash_password
    from app.models.user import User
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Estimates Export Permission Employee"}).json()
    user = User(
        username="estimatesexportuser", email="estimatesexportuser@example.com",
        full_name="Estimates Export User", password_hash=hash_password("EmpPass1!"),
        role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "estimatesexportuser@example.com", "password": "EmpPass1!"})

    resp = client.get("/api/reports/estimates.xlsx")
    assert resp.status_code == 403


def test_estimates_export_works_for_master(client, test_user):
    _login(client, test_user)
    client_id = client.post("/api/clients/", json={"name": "Estimates Export Master Client", "phone": "9000010174"}).json()["id"]
    client.post("/api/estimates/", json={"client_id": client_id, "material_cost": "5000"})

    resp = client.get("/api/reports/estimates.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert "Estimates" in wb.sheetnames


def test_chatbot_pending_estimates(client, test_user):
    _login(client, test_user)
    client_id = client.post("/api/clients/", json={"name": "Chatbot Pending Estimate Client", "phone": "9000010175"}).json()["id"]
    estimate = client.post("/api/estimates/", json={"client_id": client_id, "material_cost": "8000"}).json()
    client.put(f"/api/estimates/{estimate['id']}", json={"status": "sent"})

    resp = client.post("/api/chat/", json={"message": "show pending estimates"})
    assert resp.status_code == 200
    assert estimate["estimate_code"] in resp.json()["response"] or any(
        r["label"] == estimate["estimate_code"] for r in resp.json()["records"]
    )
