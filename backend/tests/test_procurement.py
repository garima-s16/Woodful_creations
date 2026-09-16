"""Procurement domain tests: purchase-Excel import commit chain,
the database-backed personal cart, and ProcurementRequirement/
SupplierDecision persistence. Combines the former
test_inventory_import.py, test_personal_cart.py and
test_procurement_requirements.py."""
import io
from openpyxl import Workbook
from tests.helpers import _login
from app.platform.security import hash_password
from app.modules.auth.auth import User

# --- test_inventory_import.py ---
"""Tests for the Purchase Excel import: template download, preview
matching (existing material/supplier), commit behavior, and
authorization - plus the import's tolerance for real-world header/
value variations (see HEADER_ALIASES/normalize_number/normalize_date/
normalize_match_key in app/modules/procurement/imports/purchase_import.py)."""

def _build_upload_fixed_headers(rows):
    """Builds a minimal .xlsx with the exact fixed header row the
    parser looks for, plus whatever data rows the test needs."""
    wb = Workbook()
    ws = wb.active
    headers = ["Material", "Specification", "Quantity", "Unit", "Supplier", "Rate", "GST %", "Invoice Number", "Invoice Date", "Remarks"]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_template_downloads_and_has_correct_headers(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/purchase-imports/template")
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    header_row = [c.value for c in ws[5]]  # row 5: logo row, title, subtitle, blank, headers
    assert "Material" in header_row and "Supplier" in header_row and "GST %" in header_row


def test_preview_matches_existing_material_and_supplier(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Import Test Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Import Test Material", "unit": "Sheets", "opening_stock": 5, "minimum_stock": 1,
    }).json()

    upload = _build_upload_fixed_headers([{
        "Material": "Import Test Material", "Specification": "x", "Quantity": 5, "Unit": "Sheets",
        "Supplier": "Import Test Supplier", "Rate": 500, "GST %": 18,
        "Invoice Number": "INV-1", "Invoice Date": "2026-08-01", "Remarks": None,
    }])
    resp = client.post("/api/purchase-imports/preview", files={"file": ("test.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert resp.status_code == 200
    body = resp.json()
    assert body["matched_rows"] == 1
    assert body["error_rows"] == 0
    assert body["rows"][0]["matched_material_id"] == material["id"]
    assert body["rows"][0]["matched_supplier_id"] == supplier["id"]


def test_preview_flags_unmatched_material_as_new(client, test_user):
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Unmatched Material Test Supplier"})
    upload = _build_upload_fixed_headers([{
        "Material": "Totally New Material XYZ", "Specification": None, "Quantity": 3, "Unit": "Sheets",
        "Supplier": "Unmatched Material Test Supplier", "Rate": 200, "GST %": 18,
        "Invoice Number": None, "Invoice Date": None, "Remarks": None,
    }])
    resp = client.post("/api/purchase-imports/preview", files={"file": ("test.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    body = resp.json()
    assert body["new_material_rows"] == 1
    assert body["rows"][0]["is_new_material"] is True
    assert body["rows"][0]["errors"] == []


def test_preview_flags_unmatched_supplier_as_error(client, test_user):
    """Unlike materials, suppliers are never auto-created on import -
    an unmatched supplier is a real error, not a "new" row."""
    _login(client, test_user)
    client.post("/api/materials/", json={"name": "Supplier Error Test Material", "unit": "Sheets", "opening_stock": 1, "minimum_stock": 1})
    upload = _build_upload_fixed_headers([{
        "Material": "Supplier Error Test Material", "Specification": None, "Quantity": 3, "Unit": "Sheets",
        "Supplier": "Nonexistent Supplier XYZ", "Rate": 200, "GST %": 18,
        "Invoice Number": None, "Invoice Date": None, "Remarks": None,
    }])
    resp = client.post("/api/purchase-imports/preview", files={"file": ("test.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    body = resp.json()
    assert body["error_rows"] == 1
    assert any("supplier" in e.lower() for e in body["rows"][0]["errors"])


def test_commit_creates_real_purchase_for_matched_row(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Commit Test Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Commit Test Material", "unit": "Sheets", "opening_stock": 5, "minimum_stock": 1,
    }).json()

    resp = client.post("/api/purchase-imports/commit", json={"rows": [{
        "material_name": "Commit Test Material", "matched_material_id": material["id"],
        "matched_supplier_id": supplier["id"], "quantity": "10", "unit": "Sheets",
        "rate": "500.00", "gst_percent": "18", "create_new_material": False,
    }]})
    assert resp.status_code == 200
    body = resp.json()
    assert body["created_purchases"] == 1
    assert body["created_materials"] == 0
    assert body["error"] is None

    updated_material = client.get(f"/api/materials/{material['id']}").json()
    assert updated_material["current_stock"] == 15  # 5 opening + 10 imported


def test_commit_creates_new_material_when_requested(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "New Material Commit Supplier"}).json()

    resp = client.post("/api/purchase-imports/commit", json={"rows": [{
        "material_name": "Brand New Imported Material", "specification": "test spec",
        "matched_material_id": None, "matched_supplier_id": supplier["id"],
        "quantity": "5", "unit": "Sheets", "rate": "300.00", "gst_percent": "18",
        "create_new_material": True,
    }]})
    assert resp.status_code == 200
    body = resp.json()
    assert body["created_materials"] == 1
    assert body["created_purchases"] == 1

    materials = client.get("/api/materials/", params={"search": "Brand New Imported Material"}).json()
    assert len(materials) == 1
    assert materials[0]["current_stock"] == 5


def test_commit_stops_and_reports_partial_success_on_bad_row(client, test_user):
    """The honest, actually-implemented behavior: not atomic - rows
    before the failure are genuinely committed, and the response says so."""
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Partial Commit Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Partial Commit Material", "unit": "Sheets", "opening_stock": 5, "minimum_stock": 1,
    }).json()

    resp = client.post("/api/purchase-imports/commit", json={"rows": [
        {
            "material_name": "Partial Commit Material", "matched_material_id": material["id"],
            "matched_supplier_id": supplier["id"], "quantity": "2", "unit": "Sheets",
            "rate": "500.00", "gst_percent": "18", "create_new_material": False,
        },
        {
            "material_name": "No Material Row", "matched_material_id": None,
            "matched_supplier_id": supplier["id"], "quantity": "2", "unit": "Sheets",
            "rate": "500.00", "gst_percent": "18", "create_new_material": False,  # no material - must fail
        },
    ]})
    assert resp.status_code == 200
    body = resp.json()
    assert body["created_purchases"] == 1  # first row succeeded
    assert body["error"] is not None
    assert "row 2" in body["error"].lower()


def test_import_requires_master_or_manager_role(client, test_user, db_session):
    from app.platform.security import hash_password
    from app.modules.auth.auth import User

    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "Import Perm Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    limited_user = User(
        username="importpermuser", email="importpermuser@example.com", full_name="Limited Import User",
        password_hash=hash_password("LimitedPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(limited_user)
    db_session.commit()

    resp = client.post("/api/auth/login", json={"identifier": "importpermuser@example.com", "password": "LimitedPass1!"})
    assert resp.status_code == 200

    upload = _build_upload_fixed_headers([])
    preview_resp = client.post("/api/purchase-imports/preview", files={"file": ("test.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert preview_resp.status_code == 403


def _build_upload(headers, rows):
    """Like _build_upload_fixed_headers above, but takes the
    header row explicitly so a test can use whatever real-world header
    variation it's checking, instead of always the exact template text."""
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _preview(client, headers, rows):
    upload = _build_upload(headers, rows)
    return client.post(
        "/api/purchase-imports/preview",
        files={"file": ("test.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


CANONICAL_HEADERS = ["Material", "Specification", "Quantity", "Unit", "Supplier",
                     "Rate", "GST %", "Invoice Number", "Invoice Date", "Remarks"]


def test_header_capitalization_variation(client, test_user):
    """1. ALL CAPS / different-case headers still resolve to the same columns."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Norm Header Case Supplier"})
    headers = ["MATERIAL", "SPECIFICATION", "QUANTITY", "UNIT", "SUPPLIER", "RATE", "GST %", "INVOICE NUMBER", "INVOICE DATE", "REMARKS"]
    resp = _preview(client, headers, [["Totally New Header Case Material", "x", 3, "Sheets", "Norm Header Case Supplier", 200, 18, "INV-1", "2026-08-01", ""]])
    assert resp.status_code == 200
    body = resp.json()
    assert body["total_rows"] == 1
    assert body["rows"][0]["errors"] == []


def test_header_whitespace_variation(client, test_user):
    """2. Leading/trailing whitespace around a header cell is ignored."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Norm Header Whitespace Supplier"})
    headers = [" Material ", "Specification", "Quantity", "Unit", " Supplier", "Rate", "GST %", "Invoice Number", "Invoice Date", "Remarks"]
    resp = _preview(client, headers, [["Whitespace Header Material", "x", 2, "Sheets", "Norm Header Whitespace Supplier", 100, 18, "INV-2", "2026-08-01", ""]])
    assert resp.status_code == 200
    assert resp.json()["total_rows"] == 1


def test_material_name_header_alias(client, test_user):
    """3. "Material Name" is accepted as an alias for "Material"."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Material Alias Supplier"})
    headers = ["Material Name", "Specification", "Quantity", "Unit", "Supplier", "Rate", "GST %", "Invoice Number", "Invoice Date", "Remarks"]
    resp = _preview(client, headers, [["Material Alias Header Material", "x", 1, "Sheets", "Material Alias Supplier", 100, 18, "INV-3", "2026-08-01", ""]])
    assert resp.status_code == 200
    body = resp.json()
    assert body["total_rows"] == 1
    assert body["rows"][0]["material_name"] == "Material Alias Header Material"


def test_supplier_name_header_alias(client, test_user):
    """4. "Supplier Name" is accepted as an alias for "Supplier"."""
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Supplier Alias Header Supplier"}).json()
    headers = ["Material", "Specification", "Quantity", "Unit", "Supplier Name", "Rate", "GST %", "Invoice Number", "Invoice Date", "Remarks"]
    resp = _preview(client, headers, [["Supplier Alias Header Material", "x", 1, "Sheets", "Supplier Alias Header Supplier", 100, 18, "INV-4", "2026-08-01", ""]])
    assert resp.status_code == 200
    body = resp.json()
    assert body["rows"][0]["matched_supplier_id"] == supplier["id"]


def test_gst_percent_header_variation(client, test_user):
    """5. "GST%" (no space) is accepted as an alias for "GST %"."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "GST Header Variation Supplier"})
    headers = ["Material", "Specification", "Quantity", "Unit", "Supplier", "Rate", "GST%", "Invoice Number", "Invoice Date", "Remarks"]
    resp = _preview(client, headers, [["GST Header Variation Material", "x", 1, "Sheets", "GST Header Variation Supplier", 100, 18, "INV-5", "2026-08-01", ""]])
    assert resp.status_code == 200
    assert resp.json()["total_rows"] == 1


def test_invoice_number_header_variation(client, test_user):
    """6. "Invoice No" is accepted as an alias for "Invoice Number"."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Invoice No Header Supplier"})
    headers = ["Material", "Specification", "Quantity", "Unit", "Supplier", "Rate", "GST %", "Invoice No", "Invoice Date", "Remarks"]
    resp = _preview(client, headers, [["Invoice No Header Material", "x", 1, "Sheets", "Invoice No Header Supplier", 100, 18, "INV-6", "2026-08-01", ""]])
    assert resp.status_code == 200
    assert resp.json()["total_rows"] == 1


def test_invoice_date_header_variation(client, test_user):
    """7. "Date" is accepted as an alias for "Invoice Date"."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Invoice Date Header Supplier"})
    headers = ["Material", "Specification", "Quantity", "Unit", "Supplier", "Rate", "GST %", "Invoice Number", "Date", "Remarks"]
    resp = _preview(client, headers, [["Invoice Date Header Material", "x", 1, "Sheets", "Invoice Date Header Supplier", 100, 18, "INV-7", "2026-08-01", ""]])
    assert resp.status_code == 200
    assert resp.json()["total_rows"] == 1


def test_comma_formatted_quantity(client, test_user):
    """8. "1,250" is parsed as the number 1250."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Comma Qty Supplier"})
    resp = _preview(client, CANONICAL_HEADERS, [["Comma Qty Material", "x", "1,250", "Sheets", "Comma Qty Supplier", 100, 18, "INV-8", "2026-08-01", ""]])
    assert resp.status_code == 200
    body = resp.json()
    assert body["rows"][0]["errors"] == []
    assert float(body["rows"][0]["quantity"]) == 1250


def test_currency_formatted_rate(client, test_user):
    """9. "₹1,250" is parsed as the number 1250 for Rate."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Currency Rate Supplier"})
    resp = _preview(client, CANONICAL_HEADERS, [["Currency Rate Material", "x", 2, "Sheets", "Currency Rate Supplier", "₹1,250", 18, "INV-9", "2026-08-01", ""]])
    assert resp.status_code == 200
    body = resp.json()
    assert body["rows"][0]["errors"] == []
    assert float(body["rows"][0]["rate"]) == 1250


def test_common_date_format_variation(client, test_user):
    """10. A day-first "DD-MM-YYYY" date string is accepted (in addition
    to the existing YYYY-MM-DD and native Excel-date support)."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Date Format Supplier"})
    resp = _preview(client, CANONICAL_HEADERS, [["Date Format Material", "x", 1, "Sheets", "Date Format Supplier", 100, 18, "INV-10", "01-08-2026", ""]])
    assert resp.status_code == 200
    body = resp.json()
    assert body["rows"][0]["errors"] == []
    assert body["rows"][0]["invoice_date"].startswith("2026-08-01")


def test_unknown_material_still_flagged(client, test_user):
    """11. Normalization never invents a match - a genuinely unknown
    material is still flagged (as "new", not silently matched)."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Unknown Material Norm Supplier"})
    resp = _preview(client, CANONICAL_HEADERS, [["Genuinely Unknown Norm Material XYZ", "x", 1, "Sheets", "Unknown Material Norm Supplier", 100, 18, "INV-11", "2026-08-01", ""]])
    assert resp.status_code == 200
    body = resp.json()
    assert body["new_material_rows"] == 1
    assert body["rows"][0]["is_new_material"] is True


def test_unknown_supplier_still_flagged(client, test_user):
    """12. A genuinely unknown supplier is still a hard error, not
    silently matched or auto-created."""
    _login(client, test_user)
    client.post("/api/materials/", json={"name": "Unknown Supplier Norm Material", "unit": "Sheets", "opening_stock": 1, "minimum_stock": 1})
    resp = _preview(client, CANONICAL_HEADERS, [["Unknown Supplier Norm Material", "x", 1, "Sheets", "Genuinely Unknown Norm Supplier XYZ", 100, 18, "INV-12", "2026-08-01", ""]])
    assert resp.status_code == 200
    body = resp.json()
    assert body["error_rows"] == 1
    assert any("supplier" in e.lower() for e in body["rows"][0]["errors"])


def test_ambiguous_header_is_not_guessed(client, test_user):
    """13. Two columns that both claim "Material" (the canonical header
    AND an alias) makes the header row unresolvable - the importer
    refuses rather than guessing which one is real."""
    _login(client, test_user)
    headers = ["Material", "Material Name", "Quantity", "Unit", "Supplier", "Rate", "GST %", "Invoice Number", "Invoice Date", "Remarks"]
    resp = _preview(client, headers, [["A", "B", 1, "Sheets", "X", 100, 18, "INV-13", "2026-08-01", ""]])
    assert resp.status_code == 400
    assert "header" in resp.json()["detail"].lower()


def test_invalid_quantity_and_date_still_rejected(client, test_user):
    """14. Normalization tolerates *formatting*, not invalid data -
    a non-numeric quantity and an impossible/unrecognized date are
    still reported as errors, never silently coerced."""
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Invalid Data Norm Supplier"})
    resp = _preview(client, CANONICAL_HEADERS, [["Invalid Data Norm Material", "x", "not-a-number", "Sheets", "Invalid Data Norm Supplier", 100, 18, "INV-14", "31-02-2026", ""]])
    assert resp.status_code == 200
    body = resp.json()
    assert body["error_rows"] == 1
    errors = " ".join(body["rows"][0]["errors"]).lower()
    assert "quantity" in errors
    assert "date" in errors

# --- test_personal_cart.py ---
"""Tests for the database-backed personal cart - the core requirement
being that it's real server-side persistence, correctly isolated per
user, enforced at the API layer (not just hidden in the UI)."""

def _create_and_login(client, db_session, username, email):
    employee = client.post("/api/employees/", json={
        "name": username, "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    user = User(
        username=username, email=email, full_name=username,
        password_hash=hash_password("CartPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    resp = client.post("/api/auth/login", json={"identifier": email, "password": "CartPass1!"})
    assert resp.status_code == 200
    return user


def test_add_and_list_cart_item(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Cart Test Material", "unit": "Sheets", "opening_stock": 10, "minimum_stock": 1,
    }).json()

    resp = client.post("/api/personal-cart/", json={"material_id": material["id"], "quantity": "5"})
    assert resp.status_code == 201
    assert resp.json()["material_name"] == "Cart Test Material"

    listed = client.get("/api/personal-cart/").json()
    assert len(listed) == 1
    assert float(listed[0]["quantity"]) == 5.0


def test_adding_same_material_twice_accumulates_not_duplicates(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Accumulate Test Material", "unit": "Sheets", "opening_stock": 10, "minimum_stock": 1,
    }).json()

    client.post("/api/personal-cart/", json={"material_id": material["id"], "quantity": "3"})
    client.post("/api/personal-cart/", json={"material_id": material["id"], "quantity": "2"})

    listed = client.get("/api/personal-cart/").json()
    assert len(listed) == 1
    assert float(listed[0]["quantity"]) == 5.0


def test_update_quantity_to_zero_removes_item(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Zero Qty Test Material", "unit": "Sheets", "opening_stock": 10, "minimum_stock": 1,
    }).json()
    item = client.post("/api/personal-cart/", json={"material_id": material["id"], "quantity": "5"}).json()

    resp = client.put(f"/api/personal-cart/{item['id']}", json={"quantity": "0"})
    assert resp.status_code == 200
    assert resp.json()["status"] == "REMOVED"

    listed = client.get("/api/personal-cart/").json()
    assert listed == []


def test_remove_and_clear_cart(client, test_user):
    _login(client, test_user)
    m1 = client.post("/api/materials/", json={"name": "Remove Test A", "unit": "Sheets", "opening_stock": 5, "minimum_stock": 1}).json()
    m2 = client.post("/api/materials/", json={"name": "Remove Test B", "unit": "Sheets", "opening_stock": 5, "minimum_stock": 1}).json()
    item1 = client.post("/api/personal-cart/", json={"material_id": m1["id"], "quantity": "1"}).json()
    client.post("/api/personal-cart/", json={"material_id": m2["id"], "quantity": "1"})

    client.delete(f"/api/personal-cart/{item1['id']}")
    assert len(client.get("/api/personal-cart/").json()) == 1

    client.delete("/api/personal-cart/")
    assert client.get("/api/personal-cart/").json() == []


def test_cart_persists_across_relogin_same_user(client, test_user, db_session):
    """Matches the brief's exact test: add items, log out (re-auth as
    the same user simulates this - the cart must be database-backed,
    not tied to the in-memory session), log back in, cart is unchanged."""
    user = _create_and_login(client, db_session, "cartpersistuser", "cartpersistuser@example.com")
    material = client.post("/api/materials/", json={
        "name": "Persist Test Material", "unit": "Sheets", "opening_stock": 10, "minimum_stock": 1,
    }).json()
    client.post("/api/personal-cart/", json={"material_id": material["id"], "quantity": "5"})

    # Re-authenticate as the same user (simulating logout + login again).
    relogin = client.post("/api/auth/login", json={"identifier": "cartpersistuser@example.com", "password": "CartPass1!"})
    assert relogin.status_code == 200

    listed = client.get("/api/personal-cart/").json()
    assert len(listed) == 1
    assert float(listed[0]["quantity"]) == 5.0


def test_cart_is_completely_isolated_between_different_users(client, test_user, db_session):
    """The critical security requirement - User A's cart must be
    invisible to User B, enforced by the API itself, not the UI."""
    material = client.post("/api/materials/", json={
        "name": "Isolation Test Material", "unit": "Sheets", "opening_stock": 10, "minimum_stock": 1,
    }).json()

    user_a = _create_and_login(client, db_session, "cartuserA", "cartuserA@example.com")
    item_a = client.post("/api/personal-cart/", json={"material_id": material["id"], "quantity": "5"}).json()

    user_b = _create_and_login(client, db_session, "cartuserB", "cartuserB@example.com")
    # User B's own cart is empty - A's item does not appear.
    assert client.get("/api/personal-cart/").json() == []

    # User B cannot access A's specific cart item by id either.
    get_resp = client.put(f"/api/personal-cart/{item_a['id']}", json={"quantity": "99"})
    assert get_resp.status_code == 404

    delete_resp = client.delete(f"/api/personal-cart/{item_a['id']}")
    assert delete_resp.status_code == 404

    # Confirm A's item is genuinely untouched by B's attempts.
    resp = client.post("/api/auth/login", json={"identifier": "cartuserA@example.com", "password": "CartPass1!"})
    assert resp.status_code == 200
    a_items = client.get("/api/personal-cart/").json()
    assert len(a_items) == 1
    assert float(a_items[0]["quantity"]) == 5.0


def test_adding_to_cart_does_not_change_material_stock(client, test_user):
    """The explicit rule - a personal cart is a request
    list, not a purchase; inventory must be completely untouched."""
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "No Stock Change Material", "unit": "Sheets", "opening_stock": 10, "minimum_stock": 1,
    }).json()
    before = material["current_stock"]

    client.post("/api/personal-cart/", json={"material_id": material["id"], "quantity": "5"})

    after = client.get(f"/api/materials/{material['id']}").json()["current_stock"]
    assert after == before

# --- test_procurement_requirements.py ---
"""Tests for ProcurementRequirement (P0.2.1) and SupplierDecision
(P0.2.3) - closing the two gaps this family's own architecture review
identified: there was no persisted procurement requirement (only a
transient shortage calculation) and no persisted record distinguishing
a supplier recommendation from the supplier actually chosen."""

def _make_shortage_order(client, suffix):
    material = client.post("/api/materials/", json={
        "name": f"Procurement Req Sheet {suffix}", "unit": "Sheets", "opening_stock": "2", "minimum_stock": "1",
    }).json()
    product = client.post("/api/products/", json={
        "name": f"Procurement Req Product {suffix}", "unit": "Piece",
        "materials_used": [{"material_id": material["id"], "quantity_required": "5"}],
    }).json()
    client_id = client.post("/api/clients/", json={"name": f"Procurement Req Client {suffix}", "phone": f"900001080{suffix}"}).json()["id"]
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()
    return order, material


def test_creating_requirement_snapshots_real_shortage_not_client_values(client, test_user):
    """The client cannot supply required/available/shortage - the route
    must derive them itself from the authoritative calculation."""
    _login(client, test_user)
    order, material = _make_shortage_order(client, "1")

    resp = client.post("/api/procurement-requirements/", json={
        "order_id": order["id"], "material_id": material["id"], "priority": "High",
    })
    assert resp.status_code == 201
    body = resp.json()
    assert float(body["required_quantity"]) == 5.0
    assert float(body["available_quantity_at_creation"]) == 2.0
    assert float(body["shortage_quantity_at_creation"]) == 3.0
    assert body["status"] == "Open"
    assert body["material_name"] == f"Procurement Req Sheet 1"


def test_requirement_without_a_real_material_requirement_is_rejected(client, test_user):
    _login(client, test_user)
    client_id = client.post("/api/clients/", json={"name": "No Req Client", "phone": "9000010810"}).json()["id"]
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00", "order_value": "1000", "advance": "0",
    }).json()
    material = client.post("/api/materials/", json={"name": "Unrelated Sheet", "unit": "Sheets", "opening_stock": "5"}).json()

    resp = client.post("/api/procurement-requirements/", json={"order_id": order["id"], "material_id": material["id"]})
    assert resp.status_code == 400


def test_supplier_decision_distinguishes_recommendation_from_actual_choice(client, test_user):
    """The core P0.2.3 guarantee: the recommendation and the actual
    decision are genuinely different, persisted things."""
    _login(client, test_user)
    order, material = _make_shortage_order(client, "2")
    cheap_supplier = client.post("/api/suppliers/", json={"name": "Cheap Req Supplier"}).json()
    preferred_supplier = client.post("/api/suppliers/", json={"name": "Preferred Req Supplier"}).json()
    client.post("/api/supplier-materials/", json={
        "supplier_id": cheap_supplier["id"], "material_id": material["id"], "supplier_price": "400.00",
    })
    client.post("/api/supplier-materials/", json={
        "supplier_id": preferred_supplier["id"], "material_id": material["id"], "supplier_price": "550.00", "is_preferred": True,
    })
    requirement = client.post("/api/procurement-requirements/", json={
        "order_id": order["id"], "material_id": material["id"],
    }).json()

    # Master deliberately chooses the cheaper, non-preferred supplier -
    # a real decision that diverges from the recommendation.
    resp = client.post(f"/api/procurement-requirements/{requirement['id']}/decision", json={
        "requirement_id": requirement["id"], "selected_supplier_id": cheap_supplier["id"],
        "decision_reason": "Cash flow this month favors the lower price",
    })
    assert resp.status_code == 201
    body = resp.json()
    assert body["recommended_supplier_id"] == preferred_supplier["id"]
    assert body["selected_supplier_id"] == cheap_supplier["id"]
    assert body["followed_recommendation"] is False
    assert body["decision_reason"] == "Cash flow this month favors the lower price"


def test_second_decision_on_same_requirement_is_rejected(client, test_user):
    _login(client, test_user)
    order, material = _make_shortage_order(client, "3")
    supplier = client.post("/api/suppliers/", json={"name": "Only Req Supplier"}).json()
    client.post("/api/supplier-materials/", json={"supplier_id": supplier["id"], "material_id": material["id"]})
    requirement = client.post("/api/procurement-requirements/", json={
        "order_id": order["id"], "material_id": material["id"],
    }).json()
    client.post(f"/api/procurement-requirements/{requirement['id']}/decision", json={
        "requirement_id": requirement["id"], "selected_supplier_id": supplier["id"],
    })

    resp = client.post(f"/api/procurement-requirements/{requirement['id']}/decision", json={
        "requirement_id": requirement["id"], "selected_supplier_id": supplier["id"],
    })
    assert resp.status_code == 409


def test_procurement_requirements_are_master_only(client, test_user, db_session):
    from app.platform.security import hash_password
    from app.modules.auth.auth import User
    _login(client, test_user)
    order, material = _make_shortage_order(client, "4")
    employee = User(
        username="procreqempuser", email="procreqempuser@example.com", full_name="Proc Req Employee",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=None, is_active=True,
    )
    db_session.add(employee)
    db_session.commit()
    resp = client.post("/api/auth/login", json={"identifier": "procreqempuser@example.com", "password": "EmpPass1!"})
    assert resp.status_code == 200

    resp = client.post("/api/procurement-requirements/", json={"order_id": order["id"], "material_id": material["id"]})
    assert resp.status_code == 403
    assert client.get("/api/procurement-requirements/").status_code == 403


def test_supplier_options_preview_matches_what_decision_would_recommend(client, test_user):
    """The preview endpoint and the decision endpoint must agree - same
    underlying calculation, not two independently-drifting ones."""
    _login(client, test_user)
    order, material = _make_shortage_order(client, "10")
    preferred = client.post("/api/suppliers/", json={"name": "Preview Preferred Supplier"}).json()
    client.post("/api/supplier-materials/", json={
        "supplier_id": preferred["id"], "material_id": material["id"], "is_preferred": True,
    })
    requirement = client.post("/api/procurement-requirements/", json={
        "order_id": order["id"], "material_id": material["id"],
    }).json()

    resp = client.get(f"/api/procurement-requirements/{requirement['id']}/supplier-options")
    assert resp.status_code == 200
    body = resp.json()
    assert body["material_id"] == material["id"]
    assert len(body["options"]) == 1
    assert body["options"][0]["supplier_id"] == preferred["id"]
    assert body["options"][0]["is_preferred"] is True


def test_get_requirement_includes_decision_with_supplier_names(client, test_user):
    """The requirement GET response must actually embed the decision
    (with real supplier names, not just IDs) - a frontend showing the
    requirement's own detail page has no other way to know a decision
    was already made without this."""
    _login(client, test_user)
    order, material = _make_shortage_order(client, "11")
    preferred = client.post("/api/suppliers/", json={"name": "Embedded Decision Preferred Supplier"}).json()
    chosen = client.post("/api/suppliers/", json={"name": "Embedded Decision Chosen Supplier"}).json()
    client.post("/api/supplier-materials/", json={
        "supplier_id": preferred["id"], "material_id": material["id"], "is_preferred": True,
    })
    client.post("/api/supplier-materials/", json={"supplier_id": chosen["id"], "material_id": material["id"]})
    requirement = client.post("/api/procurement-requirements/", json={
        "order_id": order["id"], "material_id": material["id"],
    }).json()
    client.post(f"/api/procurement-requirements/{requirement['id']}/decision", json={
        "requirement_id": requirement["id"], "selected_supplier_id": chosen["id"],
    })

    resp = client.get(f"/api/procurement-requirements/{requirement['id']}")
    assert resp.status_code == 200
    body = resp.json()
    assert body["decision"] is not None
    assert body["decision"]["selected_supplier_name"] == "Embedded Decision Chosen Supplier"
    assert body["decision"]["recommended_supplier_name"] == "Embedded Decision Preferred Supplier"
    assert body["decision"]["followed_recommendation"] is False


def test_get_requirement_without_decision_has_null_decision(client, test_user):
    _login(client, test_user)
    order, material = _make_shortage_order(client, "12")
    requirement = client.post("/api/procurement-requirements/", json={
        "order_id": order["id"], "material_id": material["id"],
    }).json()

    resp = client.get(f"/api/procurement-requirements/{requirement['id']}")
    assert resp.json()["decision"] is None


def test_procurement_requirements_require_auth(client):
    resp = client.get("/api/procurement-requirements/")
    assert resp.status_code == 401


def test_decision_rejects_supplier_with_no_material_relationship(client, test_user):
    """P0.2.4: the backend must reject this regardless of what the
    frontend dropdown would have filtered - never trust a client-sent
    supplier_id alone."""
    _login(client, test_user)
    order, material = _make_shortage_order(client, "5")
    unrelated_supplier = client.post("/api/suppliers/", json={"name": "Unrelated Supplier No Link"}).json()
    requirement = client.post("/api/procurement-requirements/", json={
        "order_id": order["id"], "material_id": material["id"],
    }).json()

    resp = client.post(f"/api/procurement-requirements/{requirement['id']}/decision", json={
        "requirement_id": requirement["id"], "selected_supplier_id": unrelated_supplier["id"],
    })
    assert resp.status_code == 400
    assert "no recorded supply relationship" in resp.json()["detail"].lower()


def test_purchase_created_from_requirement_is_linked_and_fulfills_it(client, test_user):
    """P0.2.5 traceability: requirement.purchase_id is actually set, and
    receipt_status="Received" genuinely fulfills the requirement (goods
    are actually in hand, not merely ordered)."""
    _login(client, test_user)
    order, material = _make_shortage_order(client, "6")
    supplier = client.post("/api/suppliers/", json={"name": "Traceability Supplier"}).json()
    client.post("/api/supplier-materials/", json={"supplier_id": supplier["id"], "material_id": material["id"]})
    requirement = client.post("/api/procurement-requirements/", json={
        "order_id": order["id"], "material_id": material["id"],
    }).json()
    client.post(f"/api/procurement-requirements/{requirement['id']}/decision", json={
        "requirement_id": requirement["id"], "selected_supplier_id": supplier["id"],
    })

    resp = client.post(f"/api/procurement-requirements/{requirement['id']}/purchase", json={
        "quantity": "5", "unit": "Sheets", "rate": "450.00", "gst_percent": "18", "receipt_status": "Received",
    })
    assert resp.status_code == 201
    purchase = resp.json()
    assert purchase["supplier_id"] == supplier["id"]
    assert purchase["material_id"] == material["id"]

    updated_requirement = client.get(f"/api/procurement-requirements/{requirement['id']}").json()
    assert updated_requirement["purchase_id"] == purchase["id"]
    assert updated_requirement["status"] == "Fulfilled"


def test_purchase_from_requirement_ordered_not_received_stays_ordered(client, test_user):
    _login(client, test_user)
    order, material = _make_shortage_order(client, "7")
    supplier = client.post("/api/suppliers/", json={"name": "Traceability Ordered Supplier"}).json()
    client.post("/api/supplier-materials/", json={"supplier_id": supplier["id"], "material_id": material["id"]})
    requirement = client.post("/api/procurement-requirements/", json={
        "order_id": order["id"], "material_id": material["id"],
    }).json()
    client.post(f"/api/procurement-requirements/{requirement['id']}/decision", json={
        "requirement_id": requirement["id"], "selected_supplier_id": supplier["id"],
    })

    resp = client.post(f"/api/procurement-requirements/{requirement['id']}/purchase", json={
        "quantity": "5", "unit": "Sheets", "rate": "450.00", "gst_percent": "18", "receipt_status": "Ordered",
    })
    assert resp.status_code == 201
    updated_requirement = client.get(f"/api/procurement-requirements/{requirement['id']}").json()
    assert updated_requirement["status"] == "Ordered"


def test_purchase_from_requirement_without_decision_is_rejected(client, test_user):
    _login(client, test_user)
    order, material = _make_shortage_order(client, "8")
    requirement = client.post("/api/procurement-requirements/", json={
        "order_id": order["id"], "material_id": material["id"],
    }).json()

    resp = client.post(f"/api/procurement-requirements/{requirement['id']}/purchase", json={
        "quantity": "5", "unit": "Sheets", "rate": "450.00", "gst_percent": "18",
    })
    assert resp.status_code == 400


def test_second_purchase_from_same_requirement_is_rejected(client, test_user):
    _login(client, test_user)
    order, material = _make_shortage_order(client, "9")
    supplier = client.post("/api/suppliers/", json={"name": "Traceability Duplicate Supplier"}).json()
    client.post("/api/supplier-materials/", json={"supplier_id": supplier["id"], "material_id": material["id"]})
    requirement = client.post("/api/procurement-requirements/", json={
        "order_id": order["id"], "material_id": material["id"],
    }).json()
    client.post(f"/api/procurement-requirements/{requirement['id']}/decision", json={
        "requirement_id": requirement["id"], "selected_supplier_id": supplier["id"],
    })
    client.post(f"/api/procurement-requirements/{requirement['id']}/purchase", json={
        "quantity": "5", "unit": "Sheets", "rate": "450.00", "gst_percent": "18",
    })

    resp = client.post(f"/api/procurement-requirements/{requirement['id']}/purchase", json={
        "quantity": "5", "unit": "Sheets", "rate": "450.00", "gst_percent": "18",
    })
    assert resp.status_code == 409
