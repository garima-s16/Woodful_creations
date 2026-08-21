import io
from openpyxl import Workbook


def _login(client, test_user):
    resp = client.post("/api/auth/login", json={"identifier": "test@example.com", "password": "TestPass123!"})
    assert resp.status_code == 200


def _build_upload(rows):
    """Builds a minimal .xlsx with the exact fixed header row the
    parser looks for, plus whatever data rows the test needs."""
    wb = Workbook()
    ws = wb.active
    headers = ["Material", "Specification", "Quantity", "Unit", "Supplier", "Rate", "GST %", "Invoice Number", "Invoice Date", "Remarks"]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_template_downloads_and_has_correct_headers(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/purchase-imports/template")
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    header_row = [c.value for c in ws[5]]  # row 5: logo row, title, subtitle, blank, headers
    assert "Material" in header_row and "Supplier" in header_row and "GST %" in header_row


def test_preview_matches_existing_material_and_supplier(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Import Test Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Import Test Material", "unit": "Sheets", "opening_stock": 5, "minimum_stock": 1,
    }).json()

    upload = _build_upload([{
        "Material": "Import Test Material", "Specification": "x", "Quantity": 5, "Unit": "Sheets",
        "Supplier": "Import Test Supplier", "Rate": 500, "GST %": 18,
        "Invoice Number": "INV-1", "Invoice Date": "2026-08-01", "Remarks": None,
    }])
    resp = client.post("/api/purchase-imports/preview", files={"file": ("test.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert resp.status_code == 200
    body = resp.json()
    assert body["matched_rows"] == 1
    assert body["error_rows"] == 0
    assert body["rows"][0]["matched_material_id"] == material["id"]
    assert body["rows"][0]["matched_supplier_id"] == supplier["id"]


def test_preview_flags_unmatched_material_as_new(client, test_user):
    _login(client, test_user)
    client.post("/api/suppliers/", json={"name": "Unmatched Material Test Supplier"})
    upload = _build_upload([{
        "Material": "Totally New Material XYZ", "Specification": None, "Quantity": 3, "Unit": "Sheets",
        "Supplier": "Unmatched Material Test Supplier", "Rate": 200, "GST %": 18,
        "Invoice Number": None, "Invoice Date": None, "Remarks": None,
    }])
    resp = client.post("/api/purchase-imports/preview", files={"file": ("test.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    body = resp.json()
    assert body["new_material_rows"] == 1
    assert body["rows"][0]["is_new_material"] is True
    assert body["rows"][0]["errors"] == []


def test_preview_flags_unmatched_supplier_as_error(client, test_user):
    """Unlike materials, suppliers are never auto-created on import -
    an unmatched supplier is a real error, not a "new" row."""
    _login(client, test_user)
    client.post("/api/materials/", json={"name": "Supplier Error Test Material", "unit": "Sheets", "opening_stock": 1, "minimum_stock": 1})
    upload = _build_upload([{
        "Material": "Supplier Error Test Material", "Specification": None, "Quantity": 3, "Unit": "Sheets",
        "Supplier": "Nonexistent Supplier XYZ", "Rate": 200, "GST %": 18,
        "Invoice Number": None, "Invoice Date": None, "Remarks": None,
    }])
    resp = client.post("/api/purchase-imports/preview", files={"file": ("test.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    body = resp.json()
    assert body["error_rows"] == 1
    assert any("supplier" in e.lower() for e in body["rows"][0]["errors"])


def test_commit_creates_real_purchase_for_matched_row(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Commit Test Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Commit Test Material", "unit": "Sheets", "opening_stock": 5, "minimum_stock": 1,
    }).json()

    resp = client.post("/api/purchase-imports/commit", json={"rows": [{
        "material_name": "Commit Test Material", "matched_material_id": material["id"],
        "matched_supplier_id": supplier["id"], "quantity": "10", "unit": "Sheets",
        "rate": "500.00", "gst_percent": "18", "create_new_material": False,
    }]})
    assert resp.status_code == 200
    body = resp.json()
    assert body["created_purchases"] == 1
    assert body["created_materials"] == 0
    assert body["error"] is None

    updated_material = client.get(f"/api/materials/{material['id']}").json()
    assert updated_material["current_stock"] == 15  # 5 opening + 10 imported


def test_commit_creates_new_material_when_requested(client, test_user):
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "New Material Commit Supplier"}).json()

    resp = client.post("/api/purchase-imports/commit", json={"rows": [{
        "material_name": "Brand New Imported Material", "specification": "test spec",
        "matched_material_id": None, "matched_supplier_id": supplier["id"],
        "quantity": "5", "unit": "Sheets", "rate": "300.00", "gst_percent": "18",
        "create_new_material": True,
    }]})
    assert resp.status_code == 200
    body = resp.json()
    assert body["created_materials"] == 1
    assert body["created_purchases"] == 1

    materials = client.get("/api/materials/", params={"search": "Brand New Imported Material"}).json()
    assert len(materials) == 1
    assert materials[0]["current_stock"] == 5


def test_commit_stops_and_reports_partial_success_on_bad_row(client, test_user):
    """The honest, actually-implemented behavior: not atomic - rows
    before the failure are genuinely committed, and the response says so."""
    _login(client, test_user)
    supplier = client.post("/api/suppliers/", json={"name": "Partial Commit Supplier"}).json()
    material = client.post("/api/materials/", json={
        "name": "Partial Commit Material", "unit": "Sheets", "opening_stock": 5, "minimum_stock": 1,
    }).json()

    resp = client.post("/api/purchase-imports/commit", json={"rows": [
        {
            "material_name": "Partial Commit Material", "matched_material_id": material["id"],
            "matched_supplier_id": supplier["id"], "quantity": "2", "unit": "Sheets",
            "rate": "500.00", "gst_percent": "18", "create_new_material": False,
        },
        {
            "material_name": "No Material Row", "matched_material_id": None,
            "matched_supplier_id": supplier["id"], "quantity": "2", "unit": "Sheets",
            "rate": "500.00", "gst_percent": "18", "create_new_material": False,  # no material - must fail
        },
    ]})
    assert resp.status_code == 200
    body = resp.json()
    assert body["created_purchases"] == 1  # first row succeeded
    assert body["error"] is not None
    assert "row 2" in body["error"].lower()


def test_import_requires_master_or_manager_role(client, test_user, db_session):
    from app.core.security import hash_password
    from app.models.user import User

    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "Import Perm Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    limited_user = User(
        username="importpermuser", email="importpermuser@example.com", full_name="Limited Import User",
        password_hash=hash_password("LimitedPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(limited_user)
    db_session.commit()

    resp = client.post("/api/auth/login", json={"identifier": "importpermuser@example.com", "password": "LimitedPass1!"})
    assert resp.status_code == 200

    upload = _build_upload([])
    preview_resp = client.post("/api/purchase-imports/preview", files={"file": ("test.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert preview_resp.status_code == 403
