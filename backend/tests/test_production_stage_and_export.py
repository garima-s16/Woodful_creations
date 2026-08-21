"""Tests for production_jobs' stage/completion_date fields and the
production.xlsx export (Part 2/10 brief - production pipeline
stages, and Excel exports respecting filters)."""
import io
from openpyxl import load_workbook


def _login(client, test_user):
    resp = client.post("/api/auth/login", json={"identifier": "test@example.com", "password": "TestPass123!"})
    assert resp.status_code == 200


def test_production_stage_can_be_set_on_create(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Panel cutting", "stage": "Cutting",
    }).json()
    assert job["stage"] == "Cutting"


def test_completion_date_set_on_transition_to_completed(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Assembly"}).json()
    assert job["completion_date"] is None

    resp = client.put(f"/api/production-jobs/{job['id']}", json={"status": "Completed"})
    assert resp.status_code == 200
    assert resp.json()["completion_date"] is not None


def test_completion_date_does_not_change_on_resave(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Finishing"}).json()
    client.put(f"/api/production-jobs/{job['id']}", json={"status": "Completed"})
    first = client.get("/api/production-jobs/").json()
    first_match = next(j for j in first if j["id"] == job["id"])

    client.put(f"/api/production-jobs/{job['id']}", json={"status": "Completed", "remarks": "verified"})
    second = client.get("/api/production-jobs/").json()
    second_match = next(j for j in second if j["id"] == job["id"])
    assert second_match["completion_date"] == first_match["completion_date"]


def test_production_export_respects_stage_filter(client, test_user):
    _login(client, test_user)
    client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Cutting op", "stage": "Cutting",
    })
    client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "QC op", "stage": "QC",
    })

    resp = client.get("/api/reports/production.xlsx", params={"stage": "Cutting"})
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["Production"]
    operations = [cell.value for row in ws.iter_rows() for cell in row if isinstance(cell.value, str)]
    assert any("Cutting op" in o for o in operations)
    assert not any("QC op" in o for o in operations)


def test_production_export_respects_operator_filter(client, test_user):
    _login(client, test_user)
    emp_a = client.post("/api/employees/", json={"name": "Production Export Operator A"}).json()
    emp_b = client.post("/api/employees/", json={"name": "Production Export Operator B"}).json()
    client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Job for A", "employee_id": emp_a["id"],
    })
    client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Job for B", "employee_id": emp_b["id"],
    })

    resp = client.get("/api/reports/production.xlsx", params={"employee_id": emp_a["id"]})
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["Production"]
    operations = [cell.value for row in ws.iter_rows() for cell in row if isinstance(cell.value, str)]
    assert any("Job for A" in o for o in operations)
    assert not any("Job for B" in o for o in operations)
