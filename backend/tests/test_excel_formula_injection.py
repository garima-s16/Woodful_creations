"""Tests for the Excel formula-injection fix - user-controlled text
(client names, remarks) starting with =, +, -, @ must not become an
executable formula when the downloaded workbook is opened."""
import io
from openpyxl import load_workbook


def _login(client, test_user):
    resp = client.post("/api/auth/login", json={"identifier": "test@example.com", "password": "TestPass123!"})
    assert resp.status_code == 200


def test_malicious_client_name_is_escaped_in_export(client, test_user):
    _login(client, test_user)
    client.post("/api/clients/", json={"name": '=cmd|"/c calc"!A1'})

    resp = client.get("/api/reports/clients.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    found = False
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "cmd|" in cell.value:
                found = True
                # Verified empirically: openpyxl preserves the leading
                # apostrophe on read-back (it's an Excel-UI-layer
                # convention, not something openpyxl simulates). The
                # actual security guarantee is data_type == 's' -
                # this cell is stored as a string, never as type 'f'
                # (formula), so Excel can never execute it.
                assert cell.data_type == "s"
    assert found, "expected the sanitized client name to appear in the export"


def test_legitimate_currency_value_unaffected(client, test_user):
    """The sanitizer must not corrupt genuinely safe values."""
    from app.utils.exporters import _sanitize_cell_value
    assert _sanitize_cell_value("Rs 1,25,000.00") == "Rs 1,25,000.00"
    assert _sanitize_cell_value(42) == 42
    assert _sanitize_cell_value(None) is None
    assert _sanitize_cell_value("Normal Client Name") == "Normal Client Name"


def test_all_four_dangerous_prefixes_are_escaped():
    from app.utils.exporters import _sanitize_cell_value
    for dangerous in ["=SUM(A1)", "+1+1", "-2+3", "@SUM(A1)"]:
        result = _sanitize_cell_value(dangerous)
        assert result.startswith("'")
        assert result == "'" + dangerous
