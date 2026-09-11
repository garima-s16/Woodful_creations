"""HR domain tests: attendance/RBAC, leaves, payroll, and Excel
exports. Combines all former test_*.py files under tests/modules/hr/."""
from app.platform.security import hash_password
from app.modules.auth.auth import User
from tests.helpers import _login
import re
import io
from openpyxl import load_workbook


# --- test_attendance_hr_rbac.py ---
def test_employee_cannot_view_another_employees_attendance(client, test_user, db_session):
    _login(client, test_user)
    other_employee = client.post("/api/employees/", json={
        "name": "Attendance RBAC Other Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    client.post("/api/attendance/", json={
        "date": "2026-08-16T00:00:00", "employee_id": other_employee["id"], "attendance_status": "Present",
    })

    employee = client.post("/api/employees/", json={
        "name": "Attendance RBAC Self Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    user = User(
        username="attendancerbacuser", email="attendancerbacuser@example.com", full_name="Attendance RBAC User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "attendancerbacuser@example.com", "password": "EmpPass1!"})

    resp = client.get("/api/attendance/", params={"employee_id": other_employee["id"]})
    assert resp.status_code == 403


def test_employee_can_view_own_attendance(client, test_user, db_session):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "Attendance RBAC Own Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    client.post("/api/attendance/", json={
        "date": "2026-08-16T00:00:00", "employee_id": employee["id"], "attendance_status": "Present",
    })
    user = User(
        username="attendanceownrbacuser", email="attendanceownrbacuser@example.com", full_name="Attendance Own RBAC User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "attendanceownrbacuser@example.com", "password": "EmpPass1!"})

    resp = client.get("/api/attendance/")
    assert resp.status_code == 200
    assert all(a["employee_id"] == employee["id"] for a in resp.json())


def test_employee_cannot_mark_attendance_for_someone_else(client, test_user, db_session):
    _login(client, test_user)
    other_employee = client.post("/api/employees/", json={
        "name": "Attendance RBAC Impersonation Target", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    employee = client.post("/api/employees/", json={
        "name": "Attendance RBAC Impersonator", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    user = User(
        username="attendanceimpersonateuser", email="attendanceimpersonateuser@example.com", full_name="Impersonator",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "attendanceimpersonateuser@example.com", "password": "EmpPass1!"})

    resp = client.post("/api/attendance/", json={
        "date": "2026-08-16T00:00:00", "employee_id": other_employee["id"], "attendance_status": "Present",
    })
    assert resp.status_code == 403


def test_employee_can_mark_own_attendance(client, test_user, db_session):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "Attendance RBAC Self Mark Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    user = User(
        username="attendanceselfmarkuser", email="attendanceselfmarkuser@example.com", full_name="Self Mark",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "attendanceselfmarkuser@example.com", "password": "EmpPass1!"})

    resp = client.post("/api/attendance/", json={
        "date": "2026-08-16T00:00:00", "employee_id": employee["id"], "attendance_status": "Present",
    })
    assert resp.status_code == 201


def test_employee_cannot_correct_attendance_record(client, test_user, db_session):
    """Correction is a supervisory action, master-only."""
    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "Attendance RBAC Correction Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    record = client.post("/api/attendance/", json={
        "date": "2026-08-16T00:00:00", "employee_id": employee["id"], "attendance_status": "Present",
    }).json()
    user = User(
        username="attendancecorrectuser", email="attendancecorrectuser@example.com", full_name="Correction User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "attendancecorrectuser@example.com", "password": "EmpPass1!"})

    resp = client.put(f"/api/attendance/{record['id']}", json={"attendance_status": "Half Day"})
    assert resp.status_code == 403


def test_master_full_attendance_access_unaffected(client, test_user):
    """Backward compatibility - master retains full functionality."""
    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "Attendance RBAC Master Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    record = client.post("/api/attendance/", json={
        "date": "2026-08-16T00:00:00", "employee_id": employee["id"], "attendance_status": "Present",
    }).json()
    assert client.get("/api/attendance/").status_code == 200
    assert client.put(f"/api/attendance/{record['id']}", json={"attendance_status": "Half Day"}).status_code == 200


def test_salary_days_matches_working_days_when_no_leave(client, test_user):
    """Family P0.43: Salary Days = actual configured working days in
    the month minus approved leave - with no leave, they're equal.
    February 2026 has 24 Monday-Saturday working days (28 calendar
    days, 4 Sundays) - deliberately not August's 26, to prove this is
    genuinely computed per-month, not a fixed assumption."""
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Salary Days No Leave Employee", "monthly_salary": "20800"}).json()

    resp = client.get("/api/salary-slips/attendance-summary", params={
        "employee_id": employee["id"], "month": "February", "year": "2026",
    })
    assert resp.status_code == 200
    data = resp.json()
    assert data["calendar_days"] == 28
    assert data["suggested_working_days"] == 24
    assert data["leave_days"] == 0
    assert data["salary_days"] == 24


def test_salary_days_subtracts_approved_leave_on_working_dates(client, test_user):
    """2026-02-05 and 2026-02-06 are a Thursday and Friday (real
    working dates) - both must be subtracted."""
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Salary Days Leave Employee", "monthly_salary": "20800"}).json()
    leave = client.post("/api/leaves/", json={
        "employee_id": employee["id"], "leave_type": "CL",
        "start_date": "2026-02-05T00:00:00", "end_date": "2026-02-06T00:00:00",
    }).json()
    client.put(f"/api/leaves/{leave['id']}", json={"status": "Approved"})

    resp = client.get("/api/salary-slips/attendance-summary", params={
        "employee_id": employee["id"], "month": "February", "year": "2026",
    })
    data = resp.json()
    assert data["leave_days"] == 2
    assert data["salary_days"] == 22  # 24 - 2


def test_salary_days_ignores_pending_leave(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Salary Days Pending Leave Employee", "monthly_salary": "20800"}).json()
    client.post("/api/leaves/", json={
        "employee_id": employee["id"], "leave_type": "CL",
        "start_date": "2026-02-05T00:00:00", "end_date": "2026-02-06T00:00:00",
    })
    # Left as "Pending" - never approved.

    resp = client.get("/api/salary-slips/attendance-summary", params={
        "employee_id": employee["id"], "month": "February", "year": "2026",
    })
    data = resp.json()
    assert data["leave_days"] == 0
    assert data["salary_days"] == 24


def test_salary_days_does_not_double_subtract_leave_on_sunday(client, test_user):
    """2026-02-01 is a Sunday - already not a working day, so an
    approved leave covering it must not reduce salary_days further."""
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Salary Days Sunday Leave Employee", "monthly_salary": "20800"}).json()
    leave = client.post("/api/leaves/", json={
        "employee_id": employee["id"], "leave_type": "CL",
        "start_date": "2026-02-01T00:00:00", "end_date": "2026-02-01T00:00:00",
    }).json()
    client.put(f"/api/leaves/{leave['id']}", json={"status": "Approved"})

    resp = client.get("/api/salary-slips/attendance-summary", params={
        "employee_id": employee["id"], "month": "February", "year": "2026",
    })
    data = resp.json()
    assert data["leave_days"] == 0  # Sunday was never a working day to begin with
    assert data["salary_days"] == 24


def test_add_overtime_creates_record_when_none_exists(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Add Overtime New Record Employee", "monthly_salary": "20800"}).json()

    resp = client.post("/api/attendance/overtime", json={
        "employee_id": employee["id"], "dates": ["2026-08-10T00:00:00"], "hours": "2", "mode": "add",
    })
    assert resp.status_code == 200
    body = resp.json()
    assert len(body) == 1
    assert float(body[0]["overtime_hours"]) == 2.0


def test_add_overtime_is_additive_not_replace(client, test_user):
    """The spec's own core rule: existing 2 + add 1 = 3, never a
    silent replace to 1."""
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Add Overtime Additive Employee", "monthly_salary": "20800"}).json()
    client.post("/api/attendance/overtime", json={
        "employee_id": employee["id"], "dates": ["2026-08-10T00:00:00"], "hours": "2", "mode": "add",
    })

    resp = client.post("/api/attendance/overtime", json={
        "employee_id": employee["id"], "dates": ["2026-08-10T00:00:00"], "hours": "1", "mode": "add",
    })
    assert float(resp.json()[0]["overtime_hours"]) == 3.0


def test_set_overtime_replaces_existing_value(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Set Overtime Replace Employee", "monthly_salary": "20800"}).json()
    client.post("/api/attendance/overtime", json={
        "employee_id": employee["id"], "dates": ["2026-08-10T00:00:00"], "hours": "3", "mode": "add",
    })

    resp = client.post("/api/attendance/overtime", json={
        "employee_id": employee["id"], "dates": ["2026-08-10T00:00:00"], "hours": "1", "mode": "set",
    })
    assert float(resp.json()[0]["overtime_hours"]) == 1.0


def test_add_overtime_multiple_dates_matches_spec_worked_example(client, test_user):
    """Spec's own example: A=1, B=2, C=0, D=3. Select A, B, D. Add 2.
    Result: A=3, B=4, C=0 (untouched), D=5."""
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Add Overtime Multi Date Employee", "monthly_salary": "20800"}).json()
    client.post("/api/attendance/overtime", json={
        "employee_id": employee["id"], "dates": ["2026-08-10T00:00:00"], "hours": "1", "mode": "set",  # A
    })
    client.post("/api/attendance/overtime", json={
        "employee_id": employee["id"], "dates": ["2026-08-11T00:00:00"], "hours": "2", "mode": "set",  # B
    })
    client.post("/api/attendance/overtime", json={
        "employee_id": employee["id"], "dates": ["2026-08-13T00:00:00"], "hours": "3", "mode": "set",  # D
    })
    # C (2026-08-12) is deliberately never touched here.

    client.post("/api/attendance/overtime", json={
        "employee_id": employee["id"],
        "dates": ["2026-08-10T00:00:00", "2026-08-11T00:00:00", "2026-08-13T00:00:00"],
        "hours": "2", "mode": "add",
    })

    resp = client.get("/api/attendance/", params={"employee_id": employee["id"]})
    by_date = {r["date"][:10]: float(r["overtime_hours"]) for r in resp.json()}
    assert by_date["2026-08-10"] == 3.0  # A: 1 + 2
    assert by_date["2026-08-11"] == 4.0  # B: 2 + 2
    assert "2026-08-12" not in by_date  # C: never created, untouched
    assert by_date["2026-08-13"] == 5.0  # D: 3 + 2


def test_add_overtime_requires_master(client, test_user, db_session):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Add Overtime RBAC Employee", "monthly_salary": "20800"}).json()
    from app.platform.security import hash_password
    from app.modules.auth.auth import User
    user = User(
        username="addovertimerbacuser", email="addovertimerbacuser@example.com", full_name="Add Overtime RBAC User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "addovertimerbacuser@example.com", "password": "EmpPass1!"})

    resp = client.post("/api/attendance/overtime", json={
        "employee_id": employee["id"], "dates": ["2026-08-10T00:00:00"], "hours": "2", "mode": "add",
    })
    assert resp.status_code == 403


def test_add_overtime_rejects_negative_hours(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Add Overtime Negative Employee", "monthly_salary": "20800"}).json()

    resp = client.post("/api/attendance/overtime", json={
        "employee_id": employee["id"], "dates": ["2026-08-10T00:00:00"], "hours": "-1", "mode": "add",
    })
    assert resp.status_code == 422


def test_add_overtime_rejects_empty_dates_list(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Add Overtime Empty Dates Employee", "monthly_salary": "20800"}).json()

    resp = client.post("/api/attendance/overtime", json={
        "employee_id": employee["id"], "dates": [], "hours": "2", "mode": "add",
    })
    assert resp.status_code == 422


def test_add_overtime_rejects_unknown_employee(client, test_user):
    _login(client, test_user)
    resp = client.post("/api/attendance/overtime", json={
        "employee_id": 999999, "dates": ["2026-08-10T00:00:00"], "hours": "2", "mode": "add",
    })
    assert resp.status_code == 404


def test_add_overtime_on_sunday_still_requires_explicit_value(client, test_user):
    """2026-08-02 is a Sunday - overtime still only comes from the
    explicit hours given, never implied by the day itself."""
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Add Overtime Sunday Employee", "monthly_salary": "20800"}).json()

    resp = client.post("/api/attendance/overtime", json={
        "employee_id": employee["id"], "dates": ["2026-08-02T00:00:00"], "hours": "4", "mode": "add",
    })
    assert float(resp.json()[0]["overtime_hours"]) == 4.0


def test_payroll_summary_reflects_real_slip_and_advance_data(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Payroll Summary Real Data Employee", "monthly_salary": "20800"}).json()
    client.post("/api/salary-slips/", json={
        "employee_id": employee["id"], "month": "September", "year": "2026", "basic": "20000",
    })

    resp = client.get("/api/salary-slips/payroll-summary", params={"month": "September", "year": "2026"})
    assert resp.status_code == 200
    body = resp.json()
    assert body["month"] == "September"
    assert body["employees_with_slip"] == 1
    assert body["status_counts"]["draft"] == 1
    assert "salary_advances" in body


def test_payroll_summary_requires_master(client, test_user, db_session):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Payroll Summary RBAC Employee", "monthly_salary": "20800"}).json()
    from app.platform.security import hash_password
    from app.modules.auth.auth import User
    user = User(
        username="payrollsummaryrbacuser", email="payrollsummaryrbacuser@example.com", full_name="Payroll Summary RBAC User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "payrollsummaryrbacuser@example.com", "password": "EmpPass1!"})

    resp = client.get("/api/salary-slips/payroll-summary", params={"month": "September", "year": "2026"})
    assert resp.status_code == 403


def test_attendance_summary_computes_paid_days_and_overtime(client, test_user):
    """Family P0.43: overtime is a fact the Master explicitly records,
    not derived from in_time/out_time - the first record below has a
    10-hour clock span but overtime_hours is only counted because it
    is explicitly set to 2, not because it was computed from the
    clock times."""
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Payroll Summary Test Employee", "monthly_salary": "20800"}).json()

    client.post("/api/attendance/", json={
        "date": "2026-08-03T00:00:00", "employee_id": employee["id"],
        "in_time": "2026-08-03T09:00:00", "out_time": "2026-08-03T19:00:00",
        "standard_hours": "8", "attendance_status": "Present", "overtime_hours": "2",
    })
    client.post("/api/attendance/", json={
        "date": "2026-08-04T00:00:00", "employee_id": employee["id"],
        "in_time": "2026-08-04T09:00:00", "out_time": "2026-08-04T13:00:00",
        "standard_hours": "8", "attendance_status": "Half Day",
    })

    resp = client.get("/api/salary-slips/attendance-summary", params={
        "employee_id": employee["id"], "month": "August", "year": "2026",
    })
    assert resp.status_code == 200
    data = resp.json()
    assert data["records_found"] == 2
    assert data["suggested_paid_days"] == 1.5  # 1 (Present) + 0.5 (Half Day)
    assert data["total_overtime_hours"] == 2.0  # explicitly recorded on the first day, not derived from clock times
    assert data["suggested_overtime_amount"] > 0


def test_overtime_is_zero_by_default_even_with_long_clock_span(client, test_user):
    """The core P0.43 correction: a long in/out span alone must never
    imply overtime - only an explicit value does."""
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "No Auto Overtime Employee", "monthly_salary": "20800"}).json()

    resp = client.post("/api/attendance/", json={
        "date": "2026-08-05T00:00:00", "employee_id": employee["id"],
        "in_time": "2026-08-05T08:00:00", "out_time": "2026-08-05T22:00:00",
        "standard_hours": "8", "attendance_status": "Present",
    })
    assert resp.status_code == 201
    assert float(resp.json()["overtime_hours"]) == 0.0


def test_sunday_work_does_not_imply_overtime_without_explicit_value(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Sunday No Auto Overtime Employee", "monthly_salary": "20800"}).json()

    # 2026-08-02 is a Sunday.
    resp = client.post("/api/attendance/", json={
        "date": "2026-08-02T00:00:00", "employee_id": employee["id"],
        "in_time": "2026-08-02T09:00:00", "out_time": "2026-08-02T17:00:00",
        "standard_hours": "8", "attendance_status": "Present",
    })
    assert resp.status_code == 201
    assert float(resp.json()["overtime_hours"]) == 0.0


def test_employee_cannot_set_overtime_hours_on_own_attendance(client, test_user, db_session):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Overtime RBAC Employee", "monthly_salary": "20800"}).json()
    from app.platform.security import hash_password
    from app.modules.auth.auth import User
    user = User(
        username="overtimerbacuser", email="overtimerbacuser@example.com", full_name="Overtime RBAC User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "overtimerbacuser@example.com", "password": "EmpPass1!"})

    resp = client.post("/api/attendance/", json={
        "date": "2026-08-06T00:00:00", "employee_id": employee["id"],
        "attendance_status": "Present", "overtime_hours": "3",
    })
    assert resp.status_code == 403


def test_employee_can_mark_own_attendance_with_zero_overtime(client, test_user, db_session):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Overtime RBAC Zero Employee", "monthly_salary": "20800"}).json()
    from app.platform.security import hash_password
    from app.modules.auth.auth import User
    user = User(
        username="overtimerbaczerouser", email="overtimerbaczerouser@example.com", full_name="Overtime RBAC Zero User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "overtimerbaczerouser@example.com", "password": "EmpPass1!"})

    resp = client.post("/api/attendance/", json={
        "date": "2026-08-06T00:00:00", "employee_id": employee["id"], "attendance_status": "Present",
    })
    assert resp.status_code == 201


def test_master_can_set_overtime_hours(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Master Overtime Employee", "monthly_salary": "20800"}).json()

    resp = client.post("/api/attendance/", json={
        "date": "2026-08-06T00:00:00", "employee_id": employee["id"],
        "attendance_status": "Present", "overtime_hours": "3",
    })
    assert resp.status_code == 201
    assert float(resp.json()["overtime_hours"]) == 3.0


def test_negative_overtime_hours_rejected(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Negative Overtime Employee", "monthly_salary": "20800"}).json()

    resp = client.post("/api/attendance/", json={
        "date": "2026-08-06T00:00:00", "employee_id": employee["id"],
        "attendance_status": "Present", "overtime_hours": "-1",
    })
    assert resp.status_code == 422


def test_attendance_summary_requires_master(client, test_user, db_session):
    from app.platform.security import hash_password
    from app.modules.auth.auth import User
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Payroll Summary Permission Employee"}).json()
    user = User(
        username="payrollsummaryuser", email="payrollsummaryuser@example.com",
        full_name="Payroll Summary User", password_hash=hash_password("EmpPass1!"),
        role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "payrollsummaryuser@example.com", "password": "EmpPass1!"})

    resp = client.get("/api/salary-slips/attendance-summary", params={
        "employee_id": employee["id"], "month": "August", "year": "2026",
    })
    assert resp.status_code == 403


def test_attendance_summary_rejects_invalid_month(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Payroll Summary Invalid Month Employee"}).json()
    resp = client.get("/api/salary-slips/attendance-summary", params={
        "employee_id": employee["id"], "month": "Augsut", "year": "2026",
    })
    assert resp.status_code == 400


def test_attendance_summary_returns_zero_for_no_records(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Payroll Summary No Records Employee"}).json()
    resp = client.get("/api/salary-slips/attendance-summary", params={
        "employee_id": employee["id"], "month": "January", "year": "2026",
    })
    assert resp.status_code == 200
    assert resp.json()["records_found"] == 0
    assert resp.json()["suggested_paid_days"] == 0


def test_default_weekday_config_is_mon_to_sat_working(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/working-calendar/weekdays")
    assert resp.status_code == 200
    weekdays = {w["weekday"]: w["is_working"] for w in resp.json()}
    assert weekdays["Monday"] is True
    assert weekdays["Saturday"] is True
    assert weekdays["Sunday"] is False


def test_working_days_differ_across_months(client, test_user):
    """The core requirement - February and August must not produce the
    same working-day count just because both used to default to 26."""
    _login(client, test_user)
    feb = client.get("/api/working-calendar/working-days", params={"year": 2026, "month": 2}).json()
    aug = client.get("/api/working-calendar/working-days", params={"year": 2026, "month": 8}).json()
    assert feb["working_days"] == 24  # 2026 is not a leap year: 28 days, 4 Sundays off
    assert aug["working_days"] == 26  # 31 days, 5 Sundays off
    assert feb["working_days"] != aug["working_days"]


def test_declared_holiday_reduces_working_days(client, test_user):
    _login(client, test_user)
    before = client.get("/api/working-calendar/working-days", params={"year": 2026, "month": 10}).json()["working_days"]

    client.post("/api/working-calendar/holidays", json={
        "date": "2026-10-02", "name": "Gandhi Jayanti", "is_working": False,
    })
    after = client.get("/api/working-calendar/working-days", params={"year": 2026, "month": 10}).json()["working_days"]
    assert after == before - 1


def test_special_working_day_increases_working_days(client, test_user):
    _login(client, test_user)
    # Find a Sunday in November 2026 to declare as a special working day
    before = client.get("/api/working-calendar/working-days", params={"year": 2026, "month": 11}).json()["working_days"]
    client.post("/api/working-calendar/holidays", json={
        "date": "2026-11-01", "name": "Special working Sunday", "is_working": True,
    })
    after = client.get("/api/working-calendar/working-days", params={"year": 2026, "month": 11}).json()["working_days"]
    assert after == before + 1


def test_weekday_config_requires_master_to_edit(client, test_user, db_session):
    from app.platform.security import hash_password
    from app.modules.auth.auth import User
    _login(client, test_user)
    weekday_id = client.get("/api/working-calendar/weekdays").json()[0]["id"]

    employee = client.post("/api/employees/", json={"name": "Calendar Permission Test Employee"}).json()
    user = User(
        username="calendarpermuser", email="calendarpermuser@example.com",
        full_name="Calendar Perm User", password_hash=hash_password("EmpPass1!"),
        role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "calendarpermuser@example.com", "password": "EmpPass1!"})

    resp = client.put(f"/api/working-calendar/weekdays/{weekday_id}", json={"is_working": False})
    assert resp.status_code == 403


def test_attendance_summary_uses_actual_month_working_days(client, test_user):
    """Confirms the payroll suggestion endpoint genuinely uses the
    calendar-computed value, not a leftover fixed constant."""
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Payroll Calendar Integration Employee", "monthly_salary": "26000"}).json()

    resp = client.get("/api/salary-slips/attendance-summary", params={
        "employee_id": employee["id"], "month": "February", "year": "2026",
    })
    assert resp.json()["suggested_working_days"] == 24


def _create_employee_rbac(client, name="Directory Test Employee", **overrides):
    payload = {"name": name, "department": "Production", "monthly_salary": "30000.00"}
    payload.update(overrides)
    resp = client.post("/api/employees/", json=payload)
    assert resp.status_code == 201, resp.text
    return resp.json()


def test_employee_directory_fields_persist(client, test_user):
    _login(client, test_user)
    emp = _create_employee_rbac(
        client, name="Pankaj Sharma", designation="Site Supervisor",
        email="pankaj@example.com", manager="Garima",
    )
    assert emp["designation"] == "Site Supervisor"
    assert emp["email"] == "pankaj@example.com"
    assert emp["manager"] == "Garima"


def test_employee_email_is_validated(client, test_user):
    _login(client, test_user)
    resp = client.post("/api/employees/", json={
        "name": "Bad Email Employee", "email": "not-an-email", "monthly_salary": "20000.00",
    })
    assert resp.status_code == 422


def test_employee_search_by_name(client, test_user):
    _login(client, test_user)
    _create_employee_rbac(client, name="Pankaj Kumar")
    _create_employee_rbac(client, name="Nikhil Verma")

    resp = client.get("/api/employees/?search=Pankaj")
    assert resp.status_code == 200
    names = [e["name"] for e in resp.json()]
    assert "Pankaj Kumar" in names
    assert "Nikhil Verma" not in names


def test_employee_directory_export_returns_xlsx(client, test_user):
    _login(client, test_user)
    _create_employee_rbac(client, name="Export Test Employee")

    resp = client.get("/api/reports/employees.xlsx")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert len(resp.content) > 0


def test_unlinked_non_master_user_sees_no_attendance(client, test_user, db_session):
    _login(client, test_user)
    other_employee = client.post("/api/employees/", json={"name": "Unlinked Attendance Other Employee"}).json()
    client.post("/api/attendance/", json={
        "date": "2026-08-16T00:00:00", "employee_id": other_employee["id"], "attendance_status": "Present",
    })
    user = User(
        username="unlinkedattenduser", email="unlinkedattenduser@example.com", full_name="Unlinked Attendance User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=None, is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "unlinkedattenduser@example.com", "password": "EmpPass1!"})

    resp = client.get("/api/attendance/")
    assert resp.status_code == 200
    assert resp.json() == []


# --- test_leaves.py ---
"""Company holiday import - commit-to-DB persistence chain. Previously
untested: holiday-imports/commit had only ever been proven to return
HTTP 200, never proven to actually write real CompanyHoliday rows to
the database, or that a same-date row is genuinely skipped (not
duplicated or silently overwritten) unless overwrite_existing is
explicitly set. Mirrors test_inventory_import.py's pattern - re-fetch
via a separate client.get() call, which the app's per-request session
override makes a genuine fresh-session read-back, not just trusting
the commit response body."""


def test_holiday_import_commit_creates_real_persisted_holiday(client, test_user):
    _login(client, test_user)
    resp = client.post("/api/holiday-imports/commit", json={"rows": [{
        "date": "2026-10-02", "name": "Holiday Import Test Day", "is_working": False,
    }]})
    assert resp.status_code == 200
    body = resp.json()
    assert body["created"] == 1
    assert body["error"] is None

    holidays = client.get("/api/working-calendar/holidays").json()
    match = next((h for h in holidays if h["name"] == "Holiday Import Test Day"), None)
    assert match is not None
    assert match["is_working"] is False


def test_holiday_import_commit_skips_existing_date_without_overwrite(client, test_user):
    _login(client, test_user)
    client.post("/api/holiday-imports/commit", json={"rows": [{
        "date": "2026-10-03", "name": "Holiday Import Original Name", "is_working": False,
    }]})

    resp = client.post("/api/holiday-imports/commit", json={"rows": [{
        "date": "2026-10-03", "name": "Holiday Import Should Not Apply", "is_working": False,
    }]})
    assert resp.status_code == 200
    body = resp.json()
    assert body["skipped"] == 1
    assert body["created"] == 0

    holidays = client.get("/api/working-calendar/holidays").json()
    match = next(h for h in holidays if h["date"].startswith("2026-10-03"))
    assert match["name"] == "Holiday Import Original Name"  # unchanged, not overwritten


def test_holiday_import_commit_overwrites_when_explicitly_requested(client, test_user):
    _login(client, test_user)
    client.post("/api/holiday-imports/commit", json={"rows": [{
        "date": "2026-10-04", "name": "Holiday Import Before Overwrite", "is_working": False,
    }]})

    resp = client.post("/api/holiday-imports/commit", json={"rows": [{
        "date": "2026-10-04", "name": "Holiday Import After Overwrite", "is_working": True,
        "overwrite_existing": True,
    }]})
    assert resp.status_code == 200
    body = resp.json()
    assert body["updated"] == 1

    holidays = client.get("/api/working-calendar/holidays").json()
    match = next(h for h in holidays if h["date"].startswith("2026-10-04"))
    assert match["name"] == "Holiday Import After Overwrite"
    assert match["is_working"] is True


def test_holiday_import_commit_requires_master(client, db_session):
    employee = User(
        username="holidayimportuser", email="holidayimportuser@example.com", full_name="Holiday Import User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=None, is_active=True,
    )
    db_session.add(employee)
    db_session.commit()
    resp = client.post("/api/auth/login", json={"identifier": "holidayimportuser@example.com", "password": "EmpPass1!"})
    assert resp.status_code == 200

    resp = client.post("/api/holiday-imports/commit", json={"rows": [{
        "date": "2026-10-05", "name": "Unauthorized Holiday", "is_working": False,
    }]})
    assert resp.status_code == 403


def test_holiday_import_template_requires_master(client, db_session):
    employee = User(
        username="holidaytemplateuser", email="holidaytemplateuser@example.com", full_name="Holiday Template User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=None, is_active=True,
    )
    db_session.add(employee)
    db_session.commit()
    resp = client.post("/api/auth/login", json={"identifier": "holidaytemplateuser@example.com", "password": "EmpPass1!"})
    assert resp.status_code == 200

    resp = client.get("/api/holiday-imports/template")
    assert resp.status_code == 403


def test_holiday_import_template_unauthenticated_rejected(client):
    resp = client.get("/api/holiday-imports/template")
    assert resp.status_code in (401, 403)


def _create_employee_leave(client):
    resp = client.post("/api/employees/", json={
        "employee_code": "EMP-LEAVE", "name": "Leave Test Employee", "monthly_salary": "18000.00",
    })
    return resp.json()["id"]


def _create_logged_in_employee(client, db_session, employee_name, username, email):
    """Unlike _create_employee_leave above (which only creates the employee
    record), this also creates and logs in as a real User account tied
    to that employee - needed for RBAC tests that check "own records"
    access by actually being logged in as the employee in question."""
    employee = client.post("/api/employees/", json={
        "name": employee_name, "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    user = User(
        username=username, email=email, full_name=username,
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    resp = client.post("/api/auth/login", json={"identifier": email, "password": "EmpPass1!"})
    assert resp.status_code == 200
    return employee


def test_leave_request_and_default_status(client, test_user):
    _login(client, test_user)
    employee_id = _create_employee_leave(client)

    resp = client.post("/api/leaves/", json={
        "employee_id": employee_id, "leave_type": "CL",
        "start_date": "2026-08-20T00:00:00", "end_date": "2026-08-21T00:00:00",
        "reason": "Family function",
    })
    assert resp.status_code == 201
    assert resp.json()["status"] == "Pending"
    assert resp.json()["days"] == "2"


def test_leave_days_is_1_when_from_equals_to(client, test_user):
    _login(client, test_user)
    employee_id = _create_employee_leave(client)

    resp = client.post("/api/leaves/", json={
        "employee_id": employee_id, "leave_type": "SL",
        "start_date": "2026-08-11T00:00:00", "end_date": "2026-08-11T00:00:00",
    })
    assert resp.status_code == 201
    assert resp.json()["days"] == "1"


def test_leave_days_inclusive_multi_day(client, test_user):
    _login(client, test_user)
    employee_id = _create_employee_leave(client)

    resp = client.post("/api/leaves/", json={
        "employee_id": employee_id, "leave_type": "PL",
        "start_date": "2026-08-11T00:00:00", "end_date": "2026-08-13T00:00:00",
    })
    assert resp.status_code == 201
    assert resp.json()["days"] == "3"


def test_leave_days_from_client_is_ignored(client, test_user):
    """Even if a caller (or a stale frontend build) sends a garbage `days`
    value, the server must compute its own - never trust the client for a
    number that must never be 0/negative/NaN."""
    _login(client, test_user)
    employee_id = _create_employee_leave(client)

    resp = client.post("/api/leaves/", json={
        "employee_id": employee_id, "leave_type": "CL",
        "start_date": "2026-08-11T00:00:00", "end_date": "2026-08-11T00:00:00",
        "days": "-1",
    })
    assert resp.status_code == 201
    assert resp.json()["days"] == "1"


def test_leave_rejects_end_before_start(client, test_user):
    _login(client, test_user)
    employee_id = _create_employee_leave(client)

    resp = client.post("/api/leaves/", json={
        "employee_id": employee_id, "leave_type": "SL",
        "start_date": "2026-08-20T00:00:00", "end_date": "2026-08-18T00:00:00",
    })
    assert resp.status_code == 400


def test_leave_approval_updates_status(client, test_user):
    _login(client, test_user)
    employee_id = _create_employee_leave(client)

    create = client.post("/api/leaves/", json={
        "employee_id": employee_id, "leave_type": "PL",
        "start_date": "2026-09-01T00:00:00", "end_date": "2026-09-02T00:00:00",
    })
    leave_id = create.json()["id"]

    resp = client.put(f"/api/leaves/{leave_id}", json={"status": "Approved", "approved_by": "Nikhil"})
    assert resp.status_code == 200
    assert resp.json()["status"] == "Approved"


def test_leaves_require_auth(client):
    resp = client.get("/api/leaves/")
    assert resp.status_code == 401


def test_employee_cannot_view_another_employees_leaves(client, test_user, db_session):
    _login(client, test_user)
    other_employee = client.post("/api/employees/", json={
        "name": "Leave RBAC Other Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    client.post("/api/leaves/", json={
        "employee_id": other_employee["id"], "leave_type": "Sick", "start_date": "2026-08-01T00:00:00", "end_date": "2026-08-02T00:00:00",
    })

    _create_logged_in_employee(client, db_session, "Leave RBAC Self Employee", "leaverbacuser", "leaverbacuser@example.com")
    resp = client.get("/api/leaves/", params={"employee_id": other_employee["id"]})
    assert resp.status_code == 403


def test_employee_can_view_own_leaves(client, test_user, db_session):
    _login(client, test_user)
    employee = _create_logged_in_employee(client, db_session, "Leave RBAC Own Employee", "leaveownrbacuser", "leaveownrbacuser@example.com")
    _login(client, test_user)
    client.post("/api/leaves/", json={
        "employee_id": employee["id"], "leave_type": "Casual", "start_date": "2026-08-05T00:00:00", "end_date": "2026-08-05T00:00:00",
    })

    client.post("/api/auth/login", json={"identifier": "leaveownrbacuser@example.com", "password": "EmpPass1!"})
    resp = client.get("/api/leaves/")
    assert resp.status_code == 200
    assert all(l["employee_id"] == employee["id"] for l in resp.json())


def test_employee_cannot_request_leave_for_someone_else(client, test_user, db_session):
    _login(client, test_user)
    other_employee = client.post("/api/employees/", json={
        "name": "Leave RBAC Impersonation Target", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    _create_logged_in_employee(client, db_session, "Leave RBAC Impersonator", "leaveimpersonateuser", "leaveimpersonateuser@example.com")

    resp = client.post("/api/leaves/", json={
        "employee_id": other_employee["id"], "leave_type": "Sick", "start_date": "2026-08-10T00:00:00", "end_date": "2026-08-10T00:00:00",
    })
    assert resp.status_code == 403


def test_chatbot_employee_cannot_view_another_employees_leaves(client, test_user, db_session):
    _login(client, test_user)
    other_employee = client.post("/api/employees/", json={
        "name": "ChatLeaveRbacTarget", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    client.post("/api/leaves/", json={
        "employee_id": other_employee["id"], "leave_type": "Sick", "start_date": "2026-08-01T00:00:00", "end_date": "2026-08-02T00:00:00",
    })

    _create_logged_in_employee(client, db_session, "Chat Leave RBAC Requester", "chatleaverbacuser", "chatleaverbacuser@example.com")
    resp = client.post("/api/chat/", json={"message": "Show ChatLeaveRbacTarget's leaves"})
    assert "own leave records" in resp.json()["response"].lower()


def test_chatbot_master_can_view_named_employees_leaves(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "ChatLeaveRbacMasterTarget", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    client.post("/api/leaves/", json={
        "employee_id": employee["id"], "leave_type": "Casual", "start_date": "2026-08-03T00:00:00", "end_date": "2026-08-03T00:00:00",
    })

    resp = client.post("/api/chat/", json={"message": "Show ChatLeaveRbacMasterTarget's leaves"})
    assert "own leave records" not in resp.json()["response"].lower()


def test_chatbot_employee_can_view_own_leaves(client, test_user, db_session):
    _login(client, test_user)
    employee = _create_logged_in_employee(client, db_session, "Chat Leave RBAC Self", "chatleaveselfrbacuser", "chatleaveselfrbacuser@example.com")
    _login(client, test_user)
    client.post("/api/leaves/", json={
        "employee_id": employee["id"], "leave_type": "Casual", "start_date": "2026-08-04T00:00:00", "end_date": "2026-08-04T00:00:00",
    })

    client.post("/api/auth/login", json={"identifier": "chatleaveselfrbacuser@example.com", "password": "EmpPass1!"})
    resp = client.post("/api/chat/", json={"message": "Show my leaves"})
    assert resp.status_code == 200
    assert "own leave records" not in resp.json()["response"].lower()


def test_unlinked_non_master_user_sees_no_leaves(client, test_user, db_session):
    """Unlinked non-master users must get zero records, not an
    unfiltered query result across every employee's leave history."""
    _login(client, test_user)
    other_id = _create_employee_leave(client)
    client.post("/api/leaves/", json={
        "employee_id": other_id, "leave_type": "Casual",
        "start_date": "2026-08-10T00:00:00", "end_date": "2026-08-10T00:00:00",
    })
    user = User(
        username="unlinkedleaveuser", email="unlinkedleaveuser@example.com", full_name="Unlinked Leave User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=None, is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "unlinkedleaveuser@example.com", "password": "EmpPass1!"})

    resp = client.get("/api/leaves/")
    assert resp.status_code == 200
    assert resp.json() == []


# --- test_payroll.py ---
"""Tests for the Salary Advance workflow (Family P0.44): employee/
Master request, Master approve (at requested or a different amount)/
reject, recovery against a real SalarySlip (never a bare number), and
the safety validations (no over-recovery, no recovery from an
unapproved/rejected advance, no recovery without a matching slip)."""


def _make_employee(client, suffix, salary="26000"):
    return client.post("/api/employees/", json={
        "name": f"Salary Advance Employee {suffix}", "monthly_salary": salary,
    }).json()


def _login_as_employee(client, db_session, employee_id, suffix):
    user = User(
        username=f"salaryadvanceuser{suffix}", email=f"salaryadvanceuser{suffix}@example.com",
        full_name=f"Salary Advance User {suffix}", password_hash=hash_password("EmpPass1!"),
        role="user", employee_id=employee_id, is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": f"salaryadvanceuser{suffix}@example.com", "password": "EmpPass1!"})


def test_employee_can_request_own_advance(client, test_user, db_session):
    _login(client, test_user)
    employee = _make_employee(client, "1")
    _login_as_employee(client, db_session, employee["id"], "1")

    resp = client.post("/api/salary-advances/", json={
        "employee_id": employee["id"], "requested_amount": "5000", "request_date": "2026-09-01T00:00:00",
        "reason": "Medical expense",
    })
    assert resp.status_code == 201
    body = resp.json()
    assert body["status"] == "Pending"
    assert float(body["requested_amount"]) == 5000.0
    assert body["created_by"] == f"salaryadvanceuser1"


def test_employee_cannot_request_advance_for_another_employee(client, test_user, db_session):
    _login(client, test_user)
    other_employee = _make_employee(client, "2a")
    self_employee = _make_employee(client, "2b")
    _login_as_employee(client, db_session, self_employee["id"], "2")

    resp = client.post("/api/salary-advances/", json={
        "employee_id": other_employee["id"], "requested_amount": "5000", "request_date": "2026-09-01T00:00:00",
    })
    assert resp.status_code == 403


def test_master_can_create_advance_on_behalf_of_employee(client, test_user):
    _login(client, test_user)
    employee = _make_employee(client, "3")

    resp = client.post("/api/salary-advances/", json={
        "employee_id": employee["id"], "requested_amount": "3000", "request_date": "2026-09-01T00:00:00",
    })
    assert resp.status_code == 201
    assert resp.json()["status"] == "Pending"


def test_approval_notifies_the_employees_linked_user(client, test_user, db_session):
    """The real end-to-end path: approve as Master, then log in as the
    employee's own linked account and confirm the notification is
    genuinely visible to them, not just that notify() was called."""
    _login(client, test_user)
    employee = _make_employee(client, "17")
    user = User(
        username="salaryadvancenotifyuser", email="salaryadvancenotifyuser@example.com",
        full_name="Salary Advance Notify User", password_hash=hash_password("EmpPass1!"),
        role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()

    advance = client.post("/api/salary-advances/", json={
        "employee_id": employee["id"], "requested_amount": "4000", "request_date": "2026-09-01T00:00:00",
    }).json()
    resp = client.put(f"/api/salary-advances/{advance['id']}/approve", json={
        "recovery_month": "September", "recovery_year": "2026",
    })
    assert resp.status_code == 200

    client.post("/api/auth/login", json={"identifier": "salaryadvancenotifyuser@example.com", "password": "EmpPass1!"})
    notifications = client.get("/api/notifications/").json()
    matches = [n for n in notifications if n["related_entity_type"] == "salary_advance" and n["related_entity_id"] == advance["id"]]
    assert len(matches) == 1
    assert "approved" in matches[0]["title"].lower()


def test_rejection_notifies_the_employees_linked_user_with_reason(client, test_user, db_session):
    _login(client, test_user)
    employee = _make_employee(client, "18")
    user = User(
        username="salaryadvancerejectnotifyuser", email="salaryadvancerejectnotifyuser@example.com",
        full_name="Salary Advance Reject Notify User", password_hash=hash_password("EmpPass1!"),
        role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()

    advance = client.post("/api/salary-advances/", json={
        "employee_id": employee["id"], "requested_amount": "4000", "request_date": "2026-09-01T00:00:00",
    }).json()
    client.put(f"/api/salary-advances/{advance['id']}/reject", json={"rejection_reason": "Insufficient reason given"})

    client.post("/api/auth/login", json={"identifier": "salaryadvancerejectnotifyuser@example.com", "password": "EmpPass1!"})
    notifications = client.get("/api/notifications/").json()
    matches = [n for n in notifications if n["related_entity_type"] == "salary_advance" and n["related_entity_id"] == advance["id"]]
    assert len(matches) == 1
    assert "rejected" in matches[0]["title"].lower()
    assert "Insufficient reason given" in matches[0]["message"]


def test_approval_does_not_fail_when_employee_has_no_linked_user(client, test_user):
    """A best-effort courtesy notification - its absence must never
    break the actual approval."""
    _login(client, test_user)
    employee = _make_employee(client, "19")
    advance = client.post("/api/salary-advances/", json={
        "employee_id": employee["id"], "requested_amount": "4000", "request_date": "2026-09-01T00:00:00",
    }).json()

    resp = client.put(f"/api/salary-advances/{advance['id']}/approve", json={
        "recovery_month": "September", "recovery_year": "2026",
    })
    assert resp.status_code == 200
    assert resp.json()["status"] == "Approved"


def test_master_approve_at_requested_amount(client, test_user):
    _login(client, test_user)
    employee = _make_employee(client, "4")
    advance = client.post("/api/salary-advances/", json={
        "employee_id": employee["id"], "requested_amount": "4000", "request_date": "2026-09-01T00:00:00",
    }).json()

    resp = client.put(f"/api/salary-advances/{advance['id']}/approve", json={
        "recovery_month": "September", "recovery_year": "2026",
    })
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "Approved"
    assert float(body["approved_amount"]) == 4000.0
    assert float(body["requested_amount"]) == 4000.0


def test_master_approve_at_different_amount_preserves_requested(client, test_user):
    """Spec: requested_amount remains historical; approved_amount is
    the separate, authoritative figure."""
    _login(client, test_user)
    employee = _make_employee(client, "5")
    advance = client.post("/api/salary-advances/", json={
        "employee_id": employee["id"], "requested_amount": "10000", "request_date": "2026-09-01T00:00:00",
    }).json()

    resp = client.put(f"/api/salary-advances/{advance['id']}/approve", json={
        "approved_amount": "6000", "recovery_month": "September", "recovery_year": "2026",
    })
    body = resp.json()
    assert float(body["approved_amount"]) == 6000.0
    assert float(body["requested_amount"]) == 10000.0  # unchanged


def test_master_reject_advance(client, test_user):
    _login(client, test_user)
    employee = _make_employee(client, "6")
    advance = client.post("/api/salary-advances/", json={
        "employee_id": employee["id"], "requested_amount": "4000", "request_date": "2026-09-01T00:00:00",
    }).json()

    resp = client.put(f"/api/salary-advances/{advance['id']}/reject", json={"rejection_reason": "Not eligible"})
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "Rejected"
    assert body["rejection_reason"] == "Not eligible"
    assert float(body["outstanding_amount"]) == 0.0


def test_cannot_approve_already_approved_advance(client, test_user):
    _login(client, test_user)
    employee = _make_employee(client, "7")
    advance = client.post("/api/salary-advances/", json={
        "employee_id": employee["id"], "requested_amount": "4000", "request_date": "2026-09-01T00:00:00",
    }).json()
    client.put(f"/api/salary-advances/{advance['id']}/approve", json={
        "recovery_month": "September", "recovery_year": "2026",
    })

    resp = client.put(f"/api/salary-advances/{advance['id']}/approve", json={
        "recovery_month": "September", "recovery_year": "2026",
    })
    assert resp.status_code == 409


def test_cannot_reject_already_rejected_advance(client, test_user):
    _login(client, test_user)
    employee = _make_employee(client, "8")
    advance = client.post("/api/salary-advances/", json={
        "employee_id": employee["id"], "requested_amount": "4000", "request_date": "2026-09-01T00:00:00",
    }).json()
    client.put(f"/api/salary-advances/{advance['id']}/reject", json={})

    resp = client.put(f"/api/salary-advances/{advance['id']}/reject", json={})
    assert resp.status_code == 409


def test_recovery_updates_advance_and_salary_slip(client, test_user):
    """The core P0.44/P0.43 connection: recovery actually changes a
    real SalarySlip's advance_deduction and net_salary - not just a
    number on the advance itself."""
    _login(client, test_user)
    employee = _make_employee(client, "9")
    advance = client.post("/api/salary-advances/", json={
        "employee_id": employee["id"], "requested_amount": "5000", "request_date": "2026-09-01T00:00:00",
    }).json()
    client.put(f"/api/salary-advances/{advance['id']}/approve", json={
        "recovery_month": "September", "recovery_year": "2026",
    })
    slip = client.post("/api/salary-slips/", json={
        "employee_id": employee["id"], "month": "September", "year": "2026", "basic": "20000",
    }).json()
    assert float(slip["net_salary"]) == 20000.0

    resp = client.post(f"/api/salary-advances/{advance['id']}/recover", json={
        "amount": "2000", "month": "September", "year": "2026",
    })
    assert resp.status_code == 200
    body = resp.json()
    assert float(body["recovered_amount"]) == 2000.0
    assert float(body["outstanding_amount"]) == 3000.0

    updated_slip = client.get(f"/api/salary-slips/{slip['id']}").json()
    assert float(updated_slip["advance_deduction"]) == 2000.0
    assert float(updated_slip["net_salary"]) == 18000.0  # 20000 - 2000


def test_cannot_recover_more_than_outstanding(client, test_user):
    _login(client, test_user)
    employee = _make_employee(client, "10")
    advance = client.post("/api/salary-advances/", json={
        "employee_id": employee["id"], "requested_amount": "5000", "request_date": "2026-09-01T00:00:00",
    }).json()
    client.put(f"/api/salary-advances/{advance['id']}/approve", json={
        "approved_amount": "3000", "recovery_month": "September", "recovery_year": "2026",
    })
    client.post("/api/salary-slips/", json={"employee_id": employee["id"], "month": "September", "year": "2026", "basic": "20000"})

    resp = client.post(f"/api/salary-advances/{advance['id']}/recover", json={
        "amount": "5000", "month": "September", "year": "2026",
    })
    assert resp.status_code == 400


def test_cannot_recover_from_pending_advance(client, test_user):
    _login(client, test_user)
    employee = _make_employee(client, "11")
    advance = client.post("/api/salary-advances/", json={
        "employee_id": employee["id"], "requested_amount": "5000", "request_date": "2026-09-01T00:00:00",
    }).json()
    client.post("/api/salary-slips/", json={"employee_id": employee["id"], "month": "September", "year": "2026", "basic": "20000"})

    resp = client.post(f"/api/salary-advances/{advance['id']}/recover", json={
        "amount": "1000", "month": "September", "year": "2026",
    })
    assert resp.status_code == 409


def test_cannot_recover_from_rejected_advance(client, test_user):
    _login(client, test_user)
    employee = _make_employee(client, "12")
    advance = client.post("/api/salary-advances/", json={
        "employee_id": employee["id"], "requested_amount": "5000", "request_date": "2026-09-01T00:00:00",
    }).json()
    client.put(f"/api/salary-advances/{advance['id']}/reject", json={})
    client.post("/api/salary-slips/", json={"employee_id": employee["id"], "month": "September", "year": "2026", "basic": "20000"})

    resp = client.post(f"/api/salary-advances/{advance['id']}/recover", json={
        "amount": "1000", "month": "September", "year": "2026",
    })
    assert resp.status_code == 409


def test_recovery_without_matching_salary_slip_rejected(client, test_user):
    _login(client, test_user)
    employee = _make_employee(client, "13")
    advance = client.post("/api/salary-advances/", json={
        "employee_id": employee["id"], "requested_amount": "5000", "request_date": "2026-09-01T00:00:00",
    }).json()
    client.put(f"/api/salary-advances/{advance['id']}/approve", json={
        "recovery_month": "September", "recovery_year": "2026",
    })
    # No SalarySlip created for this employee/month at all.

    resp = client.post(f"/api/salary-advances/{advance['id']}/recover", json={
        "amount": "1000", "month": "September", "year": "2026",
    })
    assert resp.status_code == 404


def test_employee_cannot_view_another_employees_advance(client, test_user, db_session):
    _login(client, test_user)
    other_employee = _make_employee(client, "14a")
    other_advance = client.post("/api/salary-advances/", json={
        "employee_id": other_employee["id"], "requested_amount": "4000", "request_date": "2026-09-01T00:00:00",
    }).json()
    self_employee = _make_employee(client, "14b")
    _login_as_employee(client, db_session, self_employee["id"], "14")

    resp = client.get(f"/api/salary-advances/{other_advance['id']}")
    assert resp.status_code == 403

    resp = client.get("/api/salary-advances/", params={"employee_id": other_employee["id"]})
    assert resp.status_code == 403


def test_employee_can_view_own_advance(client, test_user, db_session):
    _login(client, test_user)
    employee = _make_employee(client, "15")
    advance = client.post("/api/salary-advances/", json={
        "employee_id": employee["id"], "requested_amount": "4000", "request_date": "2026-09-01T00:00:00",
    }).json()
    _login_as_employee(client, db_session, employee["id"], "15")

    resp = client.get(f"/api/salary-advances/{advance['id']}")
    assert resp.status_code == 200


def test_employee_cannot_approve_or_reject(client, test_user, db_session):
    _login(client, test_user)
    employee = _make_employee(client, "16")
    advance = client.post("/api/salary-advances/", json={
        "employee_id": employee["id"], "requested_amount": "4000", "request_date": "2026-09-01T00:00:00",
    }).json()
    _login_as_employee(client, db_session, employee["id"], "16")

    resp = client.put(f"/api/salary-advances/{advance['id']}/approve", json={
        "recovery_month": "September", "recovery_year": "2026",
    })
    assert resp.status_code == 403


def test_salary_advances_require_auth(client):
    resp = client.get("/api/salary-advances/")
    assert resp.status_code == 401


"""Salary slip tests: creation/computation (net salary, working/paid
days validation, duplicate rejection), PDF generation across various
field-completeness scenarios, ownership-based RBAC (an employee sees
only their own slip, master sees all), and the unlinked-non-master
defense-in-depth regression guard."""


BUSINESS_ID_PATTERN = re.compile(r"^[A-Z0-9]{10}$")



def _create_employee_payroll(client, name="Salary Slip Test Employee"):
    resp = client.post("/api/employees/", json={
        "name": name, "designation": "Carpenter", "department": "Production", "monthly_salary": "26000.00",
    })
    assert resp.status_code == 201
    return resp.json()["id"]


def test_create_salary_slip_end_to_end(client, test_user):
    _login(client, test_user)
    employee_id = _create_employee_payroll(client)

    resp = client.post("/api/salary-slips/", json={
        "employee_id": employee_id, "month": "August", "year": "2026",
        "working_days": "26", "paid_days": "26",
        "basic": "20000.00", "da": "2000.00", "hra": "3000.00", "overtime_amount": "500.00",
        "pf_deduction": "1200.00", "tds_deduction": "0.00", "other_deductions": "0.00",
    })
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert BUSINESS_ID_PATTERN.match(body["business_id"])
    # 20000 + 2000 + 3000 + 500 - 1200 = 24300
    assert float(body["net_salary"]) == 24300.0


def test_salary_slip_defaults_working_and_paid_days(client, test_user):
    _login(client, test_user)
    employee_id = _create_employee_payroll(client)

    resp = client.post("/api/salary-slips/", json={
        "employee_id": employee_id, "month": "September", "year": "2026", "basic": "20000.00",
    })
    assert resp.status_code == 201
    body = resp.json()
    assert float(body["working_days"]) == 26.0
    assert float(body["paid_days"]) == 26.0


def test_salary_slip_rejects_paid_days_exceeding_working_days(client, test_user):
    _login(client, test_user)
    employee_id = _create_employee_payroll(client)

    resp = client.post("/api/salary-slips/", json={
        "employee_id": employee_id, "month": "October", "year": "2026",
        "working_days": "20", "paid_days": "25", "basic": "20000.00",
    })
    assert resp.status_code == 422


def test_salary_slip_rejects_zero_or_negative_days(client, test_user):
    _login(client, test_user)
    employee_id = _create_employee_payroll(client)

    resp = client.post("/api/salary-slips/", json={
        "employee_id": employee_id, "month": "November", "year": "2026",
        "working_days": "0", "paid_days": "0", "basic": "20000.00",
    })
    assert resp.status_code == 422


def test_duplicate_salary_slip_rejected(client, test_user):
    _login(client, test_user)
    employee_id = _create_employee_payroll(client)
    payload = {"employee_id": employee_id, "month": "December", "year": "2026", "basic": "20000.00"}

    first = client.post("/api/salary-slips/", json=payload)
    assert first.status_code == 201
    second = client.post("/api/salary-slips/", json=payload)
    assert second.status_code == 400


def test_salary_slip_update_rejects_paid_days_exceeding_working_days(client, test_user):
    _login(client, test_user)
    employee_id = _create_employee_payroll(client)
    create = client.post("/api/salary-slips/", json={
        "employee_id": employee_id, "month": "January", "year": "2027",
        "working_days": "26", "paid_days": "26", "basic": "20000.00",
    })
    slip_id = create.json()["id"]

    resp = client.put(f"/api/salary-slips/{slip_id}", json={"working_days": "10"})
    assert resp.status_code == 400


def test_salary_slip_pdf_generates_successfully_with_minimal_fields(client, test_user):
    _login(client, test_user)
    employee_id = _create_employee_payroll(client)
    create = client.post("/api/salary-slips/", json={
        "employee_id": employee_id, "month": "February", "year": "2027", "basic": "20000.00",
    })
    slip_id = create.json()["id"]

    resp = client.get(f"/api/reports/salary-slips/{slip_id}.pdf")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert len(resp.content) > 0


def test_employee_can_view_own_salary_slip(client, test_user, db_session):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "Salary RBAC Own Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    slip = client.post("/api/salary-slips/", json={
        "employee_id": employee["id"], "month": "8", "year": "2026", "working_days": "26", "paid_days": "26",
        "basic": "15000", "da": "0", "hra": "0", "overtime_amount": "0",
        "pf_deduction": "0", "tds_deduction": "0", "other_deductions": "0",
    }).json()

    user = User(
        username="salaryownrbacuser", email="salaryownrbacuser@example.com", full_name="Salary Own",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "salaryownrbacuser@example.com", "password": "EmpPass1!"})

    resp = client.get(f"/api/salary-slips/{slip['id']}")
    assert resp.status_code == 200


def test_employee_cannot_view_another_employees_salary_slip(client, test_user, db_session):
    _login(client, test_user)
    other_employee = client.post("/api/employees/", json={
        "name": "Salary RBAC Other Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()
    slip = client.post("/api/salary-slips/", json={
        "employee_id": other_employee["id"], "month": "8", "year": "2026", "working_days": "26", "paid_days": "26",
        "basic": "15000", "da": "0", "hra": "0", "overtime_amount": "0",
        "pf_deduction": "0", "tds_deduction": "0", "other_deductions": "0",
    }).json()

    _create_logged_in_employee(client, db_session, "Salary RBAC Requester", "salaryotherrbacuser", "salaryotherrbacuser@example.com")
    resp = client.get(f"/api/salary-slips/{slip['id']}")
    assert resp.status_code == 403


def test_employee_cannot_create_salary_slip(client, test_user, db_session):
    """Creation stays master-only - viewing opened up, editing did not."""
    _login(client, test_user)
    employee = _create_logged_in_employee(client, db_session, "Salary RBAC Create Employee", "salarycreaterbacuser", "salarycreaterbacuser@example.com")
    resp = client.post("/api/salary-slips/", json={
        "employee_id": employee["id"], "month": "9", "year": "2026", "working_days": "26", "paid_days": "26",
        "basic": "15000", "da": "0", "hra": "0", "overtime_amount": "0",
        "pf_deduction": "0", "tds_deduction": "0", "other_deductions": "0",
    })
    assert resp.status_code == 403


"""Tests that the salary slip PDF endpoint actually generates a valid
PDF response for real stored data, and that access to it is correctly
scoped: master can download any slip, an employee can download their
own, and cannot download another employee's - a real inconsistency
found during a full-backend systematic sweep, since get_salary_slip
(the API for the same underlying record) had already been widened to
let an employee view their own, but the PDF export was still
master-only."""


def test_salary_slip_pdf_generates_successfully(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "PDF Gen Test Employee", "monthly_salary": "20000",
        "designation": "Carpenter", "department": "Assembly",
        "pan": "ABCPT1234A", "uan": "100200300999",
        "bank_name": "State Bank of India", "bank_account_number": "998877661234",
        "tax_regime": "New",
    }).json()
    slip = client.post("/api/salary-slips/", json={
        "employee_id": employee["id"], "month": "August", "year": "2026",
        "working_days": "26", "paid_days": "26",
        "basic": "9500", "da": "500", "hra": "6000", "overtime_amount": "1500",
        "pf_deduction": "1140", "tds_deduction": "0", "other_deductions": "300",
    }).json()

    resp = client.get(f"/api/reports/salary-slips/{slip['id']}.pdf")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert len(resp.content) > 1000  # a genuine PDF, not an empty/error stub


def test_salary_slip_pdf_generates_with_no_leave_history(client, test_user):
    """An employee with zero leave records must not crash the PDF -
    the Leave Balance section should degrade gracefully."""
    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "PDF Gen No Leave Employee", "monthly_salary": "18000",
    }).json()
    slip = client.post("/api/salary-slips/", json={
        "employee_id": employee["id"], "month": "August", "year": "2026",
        "working_days": "26", "paid_days": "26",
        "basic": "9000", "da": "0", "hra": "0", "overtime_amount": "0",
        "pf_deduction": "0", "tds_deduction": "0", "other_deductions": "0",
    }).json()

    resp = client.get(f"/api/reports/salary-slips/{slip['id']}.pdf")
    assert resp.status_code == 200


def test_salary_slip_pdf_generates_with_missing_employee_payroll_fields(client, test_user):
    """PAN/UAN/bank details are all optional - a slip for an employee
    without them must still generate cleanly (shown as '-')."""
    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "PDF Gen Minimal Employee", "monthly_salary": "14000",
    }).json()
    assert employee.get("pan") is None
    slip = client.post("/api/salary-slips/", json={
        "employee_id": employee["id"], "month": "August", "year": "2026",
        "working_days": "26", "paid_days": "26",
        "basic": "7000", "da": "0", "hra": "3000", "overtime_amount": "0",
        "pf_deduction": "0", "tds_deduction": "0", "other_deductions": "200",
    }).json()

    resp = client.get(f"/api/reports/salary-slips/{slip['id']}.pdf")
    assert resp.status_code == 200


def test_employee_can_download_own_salary_slip_pdf(client, test_user, db_session):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "Salary PDF RBAC Own Employee", "monthly_salary": "20000",
    }).json()
    slip = client.post("/api/salary-slips/", json={
        "employee_id": employee["id"], "month": "8", "year": "2026", "working_days": "26", "paid_days": "26",
        "basic": "15000", "da": "0", "hra": "0", "overtime_amount": "0",
        "pf_deduction": "0", "tds_deduction": "0", "other_deductions": "0",
    }).json()
    user = User(
        username="salarypdfownrbacuser", email="salarypdfownrbacuser@example.com", full_name="Salary PDF Own RBAC User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "salarypdfownrbacuser@example.com", "password": "EmpPass1!"})

    resp = client.get(f"/api/reports/salary-slips/{slip['id']}.pdf")
    assert resp.status_code == 200


def test_employee_cannot_download_another_employees_salary_slip_pdf(client, test_user, db_session):
    _login(client, test_user)
    other_employee = client.post("/api/employees/", json={
        "name": "Salary PDF RBAC Other Employee", "monthly_salary": "25000",
    }).json()
    slip = client.post("/api/salary-slips/", json={
        "employee_id": other_employee["id"], "month": "8", "year": "2026", "working_days": "26", "paid_days": "26",
        "basic": "18000", "da": "0", "hra": "0", "overtime_amount": "0",
        "pf_deduction": "0", "tds_deduction": "0", "other_deductions": "0",
    }).json()
    employee = client.post("/api/employees/", json={
        "name": "Salary PDF RBAC Requester", "monthly_salary": "18000",
    }).json()
    user = User(
        username="salarypdfotherrbacuser", email="salarypdfotherrbacuser@example.com", full_name="Salary PDF Other RBAC User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "salarypdfotherrbacuser@example.com", "password": "EmpPass1!"})

    resp = client.get(f"/api/reports/salary-slips/{slip['id']}.pdf")
    assert resp.status_code == 403


def test_master_can_download_any_salary_slip_pdf(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "Salary PDF RBAC Master Employee", "monthly_salary": "20000",
    }).json()
    slip = client.post("/api/salary-slips/", json={
        "employee_id": employee["id"], "month": "8", "year": "2026", "working_days": "26", "paid_days": "26",
        "basic": "15000", "da": "0", "hra": "0", "overtime_amount": "0",
        "pf_deduction": "0", "tds_deduction": "0", "other_deductions": "0",
    }).json()

    resp = client.get(f"/api/reports/salary-slips/{slip['id']}.pdf")
    assert resp.status_code == 200


def test_master_sees_all_employee_salaries(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "Salary RBAC Master Employee", "monthly_salary": "25000",
    }).json()
    resp = client.get(f"/api/employees/{employee['id']}").json()
    assert resp["monthly_salary"] is not None
    assert resp["daily_wage"] is not None


def test_employee_cannot_see_another_employees_salary(client, test_user, db_session):
    _login(client, test_user)
    other_employee = client.post("/api/employees/", json={
        "name": "Salary RBAC Other Employee", "monthly_salary": "30000",
    }).json()
    employee = client.post("/api/employees/", json={
        "name": "Salary RBAC Self Employee", "monthly_salary": "20000",
    }).json()
    user = User(
        username="empsalaryrbacuser", email="empsalaryrbacuser@example.com", full_name="Emp Salary RBAC User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "empsalaryrbacuser@example.com", "password": "EmpPass1!"})

    resp = client.get(f"/api/employees/{other_employee['id']}").json()
    assert resp["monthly_salary"] is None
    assert resp["daily_wage"] is None
    # Non-financial directory info remains visible.
    assert resp["name"] == "Salary RBAC Other Employee"
    assert "designation" in resp


def test_employee_can_see_own_salary(client, test_user, db_session):
    """The specific nuance of this fix - own record stays visible,
    matching "employee can access only their OWN confidential HR data",
    not a blanket denial of all salary info."""
    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "Salary RBAC Own Employee", "monthly_salary": "22000",
    }).json()
    user = User(
        username="empsalaryownrbacuser", email="empsalaryownrbacuser@example.com", full_name="Emp Salary Own RBAC User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "empsalaryownrbacuser@example.com", "password": "EmpPass1!"})

    resp = client.get(f"/api/employees/{employee['id']}").json()
    assert resp["monthly_salary"] is not None
    assert float(resp["monthly_salary"]) == 22000.0
    assert resp["daily_wage"] is not None


def test_employee_list_view_redacts_others_but_shows_own_salary(client, test_user, db_session):
    _login(client, test_user)
    other_employee = client.post("/api/employees/", json={
        "name": "Salary RBAC List Other Employee", "monthly_salary": "35000",
    }).json()
    employee = client.post("/api/employees/", json={
        "name": "Salary RBAC List Self Employee", "monthly_salary": "18000",
    }).json()
    user = User(
        username="empsalarylistrbacuser", email="empsalarylistrbacuser@example.com", full_name="Emp Salary List RBAC User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "empsalarylistrbacuser@example.com", "password": "EmpPass1!"})

    employees = client.get("/api/employees/").json()
    other_row = next(e for e in employees if e["id"] == other_employee["id"])
    own_row = next(e for e in employees if e["id"] == employee["id"])
    assert other_row["monthly_salary"] is None
    assert own_row["monthly_salary"] is not None


def test_unlinked_non_master_user_sees_no_salary_slips(client, test_user, db_session):
    """Unlinked non-master users must get zero records, not every
    employee's salary slips - the most sensitive instance of this
    class of bug, since it's compensation data."""
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Unlinked Salary Other Employee"}).json()
    client.post("/api/salary-slips/", json={
        "employee_id": employee["id"], "month": "August", "year": "2026",
        "working_days": "26", "paid_days": "26",
        "basic": "20000", "da": "0", "hra": "0", "overtime_amount": "0",
        "pf_deduction": "0", "tds_deduction": "0", "other_deductions": "0",
    })
    user = User(
        username="unlinkedsalaryuser", email="unlinkedsalaryuser@example.com", full_name="Unlinked Salary User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=None, is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "unlinkedsalaryuser@example.com", "password": "EmpPass1!"})

    resp = client.get("/api/salary-slips/")
    assert resp.status_code == 200
    assert resp.json() == []


def test_salary_slip_computes_net(client, test_user):
    _login(client, test_user)
    resp = client.post("/api/employees/", json={
        "employee_code": "EMP-SAL", "name": "Salary Test Employee", "monthly_salary": "20000.00",
    })
    employee_id = resp.json()["id"]

    resp = client.post("/api/salary-slips/", json={
        "employee_id": employee_id, "month": "August", "year": "2026",
        "basic": "15000.00", "da": "2000.00", "hra": "3000.00",
        "pf_deduction": "1800.00", "tds_deduction": "0.00",
    })
    assert resp.status_code == 201
    assert float(resp.json()["net_salary"]) == 18200.0


# --- test_workforce_excel_exports.py ---
def test_attendance_export_opens_and_respects_employee_filter(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Excel Export Attendance Employee"}).json()
    client.post("/api/attendance/", json={
        "date": "2026-08-10T00:00:00", "employee_id": employee["id"],
        "in_time": "2026-08-10T09:00:00", "out_time": "2026-08-10T18:00:00", "attendance_status": "Present",
    })

    resp = client.get("/api/reports/attendance.xlsx", params={"employee_id": employee["id"]})
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert "Attendance" in wb.sheetnames


def test_overtime_export_only_includes_overtime_records(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Excel Export Overtime Employee"}).json()
    client.post("/api/attendance/", json={
        "date": "2026-08-11T00:00:00", "employee_id": employee["id"],
        "in_time": "2026-08-11T09:00:00", "out_time": "2026-08-11T19:00:00",
        "standard_hours": "8", "attendance_status": "Present",
    })
    resp = client.get("/api/reports/attendance.xlsx", params={"employee_id": employee["id"], "overtime_only": True})
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert "Overtime" in wb.sheetnames


def test_leaves_export_opens_correctly(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Excel Export Leave Employee"}).json()
    client.post("/api/leaves/", json={
        "employee_id": employee["id"], "leave_type": "CL",
        "start_date": "2026-08-15T00:00:00", "end_date": "2026-08-15T00:00:00",
    })
    resp = client.get("/api/reports/leaves.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert "Leave" in wb.sheetnames


def test_tasks_export_respects_employee_filter(client, test_user):
    """A filtered
    export must not silently return the whole database."""
    _login(client, test_user)
    emp_a = client.post("/api/employees/", json={"name": "Excel Export Tasks Employee A"}).json()
    emp_b = client.post("/api/employees/", json={"name": "Excel Export Tasks Employee B"}).json()
    client.post("/api/daily-tasks/", json={
        "date": "2026-08-19T00:00:00", "employee_id": emp_a["id"], "task_description": "Task for A only",
    })
    client.post("/api/daily-tasks/", json={
        "date": "2026-08-19T00:00:00", "employee_id": emp_b["id"], "task_description": "Task for B only",
    })

    resp = client.get("/api/reports/tasks.xlsx", params={"employee_id": emp_a["id"]})
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["Tasks"]
    task_descriptions = [cell.value for row in ws.iter_rows(min_row=1) for cell in row if isinstance(cell.value, str)]
    assert any("Task for A only" in d for d in task_descriptions)
    assert not any("Task for B only" in d for d in task_descriptions)


def test_payroll_export_requires_master(client, test_user, db_session):
    from app.platform.security import hash_password
    from app.modules.auth.auth import User
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Excel Export Payroll Permission Employee"}).json()
    user = User(
        username="payrollexportuser", email="payrollexportuser@example.com",
        full_name="Payroll Export User", password_hash=hash_password("EmpPass1!"),
        role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "payrollexportuser@example.com", "password": "EmpPass1!"})

    resp = client.get("/api/reports/payroll.xlsx")
    assert resp.status_code == 403


def test_payroll_export_works_for_master(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Excel Export Payroll Master Employee", "monthly_salary": "20000"}).json()
    client.post("/api/salary-slips/", json={
        "employee_id": employee["id"], "month": "August", "year": "2026",
        "working_days": "26", "paid_days": "26",
        "basic": "15000", "da": "0", "hra": "5000", "overtime_amount": "0",
        "pf_deduction": "0", "tds_deduction": "0", "other_deductions": "0",
    })
    resp = client.get("/api/reports/payroll.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert "Payroll" in wb.sheetnames


def test_tasks_excel_export_contains_real_data(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Tasks Excel Employee"}).json()
    client.post("/api/daily-tasks/", json={
        "date": "2026-08-18T00:00:00", "employee_id": employee["id"], "task_description": "Tasks Excel Test Task",
    })
    client.post("/api/production-jobs/", json={
        "date": "2026-08-18T00:00:00", "employee_id": employee["id"], "operation": "Tasks Excel Test Operation",
        "planned_qty": 5, "status": "Not Started",
    })

    resp = client.get("/api/reports/tasks.xlsx")
    assert resp.status_code == 200
    assert len(resp.content) > 1000  # a genuine workbook, not an empty/error stub

    wb = load_workbook(io.BytesIO(resp.content))
    assert "Tasks" in wb.sheetnames
    assert "Production" in wb.sheetnames


def test_tasks_excel_export_open_to_any_authenticated_role(client, test_user, db_session):
    """Everyone can view tasks - the export must not be restricted."""
    from app.platform.security import hash_password
    from app.modules.auth.auth import User

    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Tasks Excel RBAC Employee"}).json()
    user = User(
        username="tasksexcelrbacuser", email="tasksexcelrbacuser@example.com", full_name="Tasks Excel RBAC User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "tasksexcelrbacuser@example.com", "password": "EmpPass1!"})

    resp = client.get("/api/reports/tasks.xlsx")
    assert resp.status_code == 200


def _create_unlinked_non_master_user(client, db_session, username, email):
    user = User(
        username=username, email=email, full_name=username,
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=None, is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": email, "password": "EmpPass1!"})


def test_unlinked_user_gets_no_attendance_export(client, test_user, db_session):
    _login(client, test_user)
    client.post("/api/employees/", json={"name": "Unlinked Export Attendance Employee"})
    _create_unlinked_non_master_user(client, db_session, "unlinkedattexport", "unlinkedattexport@example.com")

    resp = client.get("/api/reports/attendance.xlsx")
    assert resp.status_code == 403


def test_unlinked_user_gets_no_leaves_export(client, test_user, db_session):
    _login(client, test_user)
    client.post("/api/employees/", json={"name": "Unlinked Export Leave Employee"})
    _create_unlinked_non_master_user(client, db_session, "unlinkedleaveexport", "unlinkedleaveexport@example.com")

    resp = client.get("/api/reports/leaves.xlsx")
    assert resp.status_code == 403


# --- Family 137 (updated) - Employee 360 / HR Command Center (section 13) ---

def test_employee_360_overview_not_found(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/employees/999999/360-overview")
    assert resp.status_code == 404


def test_employee_360_overview_master_sees_salary_and_cost(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "E360 Overview Employee", "monthly_salary": "30000"}).json()
    resp = client.get(f"/api/employees/{employee['id']}/360-overview")
    assert resp.status_code == 200
    data = resp.json()
    assert data["label"] == "FACT"
    assert data["profile"]["name"] == "E360 Overview Employee"
    assert "salary_summary" in data
    assert "cost_contribution" in data
    assert data["cost_contribution"]["disclaimer"]
    assert data["documents_count"] == 0
    from app.modules.hr.models import ONBOARDING_ITEMS
    assert data["onboarding_progress"]["total"] == len(ONBOARDING_ITEMS)


def test_employee_360_overview_other_employee_forbidden(client, test_user, db_session):
    _login(client, test_user)
    other_employee = client.post("/api/employees/", json={"name": "E360 RBAC Other Employee"}).json()
    self_employee = client.post("/api/employees/", json={"name": "E360 RBAC Self Employee"}).json()
    user = User(
        username="e360rbacuser", email="e360rbacuser@example.com", full_name="E360 RBAC User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=self_employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "e360rbacuser@example.com", "password": "EmpPass1!"})

    resp = client.get(f"/api/employees/{other_employee['id']}/360-overview")
    assert resp.status_code == 403

    own_resp = client.get(f"/api/employees/{self_employee['id']}/360-overview")
    assert own_resp.status_code == 200
    # A non-master viewing their OWN overview must not see document facts
    # (the Documents API itself is master-only for the employee parent
    # type, even for an employee's own documents).
    assert own_resp.json()["documents_count"] is None


def test_employee_workload_classification_is_deterministic(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "E360 Workload Employee"}).json()

    low = client.get(f"/api/employees/{employee['id']}/workload").json()
    assert low["workload_level"] == "Low"
    assert low["basis"]

    for i in range(5):
        client.post("/api/daily-tasks/", json={
            "date": "2026-08-20T00:00:00", "employee_id": employee["id"],
            "task_description": f"Workload task {i}",
        })
    high = client.get(f"/api/employees/{employee['id']}/workload").json()
    assert high["active_tasks"] == 5


def test_employee_calendar_shape(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "E360 Calendar Employee"}).json()
    resp = client.get(f"/api/employees/{employee['id']}/calendar", params={"year": 2026, "month": 8})
    assert resp.status_code == 200
    data = resp.json()
    assert data["year"] == 2026 and data["month"] == 8
    assert len(data["days"]) == 31
    assert all(d["event_type"] in ("present", "half_day", "absent", "leave", "holiday", "week_off", "work_day") for d in data["days"])


def test_employee_lifecycle_onboarding_seeds_and_auto_computes(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={
        "name": "E360 Lifecycle Employee", "designation": "Carpenter", "manager": "Someone",
    }).json()
    resp = client.get(f"/api/employees/{employee['id']}/lifecycle", params={"phase": "onboarding"})
    assert resp.status_code == 200
    items = {i["item_key"]: i for i in resp.json()["items"]}
    assert items["record_created"]["is_complete"] is True
    assert items["role_assigned"]["is_complete"] is True  # designation was set on create
    assert items["manager_assigned"]["is_complete"] is True
    assert items["documents_submitted"]["is_complete"] is False
    assert items["contract_completed"]["is_complete"] is False


def test_employee_lifecycle_offboarding_empty_until_inactive(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "E360 Offboarding Employee"}).json()
    resp = client.get(f"/api/employees/{employee['id']}/lifecycle", params={"phase": "offboarding"})
    assert resp.json()["items"] == []

    client.put(f"/api/employees/{employee['id']}", json={"status": "Inactive"})
    resp2 = client.get(f"/api/employees/{employee['id']}/lifecycle", params={"phase": "offboarding"})
    assert len(resp2.json()["items"]) > 0


def test_employee_lifecycle_item_toggle_is_master_only_and_audited(client, test_user, db_session):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "E360 Toggle Employee"}).json()
    items = client.get(f"/api/employees/{employee['id']}/lifecycle", params={"phase": "onboarding"}).json()["items"]
    manual_item = next(i for i in items if not i["auto_computed"])

    resp = client.put(f"/api/employees/{employee['id']}/lifecycle/{manual_item['id']}", json={"is_complete": True})
    assert resp.status_code == 200
    assert resp.json()["is_complete"] is True

    from app.platform.audit import AuditLog
    logged = db_session.query(AuditLog).filter(AuditLog.action == "update_lifecycle_item").first()
    assert logged is not None


def test_employee_lifecycle_auto_item_cannot_be_toggled_directly(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "E360 Auto Item Employee"}).json()
    items = client.get(f"/api/employees/{employee['id']}/lifecycle", params={"phase": "onboarding"}).json()["items"]
    auto_item = next(i for i in items if i["auto_computed"])

    resp = client.put(f"/api/employees/{employee['id']}/lifecycle/{auto_item['id']}", json={"is_complete": True})
    assert resp.status_code == 400


def test_employee_activity_timeline_includes_update_audit(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "E360 Timeline Employee", "designation": "Helper"}).json()
    client.put(f"/api/employees/{employee['id']}", json={"designation": "Senior Helper"})

    resp = client.get(f"/api/employees/{employee['id']}/activity-timeline")
    assert resp.status_code == 200
    data = resp.json()
    assert data["total_entries"] >= 1
    texts = [e["text"] for e in data["entries"]]
    assert any("designation" in t for t in texts)


def test_employee_needs_attention_flags_overdue_task(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "E360 Needs Attention Employee"}).json()
    client.post("/api/daily-tasks/", json={
        "date": "2020-01-01T00:00:00", "employee_id": employee["id"], "task_description": "Old overdue task",
    })
    resp = client.get(f"/api/employees/{employee['id']}/360-overview")
    needs_attention = resp.json()["needs_attention"]
    assert needs_attention["total"] >= 1
    assert any(i["type"] == "overdue_tasks" for i in needs_attention["items"])


def test_employee_relationships_direct_reports(client, test_user):
    _login(client, test_user)
    manager = client.post("/api/employees/", json={"name": "E360 Manager Employee"}).json()
    client.post("/api/employees/", json={"name": "E360 Report Employee", "manager": "E360 Manager Employee"})

    resp = client.get(f"/api/employees/{manager['id']}/relationships")
    assert resp.status_code == 200
    reports = resp.json()["direct_reports"]
    assert any(r["name"] == "E360 Report Employee" for r in reports)


def test_document_upload_accepts_type_and_expiry(client, test_user):
    import io as io_module

    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "E360 Document Employee"}).json()
    file_bytes = b"%PDF-1.4 test content"
    resp = client.post(
        f"/api/documents/employee/{employee['id']}",
        files={"file": ("id_proof.pdf", io_module.BytesIO(file_bytes), "application/pdf")},
        data={"document_type": "Identity", "expiry_date": "2099-01-01"},
    )
    assert resp.status_code == 201
    doc = resp.json()
    assert doc["document_type"] == "Identity"
    assert doc["expiry_date"] == "2099-01-01"


# --- test_attendance_overtime_command_center.py ---
"""Tests for the Attendance & Overtime Command Center redesign:
attendance_period_summary/day_detail/team_attendance_grid/
attendance_exceptions, the new date_from/date_to range filter on
GET /api/attendance/, and the new OvertimeRequest Draft -> Submitted
-> Approved/Rejected workflow (converging on the same
Attendance.overtime_hours the pre-existing bulk 'Manage Overtime'
action writes, via the shared apply_overtime_hours). Aug 2026 dates
throughout are ordinary Mon-Sat working days under the default
weekday configuration, safely in the past relative to this suite's
'today' (2026-08-17 is a Monday)."""


def _make_cc_employee(client, suffix):
    return client.post("/api/employees/", json={"name": f"Command Center Employee {suffix}", "monthly_salary": "26000"}).json()


def _login_as_cc_employee(client, db_session, employee_id, suffix):
    user = User(
        username=f"ccuser{suffix}", email=f"ccuser{suffix}@example.com", full_name=f"Command Center User {suffix}",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee_id, is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": f"ccuser{suffix}@example.com", "password": "EmpPass1!"})
    return user


def test_period_summary_present_day_with_overtime(client, test_user):
    _login(client, test_user)
    employee = _make_cc_employee(client, "1")
    client.post("/api/attendance/", json={
        "date": "2026-08-17T00:00:00", "employee_id": employee["id"], "attendance_status": "Present",
        "in_time": "2026-08-17T09:00:00", "out_time": "2026-08-17T18:00:00",
    })
    client.post("/api/attendance/overtime", json={"employee_id": employee["id"], "dates": ["2026-08-17T00:00:00"], "hours": "2"})

    resp = client.get("/api/attendance/period-summary", params={
        "employee_id": employee["id"], "start": "2026-08-17T00:00:00", "end": "2026-08-17T00:00:00",
    })
    assert resp.status_code == 200
    body = resp.json()
    day = body["days"][0]
    assert day["scheduled_hours"] == 8.0
    assert day["worked_hours"] == 9.0
    assert day["overtime_hours"] == 2.0
    assert day["payable_hours"] == 10.0  # 1 (Present day-value) * 8 scheduled + 2 overtime
    assert day["shortfall_hours"] == 0.0  # worked (9h) already exceeds scheduled (8h)
    assert body["summary"]["payable_hours"] == 10.0


def test_period_summary_half_day_and_absent_have_shortfall(client, test_user):
    _login(client, test_user)
    employee = _make_cc_employee(client, "2")
    client.post("/api/attendance/", json={"date": "2026-08-18T00:00:00", "employee_id": employee["id"], "attendance_status": "Half Day"})
    client.post("/api/attendance/", json={"date": "2026-08-19T00:00:00", "employee_id": employee["id"], "attendance_status": "Absent"})

    resp = client.get("/api/attendance/period-summary", params={
        "employee_id": employee["id"], "start": "2026-08-18T00:00:00", "end": "2026-08-19T00:00:00",
    })
    body = resp.json()
    half_day, absent = body["days"]
    assert half_day["payable_hours"] == 4.0  # 0.5 * 8 scheduled hours
    assert half_day["shortfall_hours"] == 8.0  # no clock time recorded at all
    assert absent["payable_hours"] == 0.0
    assert absent["shortfall_hours"] == 8.0
    assert body["summary"]["attendance_percent"] == 25.0  # (0.5 + 0) / 2 marked records * 100


def test_period_summary_requires_own_or_master(client, test_user, db_session):
    _login(client, test_user)
    employee = _make_cc_employee(client, "3")
    other = _make_cc_employee(client, "3b")
    _login_as_cc_employee(client, db_session, employee["id"], "3")

    forbidden = client.get("/api/attendance/period-summary", params={
        "employee_id": other["id"], "start": "2026-08-17T00:00:00", "end": "2026-08-17T00:00:00",
    })
    assert forbidden.status_code == 403

    allowed = client.get("/api/attendance/period-summary", params={
        "employee_id": employee["id"], "start": "2026-08-17T00:00:00", "end": "2026-08-17T00:00:00",
    })
    assert allowed.status_code == 200


def test_day_detail_places_task_with_planned_time_on_timeline(client, test_user):
    _login(client, test_user)
    employee = _make_cc_employee(client, "4")
    client.post("/api/daily-tasks/", json={
        "date": "2026-08-17T00:00:00", "employee_id": employee["id"], "task_description": "Cut panels",
        "planned_start": "09:00:00", "planned_end": "11:00:00",
    })

    resp = client.get("/api/attendance/day-detail", params={"employee_id": employee["id"], "date": "2026-08-17T00:00:00"})
    assert resp.status_code == 200
    body = resp.json()
    assert len(body["timeline"]) == 1
    assert body["timeline"][0]["description"] == "Cut panels"
    assert body["timeline"][0]["start"] == "2026-08-17T09:00:00"
    assert body["timeline"][0]["end"] == "2026-08-17T11:00:00"
    assert body["unscheduled_work"] == []


def test_day_detail_lists_task_without_planned_time_as_unscheduled(client, test_user):
    _login(client, test_user)
    employee = _make_cc_employee(client, "5")
    client.post("/api/daily-tasks/", json={
        "date": "2026-08-17T00:00:00", "employee_id": employee["id"], "task_description": "General assembly work",
    })

    resp = client.get("/api/attendance/day-detail", params={"employee_id": employee["id"], "date": "2026-08-17T00:00:00"})
    body = resp.json()
    assert body["timeline"] == []
    assert len(body["unscheduled_work"]) == 1
    assert body["unscheduled_work"][0]["description"] == "General assembly work"


def test_team_grid_requires_master(client, test_user, db_session):
    _login(client, test_user)
    employee = _make_cc_employee(client, "6")
    _login_as_cc_employee(client, db_session, employee["id"], "6")

    resp = client.get("/api/attendance/team-grid", params={"start": "2026-08-17T00:00:00", "end": "2026-08-18T00:00:00"})
    assert resp.status_code == 403


def test_team_grid_covers_every_active_employee(client, test_user):
    _login(client, test_user)
    e1 = _make_cc_employee(client, "7a")
    e2 = _make_cc_employee(client, "7b")
    client.post("/api/attendance/", json={"date": "2026-08-17T00:00:00", "employee_id": e1["id"], "attendance_status": "Present"})
    client.post("/api/attendance/", json={"date": "2026-08-17T00:00:00", "employee_id": e2["id"], "attendance_status": "Absent"})

    resp = client.get("/api/attendance/team-grid", params={"start": "2026-08-17T00:00:00", "end": "2026-08-17T00:00:00"})
    assert resp.status_code == 200
    body = resp.json()
    ids = {row["employee_id"] for row in body["employees"]}
    assert e1["id"] in ids and e2["id"] in ids
    row1 = next(r for r in body["employees"] if r["employee_id"] == e1["id"])
    assert row1["days"][0]["event_type"] == "present"


def test_exceptions_flags_missing_attendance_on_a_working_day(client, test_user):
    _login(client, test_user)
    employee = _make_cc_employee(client, "8")
    # No attendance record created at all for this ordinary Monday.
    resp = client.get("/api/attendance/exceptions", params={
        "employee_id": employee["id"], "start": "2026-08-17T00:00:00", "end": "2026-08-17T00:00:00",
    })
    assert resp.status_code == 200
    types = {e["type"] for e in resp.json()["exceptions"]}
    assert "missing_attendance" in types


def test_exceptions_flags_missing_checkout(client, test_user):
    _login(client, test_user)
    employee = _make_cc_employee(client, "9")
    client.post("/api/attendance/", json={
        "date": "2026-08-17T00:00:00", "employee_id": employee["id"], "attendance_status": "Present",
        "in_time": "2026-08-17T09:00:00",
    })
    resp = client.get("/api/attendance/exceptions", params={
        "employee_id": employee["id"], "start": "2026-08-17T00:00:00", "end": "2026-08-17T00:00:00",
    })
    types = {e["type"] for e in resp.json()["exceptions"]}
    assert "missing_checkout" in types


def test_exceptions_employee_scoped_to_self(client, test_user, db_session):
    _login(client, test_user)
    employee = _make_cc_employee(client, "10")
    other = _make_cc_employee(client, "10b")
    _login_as_cc_employee(client, db_session, employee["id"], "10")

    forbidden = client.get("/api/attendance/exceptions", params={
        "employee_id": other["id"], "start": "2026-08-17T00:00:00", "end": "2026-08-17T00:00:00",
    })
    assert forbidden.status_code == 403

    own = client.get("/api/attendance/exceptions", params={"start": "2026-08-17T00:00:00", "end": "2026-08-17T00:00:00"})
    assert own.status_code == 200
    assert all(e["employee_id"] == employee["id"] for e in own.json()["exceptions"])


def test_attendance_list_supports_date_range_filter(client, test_user):
    _login(client, test_user)
    employee = _make_cc_employee(client, "11")
    client.post("/api/attendance/", json={"date": "2026-08-17T00:00:00", "employee_id": employee["id"], "attendance_status": "Present"})
    client.post("/api/attendance/", json={"date": "2026-08-20T00:00:00", "employee_id": employee["id"], "attendance_status": "Present"})
    client.post("/api/attendance/", json={"date": "2026-08-25T00:00:00", "employee_id": employee["id"], "attendance_status": "Present"})

    resp = client.get("/api/attendance/", params={
        "employee_id": employee["id"], "date_from": "2026-08-17T00:00:00", "date_to": "2026-08-21T00:00:00",
    })
    assert resp.status_code == 200
    dates = {a["date"][:10] for a in resp.json()}
    assert dates == {"2026-08-17", "2026-08-20"}


# --- Overtime request workflow (Draft -> Submitted -> Approved/Rejected) ---
def test_employee_can_create_and_submit_own_overtime_request(client, test_user, db_session):
    _login(client, test_user)
    employee = _make_cc_employee(client, "12")
    _login_as_cc_employee(client, db_session, employee["id"], "12")

    create = client.post("/api/overtime-requests/", json={
        "employee_id": employee["id"], "date": "2026-08-17T00:00:00", "requested_hours": "2.5", "reason": "Urgent order",
    })
    assert create.status_code == 201
    request = create.json()
    assert request["status"] == "Draft"

    submitted = client.post(f"/api/overtime-requests/{request['id']}/submit")
    assert submitted.status_code == 200
    assert submitted.json()["status"] == "Submitted"
    assert submitted.json()["submitted_at"] is not None


def test_employee_cannot_create_overtime_request_for_another_employee(client, test_user, db_session):
    _login(client, test_user)
    employee = _make_cc_employee(client, "13")
    other = _make_cc_employee(client, "13b")
    _login_as_cc_employee(client, db_session, employee["id"], "13")

    resp = client.post("/api/overtime-requests/", json={
        "employee_id": other["id"], "date": "2026-08-17T00:00:00", "requested_hours": "2",
    })
    assert resp.status_code == 403


def test_master_approving_overtime_request_applies_to_attendance(client, test_user):
    _login(client, test_user)
    employee = _make_cc_employee(client, "14")
    request = client.post("/api/overtime-requests/", json={
        "employee_id": employee["id"], "date": "2026-08-17T00:00:00", "requested_hours": "3",
    }).json()
    client.post(f"/api/overtime-requests/{request['id']}/submit")

    resp = client.put(f"/api/overtime-requests/{request['id']}/approve", json={})
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "Approved"
    assert float(body["approved_hours"]) == 3.0

    attendance = client.get("/api/attendance/", params={"employee_id": employee["id"], "date": "2026-08-17T00:00:00"}).json()
    assert float(attendance[0]["overtime_hours"]) == 3.0


def test_approving_overtime_request_is_additive_with_existing_overtime(client, test_user):
    """Spec: the approval path converges on the same apply_overtime_hours
    the bulk Master action uses, mode="add" - an approval adds to
    whatever overtime that date already carries rather than silently
    replacing a prior Master correction."""
    _login(client, test_user)
    employee = _make_cc_employee(client, "15")
    client.post("/api/attendance/overtime", json={"employee_id": employee["id"], "dates": ["2026-08-17T00:00:00"], "hours": "1.5"})

    request = client.post("/api/overtime-requests/", json={
        "employee_id": employee["id"], "date": "2026-08-17T00:00:00", "requested_hours": "2",
    }).json()
    client.post(f"/api/overtime-requests/{request['id']}/submit")
    client.put(f"/api/overtime-requests/{request['id']}/approve", json={})

    attendance = client.get("/api/attendance/", params={"employee_id": employee["id"], "date": "2026-08-17T00:00:00"}).json()
    assert float(attendance[0]["overtime_hours"]) == 3.5  # 1.5 existing + 2 approved


def test_master_can_approve_at_different_amount_than_requested(client, test_user):
    _login(client, test_user)
    employee = _make_cc_employee(client, "16")
    request = client.post("/api/overtime-requests/", json={
        "employee_id": employee["id"], "date": "2026-08-17T00:00:00", "requested_hours": "4",
    }).json()
    client.post(f"/api/overtime-requests/{request['id']}/submit")

    resp = client.put(f"/api/overtime-requests/{request['id']}/approve", json={"approved_hours": "2.5"})
    body = resp.json()
    assert float(body["approved_hours"]) == 2.5
    assert float(body["requested_hours"]) == 4.0  # unchanged, stays historical


def test_master_can_reject_overtime_request(client, test_user):
    _login(client, test_user)
    employee = _make_cc_employee(client, "17")
    request = client.post("/api/overtime-requests/", json={
        "employee_id": employee["id"], "date": "2026-08-17T00:00:00", "requested_hours": "2",
    }).json()
    client.post(f"/api/overtime-requests/{request['id']}/submit")

    resp = client.put(f"/api/overtime-requests/{request['id']}/reject", json={"rejection_reason": "Not required"})
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "Rejected"
    assert body["rejection_reason"] == "Not required"

    attendance = client.get("/api/attendance/", params={"employee_id": employee["id"], "date": "2026-08-17T00:00:00"}).json()
    assert attendance == []  # rejection never touches Attendance


def test_cannot_approve_a_draft_overtime_request(client, test_user):
    _login(client, test_user)
    employee = _make_cc_employee(client, "18")
    request = client.post("/api/overtime-requests/", json={
        "employee_id": employee["id"], "date": "2026-08-17T00:00:00", "requested_hours": "2",
    }).json()

    resp = client.put(f"/api/overtime-requests/{request['id']}/approve", json={})
    assert resp.status_code == 409


def test_cannot_edit_a_submitted_overtime_request(client, test_user):
    _login(client, test_user)
    employee = _make_cc_employee(client, "19")
    request = client.post("/api/overtime-requests/", json={
        "employee_id": employee["id"], "date": "2026-08-17T00:00:00", "requested_hours": "2",
    }).json()
    client.post(f"/api/overtime-requests/{request['id']}/submit")

    resp = client.put(f"/api/overtime-requests/{request['id']}", json={"requested_hours": "5"})
    assert resp.status_code == 409


def test_draft_overtime_request_can_be_cancelled(client, test_user):
    _login(client, test_user)
    employee = _make_cc_employee(client, "20")
    request = client.post("/api/overtime-requests/", json={
        "employee_id": employee["id"], "date": "2026-08-17T00:00:00", "requested_hours": "2",
    }).json()

    resp = client.delete(f"/api/overtime-requests/{request['id']}")
    assert resp.status_code == 204
    assert client.get("/api/overtime-requests/", params={"employee_id": employee["id"]}).json() == []


def test_cannot_cancel_a_submitted_overtime_request(client, test_user):
    _login(client, test_user)
    employee = _make_cc_employee(client, "21")
    request = client.post("/api/overtime-requests/", json={
        "employee_id": employee["id"], "date": "2026-08-17T00:00:00", "requested_hours": "2",
    }).json()
    client.post(f"/api/overtime-requests/{request['id']}/submit")

    resp = client.delete(f"/api/overtime-requests/{request['id']}")
    assert resp.status_code == 409


def test_employee_cannot_approve_own_overtime_request(client, test_user, db_session):
    _login(client, test_user)
    employee = _make_cc_employee(client, "22")
    request = client.post("/api/overtime-requests/", json={
        "employee_id": employee["id"], "date": "2026-08-17T00:00:00", "requested_hours": "2",
    }).json()
    client.post(f"/api/overtime-requests/{request['id']}/submit")
    _login_as_cc_employee(client, db_session, employee["id"], "22")

    resp = client.put(f"/api/overtime-requests/{request['id']}/approve", json={})
    assert resp.status_code == 403


def test_employee_cannot_view_another_employees_overtime_requests(client, test_user, db_session):
    _login(client, test_user)
    other = _make_cc_employee(client, "23a")
    employee = _make_cc_employee(client, "23b")
    client.post("/api/overtime-requests/", json={
        "employee_id": other["id"], "date": "2026-08-17T00:00:00", "requested_hours": "2",
    })
    _login_as_cc_employee(client, db_session, employee["id"], "23")

    resp = client.get("/api/overtime-requests/", params={"employee_id": other["id"]})
    assert resp.status_code == 403


def test_overtime_request_approval_notifies_employee(client, test_user, db_session):
    _login(client, test_user)
    employee = _make_cc_employee(client, "24")
    user = User(
        username="overtimenotifyuser", email="overtimenotifyuser@example.com", full_name="Overtime Notify User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()

    request = client.post("/api/overtime-requests/", json={
        "employee_id": employee["id"], "date": "2026-08-17T00:00:00", "requested_hours": "2",
    }).json()
    client.post(f"/api/overtime-requests/{request['id']}/submit")
    client.put(f"/api/overtime-requests/{request['id']}/approve", json={})

    client.post("/api/auth/login", json={"identifier": "overtimenotifyuser@example.com", "password": "EmpPass1!"})
    notifications = client.get("/api/notifications/").json()
    matches = [n for n in notifications if n["related_entity_type"] == "overtime_request" and n["related_entity_id"] == request["id"]]
    assert len(matches) == 1
    assert "approved" in matches[0]["title"].lower()


def test_overtime_request_rejection_notifies_employee_with_reason(client, test_user, db_session):
    _login(client, test_user)
    employee = _make_cc_employee(client, "25")
    user = User(
        username="overtimerejectnotifyuser", email="overtimerejectnotifyuser@example.com", full_name="Overtime Reject Notify User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()

    request = client.post("/api/overtime-requests/", json={
        "employee_id": employee["id"], "date": "2026-08-17T00:00:00", "requested_hours": "2",
    }).json()
    client.post(f"/api/overtime-requests/{request['id']}/submit")
    client.put(f"/api/overtime-requests/{request['id']}/reject", json={"rejection_reason": "Budget exceeded"})

    client.post("/api/auth/login", json={"identifier": "overtimerejectnotifyuser@example.com", "password": "EmpPass1!"})
    notifications = client.get("/api/notifications/").json()
    matches = [n for n in notifications if n["related_entity_type"] == "overtime_request" and n["related_entity_id"] == request["id"]]
    assert len(matches) == 1
    assert "rejected" in matches[0]["title"].lower()
    assert "Budget exceeded" in matches[0]["message"]


def test_overtime_request_hours_must_be_positive_and_at_most_24(client, test_user):
    _login(client, test_user)
    employee = _make_cc_employee(client, "26")
    too_much = client.post("/api/overtime-requests/", json={
        "employee_id": employee["id"], "date": "2026-08-17T00:00:00", "requested_hours": "25",
    })
    assert too_much.status_code == 422
    zero = client.post("/api/overtime-requests/", json={
        "employee_id": employee["id"], "date": "2026-08-17T00:00:00", "requested_hours": "0",
    })
    assert zero.status_code == 422
