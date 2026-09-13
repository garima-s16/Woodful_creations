"""Operations domain tests: production jobs, work centres,
cutting requirements, and general operations. Combines all former
test_*.py files under tests/modules/operations/."""
import io
from openpyxl import load_workbook
from tests.helpers import _login
from datetime import datetime, timedelta
from sqlalchemy import event
from app.platform.database import engine
from app.platform.security import hash_password
from app.modules.auth.auth import User


# --- test_production_jobs.py ---
"""Tests for production_jobs: the security gap in update_production_job
(every field was previously open to any authenticated user, unlike the
equivalent DailyTask endpoint), the new production-bottlenecks chatbot
handler, the stage/completion_date fields and production.xlsx export
(filters respected), and the staff dashboard's production status
summary (a real gap - it had zero production job data despite the
brief requiring "Production/job status where applicable")."""


def test_employee_can_update_status_and_completed_qty(client, test_user, db_session):
    """The genuinely-permitted self-service fields must still work for
    a non-master employee."""
    from app.platform.security import hash_password
    from app.modules.auth.auth import User
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Cutting"}).json()
    employee = client.post("/api/employees/", json={"name": "Production Update Self Service Employee"}).json()
    user = User(
        username="prodselfserviceuser", email="prodselfserviceuser@example.com",
        full_name="Prod Self Service User", password_hash=hash_password("EmpPass1!"),
        role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "prodselfserviceuser@example.com", "password": "EmpPass1!"})

    resp = client.put(f"/api/production-jobs/{job['id']}", json={"status": "In Progress", "completed_qty": 5})
    assert resp.status_code == 200


def test_employee_cannot_change_stage_or_remarks(client, test_user, db_session):
    """The real fix - planning fields must be blocked for non-master,
    which was not previously enforced at all."""
    from app.platform.security import hash_password
    from app.modules.auth.auth import User
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Assembly"}).json()
    employee = client.post("/api/employees/", json={"name": "Production Update Restricted Employee"}).json()
    user = User(
        username="prodrestricteduser", email="prodrestricteduser@example.com",
        full_name="Prod Restricted User", password_hash=hash_password("EmpPass1!"),
        role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "prodrestricteduser@example.com", "password": "EmpPass1!"})

    resp = client.put(f"/api/production-jobs/{job['id']}", json={"stage": "QC"})
    assert resp.status_code == 403

    resp2 = client.put(f"/api/production-jobs/{job['id']}", json={"remarks": "trying to sneak this in"})
    assert resp2.status_code == 403


def test_employee_can_set_blocker_reason(client, test_user, db_session):
    from app.platform.security import hash_password
    from app.modules.auth.auth import User
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Finishing"}).json()
    employee = client.post("/api/employees/", json={"name": "Production Blocker Reason Employee"}).json()
    user = User(
        username="prodblockeruser", email="prodblockeruser@example.com",
        full_name="Prod Blocker User", password_hash=hash_password("EmpPass1!"),
        role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "prodblockeruser@example.com", "password": "EmpPass1!"})

    resp = client.put(f"/api/production-jobs/{job['id']}", json={
        "status": "Blocked", "blocker_reason": "Waiting on CNC router availability",
    })
    assert resp.status_code == 200
    assert resp.json()["blocker_reason"] == "Waiting on CNC router availability"


def test_master_can_still_update_any_field(client, test_user):
    """Regression guard - master must remain unrestricted."""
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Packing"}).json()
    resp = client.put(f"/api/production-jobs/{job['id']}", json={"stage": "Packing", "remarks": "master note"})
    assert resp.status_code == 200


def test_chatbot_production_bottlenecks(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Blocked bottleneck job", "machine": "CNC Router",
    }).json()
    client.put(f"/api/production-jobs/{job['id']}", json={"status": "Blocked", "blocker_reason": "No material"})

    resp = client.post("/api/chat/", json={"message": "show production bottlenecks"})
    assert resp.status_code == 200
    data = resp.json()
    assert "1 job(s) blocked" in data["response"]
    assert any(r["label"] == job["job_code"] for r in data["records"])


def test_chatbot_no_bottlenecks_reports_honestly(client, test_user):
    _login(client, test_user)
    resp = client.post("/api/chat/", json={"message": "any bottleneck in production"})
    assert resp.status_code == 200


def test_production_stage_can_be_set_on_create(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Panel cutting", "stage": "Cutting",
    }).json()
    assert job["stage"] == "Cutting"


def test_completion_date_set_on_transition_to_completed(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Assembly"}).json()
    assert job["completion_date"] is None

    resp = client.put(f"/api/production-jobs/{job['id']}", json={"status": "Completed"})
    assert resp.status_code == 200
    assert resp.json()["completion_date"] is not None


def test_completion_date_does_not_change_on_resave(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Finishing"}).json()
    client.put(f"/api/production-jobs/{job['id']}", json={"status": "Completed"})
    first = client.get("/api/production-jobs/").json()
    first_match = next(j for j in first if j["id"] == job["id"])

    client.put(f"/api/production-jobs/{job['id']}", json={"status": "Completed", "remarks": "verified"})
    second = client.get("/api/production-jobs/").json()
    second_match = next(j for j in second if j["id"] == job["id"])
    assert second_match["completion_date"] == first_match["completion_date"]


def test_production_export_respects_stage_filter(client, test_user):
    _login(client, test_user)
    client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Cutting op", "stage": "Cutting",
    })
    client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "QC op", "stage": "QC",
    })

    resp = client.get("/api/reports/production.xlsx", params={"stage": "Cutting"})
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["Production"]
    operations = [cell.value for row in ws.iter_rows() for cell in row if isinstance(cell.value, str)]
    assert any("Cutting op" in o for o in operations)
    assert not any("QC op" in o for o in operations)


def test_production_export_respects_operator_filter(client, test_user):
    _login(client, test_user)
    emp_a = client.post("/api/employees/", json={"name": "Production Export Operator A"}).json()
    emp_b = client.post("/api/employees/", json={"name": "Production Export Operator B"}).json()
    client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Job for A", "employee_id": emp_a["id"],
    })
    client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Job for B", "employee_id": emp_b["id"],
    })

    resp = client.get("/api/reports/production.xlsx", params={"employee_id": emp_a["id"]})
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["Production"]
    operations = [cell.value for row in ws.iter_rows() for cell in row if isinstance(cell.value, str)]
    assert any("Job for A" in o for o in operations)
    assert not any("Job for B" in o for o in operations)


def test_staff_dashboard_includes_production_status_summary(client, test_user):
    _login(client, test_user)
    order_client = client.post("/api/clients/", json={"name": "Staff Dashboard Production Client", "phone": "9000010184"}).json()
    order = client.post("/api/orders/", json={
        "client_id": order_client["id"], "order_date": "2026-08-18T00:00:00", "order_value": "50000", "advance": "0",
    }).json()
    employee = client.post("/api/employees/", json={"name": "Staff Dashboard Production Employee", "monthly_salary": "20000"}).json()
    client.post("/api/production-jobs/", json={
        "date": "2026-08-18T00:00:00", "order_id": order["id"], "employee_id": employee["id"], "operation": "Cutting",
        "planned_qty": 10, "status": "In Progress",
    })

    resp = client.get("/api/dashboard/staff").json()
    assert "production_status_summary" in resp
    match = next((p for p in resp["production_status_summary"] if p["status"] == "In Progress"), None)
    assert match is not None
    assert match["count"] >= 1


def test_staff_dashboard_production_summary_reflects_real_counts_not_hardcoded(client, test_user):
    """Adding a second job in a different status must change the
    summary - proving it's derived from real data, not a fixed value."""
    _login(client, test_user)
    before = client.get("/api/dashboard/staff").json()["production_status_summary"]

    order_client = client.post("/api/clients/", json={"name": "Staff Dashboard Production Client 2", "phone": "9000010185"}).json()
    order = client.post("/api/orders/", json={
        "client_id": order_client["id"], "order_date": "2026-08-18T00:00:00", "order_value": "30000", "advance": "0",
    }).json()
    employee = client.post("/api/employees/", json={"name": "Staff Dashboard Production Employee 2", "monthly_salary": "18000"}).json()
    client.post("/api/production-jobs/", json={
        "date": "2026-08-18T00:00:00", "order_id": order["id"], "employee_id": employee["id"], "operation": "Assembly",
        "planned_qty": 5, "status": "Pending",
    })

    after = client.get("/api/dashboard/staff").json()["production_status_summary"]
    assert after != before


"""Tests for ProductionOperation (P0.3.3/3.4) - a single manufacturing
step within a ProductionJob, with a deliberately simple, single-
predecessor dependency (not a general workflow engine). Focused on the
real business rule: a dependent operation cannot be started/completed
while its declared predecessor is still incomplete."""


def _make_job_prod(client):
    return client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "General"}).json()


def test_operation_with_incomplete_dependency_cannot_be_started(client, test_user):
    _login(client, test_user)
    job = _make_job_prod(client)
    cutting = client.post("/api/production-operations/", json={
        "production_job_id": job["id"], "sequence": 1, "operation_name": "Cutting",
    }).json()
    cnc = client.post("/api/production-operations/", json={
        "production_job_id": job["id"], "sequence": 2, "operation_name": "CNC",
        "depends_on_operation_id": cutting["id"],
    }).json()

    resp = client.put(f"/api/production-operations/{cnc['id']}", json={"status": "In Progress"})
    assert resp.status_code == 409
    assert "Cutting" in resp.json()["detail"]


def test_operation_becomes_unblocked_once_dependency_completes(client, test_user):
    _login(client, test_user)
    job = _make_job_prod(client)
    cutting = client.post("/api/production-operations/", json={
        "production_job_id": job["id"], "sequence": 1, "operation_name": "Cutting",
    }).json()
    cnc = client.post("/api/production-operations/", json={
        "production_job_id": job["id"], "sequence": 2, "operation_name": "CNC",
        "depends_on_operation_id": cutting["id"],
    }).json()

    client.put(f"/api/production-operations/{cutting['id']}", json={"status": "Completed"})

    resp = client.put(f"/api/production-operations/{cnc['id']}", json={"status": "In Progress"})
    assert resp.status_code == 200
    assert resp.json()["status"] == "In Progress"


def test_operation_list_reports_blocked_flag_without_starting_it(client, test_user):
    _login(client, test_user)
    job = _make_job_prod(client)
    cutting = client.post("/api/production-operations/", json={
        "production_job_id": job["id"], "sequence": 1, "operation_name": "Cutting",
    }).json()
    client.post("/api/production-operations/", json={
        "production_job_id": job["id"], "sequence": 2, "operation_name": "CNC",
        "depends_on_operation_id": cutting["id"],
    })

    resp = client.get("/api/production-operations/", params={"production_job_id": job["id"]})
    assert resp.status_code == 200
    rows = {r["operation_name"]: r for r in resp.json()}
    assert rows["Cutting"]["is_blocked_by_dependency"] is False
    assert rows["CNC"]["is_blocked_by_dependency"] is True


def test_operation_without_dependency_is_never_blocked(client, test_user):
    _login(client, test_user)
    job = _make_job_prod(client)
    op = client.post("/api/production-operations/", json={
        "production_job_id": job["id"], "sequence": 1, "operation_name": "Finishing",
    }).json()

    resp = client.put(f"/api/production-operations/{op['id']}", json={"status": "Completed"})
    assert resp.status_code == 200


def test_dependency_must_belong_to_the_same_production_job(client, test_user):
    _login(client, test_user)
    job_a = _make_job_prod(client)
    job_b = _make_job_prod(client)
    op_a = client.post("/api/production-operations/", json={
        "production_job_id": job_a["id"], "sequence": 1, "operation_name": "Cutting A",
    }).json()

    resp = client.post("/api/production-operations/", json={
        "production_job_id": job_b["id"], "sequence": 1, "operation_name": "CNC B",
        "depends_on_operation_id": op_a["id"],
    })
    assert resp.status_code == 400


def test_production_operations_require_auth(client):
    resp = client.get("/api/production-operations/")
    assert resp.status_code == 401


def test_employee_cannot_create_operation(client, test_user, db_session):
    from app.platform.security import hash_password
    from app.modules.auth.auth import User
    _login(client, test_user)
    job = _make_job_prod(client)
    employee = User(
        username="prodopempuser", email="prodopempuser@example.com", full_name="Prod Op Employee",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=None, is_active=True,
    )
    db_session.add(employee)
    db_session.commit()
    resp = client.post("/api/auth/login", json={"identifier": "prodopempuser@example.com", "password": "EmpPass1!"})
    assert resp.status_code == 200

    resp = client.post("/api/production-operations/", json={
        "production_job_id": job["id"], "sequence": 1, "operation_name": "Unauthorized Op",
    })
    assert resp.status_code == 403


"""Tests for GET /api/production-jobs/{id}/variance - planned vs actual
(P0.3 section 27/28). Never invents an actual value: duration variance
is computed only from operations that actually have actual_duration_minutes
recorded, and duration_complete distinguishes a partial variance from
the final one."""


def test_quantity_variance_reflects_real_planned_vs_completed(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Cutting", "planned_qty": 100, "completed_qty": 90,
    }).json()

    resp = client.get(f"/api/production-jobs/{job['id']}/variance")
    assert resp.status_code == 200
    body = resp.json()
    assert body["quantity_planned"] == 100
    assert body["quantity_completed"] == 90
    assert body["quantity_variance"] == -10
    assert body["quantity_variance_percent"] == -10.0


def test_duration_variance_none_when_no_actuals_recorded_yet(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Cutting"}).json()
    client.post("/api/production-operations/", json={
        "production_job_id": job["id"], "sequence": 1, "operation_name": "Cut Sheets",
        "estimated_duration_minutes": 120,
    })

    resp = client.get(f"/api/production-jobs/{job['id']}/variance")
    body = resp.json()
    assert body["duration_planned_minutes"] == 120
    assert body["duration_actual_minutes"] is None
    assert body["duration_variance_minutes"] is None
    assert body["duration_complete"] is False


def test_duration_variance_computed_once_actual_is_recorded(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Cutting"}).json()
    operation = client.post("/api/production-operations/", json={
        "production_job_id": job["id"], "sequence": 1, "operation_name": "Cut Sheets",
        "estimated_duration_minutes": 120,
    }).json()
    client.put(f"/api/production-operations/{operation['id']}", json={"actual_duration_minutes": 150})

    resp = client.get(f"/api/production-jobs/{job['id']}/variance")
    body = resp.json()
    assert body["duration_planned_minutes"] == 120
    assert body["duration_actual_minutes"] == 150
    assert body["duration_variance_minutes"] == 30
    assert body["duration_complete"] is True
    assert body["operations_total"] == 1
    assert body["operations_with_actual_duration"] == 1


def test_duration_incomplete_when_only_some_operations_have_actuals(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Cutting"}).json()
    op1 = client.post("/api/production-operations/", json={
        "production_job_id": job["id"], "sequence": 1, "operation_name": "Cut Sheets",
        "estimated_duration_minutes": 60,
    }).json()
    client.post("/api/production-operations/", json={
        "production_job_id": job["id"], "sequence": 2, "operation_name": "Assemble",
        "estimated_duration_minutes": 90,
    })
    client.put(f"/api/production-operations/{op1['id']}", json={"actual_duration_minutes": 70})

    resp = client.get(f"/api/production-jobs/{job['id']}/variance")
    body = resp.json()
    assert body["operations_total"] == 2
    assert body["operations_with_actual_duration"] == 1
    assert body["duration_complete"] is False
    # Only the operation with a real actual counts toward the variance -
    # not the still-unrecorded second operation.
    assert body["duration_actual_minutes"] == 70
    assert body["duration_variance_minutes"] == 10


def test_variance_requires_auth(client):
    resp = client.get("/api/production-jobs/1/variance")
    assert resp.status_code == 401


def test_variance_unknown_job_returns_404(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/production-jobs/999999/variance")
    assert resp.status_code == 404


"""Production job material-status (GET
/api/production-jobs/{id}/material-status) - Phase D dependency-aware
planning: is this job blocked by a real material shortage? Reuses
StockService.calculate_order_material_requirements for the job's own
order; no separate shortage calculation exists here."""


def test_job_material_status_reports_real_shortage(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Job Shortage Test Sheet", "unit": "Sheets", "opening_stock": "2", "minimum_stock": "1",
    }).json()
    product = client.post("/api/products/", json={
        "name": "Job Shortage Test Product", "unit": "Piece",
        "materials_used": [{"material_id": material["id"], "quantity_required": "5"}],
    }).json()
    client_id = client.post("/api/clients/", json={"name": "Job Shortage Client", "phone": "9000010199"}).json()["id"]
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()
    job = client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Cutting",
        "order_id": order["id"], "material_id": material["id"],
    }).json()

    resp = client.get(f"/api/production-jobs/{job['id']}/material-status")
    assert resp.status_code == 200
    body = resp.json()
    assert body["has_shortage_data"] is True
    assert body["is_blocked_by_shortage"] is True
    assert float(body["shortage"]) == 3.0  # 5 required - 2 available


def test_job_material_status_no_shortage_when_stock_covers_it(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Job No Shortage Test Sheet", "unit": "Sheets", "opening_stock": "20", "minimum_stock": "1",
    }).json()
    product = client.post("/api/products/", json={
        "name": "Job No Shortage Test Product", "unit": "Piece",
        "materials_used": [{"material_id": material["id"], "quantity_required": "5"}],
    }).json()
    client_id = client.post("/api/clients/", json={"name": "Job No Shortage Client", "phone": "9000010198"}).json()["id"]
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()
    job = client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Cutting",
        "order_id": order["id"], "material_id": material["id"],
    }).json()

    resp = client.get(f"/api/production-jobs/{job['id']}/material-status")
    body = resp.json()
    assert body["has_shortage_data"] is True
    assert body["is_blocked_by_shortage"] is False


def test_job_without_order_or_material_has_no_shortage_data(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Assembly"}).json()

    resp = client.get(f"/api/production-jobs/{job['id']}/material-status")
    assert resp.status_code == 200
    body = resp.json()
    assert body["has_shortage_data"] is False
    assert body["is_blocked_by_shortage"] is False


def test_job_material_status_requires_auth(client):
    resp = client.get("/api/production-jobs/1/material-status")
    assert resp.status_code == 401


def test_job_material_status_unknown_job_returns_404(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/production-jobs/999999/material-status")
    assert resp.status_code == 404


def test_job_material_status_includes_supplier_options_when_blocked(client, test_user):
    """Family 130 section 8: a blocked job's material-status must show
    which supplier can actually resolve the shortage, not just the raw
    numbers - preferred supplier first regardless of price."""
    _login(client, test_user)
    material = client.post("/api/materials/", json={
        "name": "Job Shortage Supplier Test Sheet", "unit": "Sheets", "opening_stock": "2", "minimum_stock": "1",
    }).json()
    cheap_supplier = client.post("/api/suppliers/", json={"name": "Cheap Sheet Supplier"}).json()
    preferred_supplier = client.post("/api/suppliers/", json={"name": "Preferred Sheet Supplier"}).json()
    client.post("/api/supplier-materials/", json={
        "supplier_id": cheap_supplier["id"], "material_id": material["id"],
        "supplier_price": "400.00", "lead_time_days": 5, "is_preferred": False,
    })
    client.post("/api/supplier-materials/", json={
        "supplier_id": preferred_supplier["id"], "material_id": material["id"],
        "supplier_price": "550.00", "lead_time_days": 2, "is_preferred": True,
    })
    product = client.post("/api/products/", json={
        "name": "Job Shortage Supplier Test Product", "unit": "Piece",
        "materials_used": [{"material_id": material["id"], "quantity_required": "5"}],
    }).json()
    client_id = client.post("/api/clients/", json={"name": "Job Shortage Supplier Client", "phone": "9000010200"}).json()["id"]
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()
    job = client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Cutting",
        "order_id": order["id"], "material_id": material["id"],
    }).json()

    resp = client.get(f"/api/production-jobs/{job['id']}/material-status")
    assert resp.status_code == 200
    body = resp.json()
    assert body["is_blocked_by_shortage"] is True
    assert float(body["recommended_purchase_quantity"]) == 3.0
    options = body["supplier_options"]
    assert len(options) == 2
    # Preferred supplier ranks first even though it is not the cheapest.
    assert options[0]["supplier_id"] == preferred_supplier["id"]
    assert options[0]["is_preferred"] is True
    assert options[0]["lead_time_days"] == 2
    assert options[1]["supplier_id"] == cheap_supplier["id"]
    assert options[1]["price"] == 400.0


"""Tests for GET /api/production-jobs/{id}/risks - Production Risk
(P0.3 section 29). Always a list of individually-explained findings,
never a single opaque score; an empty list is a real "nothing found
from checkable data", not a claim of guaranteed safety."""


def test_no_risks_when_everything_is_fine(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Assembly"}).json()

    resp = client.get(f"/api/production-jobs/{job['id']}/risks")
    assert resp.status_code == 200
    assert resp.json()["risks"] == []


def test_completed_job_has_no_risks_regardless_of_material(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={"name": "Risk Test Sheet A", "unit": "Sheets", "opening_stock": "0"}).json()
    product = client.post("/api/products/", json={
        "name": "Risk Test Product A", "unit": "Piece",
        "materials_used": [{"material_id": material["id"], "quantity_required": "5"}],
    }).json()
    client_id = client.post("/api/clients/", json={"name": "Risk Test Client A", "phone": "9000010801"}).json()["id"]
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()
    job = client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Cutting", "order_id": order["id"], "material_id": material["id"],
    }).json()
    client.put(f"/api/production-jobs/{job['id']}", json={"status": "Completed"})

    resp = client.get(f"/api/production-jobs/{job['id']}/risks")
    assert resp.json()["risks"] == []


def test_material_shortage_produces_a_high_severity_risk(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={"name": "Risk Test Sheet B", "unit": "Sheets", "opening_stock": "0"}).json()
    product = client.post("/api/products/", json={
        "name": "Risk Test Product B", "unit": "Piece",
        "materials_used": [{"material_id": material["id"], "quantity_required": "5"}],
    }).json()
    client_id = client.post("/api/clients/", json={"name": "Risk Test Client B", "phone": "9000010802"}).json()["id"]
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()
    job = client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Cutting", "order_id": order["id"], "material_id": material["id"],
    }).json()

    resp = client.get(f"/api/production-jobs/{job['id']}/risks")
    risks = resp.json()["risks"]
    shortage_risks = [r for r in risks if r["type"] == "material_shortage"]
    assert len(shortage_risks) == 1
    assert shortage_risks[0]["severity"] == "high"
    assert "Risk Test Sheet B" in shortage_risks[0]["why"]


def test_dependency_blockage_produces_a_risk_per_blocked_operation(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Assembly"}).json()
    cutting = client.post("/api/production-operations/", json={
        "production_job_id": job["id"], "sequence": 1, "operation_name": "Risk Test Cutting",
    }).json()
    client.post("/api/production-operations/", json={
        "production_job_id": job["id"], "sequence": 2, "operation_name": "Risk Test CNC",
        "depends_on_operation_id": cutting["id"],
    })

    resp = client.get(f"/api/production-jobs/{job['id']}/risks")
    risks = resp.json()["risks"]
    dep_risks = [r for r in risks if r["type"] == "dependency_blockage"]
    assert len(dep_risks) == 1
    assert "Risk Test CNC" in dep_risks[0]["what"]
    assert "Risk Test Cutting" in dep_risks[0]["why"]


def test_capacity_overload_produces_a_risk(client, test_user):
    _login(client, test_user)
    centre = client.post("/api/work-centres/", json={"name": "Risk Test CNC Centre", "capacity_hours_per_day": "8"}).json()
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Cutting"}).json()
    client.post("/api/production-operations/", json={
        "production_job_id": job["id"], "sequence": 1, "operation_name": "Risk Test Overload Op",
        "work_centre_id": centre["id"], "estimated_duration_minutes": 600, "start_time": "2026-08-19T09:00:00",
    })

    resp = client.get(f"/api/production-jobs/{job['id']}/risks")
    risks = resp.json()["risks"]
    capacity_risks = [r for r in risks if r["type"] == "capacity_overload"]
    assert len(capacity_risks) == 1
    assert "Risk Test CNC Centre" in capacity_risks[0]["what"]
    assert capacity_risks[0]["when"] == "2026-08-19"


def test_schedule_risk_when_delivery_date_already_passed(client, test_user):
    _login(client, test_user)
    client_id = client.post("/api/clients/", json={"name": "Risk Test Client C", "phone": "9000010803"}).json()["id"]
    past_date = (datetime.utcnow() - timedelta(days=5)).isoformat()
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-01-01T00:00:00", "delivery_date": past_date,
    }).json()
    job = client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Assembly", "order_id": order["id"],
    }).json()

    resp = client.get(f"/api/production-jobs/{job['id']}/risks")
    risks = resp.json()["risks"]
    schedule_risks = [r for r in risks if r["type"] == "schedule_risk"]
    assert len(schedule_risks) == 1
    assert schedule_risks[0]["severity"] == "high"
    assert "already passed" in schedule_risks[0]["what"]


def test_schedule_risk_when_delivery_date_imminent(client, test_user):
    _login(client, test_user)
    client_id = client.post("/api/clients/", json={"name": "Risk Test Client D", "phone": "9000010804"}).json()["id"]
    soon_date = (datetime.utcnow() + timedelta(days=2)).isoformat()
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-01-01T00:00:00", "delivery_date": soon_date,
    }).json()
    job = client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Assembly", "order_id": order["id"],
    }).json()

    resp = client.get(f"/api/production-jobs/{job['id']}/risks")
    risks = resp.json()["risks"]
    schedule_risks = [r for r in risks if r["type"] == "schedule_risk"]
    assert len(schedule_risks) == 1
    assert schedule_risks[0]["severity"] == "medium"
    assert "imminent" in schedule_risks[0]["what"]


def test_risks_requires_auth(client):
    resp = client.get("/api/production-jobs/1/risks")
    assert resp.status_code == 401


def test_risks_unknown_job_returns_404(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/production-jobs/999999/risks")
    assert resp.status_code == 404


"""Tests for the Production Readiness Engine
(GET /api/production-jobs/{id}/readiness) - P0.3 section 2.
READY/PARTIALLY_READY/BLOCKED, reusing
StockService.calculate_order_material_requirements exactly (the same
calculation the material-status endpoint, dashboard, daily-tasks list
and AI chatbot all already use) - no separate, independently-drifting
shortage logic here."""


def _make_job_against_order(client, material_id, product_qty_required, opening_stock, suffix, pending_qty=None, supplier_name=None):
    material_kwargs = {
        "name": f"Readiness Sheet {suffix}", "unit": "Sheets", "opening_stock": str(opening_stock), "minimum_stock": "1",
    }
    if material_id is None:
        material = client.post("/api/materials/", json=material_kwargs).json()
        material_id = material["id"]
    if pending_qty is not None:
        supplier = client.post("/api/suppliers/", json={"name": supplier_name or f"Readiness Supplier {suffix}"}).json()
        client.post("/api/purchases/", json={
            "date": "2026-08-01T00:00:00", "supplier_id": supplier["id"], "material_id": material_id,
            "quantity": str(pending_qty), "unit": "Sheets", "rate": "500.00", "gst_percent": "18",
            "receipt_status": "Ordered",
        })
    product = client.post("/api/products/", json={
        "name": f"Readiness Product {suffix}", "unit": "Piece",
        "materials_used": [{"material_id": material_id, "quantity_required": str(product_qty_required)}],
    }).json()
    client_id = client.post("/api/clients/", json={"name": f"Readiness Client {suffix}", "phone": f"900001070{suffix}"}).json()["id"]
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()
    job = client.post("/api/production-jobs/", json={
        "date": "2026-08-19T00:00:00", "operation": "Cutting", "order_id": order["id"], "material_id": material_id,
    }).json()
    return job, material_id


def test_readiness_is_ready_when_material_is_fully_available(client, test_user):
    _login(client, test_user)
    job, _ = _make_job_against_order(client, None, product_qty_required=2, opening_stock=100, suffix="1")

    resp = client.get(f"/api/production-jobs/{job['id']}/readiness")
    assert resp.status_code == 200
    body = resp.json()
    assert body["readiness"] == "READY"
    assert "fully available" in body["reason"]


def test_readiness_is_blocked_when_shortage_has_no_pending_coverage(client, test_user):
    _login(client, test_user)
    job, _ = _make_job_against_order(client, None, product_qty_required=5, opening_stock=2, suffix="2")

    resp = client.get(f"/api/production-jobs/{job['id']}/readiness")
    assert resp.status_code == 200
    body = resp.json()
    assert body["readiness"] == "BLOCKED"
    assert "Readiness Sheet 2" in body["reason"]


def test_readiness_is_partially_ready_when_pending_purchase_covers_the_gap(client, test_user):
    """The key distinction this engine must get right: shortage is
    computationally 0 once a covering purchase is on order, but the
    material is not physically in stock yet - the job cannot actually
    start, so this must be PARTIALLY_READY, not READY."""
    _login(client, test_user)
    job, _ = _make_job_against_order(
        client, None, product_qty_required=5, opening_stock=2, suffix="3",
        pending_qty=3, supplier_name="Readiness Covering Supplier",
    )

    resp = client.get(f"/api/production-jobs/{job['id']}/readiness")
    assert resp.status_code == 200
    body = resp.json()
    assert body["readiness"] == "PARTIALLY_READY"
    assert "already on order" in body["reason"]


def test_readiness_blocked_by_employee_reported_blocker_overrides_material_check(client, test_user):
    """A job the employee has explicitly flagged as blocked must report
    that reason even if the material itself is fully available - the
    blocker is real, reported information, and must not be silently
    overridden by an unrelated material check."""
    _login(client, test_user)
    job, _ = _make_job_against_order(client, None, product_qty_required=2, opening_stock=100, suffix="4")
    client.put(f"/api/production-jobs/{job['id']}", json={"blocker_reason": "CNC machine is down for maintenance"})

    resp = client.get(f"/api/production-jobs/{job['id']}/readiness")
    assert resp.status_code == 200
    body = resp.json()
    assert body["readiness"] == "BLOCKED"
    assert body["reason"] == "CNC machine is down for maintenance"


def test_readiness_is_ready_for_a_completed_job_regardless_of_material(client, test_user):
    _login(client, test_user)
    job, _ = _make_job_against_order(client, None, product_qty_required=5, opening_stock=0, suffix="5")
    client.put(f"/api/production-jobs/{job['id']}", json={"status": "Completed"})

    resp = client.get(f"/api/production-jobs/{job['id']}/readiness")
    assert resp.status_code == 200
    assert resp.json()["readiness"] == "READY"


def test_readiness_is_ready_when_job_has_no_order_or_material_link(client, test_user):
    _login(client, test_user)
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Assembly"}).json()

    resp = client.get(f"/api/production-jobs/{job['id']}/readiness")
    assert resp.status_code == 200
    body = resp.json()
    assert body["readiness"] == "READY"
    assert "no linked order/material" in body["reason"]


def test_readiness_requires_auth(client):
    resp = client.get("/api/production-jobs/1/readiness")
    assert resp.status_code == 401


def test_readiness_unknown_job_returns_404(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/production-jobs/999999/readiness")
    assert resp.status_code == 404


# --- test_work_centres.py ---
def test_create_and_list_work_centre(client, test_user):
    _login(client, test_user)
    resp = client.post("/api/work-centres/", json={"name": "Test CNC", "type": "CNC", "capacity_hours_per_day": "8"})
    assert resp.status_code == 201
    centre = resp.json()
    assert float(centre["capacity_hours_per_day"]) == 8.0

    resp = client.get("/api/work-centres/")
    assert resp.status_code == 200
    assert any(c["name"] == "Test CNC" for c in resp.json())


def test_duplicate_work_centre_name_rejected(client, test_user):
    _login(client, test_user)
    client.post("/api/work-centres/", json={"name": "Duplicate CNC"})
    resp = client.post("/api/work-centres/", json={"name": "Duplicate CNC"})
    assert resp.status_code == 409


def test_capacity_conflict_detected_when_scheduled_work_exceeds_capacity(client, test_user):
    _login(client, test_user)
    centre = client.post("/api/work-centres/", json={"name": "Capacity CNC", "capacity_hours_per_day": "8"}).json()
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Cutting"}).json()
    client.post("/api/production-operations/", json={
        "production_job_id": job["id"], "sequence": 1, "operation_name": "Big Cut",
        "work_centre_id": centre["id"], "estimated_duration_minutes": 600,
        "start_time": "2026-08-19T09:00:00",
    })

    resp = client.get(f"/api/work-centres/{centre['id']}/capacity", params={"date": "2026-08-19"})
    assert resp.status_code == 200
    body = resp.json()
    assert body["has_capacity_data"] is True
    assert body["scheduled_minutes"] == 600
    # 8 hours capacity = 480 minutes; 600 scheduled minutes exceeds it.
    assert body["status"] == "CAPACITY_CONFLICT"
    assert body["remaining_minutes"] == -120


def test_capacity_is_ok_when_scheduled_work_fits(client, test_user):
    _login(client, test_user)
    centre = client.post("/api/work-centres/", json={"name": "Fits CNC", "capacity_hours_per_day": "8"}).json()
    job = client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Cutting"}).json()
    client.post("/api/production-operations/", json={
        "production_job_id": job["id"], "sequence": 1, "operation_name": "Small Cut",
        "work_centre_id": centre["id"], "estimated_duration_minutes": 240,
        "start_time": "2026-08-19T09:00:00",
    })

    resp = client.get(f"/api/work-centres/{centre['id']}/capacity", params={"date": "2026-08-19"})
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "OK"
    assert body["remaining_minutes"] == 240


def test_capacity_reports_no_data_when_unconfigured(client, test_user):
    _login(client, test_user)
    centre = client.post("/api/work-centres/", json={"name": "Unconfigured CNC"}).json()
    resp = client.get(f"/api/work-centres/{centre['id']}/capacity")
    assert resp.status_code == 200
    assert resp.json()["has_capacity_data"] is False


def test_work_centres_require_auth(client):
    resp = client.get("/api/work-centres/")
    assert resp.status_code == 401


# --- test_cutting_requirements.py ---
def _make_job_cutting(client):
    return client.post("/api/production-jobs/", json={"date": "2026-08-19T00:00:00", "operation": "Cutting"}).json()


def test_create_and_get_cutting_requirement(client, test_user):
    _login(client, test_user)
    job = _make_job_cutting(client)
    material = client.post("/api/materials/", json={"name": "Cutting Req Sheet", "unit": "Sheets", "opening_stock": "10"}).json()

    resp = client.post("/api/cutting-requirements/", json={
        "production_job_id": job["id"], "material_id": material["id"], "part_name": "Side Panel",
        "quantity": 4, "length_mm": "600", "width_mm": "400", "thickness_mm": "18",
        "grain_direction": "Length", "rotation_allowed": False, "kerf_mm": "3.2",
    })
    assert resp.status_code == 201
    body = resp.json()
    assert body["part_name"] == "Side Panel"
    assert body["material_name"] == "Cutting Req Sheet"
    assert body["rotation_allowed"] is False
    assert float(body["length_mm"]) == 600.0

    resp = client.get(f"/api/cutting-requirements/{body['id']}")
    assert resp.status_code == 200
    assert resp.json()["part_name"] == "Side Panel"


def test_creating_cutting_requirement_never_touches_stock(client, test_user):
    """The explicit brief requirement: this is a plan, not a
    consumption event."""
    _login(client, test_user)
    job = _make_job_cutting(client)
    material = client.post("/api/materials/", json={"name": "Cutting Req Stock Sheet", "unit": "Sheets", "opening_stock": "10"}).json()

    client.post("/api/cutting-requirements/", json={
        "production_job_id": job["id"], "material_id": material["id"], "part_name": "Top",
        "quantity": 2, "length_mm": "500", "width_mm": "300",
    })

    unchanged = client.get(f"/api/materials/{material['id']}").json()
    assert float(unchanged["current_stock"]) == 10.0


def test_dimensions_must_be_positive(client, test_user):
    _login(client, test_user)
    job = _make_job_cutting(client)
    material = client.post("/api/materials/", json={"name": "Cutting Req Bad Dim Sheet", "unit": "Sheets", "opening_stock": "10"}).json()

    resp = client.post("/api/cutting-requirements/", json={
        "production_job_id": job["id"], "material_id": material["id"], "part_name": "Bad Part",
        "length_mm": "0", "width_mm": "300",
    })
    assert resp.status_code == 422


def test_unknown_production_job_rejected(client, test_user):
    _login(client, test_user)
    material = client.post("/api/materials/", json={"name": "Cutting Req Orphan Sheet", "unit": "Sheets", "opening_stock": "10"}).json()

    resp = client.post("/api/cutting-requirements/", json={
        "production_job_id": 999999, "material_id": material["id"], "part_name": "Orphan Part",
        "length_mm": "500", "width_mm": "300",
    })
    assert resp.status_code == 404


def test_list_filters_by_production_job(client, test_user):
    _login(client, test_user)
    job_a = _make_job_cutting(client)
    job_b = _make_job_cutting(client)
    material = client.post("/api/materials/", json={"name": "Cutting Req Filter Sheet", "unit": "Sheets", "opening_stock": "10"}).json()
    client.post("/api/cutting-requirements/", json={
        "production_job_id": job_a["id"], "material_id": material["id"], "part_name": "Job A Part",
        "length_mm": "500", "width_mm": "300",
    })
    client.post("/api/cutting-requirements/", json={
        "production_job_id": job_b["id"], "material_id": material["id"], "part_name": "Job B Part",
        "length_mm": "500", "width_mm": "300",
    })

    resp = client.get("/api/cutting-requirements/", params={"production_job_id": job_a["id"]})
    parts = [r["part_name"] for r in resp.json()]
    assert "Job A Part" in parts
    assert "Job B Part" not in parts


def test_update_and_delete_cutting_requirement(client, test_user):
    _login(client, test_user)
    job = _make_job_cutting(client)
    material = client.post("/api/materials/", json={"name": "Cutting Req Update Sheet", "unit": "Sheets", "opening_stock": "10"}).json()
    requirement = client.post("/api/cutting-requirements/", json={
        "production_job_id": job["id"], "material_id": material["id"], "part_name": "Original Name",
        "length_mm": "500", "width_mm": "300",
    }).json()

    resp = client.put(f"/api/cutting-requirements/{requirement['id']}", json={"part_name": "Renamed Part", "quantity": 3})
    assert resp.status_code == 200
    assert resp.json()["part_name"] == "Renamed Part"
    assert resp.json()["quantity"] == 3

    resp = client.delete(f"/api/cutting-requirements/{requirement['id']}")
    assert resp.status_code == 204
    assert client.get(f"/api/cutting-requirements/{requirement['id']}").status_code == 404


def test_cutting_requirements_are_master_only_to_write(client, test_user, db_session):
    from app.platform.security import hash_password
    from app.modules.auth.auth import User
    _login(client, test_user)
    job = _make_job_cutting(client)
    material = client.post("/api/materials/", json={"name": "Cutting Req RBAC Sheet", "unit": "Sheets", "opening_stock": "10"}).json()
    employee = User(
        username="cuttingrequser", email="cuttingrequser@example.com", full_name="Cutting Req Employee",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=None, is_active=True,
    )
    db_session.add(employee)
    db_session.commit()
    resp = client.post("/api/auth/login", json={"identifier": "cuttingrequser@example.com", "password": "EmpPass1!"})
    assert resp.status_code == 200

    resp = client.post("/api/cutting-requirements/", json={
        "production_job_id": job["id"], "material_id": material["id"], "part_name": "Employee Part",
        "length_mm": "500", "width_mm": "300",
    })
    assert resp.status_code == 403
    # Reading is fine for any authenticated employee.
    assert client.get("/api/cutting-requirements/").status_code == 200


def test_cutting_requirements_require_auth(client):
    resp = client.get("/api/cutting-requirements/")
    assert resp.status_code == 401


def test_cut_list_groups_by_material(client, test_user):
    """Family 131 section 12 - the cut list must group real
    CuttingRequirement rows by material, using the exact same data
    already entered against the job, never a second data source."""
    _login(client, test_user)
    job = _make_job_cutting(client)
    plywood = client.post("/api/materials/", json={"name": "Cut List Plywood", "unit": "Sheets", "opening_stock": "10"}).json()
    mdf = client.post("/api/materials/", json={"name": "Cut List MDF", "unit": "Sheets", "opening_stock": "10"}).json()

    client.post("/api/cutting-requirements/", json={
        "production_job_id": job["id"], "material_id": plywood["id"], "part_name": "Side Panel",
        "quantity": 2, "length_mm": "600", "width_mm": "400", "thickness_mm": "18",
    })
    client.post("/api/cutting-requirements/", json={
        "production_job_id": job["id"], "material_id": plywood["id"], "part_name": "Shelf",
        "quantity": 3, "length_mm": "500", "width_mm": "300", "thickness_mm": "18",
    })
    client.post("/api/cutting-requirements/", json={
        "production_job_id": job["id"], "material_id": mdf["id"], "part_name": "Door",
        "quantity": 1, "length_mm": "700", "width_mm": "450", "thickness_mm": "6",
    })

    resp = client.get(f"/api/production-jobs/{job['id']}/cut-list")
    assert resp.status_code == 200
    body = resp.json()
    assert body["total_parts"] == 6  # 2 + 3 + 1
    assert len(body["material_groups"]) == 2
    plywood_group = next(g for g in body["material_groups"] if g["material_id"] == plywood["id"])
    assert len(plywood_group["parts"]) == 2
    part_names = {p["part_name"] for p in plywood_group["parts"]}
    assert part_names == {"Side Panel", "Shelf"}


def test_nesting_calculates_sheets_required(client, test_user):
    """Family 131 section 13 - nesting must run against real part
    dimensions and report an honest, deterministic result, never a
    fake 'optimized' number."""
    _login(client, test_user)
    job = _make_job_cutting(client)
    material = client.post("/api/materials/", json={"name": "Nesting Plywood", "unit": "Sheets", "opening_stock": "10"}).json()
    client.post("/api/cutting-requirements/", json={
        "production_job_id": job["id"], "material_id": material["id"], "part_name": "Panel",
        "quantity": 4, "length_mm": "600", "width_mm": "400", "rotation_allowed": True,
    })

    resp = client.get(f"/api/production-jobs/{job['id']}/nesting")
    assert resp.status_code == 200
    body = resp.json()
    assert len(body["material_results"]) == 1
    result = body["material_results"][0]
    assert result["material_id"] == material["id"]
    assert result["sheets_required"] == 1
    assert result["parts_placed"] == 4
    assert 0 < result["estimated_utilization_percent"] < 100


def test_nesting_reports_error_for_oversized_part(client, test_user):
    """A part larger than the sheet (even with rotation, since it's
    disallowed here) must be reported as a real, explainable error -
    never silently dropped or given a wrong sheet count."""
    _login(client, test_user)
    job = _make_job_cutting(client)
    material = client.post("/api/materials/", json={"name": "Oversized Part Material", "unit": "Sheets", "opening_stock": "10"}).json()
    client.post("/api/cutting-requirements/", json={
        "production_job_id": job["id"], "material_id": material["id"], "part_name": "Too Big",
        "quantity": 1, "length_mm": "3000", "width_mm": "400", "rotation_allowed": False,
    })

    resp = client.get(f"/api/production-jobs/{job['id']}/nesting", params={"sheet_length_mm": 2440, "sheet_width_mm": 1220})
    assert resp.status_code == 200
    result = resp.json()["material_results"][0]
    assert result["error"] == "parts_too_large_for_sheet"


def test_nesting_needs_multiple_sheets_when_parts_dont_fit_on_one(client, test_user):
    _login(client, test_user)
    job = _make_job_cutting(client)
    material = client.post("/api/materials/", json={"name": "Multi Sheet Material", "unit": "Sheets", "opening_stock": "10"}).json()
    client.post("/api/cutting-requirements/", json={
        "production_job_id": job["id"], "material_id": material["id"], "part_name": "Large Panel",
        "quantity": 5, "length_mm": "600", "width_mm": "1220", "rotation_allowed": True,
    })

    resp = client.get(f"/api/production-jobs/{job['id']}/nesting")
    assert resp.status_code == 200
    result = resp.json()["material_results"][0]
    assert result["sheets_required"] >= 2


def test_cut_list_and_nesting_require_auth(client):
    assert client.get("/api/production-jobs/1/cut-list").status_code == 401
    assert client.get("/api/production-jobs/1/nesting").status_code == 401


def test_cut_list_unknown_job_returns_404(client, test_user):
    _login(client, test_user)
    resp = client.get("/api/production-jobs/999999/cut-list")
    assert resp.status_code == 404


# --- test_operations.py ---
def test_create_milestone(client, test_user):
    _login(client, test_user)
    client_id = client.post("/api/clients/", json={"name": "Milestone Test Client", "phone": "9000010114"}).json()["id"]
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00", "order_value": "10000", "advance": "0",
    }).json()
    resp = client.post("/api/milestones/", json={
        "order_id": order["id"], "name": "Design approval", "target_date": "2026-09-01T00:00:00",
    })
    assert resp.status_code == 201
    assert resp.json()["name"] == "Design approval"
    assert resp.json()["completed_date"] is None


def test_milestone_creation_requires_master(client, test_user, db_session):
    _login(client, test_user)
    client_id = client.post("/api/clients/", json={"name": "Milestone RBAC Client", "phone": "9000010115"}).json()["id"]
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00", "order_value": "10000", "advance": "0",
    }).json()
    employee = client.post("/api/employees/", json={"name": "Milestone RBAC Employee"}).json()
    user = User(
        username="milestonerbacuser", email="milestonerbacuser@example.com", full_name="Milestone RBAC User",
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee["id"], is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    client.post("/api/auth/login", json={"identifier": "milestonerbacuser@example.com", "password": "EmpPass1!"})

    resp = client.post("/api/milestones/", json={
        "order_id": order["id"], "name": "Employee Attempt", "target_date": "2026-09-01T00:00:00",
    })
    assert resp.status_code == 403


def test_milestone_can_be_marked_complete(client, test_user):
    _login(client, test_user)
    client_id = client.post("/api/clients/", json={"name": "Milestone Complete Client", "phone": "9000010116"}).json()["id"]
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00", "order_value": "10000", "advance": "0",
    }).json()
    milestone = client.post("/api/milestones/", json={
        "order_id": order["id"], "name": "Site measurement", "target_date": "2026-08-25T00:00:00",
    }).json()

    resp = client.put(f"/api/milestones/{milestone['id']}", json={"completed_date": "2026-08-24T00:00:00"})
    assert resp.status_code == 200
    assert resp.json()["completed_date"] is not None


def test_milestones_list_filters_by_order(client, test_user):
    _login(client, test_user)
    client_id = client.post("/api/clients/", json={"name": "Milestone Filter Client", "phone": "9000010117"}).json()["id"]
    order_a = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00", "order_value": "10000", "advance": "0",
    }).json()
    order_b = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00", "order_value": "12000", "advance": "0",
    }).json()
    client.post("/api/milestones/", json={"order_id": order_a["id"], "name": "A Milestone", "target_date": "2026-09-01T00:00:00"})
    client.post("/api/milestones/", json={"order_id": order_b["id"], "name": "B Milestone", "target_date": "2026-09-01T00:00:00"})

    resp = client.get("/api/milestones/", params={"order_id": order_a["id"]})
    assert resp.status_code == 200
    names = [m["name"] for m in resp.json()]
    assert "A Milestone" in names
    assert "B Milestone" not in names


def test_chatbot_next_action_via_client_name(client, test_user):
    _login(client, test_user)
    client_id = client.post("/api/clients/", json={"name": "Next Action Client", "phone": "9000010118"}).json()["id"]
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00", "order_value": "10000", "advance": "0",
    }).json()
    client.post("/api/daily-tasks/", json={
        "order_id": order["id"], "date": "2026-08-20T00:00:00",
        "task_description": "Site measurement", "status": "TO DO",
    })

    resp = client.post("/api/chat/", json={"message": "what is the next action for Next Action Client"})
    assert resp.status_code == 200
    assert "Site measurement" in resp.json()["response"]


def test_chatbot_next_action_flags_blocked_task(client, test_user):
    _login(client, test_user)
    client_id = client.post("/api/clients/", json={"name": "Next Action Blocked Client", "phone": "9000010119"}).json()["id"]
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00", "order_value": "10000", "advance": "0",
    }).json()
    client.post("/api/daily-tasks/", json={
        "order_id": order["id"], "date": "2026-08-20T00:00:00",
        "task_description": "Blocked cutting", "status": "BLOCKED", "delay_reason": "Waiting for material",
    })

    resp = client.post("/api/chat/", json={"message": "what is the next action for Next Action Blocked Client"})
    assert resp.status_code == 200
    assert "Blocked cutting" in resp.json()["response"]


def test_chatbot_next_action_no_open_tasks(client, test_user):
    _login(client, test_user)
    client_id = client.post("/api/clients/", json={"name": "Next Action Empty Client", "phone": "9000010120"}).json()["id"]
    client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00", "order_value": "10000", "advance": "0",
    })

    resp = client.post("/api/chat/", json={"message": "what is the next action for Next Action Empty Client"})
    assert resp.status_code == 200


def _create_employee_user_for_task_auth(client, db_session, employee_id, username, email):
    """Creates a real 'user'-role account linked to a specific employee,
    the same way a master would via the Users admin page."""
    user = User(
        username=username, email=email, full_name=username,
        password_hash=hash_password("EmployeePass1!"), role="user", employee_id=employee_id, is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    return user


def test_employee_can_update_own_task_status(client, test_user, db_session):
    _login(client, test_user)
    employee_id = client.post("/api/employees/", json={
        "name": "Task Auth Employee One", "monthly_salary": "20000", "daily_wage": "800",
    }).json()["id"]
    task = client.post("/api/daily-tasks/", json={
        "date": "2026-08-13T00:00:00", "employee_id": employee_id, "task_description": "Sand panels",
    }).json()
    _create_employee_user_for_task_auth(client, db_session, employee_id, "taskauthemp1", "taskauthemp1@example.com")

    resp = client.post("/api/auth/login", json={"identifier": "taskauthemp1@example.com", "password": "EmployeePass1!"})
    assert resp.status_code == 200

    update_resp = client.put(f"/api/daily-tasks/{task['id']}", json={"status": "DOING"})
    assert update_resp.status_code == 200
    assert update_resp.json()["status"] == "DOING"


def test_employee_cannot_update_another_employees_task_status(client, test_user, db_session):
    """The core rule: ownership is resolved via
    employee_id, not open to anyone - an employee cannot touch a task
    assigned to someone else, even just its status."""
    _login(client, test_user)
    employee_a = client.post("/api/employees/", json={
        "name": "Task Auth Employee A", "monthly_salary": "20000", "daily_wage": "800",
    }).json()["id"]
    employee_b = client.post("/api/employees/", json={
        "name": "Task Auth Employee B", "monthly_salary": "20000", "daily_wage": "800",
    }).json()["id"]
    task_for_b = client.post("/api/daily-tasks/", json={
        "date": "2026-08-13T00:00:00", "employee_id": employee_b, "task_description": "B's task",
    }).json()
    _create_employee_user_for_task_auth(client, db_session, employee_a, "taskauthempa", "taskauthempa@example.com")

    resp = client.post("/api/auth/login", json={"identifier": "taskauthempa@example.com", "password": "EmployeePass1!"})
    assert resp.status_code == 200

    update_resp = client.put(f"/api/daily-tasks/{task_for_b['id']}", json={"status": "DONE"})
    assert update_resp.status_code == 403


def test_employee_can_update_completion_percent_on_own_task(client, test_user, db_session):
    """completion_percent IS in the self-service field set
    (EMPLOYEE_SELF_SERVICE_FIELDS in daily_tasks.py) - an employee can
    update it on their OWN task."""
    _login(client, test_user)
    employee_id = client.post("/api/employees/", json={
        "name": "Task Auth Employee C", "monthly_salary": "20000", "daily_wage": "800",
    }).json()["id"]
    task = client.post("/api/daily-tasks/", json={
        "date": "2026-08-13T00:00:00", "employee_id": employee_id, "task_description": "C's task",
    }).json()
    _create_employee_user_for_task_auth(client, db_session, employee_id, "taskauthempc", "taskauthempc@example.com")

    resp = client.post("/api/auth/login", json={"identifier": "taskauthempc@example.com", "password": "EmployeePass1!"})
    assert resp.status_code == 200

    update_resp = client.put(f"/api/daily-tasks/{task['id']}", json={"completion_percent": 40})
    assert update_resp.status_code == 200
    assert update_resp.json()["completion_percent"] == 40


def test_employee_cannot_reassign_task_via_self_service_fields(client, test_user, db_session):
    """Even on their own task, an employee still can't change fields
    outside the self-service allowlist - e.g. checked_by."""
    _login(client, test_user)
    employee_id = client.post("/api/employees/", json={
        "name": "Task Auth Employee D", "monthly_salary": "20000", "daily_wage": "800",
    }).json()["id"]
    task = client.post("/api/daily-tasks/", json={
        "date": "2026-08-13T00:00:00", "employee_id": employee_id, "task_description": "D's task",
    }).json()
    _create_employee_user_for_task_auth(client, db_session, employee_id, "taskauthempd", "taskauthempd@example.com")

    resp = client.post("/api/auth/login", json={"identifier": "taskauthempd@example.com", "password": "EmployeePass1!"})
    assert resp.status_code == 200

    update_resp = client.put(f"/api/daily-tasks/{task['id']}", json={"checked_by": "Someone Else"})
    assert update_resp.status_code == 403


def test_employee_with_no_linked_employee_record_cannot_update_any_task(client, test_user, db_session):
    """Ownership is resolved via authenticated user -> employee_id ->
    DailyTask.employee_id (B13). A 'user'-role account with no linked
    Employee record has employee_id=None, which can never equal a
    real task's employee_id - so it can't update any task, including
    its own status, since it has no "own" task to begin with."""
    _login(client, test_user)
    employee_id = client.post("/api/employees/", json={
        "name": "Task Auth Employee E", "monthly_salary": "20000", "daily_wage": "800",
    }).json()["id"]
    task = client.post("/api/daily-tasks/", json={
        "date": "2026-08-13T00:00:00", "employee_id": employee_id, "task_description": "E's task",
    }).json()
    unlinked = User(
        username="taskauthunlinked", email="taskauthunlinked@example.com", full_name="Unlinked",
        password_hash=hash_password("EmployeePass1!"), role="user", employee_id=None, is_active=True,
    )
    db_session.add(unlinked)
    db_session.commit()

    resp = client.post("/api/auth/login", json={"identifier": "taskauthunlinked@example.com", "password": "EmployeePass1!"})
    assert resp.status_code == 200

    update_resp = client.put(f"/api/daily-tasks/{task['id']}", json={"status": "DOING"})
    assert update_resp.status_code == 403


def test_master_can_edit_all_task_fields(client, test_user):
    _login(client, test_user)
    employee_id = client.post("/api/employees/", json={
        "name": "Task Auth Master Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()["id"]
    task = client.post("/api/daily-tasks/", json={
        "date": "2026-08-13T00:00:00", "employee_id": employee_id, "task_description": "Master editable task",
    }).json()

    update_resp = client.put(f"/api/daily-tasks/{task['id']}", json={
        "status": "DOING", "completion_percent": 40, "checked_by": "Garima",
    })
    assert update_resp.status_code == 200
    assert update_resp.json()["completion_percent"] == 40
    assert update_resp.json()["checked_by"] == "Garima"


def test_master_can_reassign_another_employees_task(client, test_user):
    """Master is not subject to the ownership restriction at all -
    reassignment (and every other field) is a management operation."""
    _login(client, test_user)
    employee_id = client.post("/api/employees/", json={
        "name": "Task Auth Master Reassign Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()["id"]
    other_employee_id = client.post("/api/employees/", json={
        "name": "Task Auth Master Reassign Other Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()["id"]
    task = client.post("/api/daily-tasks/", json={
        "date": "2026-08-13T00:00:00", "employee_id": employee_id, "task_description": "Reassignable task",
    }).json()

    update_resp = client.put(f"/api/daily-tasks/{task['id']}", json={"employee_id": other_employee_id})
    assert update_resp.status_code == 200
    assert update_resp.json()["employee_id"] == other_employee_id


def test_completed_status_forces_completion_percent_to_100(client, test_user):
    _login(client, test_user)
    employee_id = client.post("/api/employees/", json={
        "name": "Task Auth Employee F", "monthly_salary": "20000", "daily_wage": "800",
    }).json()["id"]
    task = client.post("/api/daily-tasks/", json={
        "date": "2026-08-13T00:00:00", "employee_id": employee_id, "task_description": "F's task",
        "completion_percent": 30,
    }).json()

    update_resp = client.put(f"/api/daily-tasks/{task['id']}", json={"status": "DONE"})
    assert update_resp.status_code == 200
    assert update_resp.json()["completion_percent"] == 100
    assert update_resp.json()["actual_completed_at"] is not None


def test_any_authenticated_role_can_view_all_tasks(client, test_user, db_session):
    """Viewing (GET list, without mine=True) is unrestricted for any
    authenticated role - the ownership restriction in B13 applies to
    updating a task, not viewing the list."""
    _login(client, test_user)
    employee_id = client.post("/api/employees/", json={
        "name": "Task Auth View Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()["id"]
    other_employee_id = client.post("/api/employees/", json={
        "name": "Task Auth View Other Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()["id"]
    client.post("/api/daily-tasks/", json={
        "date": "2026-08-13T00:00:00", "employee_id": other_employee_id, "task_description": "Someone else's task",
    })
    _create_employee_user_for_task_auth(client, db_session, employee_id, "taskauthviewer", "taskauthviewer@example.com")

    resp = client.post("/api/auth/login", json={"identifier": "taskauthviewer@example.com", "password": "EmployeePass1!"})
    assert resp.status_code == 200

    tasks = client.get("/api/daily-tasks/").json()
    assert any(t["task_description"] == "Someone else's task" for t in tasks)


def test_mine_query_still_filters_to_the_linked_employees_own_tasks(client, test_user, db_session):
    """"mine" is a convenience filter for an employee's own task list -
    a separate concern from B13's update-ownership restriction, but
    consistent with it (both resolve via the same employee_id link)."""
    _login(client, test_user)
    employee_id = client.post("/api/employees/", json={
        "name": "Task Auth Mine Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()["id"]
    other_employee_id = client.post("/api/employees/", json={
        "name": "Task Auth Mine Other Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()["id"]
    client.post("/api/daily-tasks/", json={
        "date": "2026-08-13T00:00:00", "employee_id": employee_id, "task_description": "My own task",
    })
    client.post("/api/daily-tasks/", json={
        "date": "2026-08-13T00:00:00", "employee_id": other_employee_id, "task_description": "Not my task",
    })
    _create_employee_user_for_task_auth(client, db_session, employee_id, "taskauthmine", "taskauthmine@example.com")

    resp = client.post("/api/auth/login", json={"identifier": "taskauthmine@example.com", "password": "EmployeePass1!"})
    assert resp.status_code == 200

    mine_tasks = client.get("/api/daily-tasks/", params={"mine": True}).json()
    descriptions = [t["task_description"] for t in mine_tasks]
    assert "My own task" in descriptions
    assert "Not my task" not in descriptions


def _create_and_login_employee_user(client, db_session, employee_id, username, email):
    user = User(
        username=username, email=email, full_name=username,
        password_hash=hash_password("EmpPass1!"), role="user", employee_id=employee_id, is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    resp = client.post("/api/auth/login", json={"identifier": email, "password": "EmpPass1!"})
    assert resp.status_code == 200


def test_complete_and_assign_next_creates_linked_task(client, test_user):
    _login(client, test_user)
    emp_a = client.post("/api/employees/", json={"name": "Handoff Employee A"}).json()
    emp_b = client.post("/api/employees/", json={"name": "Handoff Employee B"}).json()
    task = client.post("/api/daily-tasks/", json={
        "date": "2026-08-19T00:00:00", "employee_id": emp_a["id"], "task_description": "Measure wardrobe",
    }).json()

    resp = client.post(f"/api/daily-tasks/{task['id']}/complete-and-assign-next", json={
        "next_employee_id": emp_b["id"], "next_task_description": "Prepare cutting drawing",
        "next_due_date": "2026-08-20T00:00:00",
    })
    assert resp.status_code == 201
    next_task = resp.json()
    assert next_task["previous_task_id"] == task["id"]
    assert next_task["employee_id"] == emp_b["id"]

    original = client.get(f"/api/daily-tasks/{task['id']}").json()
    assert original["status"] == "DONE"
    assert original["completion_percent"] == 100


def test_original_task_not_overwritten_by_handoff(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Handoff No Overwrite Employee"}).json()
    next_employee = client.post("/api/employees/", json={"name": "Handoff No Overwrite Next"}).json()
    task = client.post("/api/daily-tasks/", json={
        "date": "2026-08-19T00:00:00", "employee_id": employee["id"], "task_description": "Original description",
    }).json()

    client.post(f"/api/daily-tasks/{task['id']}/complete-and-assign-next", json={
        "next_employee_id": next_employee["id"], "next_task_description": "Different description",
        "next_due_date": "2026-08-20T00:00:00",
    })

    original = client.get(f"/api/daily-tasks/{task['id']}").json()
    assert original["task_description"] == "Original description"


def test_add_and_list_task_comments(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Comment Test Employee"}).json()
    task = client.post("/api/daily-tasks/", json={
        "date": "2026-08-19T00:00:00", "employee_id": employee["id"], "task_description": "Cut panels",
    }).json()

    resp = client.post(f"/api/daily-tasks/{task['id']}/comments", json={"text": "Ready for cutting."})
    assert resp.status_code == 201

    comments = client.get(f"/api/daily-tasks/{task['id']}/comments").json()
    assert len(comments) == 1
    assert comments[0]["text"] == "Ready for cutting."


def test_subtask_links_to_parent(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Subtask Test Employee"}).json()
    parent = client.post("/api/daily-tasks/", json={
        "date": "2026-08-19T00:00:00", "employee_id": employee["id"], "task_description": "Wardrobe - full build",
    }).json()
    sub = client.post("/api/daily-tasks/", json={
        "date": "2026-08-19T00:00:00", "employee_id": employee["id"], "task_description": "Assembly",
        "parent_task_id": parent["id"],
    }).json()
    assert sub["parent_task_id"] == parent["id"]


def test_employee_can_set_status_and_block_reason_together(client, test_user, db_session):
    """The explicit brief requirement - mark BLOCKED with a reason -
    previously impossible since only status was self-service."""
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Block Reason Employee"}).json()
    task = client.post("/api/daily-tasks/", json={
        "date": "2026-08-19T00:00:00", "employee_id": employee["id"], "task_description": "Apply laminate",
    }).json()
    _create_and_login_employee_user(client, db_session, employee["id"], "blockreasonuser", "blockreasonuser@example.com")

    resp = client.put(f"/api/daily-tasks/{task['id']}", json={
        "status": "BLOCKED", "delay_reason": "Waiting for laminate",
    })
    assert resp.status_code == 200
    assert resp.json()["delay_reason"] == "Waiting for laminate"


def test_employee_notified_when_task_blocked(client, test_user, db_session):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Block Notif Employee"}).json()
    task = client.post("/api/daily-tasks/", json={
        "date": "2026-08-19T00:00:00", "employee_id": employee["id"], "task_description": "Sand panels",
    }).json()
    _create_and_login_employee_user(client, db_session, employee["id"], "blocknotifuser", "blocknotifuser@example.com")

    client.put(f"/api/daily-tasks/{task['id']}", json={"status": "BLOCKED", "delay_reason": "Material unavailable"})

    notifications = client.get("/api/notifications/").json()
    assert any(n["notification_type"] == "TASK_BLOCKED" for n in notifications)


def test_created_by_set_on_task_creation(client, test_user):
    _login(client, test_user)
    employee = client.post("/api/employees/", json={"name": "Created By Employee"}).json()
    task = client.post("/api/daily-tasks/", json={
        "date": "2026-08-19T00:00:00", "employee_id": employee["id"], "task_description": "Install hardware",
    }).json()
    assert task["created_by"] == "test@example.com"


def test_employee_cannot_complete_unrelated_task(client, test_user, db_session):
    """The explicit brief requirement - an employee must not be able to
    use Complete & Assign Next on a task assigned to someone else."""
    _login(client, test_user)
    owner = client.post("/api/employees/", json={"name": "Handoff Owner Employee"}).json()
    other_employee = client.post("/api/employees/", json={"name": "Handoff Unrelated Employee"}).json()
    next_employee = client.post("/api/employees/", json={"name": "Handoff Unrelated Next"}).json()
    task = client.post("/api/daily-tasks/", json={
        "date": "2026-08-19T00:00:00", "employee_id": owner["id"], "task_description": "Owner's task",
    }).json()
    _create_and_login_employee_user(client, db_session, other_employee["id"], "handoffunrelateduser", "handoffunrelateduser@example.com")

    resp = client.post(f"/api/daily-tasks/{task['id']}/complete-and-assign-next", json={
        "next_employee_id": next_employee["id"], "next_task_description": "Should not be created",
        "next_due_date": "2026-08-20T00:00:00",
    })
    assert resp.status_code == 403

    unchanged = client.get(f"/api/daily-tasks/{task['id']}")
    assert unchanged.status_code == 200
    assert unchanged.json()["status"] != "DONE"


def test_employee_can_complete_and_assign_next_for_own_task(client, test_user, db_session):
    _login(client, test_user)
    owner = client.post("/api/employees/", json={"name": "Handoff Self Owner"}).json()
    next_employee = client.post("/api/employees/", json={"name": "Handoff Self Next"}).json()
    task = client.post("/api/daily-tasks/", json={
        "date": "2026-08-19T00:00:00", "employee_id": owner["id"], "task_description": "My own task",
    }).json()
    _create_and_login_employee_user(client, db_session, owner["id"], "handoffselfuser", "handoffselfuser@example.com")

    resp = client.post(f"/api/daily-tasks/{task['id']}/complete-and-assign-next", json={
        "next_employee_id": next_employee["id"], "next_task_description": "Handed off task",
        "next_due_date": "2026-08-20T00:00:00",
    })
    assert resp.status_code == 201


class _QueryCounter:
    def __init__(self):
        self.count = 0

    def __enter__(self):
        event.listen(engine, "before_cursor_execute", self._callback)
        return self

    def __exit__(self, *args):
        event.remove(engine, "before_cursor_execute", self._callback)

    def _callback(self, conn, cursor, statement, parameters, context, executemany):
        self.count += 1


def test_tasks_export_query_count_does_not_scale_with_row_count(client, test_user):
    """The core N+1 regression guard - export a growing number of
    tasks, each with a DIFFERENT employee and order, and confirm the
    query count stays bounded rather than growing linearly."""
    _login(client, test_user)
    for i in range(8):
        client_id = client.post("/api/clients/", json={"name": f"N+1 Test Client {i}", "phone": "9000010077"}).json()["id"]
        order = client.post("/api/orders/", json={
            "client_id": client_id, "order_date": "2026-08-21T00:00:00", "order_value": "10000", "advance": "0",
        }).json()
        employee = client.post("/api/employees/", json={"name": f"N+1 Test Employee {i}"}).json()
        client.post("/api/daily-tasks/", json={
            "date": "2026-08-21T00:00:00", "employee_id": employee["id"], "order_id": order["id"],
            "task_description": f"N+1 test task {i}",
        })

    with _QueryCounter() as counter:
        resp = client.get("/api/reports/tasks.xlsx")
    assert resp.status_code == 200
    # With 8 distinct employees/orders, an N+1 bug would need roughly
    # 8 (tasks) + 8 (employee lookups) + 8 (order lookups) = 24+ queries
    # for this one export alone. Eager-loading keeps it to a small,
    # fixed number regardless of row count.
    assert counter.count < 15, f"expected a bounded query count, got {counter.count} - possible N+1 regression"


def _make_at_risk_order(client, suffix):
    """Same shortage-inducing setup as test_shortage_intelligence.py -
    duplicated in miniature here since that module owns the shortage
    formula's own tests; this file only needs one at-risk order to
    prove the task-list enrichment surfaces it correctly."""
    material = client.post("/api/materials/", json={
        "name": f"Task Risk Sheet {suffix}", "unit": "Sheets", "opening_stock": "1", "minimum_stock": "1",
    }).json()
    product = client.post("/api/products/", json={
        "name": f"Task Risk Product {suffix}", "unit": "Piece",
        "materials_used": [{"material_id": material["id"], "quantity_required": "5"}],
    }).json()
    client_id = client.post("/api/clients/", json={"name": f"Task Risk Client {suffix}", "phone": f"900001040{suffix}"}).json()["id"]
    return client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()


def test_task_list_flags_material_at_risk_when_requested(client, test_user):
    _login(client, test_user)
    employee_id = client.post("/api/employees/", json={
        "name": "Task Risk Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()["id"]
    order = _make_at_risk_order(client, "1")
    task = client.post("/api/daily-tasks/", json={
        "date": "2026-08-19T00:00:00", "employee_id": employee_id, "order_id": order["id"],
        "task_description": "Cut sheets for at-risk order",
    }).json()

    resp = client.get("/api/daily-tasks/", params={"include_material_risk": "true"})
    assert resp.status_code == 200
    row = next(r for r in resp.json() if r["id"] == task["id"])
    assert row["material_at_risk"] is True


def test_task_list_omits_material_risk_by_default(client, test_user):
    """The flag is opt-in - a plain list call (e.g. the dashboard's
    small "my tasks" widget) must not pay for the bulk shortage
    calculation it never asked for, and must not report a false
    positive by defaulting to True."""
    _login(client, test_user)
    employee_id = client.post("/api/employees/", json={
        "name": "Task Risk Default Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()["id"]
    order = _make_at_risk_order(client, "2")
    task = client.post("/api/daily-tasks/", json={
        "date": "2026-08-19T00:00:00", "employee_id": employee_id, "order_id": order["id"],
        "task_description": "Cut sheets for at-risk order, no flag requested",
    }).json()

    resp = client.get("/api/daily-tasks/")
    assert resp.status_code == 200
    row = next(r for r in resp.json() if r["id"] == task["id"])
    assert row["material_at_risk"] is False


def test_task_list_material_at_risk_false_for_well_stocked_order(client, test_user):
    _login(client, test_user)
    employee_id = client.post("/api/employees/", json={
        "name": "Task No Risk Employee", "monthly_salary": "20000", "daily_wage": "800",
    }).json()["id"]
    material = client.post("/api/materials/", json={
        "name": "Task No Risk Sheet", "unit": "Sheets", "opening_stock": "100", "minimum_stock": "1",
    }).json()
    product = client.post("/api/products/", json={
        "name": "Task No Risk Product", "unit": "Piece",
        "materials_used": [{"material_id": material["id"], "quantity_required": "2"}],
    }).json()
    client_id = client.post("/api/clients/", json={"name": "Task No Risk Client", "phone": "9000010403"}).json()["id"]
    order = client.post("/api/orders/", json={
        "client_id": client_id, "order_date": "2026-08-19T00:00:00",
        "items": [{"description": "Item", "quantity": "1", "unit": "Piece", "rate": "5000", "product_id": product["id"]}],
    }).json()
    task = client.post("/api/daily-tasks/", json={
        "date": "2026-08-19T00:00:00", "employee_id": employee_id, "order_id": order["id"],
        "task_description": "Cut sheets for well-stocked order",
    }).json()

    resp = client.get("/api/daily-tasks/", params={"include_material_risk": "true"})
    assert resp.status_code == 200
    row = next(r for r in resp.json() if r["id"] == task["id"])
    assert row["material_at_risk"] is False
