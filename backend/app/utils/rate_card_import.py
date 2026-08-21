"""Excel import for the Rate Master. Same discipline as
product_import.py/client_import.py: one fixed template shared by the
download and the parser, never writes a parsed row directly to the
database, preview/commit are two separate steps.

Business rule specific to this import: uploaded rates are NEVER
allowed to blindly overwrite a historical rate. Every committed row
goes through the exact same create-new-version path as a UI edit (see
rate_cards.py's revise_rate_card) - a Rate ID in the sheet identifies
WHICH item to version, it never lets the sheet dictate the new row's
own ID.
"""
import re
from io import BytesIO
from decimal import Decimal, InvalidOperation
from typing import List, Optional

import openpyxl
from openpyxl import Workbook

from app.utils.exporters import write_sheet
from app.models.rate_card import RATE_SOURCE_TYPES, RATE_CONFIDENCE_LEVELS, STANDARD_UOMS, uom_allowed_for_category

RATE_IMPORT_COLUMNS = [
    "Rate ID", "Category", "Subcategory", "Item Name", "Specification", "Location", "UOM",
    "Market Reference Rate", "Woodful Cost Rate", "Woodful Selling Rate",
    "Overhead %", "Target Margin %", "Wastage %", "Tax %",
    "Source Type", "Source Reference", "Confidence", "Notes",
]

HEADER_ALIASES = {
    "rate id": "Rate ID", "item name": "Item Name", "item": "Item Name", "product/service": "Item Name",
    "category": "Category", "subcategory": "Subcategory", "specification": "Specification",
    "spec": "Specification", "location": "Location", "uom": "UOM", "unit": "UOM",
    "market reference rate": "Market Reference Rate", "market rate": "Market Reference Rate",
    "woodful cost rate": "Woodful Cost Rate", "cost rate": "Woodful Cost Rate",
    "woodful selling rate": "Woodful Selling Rate", "selling rate": "Woodful Selling Rate",
    "overhead %": "Overhead %", "overhead": "Overhead %",
    "target margin %": "Target Margin %", "margin": "Target Margin %", "margin %": "Target Margin %",
    "wastage %": "Wastage %", "wastage": "Wastage %",
    "tax %": "Tax %", "tax": "Tax %", "gst": "Tax %", "gst %": "Tax %",
    "source type": "Source Type", "source reference": "Source Reference", "source": "Source Reference",
    "confidence": "Confidence", "notes": "Notes",
}
for _col in RATE_IMPORT_COLUMNS:
    HEADER_ALIASES.setdefault(_col.lower(), _col)

REQUIRED_COLUMNS = ["Category", "Item Name", "UOM", "Source Type"]


def _normalize_header_cell(raw) -> Optional[str]:
    if raw is None:
        return None
    key = re.sub(r"\s+", " ", str(raw).strip().lower())
    return HEADER_ALIASES.get(key)


def _resolve_header_row(values) -> Optional[dict]:
    resolved = {}
    for idx, raw in enumerate(values):
        canonical = _normalize_header_cell(raw)
        if canonical is None:
            continue
        if canonical in resolved:
            return None
        resolved[canonical] = idx
    if all(col in resolved for col in REQUIRED_COLUMNS):
        return resolved
    return None


_NUMERIC_NOISE_RE = re.compile(r"[\u20b9$\u20ac\u00a3%,\s]")


def normalize_number(raw) -> Optional[Decimal]:
    if raw is None:
        return None
    if isinstance(raw, (int, float, Decimal)):
        return Decimal(str(raw))
    text = _NUMERIC_NOISE_RE.sub("", str(raw))
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


EXAMPLE_ROWS = [
    {
        "Rate ID": "", "Category": "Material", "Subcategory": "Plywood", "Item Name": "Plywood - MR Grade",
        "Specification": "18mm, commercial/MR", "Location": "Indore, Madhya Pradesh", "UOM": "Sq Ft",
        "Market Reference Rate": 60, "Woodful Cost Rate": "", "Woodful Selling Rate": "",
        "Overhead %": 12, "Target Margin %": 30, "Wastage %": 5, "Tax %": 18,
        "Source Type": "INDORE_SUPPLIER", "Source Reference": "Indore plywood dealer listings, Aug 2026",
        "Confidence": "MEDIUM", "Notes": "New rate - leave Rate ID blank",
    },
    {
        "Rate ID": "RATE-001", "Category": "Material", "Subcategory": "Plywood", "Item Name": "Plywood - BWP Grade",
        "Specification": "18mm, waterproof/marine", "Location": "Indore, Madhya Pradesh", "UOM": "Sq Ft",
        "Market Reference Rate": 90, "Woodful Cost Rate": "", "Woodful Selling Rate": "",
        "Overhead %": 12, "Target Margin %": 30, "Wastage %": 5, "Tax %": 18,
        "Source Type": "INDORE_SUPPLIER", "Source Reference": "Indore waterproof plywood listings, Aug 2026",
        "Confidence": "MEDIUM", "Notes": "Existing rate - update by Rate ID, creates a new version",
    },
]


def build_import_template() -> BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(
        wb, sheet_name="Rate Import", title="Woodful Creations - Rate Master Import Template",
        subtitle="Leave Rate ID blank to create a new rate. Enter an existing Rate ID to update it - this "
                  "creates a NEW version and never overwrites the old one, which stays available for "
                  "historical Estimates/Orders. Category, Item Name, UOM, and Source Type are required. "
                  f"Source Type must be one of: {', '.join(sorted(RATE_SOURCE_TYPES))}. "
                  f"Confidence must be one of: {', '.join(sorted(RATE_CONFIDENCE_LEVELS))}.",
        columns=RATE_IMPORT_COLUMNS, rows=EXAMPLE_ROWS,
    )
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    return value


# ---------------------------------------------------------------------
# Alternate real-world layout support: "Type / Category / Item-Grade /
# Thickness-Basis / UOM / Market_Low_INR / Market_High_INR /
# Source_Basis / Notes" - a genuinely different, valid shape (a market
# rate RANGE rather than one figure, no explicit Source Type/Rate ID
# columns) rather than the primary template's shape. Detected
# separately and mapped into the SAME canonical row dict validate_row()
# already expects, so validation/versioning/commit logic is never
# duplicated - only the column mapping differs.
# ---------------------------------------------------------------------
_MARKET_RANGE_HEADERS = ["Type", "Category", "Item/Grade", "Thickness/Basis", "UOM",
                          "Market_Low_INR", "Market_High_INR", "Source_Basis", "Notes"]


def _is_market_range_layout(values) -> bool:
    normalized = [str(v).strip() if v is not None else None for v in values]
    return all(h in normalized for h in _MARKET_RANGE_HEADERS)


def _infer_source_type(source_basis: Optional[str]) -> str:
    text = (source_basis or "").lower()
    if "indore" in text:
        return "INDORE_MARKET_LISTING"
    if "national" in text or "india" in text:
        return "NATIONAL_MARKET_BENCHMARK"
    return "NATIONAL_MARKET_BENCHMARK"


def parse_market_range_workbook(file_bytes: bytes) -> List[dict]:
    """Parses the Type/Category/Item-Grade/Thickness-Basis/UOM/
    Market_Low_INR/Market_High_INR/Source_Basis/Notes layout, mapping
    it to this module's canonical row shape:
        Category    <- Type            (top-level: Material/Furniture/...)
        Subcategory <- Category        (e.g. Plywood, Tables, CNC Cutting)
        Item Name   <- Item/Grade      (e.g. MR Commercial, Coffee Table)
        Specification <- Thickness/Basis
        Market Reference Rate <- midpoint of (Market_Low_INR, Market_High_INR)
        Source Type <- inferred from Source_Basis text (INDORE_* if
                        "Indore" appears, else NATIONAL_MARKET_BENCHMARK)
        Confidence  <- "LOW" (every row in this layout is an unverified
                        market reference BAND, never a fixed Woodful
                        rate - the source sheet's own Notes column says
                        so explicitly, e.g. "Do not use as fixed
                        Woodful rate", "Woodful must verify internal cost")
        Source Reference <- original Source_Basis, with the full
                        Low-High range appended so it's never lost by
                        collapsing to a midpoint.
    This mapping follows the spec's own Category -> Subcategory ->
    Product/Service -> Variant/Specification hierarchy exactly.
    """
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb.worksheets[0]

    header_row_idx = None
    for row in ws.iter_rows(min_row=1, max_row=5):
        values = [c.value for c in row]
        if _is_market_range_layout(values):
            header_row_idx = row[0].row
            break
    if header_row_idx is None:
        return []  # not this layout - caller falls back to the primary template parser

    rows = []
    for row in ws.iter_rows(min_row=header_row_idx + 1):
        values = [_clean(c.value) for c in row]
        if all(v is None for v in values):
            continue
        type_, category, item_grade, spec, uom, low, high, source_basis, notes = (values + [None] * 9)[:9]
        if not (type_ and item_grade and uom):
            continue  # skip malformed/incomplete rows rather than erroring the whole import

        low_n = normalize_number(low)
        high_n = normalize_number(high)
        midpoint = None
        if low_n is not None and high_n is not None:
            midpoint = ((low_n + high_n) / 2).quantize(Decimal("0.01"))
        elif low_n is not None:
            midpoint = low_n
        elif high_n is not None:
            midpoint = high_n

        range_text = f"₹{low}-₹{high}" if (low is not None and high is not None) else ""
        source_reference = " | ".join(filter(None, [source_basis, f"Market range {range_text}" if range_text else None]))

        rows.append({
            "Category": type_, "Subcategory": category, "Item Name": item_grade,
            "Specification": spec, "UOM": uom,
            "Market Reference Rate": midpoint,
            "Source Type": _infer_source_type(source_basis),
            "Source Reference": source_reference or None,
            "Confidence": "LOW",
            "Notes": notes,
        })
    return rows


def parse_uploaded_workbook(file_bytes: bytes) -> List[dict]:
    # Real-world supplier/market-research files (e.g. the initial
    # Woodful Rate Master dataset) use a different, valid column
    # layout - detect and map it before falling back to the primary
    # template's own header set.
    market_range_rows = parse_market_range_workbook(file_bytes)
    if market_range_rows:
        return market_range_rows

    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb.worksheets[0]

    header_row_idx = None
    col_index = None
    for row in ws.iter_rows(min_row=1, max_row=10):
        values = [c.value for c in row]
        resolved = _resolve_header_row(values)
        if resolved:
            header_row_idx = row[0].row
            col_index = resolved
            break
    if header_row_idx is None:
        raise ValueError(
            "Couldn't find the expected column headers in this file. "
            "Please use the downloaded template, or make sure Category, Item Name, UOM, and Source Type "
            "have a recognizable header and aren't duplicated."
        )

    rows = []
    for row in ws.iter_rows(min_row=header_row_idx + 1):
        values = [c.value for c in row]
        if all(_clean(v) is None for v in values):
            continue
        rows.append({name: _clean(values[idx]) if idx < len(values) else None for name, idx in col_index.items()})
    return rows


def validate_row(row: dict, existing_by_code: dict):
    """Validates one parsed row. Returns (result_dict, errors_list,
    is_update); never writes anything. rate_id present but unknown is
    an error (never silently treated as new)."""
    errors = []
    rate_code = row.get("Rate ID")
    matched = None
    if rate_code:
        matched = existing_by_code.get(rate_code.strip().upper())
        if not matched:
            errors.append(f"Rate ID not found: {rate_code!r}. Leave Rate ID blank to create a new rate.")

    for col in REQUIRED_COLUMNS:
        if not row.get(col):
            errors.append(f"{col} is required")

    uom = row.get("UOM")
    category = row.get("Category")
    if uom and uom not in STANDARD_UOMS:
        errors.append(f"'{uom}' is not a standard UOM")
    elif uom and category and not uom_allowed_for_category(category, uom):
        errors.append(f"'{uom}' is not an appropriate unit for category '{category}'")

    source_type_raw = row.get("Source Type")
    source_type = None
    if source_type_raw:
        normalized = source_type_raw.upper().replace(" ", "_")
        if normalized not in RATE_SOURCE_TYPES:
            errors.append(f"'{source_type_raw}' is not a recognized Source Type")
        else:
            source_type = normalized

    confidence_raw = row.get("Confidence")
    confidence = "NOT_VERIFIED"
    if confidence_raw:
        if confidence_raw.upper() not in RATE_CONFIDENCE_LEVELS:
            errors.append(f"'{confidence_raw}' is not a recognized Confidence level")
        else:
            confidence = confidence_raw.upper()

    numeric_fields = {}
    for col in ("Market Reference Rate", "Woodful Cost Rate", "Woodful Selling Rate",
                "Overhead %", "Target Margin %", "Wastage %", "Tax %"):
        raw = row.get(col)
        if raw is None:
            numeric_fields[col] = None
            continue
        val = normalize_number(raw)
        if val is None:
            errors.append(f"Invalid {col}: {raw!r}")
        numeric_fields[col] = val

    result = {
        "matched_rate_id": matched.id if matched else None,
        "category": row.get("Category"), "subcategory": row.get("Subcategory"),
        "item_name": row.get("Item Name"), "specification": row.get("Specification"),
        "location": row.get("Location") or "Indore, Madhya Pradesh", "uom": uom,
        "market_reference_rate": numeric_fields["Market Reference Rate"],
        "woodful_cost_rate": numeric_fields["Woodful Cost Rate"],
        "woodful_selling_rate": numeric_fields["Woodful Selling Rate"],
        "overhead_percent": numeric_fields["Overhead %"], "target_margin_percent": numeric_fields["Target Margin %"],
        "wastage_percent": numeric_fields["Wastage %"], "tax_percent": numeric_fields["Tax %"],
        "source_type": source_type, "source_reference": row.get("Source Reference"),
        "confidence": confidence, "notes": row.get("Notes"),
    }
    return result, errors, matched is not None
