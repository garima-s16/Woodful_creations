"""Reusable Excel bulk-import infrastructure (Family 21 - "Download
Template -> Fill Excel -> Upload -> Validate -> Preview -> Import").

This is the shared, entity-agnostic foundation the brief asks for:
"Prepare reusable infrastructure ... Do not make users enter IDs
manually." Every entity-specific import (currently: Product Master via
api/routes/product_imports.py + utils/product_import.py) builds its own
column list/validation on top of these functions rather than
reimplementing header parsing, number/date normalization, or the
template workbook shape each time.

The existing Purchase Import (utils/purchase_import.py, Family 6) is
left exactly as it was - it already implements this same
template -> upload -> validate -> preview -> commit shape correctly and
has real tests against it; this module is not a rewrite of that
feature, it is the generalized version of the same pattern for
everything built after it. Nothing here ever writes to the database or
generates a business ID - that stays entirely the caller's
responsibility (via the centralized utils/id_generator.py), matching
the standing rule that an import can validate/preview freely but only
an explicit commit step, reviewed by the user, ever creates records.
"""
import re
from io import BytesIO
from decimal import Decimal, InvalidOperation
from datetime import datetime
from typing import List, Optional, Dict

import openpyxl
from openpyxl import Workbook

from app.utils.exporters import write_sheet

_NUMERIC_NOISE_RE = re.compile(r"[₹$€£%,\s]")

_DATE_FORMATS = [
    "%Y-%m-%d", "%Y/%m/%d",
    "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y",
    "%d-%b-%Y", "%d %b %Y", "%d %B %Y",
    "%b %d, %Y", "%B %d, %Y", "%b %d %Y", "%B %d %Y",
]


def normalize_number(raw) -> Optional[Decimal]:
    """Real-world numeric-cell formatting -> Decimal, or None if nothing
    usable is left ("couldn't parse" - never silently defaulted to 0)."""
    if raw is None:
        return None
    if isinstance(raw, (int, float, Decimal)):
        return Decimal(str(raw))
    text = _NUMERIC_NOISE_RE.sub("", str(raw))
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def normalize_date(raw) -> Optional[datetime]:
    """Real-world date-cell formatting -> datetime, or None if nothing in
    _DATE_FORMATS matches (never silently coerced to "today")."""
    if raw is None:
        return None
    if isinstance(raw, datetime):
        return raw
    text = str(raw).strip()
    if not text:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def normalize_match_key(name: Optional[str]) -> str:
    """Case/whitespace-insensitive lookup key, used consistently on both
    sides of any "match this typed name against a real record" step -
    never fuzzy matching, per the standing project principle against
    silently merging records on a guess."""
    return re.sub(r"\s+", " ", (name or "").strip().lower())


def clean_cell(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    return value


def build_header_aliases(required_columns: List[str], extra_aliases: Optional[Dict[str, str]] = None) -> Dict[str, str]:
    """Every canonical column also maps to itself (case-insensitive) -
    callers only need to list genuine *variations* in extra_aliases."""
    aliases = dict(extra_aliases or {})
    for col in required_columns:
        aliases.setdefault(col.lower(), col)
    return aliases


def resolve_header_row(values, required_columns: List[str], aliases: Dict[str, str]) -> Optional[dict]:
    """One worksheet row's raw cell values -> {canonical_name: column_index},
    only if every required column resolves to exactly one cell (never
    guessed, never fuzzy - see module docstring)."""
    resolved = {}
    for idx, raw in enumerate(values):
        if raw is None:
            continue
        key = re.sub(r"\s+", " ", str(raw).strip().lower())
        canonical = aliases.get(key)
        if canonical is None:
            continue
        if canonical in resolved:
            return None
        resolved[canonical] = idx
    if all(col in resolved for col in required_columns):
        return resolved
    return None


def parse_uploaded_workbook(file_bytes: bytes, required_columns: List[str], aliases: Dict[str, str]) -> List[dict]:
    """Reads an uploaded .xlsx against a fixed column list (tolerating
    reasonable header variations via `aliases`), returning one raw dict
    per data row (blank rows skipped). Pure parsing - no validation or
    business-rule matching, so a malformed-file error is always
    distinguishable from a bad-data error."""
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb.worksheets[0]

    header_row_idx = None
    col_index = None
    for row in ws.iter_rows(min_row=1, max_row=10):
        values = [c.value for c in row]
        resolved = resolve_header_row(values, required_columns, aliases)
        if resolved:
            header_row_idx = row[0].row
            col_index = resolved
            break
    if header_row_idx is None:
        raise ValueError(
            "Couldn't find the expected column headers in this file. "
            "Please use the downloaded template, or make sure every required "
            f"column ({', '.join(required_columns)}) has a recognizable "
            "header and isn't duplicated."
        )

    rows = []
    for row in ws.iter_rows(min_row=header_row_idx + 1):
        values = [c.value for c in row]
        if all(clean_cell(v) is None for v in values):
            continue
        rows.append({name: clean_cell(values[idx]) if idx < len(values) else None for name, idx in col_index.items()})
    return rows


def build_import_template(sheet_name: str, title: str, subtitle: str,
                           columns: List[str], example_rows: List[dict]) -> BytesIO:
    """The downloadable .xlsx a user fills in - real Woodful branding
    (via write_sheet, the same helper every export in this app already
    uses), example rows showing the expected format, no totals row
    (this is an input template, not a report)."""
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, sheet_name=sheet_name, title=title, subtitle=subtitle, columns=columns, rows=example_rows)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
