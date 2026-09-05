"""Excel import for Company Holidays. Same discipline as app/utils/material_import.py: one
fixed template shared by the download and the parser, never writes a
parsed row directly to the database (preview/commit are two separate
steps). Simpler than the Material/Product importers since there's no
catalog to fuzzy-match against - a holiday is identified by its date,
and the only duplicate concern is an exact date collision, which is
already how create_holiday/update_holiday enforce uniqueness (one
calendar override per date).
"""
import re
from datetime import datetime
from io import BytesIO
from typing import List, Optional

import openpyxl
from app.shared.import_normalize import enforce_workbook_row_limit
from openpyxl import Workbook

from app.shared.exporters import write_sheet, write_instructions_sheet

TEMPLATE_VERSION = "1.0"

HOLIDAY_IMPORT_COLUMNS = ["Date *", "Holiday Name *", "Type *", "Remarks"]

HEADER_ALIASES = {
    "date": "Date *",
    "holiday name": "Holiday Name *",
    "name": "Holiday Name *",
    "holiday": "Holiday Name *",
    "type": "Type *",
    "remarks": "Remarks",
    "notes": "Remarks",
}
for _col in HOLIDAY_IMPORT_COLUMNS:
    HEADER_ALIASES.setdefault(_col.lower(), _col)
    HEADER_ALIASES.setdefault(_col.rstrip(" *").lower(), _col)

# The Type column is deliberately plain-English, not the raw is_working
# boolean the model actually stores - "Holiday" reads naturally in a
# spreadsheet a Master fills in by hand; is_working=False is the
# implementation detail underneath it.
TYPE_TO_IS_WORKING = {"holiday": False, "special working day": True}

EXAMPLE_ROWS = [
    {"Date *": "2026-08-15", "Holiday Name *": "Independence Day", "Type *": "Holiday", "Remarks": ""},
    {"Date *": "2026-08-09", "Holiday Name *": "Special working Sunday - order backlog",
     "Type *": "Special Working Day", "Remarks": "Declared working to catch up on pending orders"},
]


def build_import_template() -> BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(
        wb, sheet_name="Company Holidays", title="Woodful Creations - Company Holiday Import Template",
        subtitle="Date, Holiday Name, and Type are required for every row (marked with *). "
                  "Each date can only have one calendar override - a date that's already a holiday "
                  "will be flagged during preview, not silently overwritten. Do not change the column headers.",
        columns=HOLIDAY_IMPORT_COLUMNS, rows=EXAMPLE_ROWS,
    )
    write_instructions_sheet(
        wb, template_name="Woodful Company Holiday Import Template", version=TEMPLATE_VERSION,
        field_docs=[
            {"name": "Date *", "mandatory": True, "meaning": "The calendar date this override applies to.",
             "format": "YYYY-MM-DD, e.g. 2026-08-15. One row per date."},
            {"name": "Holiday Name *", "mandatory": True, "meaning": "A short label for the date.",
             "format": "Free text, e.g. 'Independence Day'."},
            {"name": "Type *", "mandatory": True,
             "meaning": "Whether this date removes a working day (Holiday) or adds one back (Special Working Day).",
             "format": "Must be exactly 'Holiday' or 'Special Working Day'."},
            {"name": "Remarks", "mandatory": False, "meaning": "Optional free text note.", "format": "Free text."},
        ],
        duplicate_rule="Each date can only appear once, matching the app's own rule (one calendar "
                        "override per date) - importing a date that already exists as a holiday will "
                        "be flagged during preview, not silently overwritten.",
        extra_notes=[
            "A completely blank row is skipped. A row with only some fields filled in is still "
            "validated - if Date, Holiday Name, or Type is missing, that row will show an error.",
        ],
    )
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _normalize_header_cell(raw) -> Optional[str]:
    if raw is None:
        return None
    key = re.sub(r"\s+", " ", str(raw).strip().lower())
    return HEADER_ALIASES.get(key)


def parse_uploaded_workbook(file_bytes: bytes) -> List[dict]:
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    enforce_workbook_row_limit(wb)
    ws = wb.worksheets[0]

    header_row_idx = None
    col_map = None
    for row in ws.iter_rows(min_row=1, max_row=10):
        values = [c.value for c in row]
        candidate = {}
        ambiguous = False
        for idx, raw in enumerate(values):
            canonical = _normalize_header_cell(raw)
            if canonical:
                if canonical in candidate:
                    # A genuinely duplicated
                    # required header makes this row ambiguous, not
                    # resolvable. Matches the established convention
                    # already used by every other importer in this app
                    # (client/material/product/purchase/estimate/order/
                    # rate_card_import.py all reject the same way) -
                    # this file previously let the last occurrence
                    # silently win, which was the one inconsistent
                    # implementation, found and fixed via that
                    # cross-importer comparison.
                    ambiguous = True
                    break
                candidate[canonical] = idx
        if ambiguous:
            continue
        if {"Date *", "Holiday Name *", "Type *"}.issubset(candidate.keys()):
            header_row_idx = row[0].row
            col_map = candidate
            break

    if header_row_idx is None:
        raise ValueError(
            "Couldn't find the expected column headers in this file. "
            "Please use the downloaded template, or make sure Date, Holiday Name, and Type "
            "have a recognizable header."
        )

    rows = []
    for raw_row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if all(v is None for v in raw_row):
            continue
        row = {}
        for canonical, idx in col_map.items():
            row[canonical] = raw_row[idx] if idx < len(raw_row) else None
        rows.append(row)
    return rows


def validate_row(row: dict, existing_dates: set, seen_dates_in_file: set) -> tuple:
    """Returns (parsed_fields_dict, errors_list). parsed_fields_dict is
    always returned (even with errors) so the preview can show
    whatever was readable; errors indicate this row must not be
    committed as-is."""
    errors = []
    raw_date = row.get("Date *")
    parsed_date = None
    if raw_date in (None, ""):
        errors.append("Date is required")
    else:
        if isinstance(raw_date, datetime):
            parsed_date = raw_date.date()
        else:
            try:
                parsed_date = datetime.strptime(str(raw_date).strip(), "%Y-%m-%d").date()
            except ValueError:
                errors.append(f"Date must be in YYYY-MM-DD format, got {raw_date!r}")

    name = (row.get("Holiday Name *") or "").strip() if row.get("Holiday Name *") else ""
    if not name:
        errors.append("Holiday Name is required")

    raw_type = (row.get("Type *") or "").strip().lower() if row.get("Type *") else ""
    is_working = TYPE_TO_IS_WORKING.get(raw_type)
    if raw_type and is_working is None:
        errors.append(f"Type must be 'Holiday' or 'Special Working Day', got {row.get('Type *')!r}")
    elif not raw_type:
        errors.append("Type is required")

    is_duplicate = False
    if parsed_date is not None:
        if parsed_date in existing_dates:
            is_duplicate = True
        if parsed_date in seen_dates_in_file:
            errors.append(f"Duplicate date within this file: {parsed_date.isoformat()}")
        seen_dates_in_file.add(parsed_date)

    return (
        {
            "date": parsed_date.isoformat() if parsed_date else None,
            "name": name or None,
            "is_working": is_working,
            "remarks": (row.get("Remarks") or None),
            "is_duplicate": is_duplicate,
        },
        errors,
    )
