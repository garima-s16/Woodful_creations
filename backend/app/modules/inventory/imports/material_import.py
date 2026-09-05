"""Excel import for the Material Master. Same
discipline as app/utils/product_import.py: one fixed template shared by
the download and the parser, never writes a parsed row directly to the
database (preview/commit are two separate steps), and reuses the same
shared fuzzy-match logic as every other importer (see
app.modules.clients.matching.find_fuzzy_name_matches) so a typo like
"Fevikol" vs an existing "Fevicol" is caught the same way here as it
would be in the Material UI's own duplicate check.

Stock fields are handled the same way materials.py's own create_material
route already does: only Opening Stock is ever accepted from a row -
current_stock/total_purchased/total_issued are maintained transactionally
by StockService from real Purchase/Issue records and are never
overwritten by an import (section 6).
"""
from io import BytesIO
from decimal import Decimal
from typing import List, Optional

from app.shared.import_common import _clean, resolve_header_row_bare

import openpyxl
from openpyxl import Workbook

from app.shared.exporters import write_sheet, write_instructions_sheet

TEMPLATE_VERSION = "1.0"

MATERIAL_IMPORT_COLUMNS = [
    "Material Name *", "Category", "Subcategory", "Brand / Grade", "Thickness / Size",
    "Unit *", "Supplier", "Opening Stock", "Minimum Stock", "Location", "Active",
]

HEADER_ALIASES = {
    "material": "Material Name *",
    "material name": "Material Name *",
    "name": "Material Name *",
    "category": "Category",
    "material category": "Category",
    "subcategory": "Subcategory",
    "sub category": "Subcategory",
    "material subcategory": "Subcategory",
    "brand": "Brand / Grade",
    "brand / grade": "Brand / Grade",
    "brand/grade": "Brand / Grade",
    "grade": "Brand / Grade",
    "thickness": "Thickness / Size",
    "thickness / size": "Thickness / Size",
    "thickness/size": "Thickness / Size",
    "size": "Thickness / Size",
    "unit": "Unit *",
    "supplier": "Supplier",
    "opening stock": "Opening Stock",
    "minimum stock": "Minimum Stock",
    "min stock": "Minimum Stock",
    "reorder level": "Minimum Stock",
    "location": "Location",
    "active": "Active",
    "is active": "Active",
}
for _col in MATERIAL_IMPORT_COLUMNS:
    HEADER_ALIASES.setdefault(_col.lower(), _col)
    HEADER_ALIASES.setdefault(_col.rstrip(" *").lower(), _col)


REQUIRED_COLUMNS = ["Material Name *", "Unit *"]


from app.shared.import_normalize import normalize_match_key, normalize_number, enforce_workbook_row_limit


TRUE_STRINGS = {"true", "yes", "y", "1", "active"}
FALSE_STRINGS = {"false", "no", "n", "0", "inactive"}

# Approved Woodful demo materials - real materials Woodful actually
# stocks, never generic Test/Demo/Reference placeholders.
EXAMPLE_ROWS = [
    {
        "Material Name *": "Fevicol SH Adhesive", "Category": "", "Subcategory": "",
        "Brand / Grade": "Fevicol", "Thickness / Size": "1 kg can", "Unit *": "Nos", "Supplier": "",
        "Opening Stock": 10, "Minimum Stock": 5, "Location": "", "Active": "Yes",
    },
    {
        "Material Name *": "BWP Plywood 18mm", "Category": "", "Subcategory": "",
        "Brand / Grade": "Century Ply", "Thickness / Size": "18mm", "Unit *": "Sheet", "Supplier": "",
        "Opening Stock": 25, "Minimum Stock": 10, "Location": "", "Active": "Yes",
    },
]

MATERIAL_FIELD_DOCS = [
    {"name": "Material Name *", "mandatory": True, "meaning": "Full name of the material as it should appear in Material Master.",
     "format": "Free text."},
    {"name": "Category", "mandatory": False, "meaning": "Top-level material grouping.",
     "format": "Must match an existing Material Category name exactly, or leave blank."},
    {"name": "Subcategory", "mandatory": False, "meaning": "More specific grouping within the category.",
     "format": "Must match an existing Material Subcategory name exactly, or leave blank."},
    {"name": "Brand / Grade", "mandatory": False, "meaning": "Brand or quality grade.", "format": "Free text."},
    {"name": "Thickness / Size", "mandatory": False, "meaning": "Thickness or size specification, where applicable.", "format": "Free text."},
    {"name": "Unit *", "mandatory": True, "meaning": "The unit this material is stocked/measured in.", "format": "Free text, e.g. Sheet, Kg, Nos, Litre."},
    {"name": "Supplier", "mandatory": False, "meaning": "The material's primary supplier.",
     "format": "Must match an existing Supplier name exactly, or leave blank. Import the Supplier first if it doesn't exist yet."},
    {"name": "Opening Stock", "mandatory": False,
     "meaning": "Starting stock quantity at the time this material is added. Only used once, when the material is created.",
     "format": "Number."},
    {"name": "Minimum Stock", "mandatory": False, "meaning": "Reorder threshold - below this, the material shows as Low Stock.", "format": "Number."},
    {"name": "Location", "mandatory": False, "meaning": "Where this material is stored.",
     "format": "Must match an existing Location name exactly, or leave blank."},
    {"name": "Active", "mandatory": False, "meaning": "Whether this material is in active use.",
     "format": "Yes/No. Defaults to Yes if left blank."},
]


def build_import_template() -> BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(
        wb, sheet_name="Materials", title="Woodful Creations - Material Import Template",
        subtitle="Material Name and Unit are required for every row (marked with *). Material ID is generated "
                  "automatically - do not add a Material ID column. Current Stock is never set here - only "
                  "Opening Stock, used once when the material is created; ongoing stock is maintained through "
                  "Purchases/Issues. Do not change the column headers.",
        columns=MATERIAL_IMPORT_COLUMNS, rows=EXAMPLE_ROWS,
    )
    write_instructions_sheet(
        wb, template_name="Woodful Material Import Template", version=TEMPLATE_VERSION,
        field_docs=MATERIAL_FIELD_DOCS,
        id_rule="Material ID is system-generated and never typed in by hand. Leave any ID column out entirely - "
                "there is none in this template.",
        duplicate_rule="A row is only treated as an existing material if its name matches an existing material "
                        "exactly - it will be reused, not re-created. A row whose name closely resembles (but "
                        "doesn't exactly match) an existing material - e.g. \"Fevikol\" vs an existing \"Fevicol\" "
                        "- is flagged as a possible match during preview; you will be asked to confirm whether to "
                        "use the existing material or create a new one. Nothing is merged or created automatically "
                        "on a possible match.",
        extra_notes=[
            "A completely blank row is skipped. A row with only some fields filled in is still validated - "
            "if Material Name or Unit is missing, that row will show an error.",
            "Current Stock, Total Purchased, and Total Issued are never set through this import - they are "
            "maintained automatically from real Purchase and Issue transactions. Only Opening Stock is used, "
            "and only when the material is first created.",
            "Category, Subcategory, Supplier, and Location must match an existing name exactly if provided - "
            "this import does not create new categories, suppliers, or locations on your behalf. The sample "
            "rows leave these blank since category/subcategory names are set up per business - check Material "
            "Master for the exact names already in use before filling these columns in.",
        ],
    )
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def parse_uploaded_workbook(file_bytes: bytes) -> List[dict]:
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    enforce_workbook_row_limit(wb)
    ws = wb.worksheets[0]

    header_row_idx = None
    col_index = None
    for row in ws.iter_rows(min_row=1, max_row=10):
        values = [c.value for c in row]
        resolved = resolve_header_row_bare(values, HEADER_ALIASES, REQUIRED_COLUMNS)
        if resolved:
            header_row_idx = row[0].row
            col_index = resolved
            break
    if header_row_idx is None:
        raise ValueError(
            "Couldn't find the expected column headers in this file. "
            "Please use the downloaded template, or make sure Material Name and Unit "
            "have a recognizable header and aren't duplicated."
        )

    rows = []
    for row in ws.iter_rows(min_row=header_row_idx + 1):
        values = [c.value for c in row]
        if all(_clean(v) is None for v in values):
            continue
        rows.append({name: _clean(values[idx]) if idx < len(values) else None for name, idx in col_index.items()})
    return rows


def validate_and_match_row(row: dict, existing_by_name: dict, categories_by_name: dict,
                            subcategories_by_name: dict, suppliers_by_name: dict, locations_by_name: dict,
                            fuzzy_candidates: Optional[list] = None):
    """Validates one parsed row, resolves Category/Subcategory/Supplier/
    Location by exact (case-insensitive) name against what already
    exists - never creating new master data on the material's behalf -
    and checks for an existing/possible-match material by name using the
    same shared fuzzy logic every other importer uses. Never
    writes anything."""
    errors = []
    name = row.get("Material Name *")
    if not name:
        errors.append("Material Name is required")

    unit = row.get("Unit *")
    if not unit:
        errors.append("Unit is required")

    category_name = row.get("Category")
    subcategory_id = None
    resolved_category_name = category_name
    subcat_text = row.get("Subcategory")
    if subcat_text:
        subcat = subcategories_by_name.get(normalize_match_key(subcat_text))
        if not subcat:
            errors.append(f"Subcategory not found: {subcat_text!r} - create it in Material Master first, or leave blank")
        else:
            subcategory_id = subcat.id
            resolved_category_name = subcat.category.name
    elif category_name and normalize_match_key(category_name) not in categories_by_name:
        errors.append(f"Category not found: {category_name!r} - create it in Material Master first, or leave blank")

    supplier_id = None
    supplier_text = row.get("Supplier")
    if supplier_text:
        supplier = suppliers_by_name.get(normalize_match_key(supplier_text))
        if not supplier:
            errors.append(f"Supplier not found: {supplier_text!r} - import/create the supplier first, or leave blank")
        else:
            supplier_id = supplier.id

    location_id = None
    location_text = row.get("Location")
    if location_text:
        location = locations_by_name.get(normalize_match_key(location_text))
        if not location:
            errors.append(f"Location not found: {location_text!r} - create it in Locations first, or leave blank")
        else:
            location_id = location.id

    opening_stock, e1 = None, None
    raw_opening = row.get("Opening Stock")
    if raw_opening is not None:
        opening_stock = normalize_number(raw_opening)
        if opening_stock is None:
            e1 = f"Invalid Opening Stock: {raw_opening!r}"
    minimum_stock, e2 = None, None
    raw_min = row.get("Minimum Stock")
    if raw_min is not None:
        minimum_stock = normalize_number(raw_min)
        if minimum_stock is None:
            e2 = f"Invalid Minimum Stock: {raw_min!r}"
    for e in (e1, e2):
        if e:
            errors.append(e)

    raw_active = row.get("Active")
    is_active = True
    if raw_active is not None:
        key = str(raw_active).strip().lower()
        if key in TRUE_STRINGS:
            is_active = True
        elif key in FALSE_STRINGS:
            is_active = False
        else:
            errors.append(f"Active must be Yes/No (got {raw_active!r})")

    matched = existing_by_name.get(normalize_match_key(name)) if name else None

    possible_match = None
    if not matched and name and fuzzy_candidates is not None:
        from app.modules.clients.matching import find_fuzzy_name_matches
        fuzzy_hits = find_fuzzy_name_matches(db=None, name=name, candidates=fuzzy_candidates, limit=1)
        possible_match = fuzzy_hits[0] if fuzzy_hits else None

    result = {
        "name": name, "category": resolved_category_name, "subcategory_id": subcategory_id,
        "brand_grade": row.get("Brand / Grade"), "thickness_size": row.get("Thickness / Size"), "unit": unit,
        "supplier_id": supplier_id, "opening_stock": opening_stock or Decimal("0"),
        "minimum_stock": minimum_stock or Decimal("0"), "location_id": location_id, "is_active": is_active,
        "matched_material_id": matched.id if matched else None,
        "is_duplicate": matched is not None,
        "possible_match_material_id": possible_match.id if possible_match else None,
        "possible_match_name": possible_match.name if possible_match else None,
    }
    return result, errors
