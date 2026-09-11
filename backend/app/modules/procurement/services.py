"""Procurement services: ProcurementService (purchase recording,
requirement/shortage tracking) and purchase Excel import (schemas,
template/workbook parsing, row validation). Combines the former
services.py, imports/purchase_schemas.py, and imports/purchase_import.py."""
from decimal import Decimal, ROUND_HALF_UP
from typing import Optional
from fastapi import HTTPException
from sqlalchemy.orm import Session, selectinload
from app.modules.inventory.models import Material
from app.modules.procurement.models import Purchase, SupplierMaterial
from app.modules.inventory.schemas import PurchaseCreate
from app.modules.communications.services import NotificationService
from app.platform.ids import generate_unique_code, generate_business_id
from pydantic import BaseModel, field_validator
from typing import Optional, List
from decimal import Decimal
from datetime import datetime
from io import BytesIO
from typing import List
import openpyxl
from openpyxl import Workbook
from app.shared import write_sheet
from app.shared_imports import _clean, resolve_header_row_bare
from app.shared_imports import normalize_match_key, normalize_number, normalize_date, enforce_workbook_row_limit


# --- services.py ---
"""ProcurementService: the purchase business record and supplier
recommendation logic that genuinely belongs to Procurement, not
Inventory (Family 130 P0.2 ownership correction, section 11).

record_purchase/mark_purchase_received create/update the Purchase
record and decide WHEN stock should move - but the actual stock
mutation math (weighted-average rate, ledger entry, location
handling) stays in StockService._apply_stock_receipt
(app.modules.inventory.services), since that is genuinely
Inventory's authority ("Inventory performs physical stock mutation").
This file calls into it rather than duplicating it.

StockService imports back from here (_supplier_options_for_materials,
used to enrich calculate_order_material_requirements/
calculate_at_risk_orders with supplier options) - both sides use a
local, function-level import of the other to avoid a circular
top-level import between the two modules.
"""

class ProcurementService:
    @staticmethod
    def record_purchase(db: Session, data: PurchaseCreate) -> Purchase:
        from app.modules.inventory.services import StockService

        material = db.query(Material).filter(Material.id == data.material_id).with_for_update().first()
        if not material:
            raise HTTPException(status_code=404, detail="Material not found")

        purchase_code = generate_unique_code(db, Purchase, "purchase_code", "PUR-")

        taxable_value = (data.quantity * data.rate).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        gst_amount = (taxable_value * data.gst_percent / Decimal("100")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        invoice_total = taxable_value + gst_amount

        purchase = Purchase(
            purchase_code=purchase_code, business_id=generate_business_id(db), date=data.date,
            expected_delivery_date=data.expected_delivery_date, supplier_id=data.supplier_id,
            material_id=data.material_id, quantity=data.quantity, unit=data.unit, rate=data.rate,
            taxable_value=taxable_value, gst_percent=data.gst_percent, gst_amount=gst_amount,
            invoice_total=invoice_total, payment_status=data.payment_status,
            receipt_status=data.receipt_status, location_id=data.location_id,
            quantity_received=data.quantity if data.receipt_status != "Ordered" else Decimal("0"),
        )
        db.add(purchase)
        db.flush()  # assigns purchase.id without committing, needed as the ledger entry's reference_id below

        # "Ordered" means the supplier has been asked for stock that
        # hasn't arrived yet - the material must stay completely
        # untouched until it's explicitly marked received (see
        # mark_purchase_received). "Received" (the default, and the
        # only option that previously existed) increases stock now,
        # exactly as before - existing callers that don't set
        # receipt_status keep their current behavior unchanged.
        if data.receipt_status != "Ordered":
            StockService._apply_stock_receipt(db, material, data.quantity, data.rate, data.supplier_id, data.material_id, reference_id=purchase.id, location_id=data.location_id)
        db.commit()
        db.refresh(purchase)
        if data.receipt_status != "Ordered":
            NotificationService.notify_purchase_received(db, purchase)
        return purchase

    @staticmethod
    def mark_purchase_received(db: Session, purchase_id: int, quantity_to_receive: Optional[Decimal] = None, location_id: Optional[int] = None) -> Purchase:
        """Transitions a purchase toward Received - the point stock
        actually increases. Supports genuine partial receipts:
        quantity_to_receive defaults to everything still outstanding
        (the original all-or-nothing behavior, unchanged for existing
        callers), but a smaller amount can be passed to receive only
        part of the order. Stock is applied only for the incremental
        amount each call - never the full purchase.quantity again -
        so receiving in two steps can never double-count. Rejects
        anything already fully Received, and rejects receiving more
        than genuinely remains outstanding (the quantity-integrity
        check this feature exists for)."""
        from app.modules.inventory.services import StockService

        purchase = db.query(Purchase).filter(Purchase.id == purchase_id).with_for_update().first()
        if not purchase:
            raise HTTPException(status_code=404, detail="Purchase not found")
        if purchase.receipt_status == "Received":
            raise HTTPException(status_code=400, detail="This purchase has already been received.")

        remaining = (purchase.quantity or Decimal("0")) - (purchase.quantity_received or Decimal("0"))
        if quantity_to_receive is None:
            quantity_to_receive = remaining
        if quantity_to_receive <= 0:
            raise HTTPException(status_code=400, detail="Quantity to receive must be positive.")
        if quantity_to_receive > remaining:
            raise HTTPException(
                status_code=400,
                detail=f"Cannot receive {quantity_to_receive} {purchase.unit} - only {remaining} {purchase.unit} remains outstanding on this purchase.",
            )

        material = db.query(Material).filter(Material.id == purchase.material_id).with_for_update().first()
        if not material:
            raise HTTPException(status_code=404, detail="Material not found")

        receive_location_id = location_id or purchase.location_id
        StockService._apply_stock_receipt(db, material, quantity_to_receive, purchase.rate, purchase.supplier_id, purchase.material_id, reference_id=purchase.id, location_id=receive_location_id)
        purchase.quantity_received = (purchase.quantity_received or Decimal("0")) + quantity_to_receive
        purchase.receipt_status = "Received" if purchase.quantity_received >= purchase.quantity else "Partially Received"
        if receive_location_id and not purchase.location_id:
            purchase.location_id = receive_location_id
        db.add(purchase)
        db.commit()
        db.refresh(purchase)
        NotificationService.notify_purchase_received(db, purchase)
        return purchase

    @staticmethod
    def _supplier_options_for_materials(db: Session, material_ids: list, limit_per_material: int = 3) -> dict:
        """Supplier price/lead-time/preference options for a set of
        materials with a genuine shortage (Family 130 section 8: consider
        supplier, lead time, recent price and availability when making a
        recommendation - show which supplier options exist). One bulk
        query regardless of how many materials are passed in, not a
        query per material. Returns {material_id: [option, ...]},
        cheapest/preferred first; a material with no supplier options at
        all is simply absent from the dict."""
        if not material_ids:
            return {}
        rows = (
            db.query(SupplierMaterial)
            .options(selectinload(SupplierMaterial.supplier))
            .filter(SupplierMaterial.material_id.in_(material_ids))
            .all()
        )
        by_material: dict = {}
        for row in rows:
            by_material.setdefault(row.material_id, []).append(row)

        options_by_material = {}
        for material_id, options in by_material.items():
            options.sort(key=lambda r: (
                not r.is_preferred,
                float(r.supplier_price) if r.supplier_price is not None else float("inf"),
            ))
            options_by_material[material_id] = [
                {
                    "supplier_id": r.supplier_id, "supplier_name": r.supplier.name if r.supplier else None,
                    "price": float(r.supplier_price) if r.supplier_price is not None else None,
                    "lead_time_days": r.lead_time_days, "moq": r.moq, "is_preferred": r.is_preferred,
                }
                for r in options[:limit_per_material]
            ]
        return options_by_material


# --- imports/purchase_schemas.py ---
class ImportRowPreview(BaseModel):
    row_number: int
    material_name: Optional[str] = None
    specification: Optional[str] = None
    quantity: Optional[Decimal] = None
    unit: Optional[str] = None
    supplier_name: Optional[str] = None
    rate: Optional[Decimal] = None
    gst_percent: Optional[Decimal] = None
    invoice_number: Optional[str] = None
    invoice_date: Optional[datetime] = None
    remarks: Optional[str] = None
    matched_material_id: Optional[int] = None
    matched_supplier_id: Optional[int] = None
    is_new_material: bool = False
    errors: List[str] = []


class ImportPreviewResponse(BaseModel):
    total_rows: int
    matched_rows: int
    new_material_rows: int
    error_rows: int
    rows: List[ImportRowPreview]


class ImportCommitRow(BaseModel):
    """What the frontend sends back after the user reviews the preview -
    each row echoes the parsed values plus the user's decision for any
    unmatched material (create it, or skip this row)."""
    material_name: Optional[str] = None
    specification: Optional[str] = None
    quantity: Decimal
    unit: str
    matched_material_id: Optional[int] = None
    matched_supplier_id: int
    rate: Decimal
    gst_percent: Decimal = Decimal("0")
    invoice_date: Optional[datetime] = None
    remarks: Optional[str] = None
    create_new_material: bool = False

    @field_validator("quantity")
    @classmethod
    def quantity_must_be_positive(cls, v: Decimal) -> Decimal:
        if v <= 0:
            raise ValueError("Quantity must be greater than zero")
        return v

    @field_validator("rate")
    @classmethod
    def rate_not_negative(cls, v: Decimal) -> Decimal:
        if v < 0:
            raise ValueError("Rate cannot be negative")
        return v

    @field_validator("gst_percent")
    @classmethod
    def gst_percent_in_range(cls, v: Decimal) -> Decimal:
        if v < 0 or v > 100:
            raise ValueError("GST percent must be between 0 and 100")
        return v


class ImportCommitRequest(BaseModel):
    rows: List[ImportCommitRow]


class ImportCommitResult(BaseModel):
    created_materials: int
    created_purchases: int
    purchase_ids: List[int]
    error: Optional[str] = None


# --- imports/purchase_import.py ---
"""Excel import for purchases - one fixed template, so the download
(what the user fills in) and the parser (what reads it back) can never
drift out of sync, since both reference this same column list.

Never writes a parsed row directly to the database - every upload goes
through preview() first (matches against real Material/Supplier
records, validates each row, never commits), and only commit() actually
creates records, and only for rows the caller explicitly confirms.
"""

PURCHASE_IMPORT_COLUMNS = [
    "Material", "Specification", "Quantity", "Unit", "Supplier",
    "Rate", "GST %", "Invoice Number", "Invoice Date", "Remarks",
]


HEADER_ALIASES = {
    "material": "Material",
    "material name": "Material",
    "specification": "Specification",
    "spec": "Specification",
    "specifications": "Specification",
    "quantity": "Quantity",
    "qty": "Quantity",
    "unit": "Unit",
    "units": "Unit",
    "supplier": "Supplier",
    "supplier name": "Supplier",
    "rate": "Rate",
    "invoice number": "Invoice Number",
    "invoice no": "Invoice Number",
    "invoice no.": "Invoice Number",
    "invoice #": "Invoice Number",
    "invoice date": "Invoice Date",
    "date": "Invoice Date",
    "invoice dt": "Invoice Date",
    "remarks": "Remarks",
    "remark": "Remarks",
    "gst": "GST %",
    "gst%": "GST %",
    "gst %": "GST %",
}


for _col in PURCHASE_IMPORT_COLUMNS:
    HEADER_ALIASES.setdefault(_col.lower(), _col)


EXAMPLE_ROWS = [
    {
        "Material": "HDHMR 18mm", "Specification": "Greenpanel, 8x4 ft", "Quantity": 5, "Unit": "Sheets",
        "Supplier": "Sanket Plywood & Boards", "Rate": 1820, "GST %": 18,
        "Invoice Number": "INV-1001", "Invoice Date": "2026-08-01", "Remarks": "",
    },
    {
        "Material": "BWP Plywood 18mm", "Specification": "Century, 8x4 ft", "Quantity": 6, "Unit": "Sheets",
        "Supplier": "Century Plywood Dealer", "Rate": 2350, "GST %": 18,
        "Invoice Number": "INV-1001", "Invoice Date": "2026-08-01", "Remarks": "Delivered with Aug 1 batch",
    },
]


def build_import_template() -> BytesIO:
    """The downloadable .xlsx a user fills in - real Woodful branding
    (via write_sheet, the same helper every other export uses), two
    example rows showing the expected format, and nothing else - no
    totals row, since this is an input template, not a report."""
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(
        wb, sheet_name="Purchase Import", title="Woodful Creations - Purchase Import Template",
        subtitle="Fill in one row per material purchased. Do not change the column headers.",
        columns=PURCHASE_IMPORT_COLUMNS, rows=EXAMPLE_ROWS,
    )
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def parse_uploaded_workbook(file_bytes: bytes) -> List[dict]:
    """Reads the uploaded file against the fixed column list (tolerating
    reasonable header variations - see HEADER_ALIASES/_resolve_header_row),
    returning one raw dict per row (empty rows skipped). Does no value
    normalization or matching/validation here - that's preview()'s job,
    kept separate so parsing failures (a genuinely malformed file, or one
    where a required column can't be confidently identified) are
    distinguishable from validation failures (a well-formed file with bad
    data)."""
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    enforce_workbook_row_limit(wb)
    ws = wb.worksheets[0]

    header_row_idx = None
    col_index = None
    for row in ws.iter_rows(min_row=1, max_row=10):
        values = [c.value for c in row]
        resolved = resolve_header_row_bare(values, HEADER_ALIASES, PURCHASE_IMPORT_COLUMNS)
        if resolved:
            header_row_idx = row[0].row
            col_index = resolved
            break
    if header_row_idx is None:
        raise ValueError(
            "Couldn't find the expected column headers in this file. "
            "Please use the downloaded template, or make sure every required "
            "column (Material, Quantity, Unit, Supplier, Rate, GST %, "
            "Invoice Number, Invoice Date, Specification, Remarks) has a "
            "recognizable header and isn't duplicated."
        )

    rows = []
    for row in ws.iter_rows(min_row=header_row_idx + 1):
        values = [c.value for c in row]
        if all(_clean(v) is None for v in values):
            continue  # skip fully blank rows
        rows.append({name: _clean(values[idx]) if idx < len(values) else None for name, idx in col_index.items()})
    return rows


def validate_and_match_row(row: dict, db, materials_by_name, suppliers_by_name):
    """Validates one parsed row and matches it against real Material/
    Supplier records (case-insensitive exact match on name - not fuzzy,
    per the standing project principle against silently merging records
    on a guess). Returns (result_dict, errors_list). Never writes
    anything - this is pure validation/matching."""
    errors = []
    material_name = row.get("Material")
    supplier_name = row.get("Supplier")

    if not material_name:
        errors.append("Material name is required")
    if not supplier_name:
        errors.append("Supplier is required")

    quantity = None
    if row.get("Quantity") is None:
        errors.append("Quantity is required")
    else:
        # normalize_number handles "1,250" / " 1250 " / currency symbols
        # (see HEADER/VALUE NORMALIZATION note above) before the
        # existing >0 check - an invalid quantity is still rejected the
        # same as before, just after stripping harmless formatting.
        quantity = normalize_number(row["Quantity"])
        if quantity is None:
            errors.append(f"Invalid quantity: {row['Quantity']!r}")
        elif quantity <= 0:
            errors.append("Quantity must be greater than zero")

    rate = None
    if row.get("Rate") is None:
        errors.append("Rate is required")
    else:
        rate = normalize_number(row["Rate"])
        if rate is None:
            errors.append(f"Invalid rate: {row['Rate']!r}")
        elif rate < 0:
            errors.append("Rate cannot be negative")

    gst_percent = Decimal("0")
    if row.get("GST %") is not None:
        normalized_gst = normalize_number(row["GST %"])
        if normalized_gst is None:
            errors.append(f"Invalid GST %: {row['GST %']!r}")
        else:
            gst_percent = normalized_gst

    invoice_date = None
    if row.get("Invoice Date"):
        raw = row["Invoice Date"]
        invoice_date = normalize_date(raw)
        if invoice_date is None:
            errors.append(f"Invoice Date must be a recognizable date (e.g. 2026-08-01 or 01-08-2026), got {raw!r}")

    matched_material = materials_by_name.get(normalize_match_key(material_name)) if material_name else None
    matched_supplier = suppliers_by_name.get(normalize_match_key(supplier_name)) if supplier_name else None
    if material_name and not matched_material:
        errors.append(f'No existing material matches "{material_name}" - it can be created on import, or create it first and re-upload.')
    if supplier_name and not matched_supplier:
        errors.append(f'No existing supplier matches "{supplier_name}" - create it on the Suppliers page first, then re-upload.')

    result = {
        "material_name": material_name, "specification": row.get("Specification"),
        "quantity": quantity, "unit": row.get("Unit") or "Nos", "supplier_name": supplier_name,
        "rate": rate, "gst_percent": gst_percent, "invoice_number": row.get("Invoice Number"),
        "invoice_date": invoice_date, "remarks": row.get("Remarks"),
        "matched_material_id": matched_material.id if matched_material else None,
        "matched_supplier_id": matched_supplier.id if matched_supplier else None,
        "is_new_material": bool(material_name) and not matched_material,
    }
    return result, errors
