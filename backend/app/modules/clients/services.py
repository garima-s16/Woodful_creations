"""Client domain services: Pydantic schemas, import processing
(header normalization, template/workbook parsing, row validation),
name/phone-based fuzzy matching, and PDF export. Combines the former
schemas (from models.py), import_utils.py, matching.py, and
pdf_generator.py."""
import re
import difflib
import hashlib
from io import BytesIO
from datetime import datetime
from decimal import Decimal
from typing import List, Optional, Tuple
import openpyxl
from openpyxl import Workbook
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError
from sqlalchemy import text as sa_text
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from pydantic import BaseModel, field_validator, model_validator
from app.shared_imports import enforce_workbook_row_limit
from app.shared import write_sheet, write_instructions_sheet, add_dropdown_validation
from app.shared import validate_phone, validate_email
from app.shared import (
    format_inr, get_styles, build_header, section_table, line_items_table, build_footer_text,
    pdf_text, fmt_date, mask_phone,
)
from app.modules.clients.models import Client, CLIENT_TYPES, CLIENT_STATUSES, ACTIVITY_TYPES
from app.platform.ids import generate_unique_code, generate_business_id


# --- Pydantic schemas (formerly in models.py) ---


class ClientBase(BaseModel):
    client_code: Optional[str] = None  # server-generated on create, ignored if supplied
    name: str
    # Mandatory in the Client form and Excel importer, which both
    # always supply it and never let a user skip past it. A default is
    # still kept here (rather than making this field itself required
    # with no default) so existing integrations/tests created before
    # this field existed keep working instead of failing with a hard
    # 422. Any value that IS supplied is still validated.
    client_type: str = "Individual"
    contact_person: Optional[str] = None
    alternate_phone: Optional[str] = None
    # Mandatory - a client record with no way to reach them isn't
    # usable. Must be EXACTLY 10 digits - no country code, no
    # separators. Every rejection reason (missing, wrong length,
    # non-numeric, formatted) uses the exact same message, by explicit
    # requirement: "Please enter valid mobile number".
    phone: str
    email: Optional[str] = None
    address: Optional[str] = None
    site_address: Optional[str] = None
    city: Optional[str] = None
    state: Optional[str] = None
    pincode: Optional[str] = None
    gstin: Optional[str] = None
    status: str = "Active"
    lead_source: Optional[str] = None
    first_contact_date: Optional[datetime] = None
    remarks: Optional[str] = None


class ClientCreate(ClientBase):
    @field_validator("client_type")
    @classmethod
    def client_type_must_be_valid(cls, v: str) -> str:
        v = (v or "").strip()
        if v not in CLIENT_TYPES:
            raise ValueError(f"Client Type must be one of: {', '.join(CLIENT_TYPES)}")
        return v

    @field_validator("phone")
    @classmethod
    def phone_must_be_valid(cls, v: str) -> str:
        v = (v or "").strip()
        if not validate_phone(v):
            raise ValueError("Please enter valid mobile number")
        return v

    @field_validator("email")
    @classmethod
    def email_must_be_valid(cls, v: Optional[str]) -> Optional[str]:
        if v and not validate_email(v):
            raise ValueError("Enter a valid email address.")
        return v

    @field_validator("gstin")
    @classmethod
    def gstin_must_be_15_characters(cls, v: Optional[str]) -> Optional[str]:
        if v and len(v) != 15:
            raise ValueError("GSTIN must contain 15 characters.")
        return v

    @field_validator("status")
    @classmethod
    def status_must_be_valid(cls, v: str) -> str:
        v = (v or "").strip()
        if v not in CLIENT_STATUSES:
            raise ValueError(f"Status must be one of: {', '.join(CLIENT_STATUSES)}")
        return v


class ClientUpdate(BaseModel):
    name: Optional[str] = None
    client_type: Optional[str] = None
    contact_person: Optional[str] = None
    phone: Optional[str] = None
    alternate_phone: Optional[str] = None
    email: Optional[str] = None
    address: Optional[str] = None
    site_address: Optional[str] = None
    city: Optional[str] = None
    state: Optional[str] = None
    pincode: Optional[str] = None
    gstin: Optional[str] = None
    status: Optional[str] = None
    lead_source: Optional[str] = None
    remarks: Optional[str] = None

    @field_validator("client_type")
    @classmethod
    def client_type_must_be_valid(cls, v: Optional[str]) -> Optional[str]:
        if v is None:
            return v
        v = v.strip()
        if v not in CLIENT_TYPES:
            raise ValueError(f"Client Type must be one of: {', '.join(CLIENT_TYPES)}")
        return v

    @field_validator("phone")
    @classmethod
    def phone_must_be_valid(cls, v: Optional[str]) -> Optional[str]:
        # Phone is mandatory on the entity - an update is allowed to
        # leave it unset (unchanged), but never to explicitly clear it
        # to empty, and any value supplied must still be exactly 10
        # digits. Same exact message for every rejection reason.
        if v is None:
            return v
        v = v.strip()
        if not validate_phone(v):
            raise ValueError("Please enter valid mobile number")
        return v

    @field_validator("email")
    @classmethod
    def email_must_be_valid(cls, v: Optional[str]) -> Optional[str]:
        if v and not validate_email(v):
            raise ValueError("Enter a valid email address.")
        return v

    @field_validator("gstin")
    @classmethod
    def gstin_must_be_15_characters(cls, v: Optional[str]) -> Optional[str]:
        if v and len(v) != 15:
            raise ValueError("GSTIN must contain 15 characters.")
        return v

    @field_validator("status")
    @classmethod
    def status_must_be_valid(cls, v: Optional[str]) -> Optional[str]:
        if v is None:
            return v
        v = v.strip()
        if v not in CLIENT_STATUSES:
            raise ValueError(f"Status must be one of: {', '.join(CLIENT_STATUSES)}")
        return v


class ClientResponse(ClientBase):
    id: int
    business_id: Optional[str] = None
    created_at: datetime
    updated_at: datetime

    class Config:
        from_attributes = True


class ClientWithStats(ClientResponse):
    total_orders: int
    total_sales: Optional[float] = None


class ClientActivityBase(BaseModel):
    client_id: int
    activity_type: str
    date: datetime
    summary: str
    logged_by: Optional[str] = None
    follow_up_date: Optional[datetime] = None


class ClientActivityCreate(ClientActivityBase):
    @field_validator("activity_type")
    @classmethod
    def activity_type_must_be_valid(cls, v: str) -> str:
        v = (v or "").strip()
        if v not in ACTIVITY_TYPES:
            raise ValueError(f"Activity Type must be one of: {', '.join(ACTIVITY_TYPES)}")
        return v


class ClientActivityUpdate(BaseModel):
    activity_type: Optional[str] = None
    date: Optional[datetime] = None
    summary: Optional[str] = None
    logged_by: Optional[str] = None
    follow_up_date: Optional[datetime] = None

    @field_validator("activity_type")
    @classmethod
    def activity_type_must_be_valid(cls, v: Optional[str]) -> Optional[str]:
        if v is None:
            return v
        v = v.strip()
        if v not in ACTIVITY_TYPES:
            raise ValueError(f"Activity Type must be one of: {', '.join(ACTIVITY_TYPES)}")
        return v


class ClientActivityResponse(ClientActivityBase):
    id: int
    created_at: datetime
    follow_up_done: bool = False

    class Config:
        from_attributes = True


class ClientDocumentResponse(BaseModel):
    id: int
    client_id: int
    original_filename: str
    content_type: Optional[str] = None
    description: Optional[str] = None
    uploaded_by: Optional[str] = None
    created_at: datetime

    class Config:
        from_attributes = True


class ClientEmailPreview(BaseModel):
    """What the frontend shows before sending - recipient/subject/body
    are all editable by the user before they confirm ("review
    recipient/subject/message before sending"). client_has_email
    distinguishes "no email on file, user must type one" from "client
    has an email, pre-filled and editable" - the frontend needs this
    to decide whether to show a warning."""
    recipient_email: Optional[str] = None
    client_has_email: bool
    subject: str
    body: str
    attachment_filename: str


class ClientEmailSendRequest(BaseModel):
    recipient_email: str
    subject: str
    body: str


class ClientEmailSendResult(BaseModel):
    sent: bool
    message: str


class ClientProductRateBase(BaseModel):
    client_id: int
    product_id: Optional[int] = None  # None = client-wide default margin, applies to every product for this client
    margin_percent: Optional[Decimal] = None
    fixed_selling_price: Optional[Decimal] = None
    notes: Optional[str] = None


class ClientProductRateCreate(ClientProductRateBase):
    @model_validator(mode="after")
    def exactly_one_override_type(self):
        # A customer override is either a margin adjustment or a
        # direct negotiated price, never presented as both at once.
        if self.margin_percent is not None and self.fixed_selling_price is not None:
            raise ValueError("Set either margin_percent or fixed_selling_price, not both.")
        if self.margin_percent is None and self.fixed_selling_price is None:
            raise ValueError("Set either margin_percent or fixed_selling_price.")
        if self.margin_percent is not None and self.margin_percent >= 100:
            raise ValueError("Margin percent must be less than 100%.")
        if self.fixed_selling_price is not None and self.fixed_selling_price < 0:
            raise ValueError("Fixed selling price cannot be negative.")
        return self


class ClientProductRateUpdate(BaseModel):
    margin_percent: Optional[Decimal] = None
    fixed_selling_price: Optional[Decimal] = None
    notes: Optional[str] = None

    @field_validator("margin_percent")
    @classmethod
    def margin_percent_must_be_under_100(cls, v: Optional[Decimal]) -> Optional[Decimal]:
        if v is not None and v >= 100:
            raise ValueError("Margin percent must be less than 100%.")
        return v

    @field_validator("fixed_selling_price")
    @classmethod
    def fixed_selling_price_not_negative(cls, v: Optional[Decimal]) -> Optional[Decimal]:
        if v is not None and v < 0:
            raise ValueError("Fixed selling price cannot be negative.")
        return v


class ClientProductRateResponse(ClientProductRateBase):
    id: int
    created_by: Optional[str] = None
    created_at: datetime
    updated_at: datetime

    class Config:
        from_attributes = True


class PricingResolveRequest(BaseModel):
    """Preview what rate would apply, and which rule would win, before
    actually saving a line item - lets the Estimate UI show "Calculated
    Price" vs "Manually Overridden Price" clearly."""
    product_id: Optional[int] = None
    client_id: Optional[int] = None
    estimate_id: Optional[int] = None
    explicit_override: Optional[Decimal] = None

    @field_validator("explicit_override")
    @classmethod
    def explicit_override_not_negative(cls, v: Optional[Decimal]) -> Optional[Decimal]:
        if v is not None and v < 0:
            raise ValueError("Explicit override cannot be negative.")
        return v


class PricingResolveResponse(BaseModel):
    selling_rate: Decimal
    pricing_rule_applied: str
    margin_percent_used: Optional[Decimal] = None
    cost_used: Optional[Decimal] = None


class ClientImportRowPreview(BaseModel):
    row_number: int
    name: Optional[str] = None
    client_type: Optional[str] = None
    contact_person: Optional[str] = None
    phone: Optional[str] = None
    alternate_phone: Optional[str] = None
    email: Optional[str] = None
    address: Optional[str] = None
    site_address: Optional[str] = None
    city: Optional[str] = None
    state: Optional[str] = None
    pincode: Optional[str] = None
    gstin: Optional[str] = None
    lead_source: Optional[str] = None
    remarks: Optional[str] = None
    matched_client_id: Optional[int] = None
    is_duplicate: bool = False
    # A name that closely resembles (but doesn't exactly match) an
    # existing client - never auto-resolved. The caller must show "Use
    # Existing" (set matched_client_id to this value on commit) /
    # "Create New" (leave as-is) / "Edit" to the user.
    possible_match_client_id: Optional[int] = None
    possible_match_name: Optional[str] = None
    errors: List[str] = []


class ClientImportPreviewResponse(BaseModel):
    total_rows: int
    new_rows: int
    duplicate_rows: int
    error_rows: int
    rows: List[ClientImportRowPreview]


class ClientImportCommitRow(BaseModel):
    """Echoes one reviewed row back for actual creation. matched_client_id
    (carried over from the preview response) means this row matched an
    existing client on BOTH name and phone - it is never re-created;
    it's counted as "existing client matched" and no new Client ID is
    generated, per the Client Recognition rule. A row without a match
    is only created if the caller explicitly leaves skip=False for it."""
    name: str
    client_type: Optional[str] = None
    contact_person: Optional[str] = None
    phone: Optional[str] = None
    alternate_phone: Optional[str] = None
    email: Optional[str] = None
    address: Optional[str] = None
    site_address: Optional[str] = None
    city: Optional[str] = None
    state: Optional[str] = None
    pincode: Optional[str] = None
    gstin: Optional[str] = None
    lead_source: Optional[str] = None
    remarks: Optional[str] = None
    matched_client_id: Optional[int] = None
    skip: bool = False  # user chose not to import this row (e.g. unresolved duplicate)

    @model_validator(mode="after")
    def validate_when_not_skipped(self):
        # A row marked skip=True is never written to the database, so
        # it doesn't need to pass the same validity checks a real
        # client creation would - the person may be intentionally
        # leaving a broken/duplicate row behind unfixed. A matched row
        # is never written either (it reuses the existing client
        # as-is), so it doesn't need re-validation of fields that
        # already belong to a valid, existing record.
        if self.skip or self.matched_client_id:
            return self
        if not self.client_type or self.client_type not in CLIENT_TYPES:
            raise ValueError(f"Client Type must be one of: {', '.join(CLIENT_TYPES)}")
        phone = (self.phone or "").strip()
        if not phone or not validate_phone(phone):
            raise ValueError("Please enter valid mobile number")
        self.phone = phone
        if self.email and not validate_email(self.email):
            raise ValueError("Enter a valid email address.")
        if self.gstin and len(self.gstin) != 15:
            raise ValueError("GSTIN must contain 15 characters.")
        return self


class ClientImportCommitRequest(BaseModel):
    rows: List[ClientImportCommitRow]


class ClientImportCommitResult(BaseModel):
    created_clients: int
    matched_existing: int
    skipped: int
    client_ids: List[int]
    error: Optional[str] = None


# --- import_utils.py ---
"""Excel import for the Client Master. Same discipline as the Product/
Purchase importers: one fixed template shared by the download and the
parser, phone/GSTIN/email validated with the exact same rules the API
enforces (so an import can't create a client the API itself would have
rejected), never writes a parsed row directly to the database -
preview and commit are two separate steps, and uploaded Client IDs are
never trusted (the column doesn't even exist in the template - IDs are
always server-generated).
"""

TEMPLATE_VERSION = "2.0"  # bumped from 1.0: adds Client Type (mandatory), State, Pincode


CLIENT_IMPORT_COLUMNS = [
    "Client Name *", "Client Type *", "Contact Person", "Phone *", "Alternate Phone", "Email",
    "Address", "Site Address", "City", "State", "Pincode", "GSTIN", "Lead Source", "Notes",
]


HEADER_ALIASES = {
    "client": "Client Name *",
    "client name": "Client Name *",
    "name": "Client Name *",
    "client type": "Client Type *",
    "type": "Client Type *",
    "contact person": "Contact Person",
    "contact": "Contact Person",
    "phone": "Phone *",
    "phone number": "Phone *",
    "mobile": "Phone *",
    "alternate phone": "Alternate Phone",
    "alt phone": "Alternate Phone",
    "email": "Email",
    "address": "Address",
    "billing address": "Address",
    "site address": "Site Address",
    "city": "City",
    "state": "State",
    "pincode": "Pincode",
    "pin code": "Pincode",
    "zip": "Pincode",
    "gstin": "GSTIN",
    "lead source": "Lead Source",
    "source": "Lead Source",
    "notes": "Notes",
    "remarks": "Notes",
}


for _col in CLIENT_IMPORT_COLUMNS:
    HEADER_ALIASES.setdefault(_col.lower(), _col)
    HEADER_ALIASES.setdefault(_col.rstrip(" *").lower(), _col)


REQUIRED_COLUMNS = ["Client Name *", "Client Type *", "Phone *"]


def _normalize_header_cell(raw) -> Optional[str]:
    if raw is None:
        return None
    key = re.sub(r"\s+", " ", str(raw).strip().lower())
    return HEADER_ALIASES.get(key)


def _resolve_header_row(values) -> Optional[dict]:
    resolved = {}
    for idx, raw in enumerate(values):
        canonical = _normalize_header_cell(raw)
        if canonical is None:
            continue
        if canonical in resolved:
            return None
        resolved[canonical] = idx
    if all(col in resolved for col in REQUIRED_COLUMNS):
        return resolved
    return None


def normalize_match_key(name: Optional[str], phone: Optional[str]) -> str:
    """A client is treated as a match ONLY if BOTH name and phone match
    an existing client - this is the same Client Recognition rule used
    for order intake (see find_or_create_client below, in this same
    file), reused here rather than re-implemented, so the two entry points can never
    silently drift apart on what "matches" means. Name alone is too
    common (multiple "Ravi"s are plausible); phone alone could be a
    shared family/office line; both together is the confirmed rule."""
    return f"{normalize_name(name)}|{normalize_phone(phone)}"


EXAMPLE_ROWS = [
    {
        "Client Name *": "Sanket", "Client Type *": "Individual", "Contact Person": "", "Phone *": "9812345670",
        "Alternate Phone": "", "Email": "sanket@example.com", "Address": "123 Vijay Nagar, Indore",
        "Site Address": "Same as address", "City": "Indore", "State": "Madhya Pradesh", "Pincode": "452010",
        "GSTIN": "", "Lead Source": "Referral", "Notes": "Sample row - delete before uploading your real data",
    },
    {
        "Client Name *": "Meenal Interiors Pvt Ltd", "Client Type *": "Business", "Contact Person": "Meenal",
        "Phone *": "9823456781", "Alternate Phone": "9834567892", "Email": "meenal@example.com",
        "Address": "456 Race Course Road, Indore", "Site Address": "789 AB Road, Indore",
        "City": "Indore", "State": "Madhya Pradesh", "Pincode": "452001", "GSTIN": "23ABCDE1234F1Z5",
        "Lead Source": "Architect", "Notes": "Sample row with GSTIN - delete before uploading your real data",
    },
]


CLIENT_FIELD_DOCS = [
    {"name": "Client Name *", "mandatory": True, "meaning": "Full name of the client or company.",
     "format": "Free text."},
    {"name": "Client Type *", "mandatory": True, "meaning": "Whether the client is a person or a company.",
     "format": "Must match exactly.", "accepted_values": list(CLIENT_TYPES)},
    {"name": "Contact Person", "mandatory": False,
     "meaning": "Who to actually call, when the client is a company or family rather than a single individual.",
     "format": "Free text."},
    {"name": "Phone *", "mandatory": True, "meaning": "Primary contact number.",
     "format": "Exactly 10 digits, no country code or separators."},
    {"name": "Alternate Phone", "mandatory": False, "meaning": "A second contact number, if any.",
     "format": "10 digits."},
    {"name": "Email", "mandatory": False, "meaning": "Contact email.", "format": "Valid email address."},
    {"name": "Address", "mandatory": False, "meaning": "Billing address.", "format": "Free text."},
    {"name": "Site Address", "mandatory": False,
     "meaning": "Where the work actually happens, if different from the billing address.", "format": "Free text."},
    {"name": "City", "mandatory": False, "meaning": "City.", "format": "Free text."},
    {"name": "State", "mandatory": False, "meaning": "State.", "format": "Free text."},
    {"name": "Pincode", "mandatory": False, "meaning": "Postal code.", "format": "Up to 10 digits."},
    {"name": "GSTIN", "mandatory": False, "meaning": "GST registration number, if the client has one.",
     "format": "Exactly 15 characters."},
    {"name": "Lead Source", "mandatory": False, "meaning": "How this client found Woodful.",
     "format": "Free text - use the same values as the Client form's Lead Source list."},
    {"name": "Notes", "mandatory": False, "meaning": "Any other remarks about the client.", "format": "Free text."},
]


def build_import_template() -> BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    ws = write_sheet(
        wb, sheet_name="Clients", title="Woodful Creations - Client Import Template",
        subtitle="Client Name, Client Type, and Phone are required for every row (marked with *). "
                  "Client ID is generated automatically - do not add a Client ID column. "
                  "Do not change the column headers.",
        columns=CLIENT_IMPORT_COLUMNS, rows=EXAMPLE_ROWS,
    )
    type_col_letter = chr(ord("A") + CLIENT_IMPORT_COLUMNS.index("Client Type *"))
    add_dropdown_validation(ws, type_col_letter, CLIENT_TYPES, first_row=ws.woodful_header_row + 1)

    write_instructions_sheet(
        wb, template_name="Woodful Clients Import Template", version=TEMPLATE_VERSION,
        field_docs=CLIENT_FIELD_DOCS,
        id_rule="Client ID is system-generated and never typed in by hand. Leave any ID column out entirely - "
                "there is none in this template.",
        duplicate_rule="A row is only treated as an existing client if BOTH Client Name and Phone match an "
                        "existing client exactly - it will be reused, not re-created. A row whose name closely "
                        "resembles (but doesn't exactly match) an existing client is flagged as a possible match "
                        "during preview; you will be asked to confirm whether to use the existing client or create "
                        "a new one. Nothing is merged or created automatically on a possible match.",
        extra_notes=[
            "A completely blank row is skipped. A row with only some fields filled in is still validated - "
            "if Client Name, Client Type, or Phone is missing, that row will show an error.",
            "This template does not use a Client ID column. An existing client is recognized only when both "
            "Client Name and Phone match exactly (see Duplicate / Typo Handling above).",
        ],
    )
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    if isinstance(value, (int, float)):
        # A phone/GSTIN typed into Excel can get read back as a number
        # (e.g. Excel auto-formats a bare digit string) - stringify it
        # rather than silently dropping it as an invalid type.
        return str(int(value)) if float(value).is_integer() else str(value)
    return value


def parse_uploaded_workbook(file_bytes: bytes) -> List[dict]:
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    enforce_workbook_row_limit(wb)
    ws = wb.worksheets[0]

    header_row_idx = None
    col_index = None
    for row in ws.iter_rows(min_row=1, max_row=10):
        values = [c.value for c in row]
        resolved = _resolve_header_row(values)
        if resolved:
            header_row_idx = row[0].row
            col_index = resolved
            break
    if header_row_idx is None:
        raise ValueError(
            "Couldn't find the expected column headers in this file. "
            "Please use the downloaded template, or make sure Client Name and Phone "
            "have a recognizable header and aren't duplicated."
        )

    rows = []
    for row in ws.iter_rows(min_row=header_row_idx + 1):
        values = [c.value for c in row]
        if all(_clean(v) is None for v in values):
            continue
        rows.append({name: _clean(values[idx]) if idx < len(values) else None for name, idx in col_index.items()})
    return rows


def validate_and_match_row(row: dict, existing_by_key: dict, fuzzy_candidates: Optional[list] = None):
    """Validates one parsed row with the exact same rules the Client
    API enforces (validate_phone/validate_email/15-char GSTIN), and
    checks whether it looks like a duplicate of an existing client -
    both an exact (name+phone) match and, when fuzzy_candidates is
    supplied, a "possible match" typo/near-match nudge,
    using the same shared logic as the UI's check-duplicates
    endpoint (see find_fuzzy_name_matches (same module)).
    Returns (result_dict, errors_list); never writes anything."""
    errors = []
    name = row.get("Client Name *")
    if not name:
        errors.append("Client Name is required")

    raw_client_type = row.get("Client Type *")
    client_type = str(raw_client_type).strip() if raw_client_type else raw_client_type
    if not client_type:
        errors.append("Client Type is required")
    elif client_type not in CLIENT_TYPES:
        errors.append(f"Client Type must be one of: {', '.join(CLIENT_TYPES)} (got {client_type!r})")

    phone = row.get("Phone *")
    if not phone or not validate_phone(str(phone).strip()):
        errors.append("Please enter valid mobile number")

    email = row.get("Email")
    if email and not validate_email(email):
        errors.append(f"Invalid email: {email!r}")

    gstin = row.get("GSTIN")
    if gstin and len(gstin) != 15:
        errors.append(f"GSTIN must contain 15 characters (got {len(gstin)}): {gstin!r}")

    matched = existing_by_key.get(normalize_match_key(name, phone)) if name and phone else None

    possible_match = None
    if not matched and name and fuzzy_candidates is not None:
        pass  # find_fuzzy_name_matches now defined in this same module
        fuzzy_hits = find_fuzzy_name_matches(db=None, name=name, candidates=fuzzy_candidates, limit=1)
        possible_match = fuzzy_hits[0] if fuzzy_hits else None

    result = {
        "name": name, "client_type": client_type, "contact_person": row.get("Contact Person"), "phone": phone,
        "alternate_phone": row.get("Alternate Phone"), "email": email, "address": row.get("Address"),
        "site_address": row.get("Site Address"), "city": row.get("City"), "state": row.get("State"),
        "pincode": row.get("Pincode"), "gstin": gstin,
        "lead_source": row.get("Lead Source"), "remarks": row.get("Notes"),
        "matched_client_id": matched.id if matched else None,
        "is_duplicate": matched is not None,
        "possible_match_client_id": possible_match.id if possible_match else None,
        "possible_match_name": possible_match.name if possible_match else None,
    }
    return result, errors


# --- matching.py ---
"""Client recognition for order intake.

Business rule: when a new order is entered with a client name + phone
(rather than an already-selected client_id), decide whether this is the
same client returning or a genuinely different person, using BOTH name
and phone - neither alone is sufficient:

    NAME matches AND PHONE matches  -> reuse the existing Client, create a new Order
    otherwise (any other combination) -> create a new Client, create a new Order

This deliberately does not use fuzzy/similarity matching (unlike
GET /api/clients/check-duplicates, which is an advisory nudge for a
human to review) - this decision is made unattended as part of order
creation, so it needs a bright-line rule a person could audit, not a
similarity score. A near-miss on either name or phone must create a
new client rather than silently guess.
"""

FUZZY_MATCH_THRESHOLD = 0.8


def normalize_name(name: Optional[str]) -> str:
    """Case-insensitive, whitespace-collapsed - "Garima", " garima ",
    and "GARIMA" all normalize to the same key. Internal punctuation is
    left alone (a name is not a phone number), since stripping it could
    quietly merge genuinely different names."""
    if not name:
        return ""
    return re.sub(r"\s+", " ", name.strip()).lower()


def normalize_phone(phone: Optional[str]) -> str:
    """Strips everything but digits, so "97578 84676", "97578-84676",
    and "(97578) 84676" all normalize to the same key. Does not strip
    or add a country code - "9757884676" and "919757884676" are treated
    as different numbers rather than guessing they're the same person
    with/without a +91 prefix."""
    if not phone:
        return ""
    return re.sub(r"\D", "", phone)


def find_matching_client(db: Session, name: Optional[str], phone: Optional[str]) -> Optional[Client]:
    """Returns the existing Client only if BOTH normalized name and
    normalized phone match some existing client - None otherwise
    (including when name or phone is empty/invalid, since a match
    can't be determined at all in that case, and Client Master
    validation will reject an empty/invalid phone on creation anyway).

    Filters by the indexed phone column at the database level first,
    since every path that creates a Client (the regular API, the Excel
    importer, and this module's own find_or_create_client) validates
    phone down to exactly 10 digits before ever storing it - a direct
    equality filter on the normalized phone is therefore a safe,
    narrow candidate set, not an approximation. The final name
    comparison runs in Python only against that small set, not the
    entire table."""
    name_key = normalize_name(name)
    phone_key = normalize_phone(phone)
    if not name_key or not phone_key:
        return None

    candidates = db.query(Client).filter(Client.phone == phone_key).all()
    for candidate in candidates:
        if normalize_name(candidate.name) == name_key and normalize_phone(candidate.phone) == phone_key:
            return candidate
    return None


def find_fuzzy_name_matches(db: Session, name: str, limit: int = 5,
                             candidates: Optional[List[Client]] = None) -> List[Client]:
    """Non-blocking "possible duplicate" nudge -
    catches typo-variants ("Fevikol" vs "Fevicol") and substrings via
    fuzzy similarity. Distinct from find_matching_client above: that
    function is a bright-line exact-match rule used to make an unattended
    decision (order intake); this one is an advisory signal for a human
    to review (both GET /api/clients/check-duplicates and the Client
    Excel importer call this same function, so the two can never
    disagree about what counts as a "possible" match) and never by
    itself creates, reuses, or rejects a record.
    """
    name_lower = (name or "").strip().lower()
    if not name_lower:
        return []
    matches = []
    for candidate in (candidates if candidates is not None else db.query(Client).all()):
        c_name_lower = candidate.name.lower()
        if name_lower in c_name_lower or c_name_lower in name_lower:
            matches.append(candidate)
            continue
        similarity = difflib.SequenceMatcher(None, name_lower, c_name_lower).ratio()
        if similarity >= FUZZY_MATCH_THRESHOLD:
            matches.append(candidate)
    return matches[:limit]


def find_or_create_client(db: Session, name: str, phone: str, **extra_fields) -> Tuple[Client, bool]:
    """Returns (client, created). Reuses an existing client only on a
    full name+phone match; otherwise creates a new one through the same
    generate_unique_code/generate_business_id path every other client
    creation uses, so a client created this way is indistinguishable
    from one created through the regular Client form.

    extra_fields are only applied when actually creating a new client -
    reusing an existing client never overwrites its stored details
    (e.g. its address) just because a new order mentioned different
    incidental details; only its own edit flow should ever change that.

    Guarded by a PostgreSQL advisory lock scoped to this specific
    (name, phone) pair for the duration of the find-then-create check -
    otherwise two simultaneous requests for the same new client could
    both find no existing match and both create a duplicate Client row.
    No DB-level uniqueness constraint protects against this the way it
    does for other duplicate-creation races this app guards against:
    unlike a client-wide product rate row, an existing Client can have
    real Estimates/Orders/Payments depending on it, so a future
    duplicate could never be safely auto-merged or deleted the way
    those simpler rows could - preventing the race here is the only
    safe option. A no-op on SQLite (tests never run genuinely
    concurrent requests against it)."""
    lock_key = None
    if db.bind is not None and db.bind.dialect.name == "postgresql":
        normalized = f"{normalize_name(name)}|{normalize_phone(phone)}"
        lock_key = int(hashlib.md5(normalized.encode()).hexdigest()[:15], 16)
        db.execute(sa_text("SELECT pg_advisory_lock(:key)"), {"key": lock_key})
    try:
        existing = find_matching_client(db, name, phone)
        if existing:
            return existing, False

        # This path constructs Client(...) directly via the ORM rather than
        # through the ClientCreate Pydantic schema, so the schema's own
        # phone validator never runs here - it has to be checked explicitly
        # before a genuinely new client can be created this way. Reusing an
        # existing client (the branch above) never needs this: an existing
        # client's phone was already validated when IT was created.
        from app.shared import validate_phone
        if not validate_phone((phone or "").strip()):
            raise ValueError("Please enter valid mobile number")

        for _ in range(5):
            code = generate_unique_code(db, Client, "client_code", "CL-")
            client = Client(
                client_code=code, business_id=generate_business_id(db),
                name=(name or "").strip(), phone=(phone or "").strip(), **extra_fields,
            )
            db.add(client)
            try:
                db.flush()
                return client, True
            except IntegrityError:
                db.rollback()
                continue
        raise RuntimeError("Unable to generate a unique client code, please try again")
    finally:
        if lock_key is not None:
            db.execute(sa_text("SELECT pg_advisory_unlock(:key)"), {"key": lock_key})


# --- pdf_generator.py ---
"""Client profile PDF generation - reportlab Platypus using the shared
Woodful document design system. Split out of the former
shared/pdf_generator.py - see modules/sales/pdf_generator.py's
docstring for why."""

def generate_client_pdf(client: Client, is_privileged: bool) -> BytesIO:
    """Client profile/summary document - contact
    information plus a real sales summary derived from the client's
    actual Orders - only real information actually
    supported by the current codebase, no fabricated
    relationships. Uses the same Woodful document header/footer as
    every other generated PDF.

    Financial figures (order values, received amounts, balances) are
    genuinely gated by is_privileged - previously this function had no
    way to know who was requesting the PDF, so it always printed the
    full sales summary and per-order values regardless of role,
    bypassing the same redaction the normal JSON client API already
    applies for non-master users."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=0.6 * inch, bottomMargin=0.7 * inch,
                             leftMargin=0.6 * inch, rightMargin=0.6 * inch)
    styles = get_styles()
    elements = build_header("CLIENT PROFILE", client.client_code, datetime.utcnow().strftime("%d %b %Y"), client.business_id)

    contact_row = [("CONTACT PERSON", client.contact_person or "-"), ("PHONE", mask_phone(client.phone))]
    if client.alternate_phone:
        contact_row = [("CONTACT PERSON", client.contact_person or "-"), ("ALTERNATE PHONE", mask_phone(client.alternate_phone))]
    elements.append(section_table(
        [contact_row,
         [("EMAIL", client.email or "-"), ("CITY", client.city or "-")],
         [("ADDRESS", client.address or "-"), ("SITE ADDRESS", client.site_address or "-")],
         [("GSTIN", client.gstin or "-"), ("STATUS", client.status)],
         [("LEAD SOURCE", client.lead_source or "-"), ("CLIENT SINCE", fmt_date(client.created_at))]],
        [3.5 * inch, 3.5 * inch],
    ))
    elements.append(Spacer(1, 16))

    orders = client.orders or []
    if orders:
        if is_privileged:
            total_value = sum(float(o.order_value or 0) for o in orders)
            total_received = sum(float(o.total_received or 0) for o in orders)
            outstanding = sum(float(o.balance or 0) for o in orders)
            elements.append(Paragraph("SALES SUMMARY", styles["section_label"]))
            elements.append(line_items_table(
                ["Metric", "Value"],
                [["Total Orders", str(len(orders))],
                 ["Total Order Value", format_inr(total_value)],
                 ["Total Received", format_inr(total_received)],
                 ["Outstanding Balance", format_inr(outstanding)]],
                [5 * inch, 2 * inch],
            ))
            elements.append(Spacer(1, 14))
        elements.append(Paragraph("ORDER HISTORY", styles["section_label"]))
        if is_privileged:
            elements.append(line_items_table(
                ["Order", "Date", "Status", "Value"],
                [[o.order_code, fmt_date(o.order_date), o.project_status, format_inr(o.order_value)]
                 for o in sorted(orders, key=lambda o: o.order_date or datetime.min, reverse=True)],
                [1.6 * inch, 1.6 * inch, 2 * inch, 1.8 * inch],
            ))
        else:
            elements.append(line_items_table(
                ["Order", "Date", "Status"],
                [[o.order_code, fmt_date(o.order_date), o.project_status]
                 for o in sorted(orders, key=lambda o: o.order_date or datetime.min, reverse=True)],
                [2.4 * inch, 2.4 * inch, 2.4 * inch],
            ))
    else:
        elements.append(Paragraph(
            '<font color="#70685D" size="9">No orders yet.</font>', styles["body"],
        ))

    if client.remarks:
        elements.append(Spacer(1, 14))
        elements.append(Paragraph(f'<font color="#70685D" size="8">NOTES</font><br/>{pdf_text(client.remarks)}', styles["body"]))

    elements.append(Spacer(1, 24))
    elements.append(Paragraph(build_footer_text(), styles["footer"]))

    doc.build(elements)
    buffer.seek(0)
    return buffer


def client_relationship_timeline(db: Session, client_id: int, limit: int = 200, is_privileged: bool = True) -> dict:
    """Family 137 feature 5 - Unified Client Relationship Timeline. A
    presentation/aggregation layer only, exactly as the spec requires
    ("do NOT create a redundant timeline table... do not duplicate
    source records") - every entry below is read from an existing
    table (ClientActivity, Estimate, Order, Payment, ClientDocument,
    GenericDocument, Notification) and merged into one chronological
    feed, the same pattern already proven for a single order's own
    activity feed (see sales/api.py's get_order_activity), just scoped
    across every estimate/order this client has ever had rather than
    one order.

    is_privileged (MASTER) includes payment amounts - a non-privileged
    caller still sees that a payment happened, never the figure,
    matching the same financial-confidentiality rule used everywhere
    else in this codebase (e.g. get_order_activity's own financial
    notification filtering)."""
    from app.modules.sales.models import Order, Estimate, Payment
    from app.modules.communications.models import Notification
    from app.modules.documents.api import GenericDocument

    client = db.query(Client).filter(Client.id == client_id).first()
    if not client:
        return None

    entries = []

    for a in db.query(ClientActivity).filter(ClientActivity.client_id == client_id).all():
        entries.append({
            "type": "activity", "subtype": a.activity_type, "date": a.date,
            "author": a.logged_by, "text": a.summary, "path": f"/clients/{client_id}",
        })

    for d in db.query(ClientDocument).filter(ClientDocument.client_id == client_id).all():
        entries.append({
            "type": "document", "subtype": "client_document", "date": d.created_at,
            "author": d.uploaded_by, "text": f"Document uploaded: {d.original_filename}",
            "path": f"/clients/{client_id}",
        })

    estimates = db.query(Estimate).filter(Estimate.client_id == client_id).all()
    for e in estimates:
        entries.append({
            "type": "estimate", "subtype": "created", "date": e.created_at, "author": None,
            "text": f"Estimate {e.estimate_code} created (v{e.version}, {e.status}).",
            "path": f"/estimates/{e.id}",
        })
        if e.approved_at:
            entries.append({
                "type": "estimate", "subtype": "approved", "date": e.approved_at,
                "author": e.approved_by,
                "text": f"Estimate {e.estimate_code} approved" + (f" by {e.approved_by}" if e.approved_by else "")
                        + (f" - {e.client_decision_comments}" if e.client_decision_comments else ""),
                "path": f"/estimates/{e.id}",
            })
        elif e.status == "changes_requested" and e.client_decision_comments:
            entries.append({
                "type": "estimate", "subtype": "changes_requested", "date": e.updated_at, "author": None,
                "text": f"Changes requested on estimate {e.estimate_code}: {e.client_decision_comments}",
                "path": f"/estimates/{e.id}",
            })

    orders = db.query(Order).filter(Order.client_id == client_id).all()
    order_ids = [o.id for o in orders]
    for o in orders:
        entries.append({
            "type": "order", "subtype": "created", "date": o.order_date, "author": None,
            "text": f"Order {o.order_code} created ({o.project_status}).",
            "path": f"/orders/{o.id}",
        })
        if o.project_status == "Completed":
            entries.append({
                "type": "order", "subtype": "status", "date": o.updated_at, "author": None,
                "text": f"Order {o.order_code} marked Completed.", "path": f"/orders/{o.id}",
            })

    if order_ids:
        payments = db.query(Payment).filter(Payment.order_id.in_(order_ids)).all()
        for p in payments:
            order = next((o for o in orders if o.id == p.order_id), None)
            order_code = order.order_code if order else p.order_id
            if is_privileged:
                text = f"Payment received against {order_code}: Rs {float(p.amount):,.2f} ({p.payment_mode})."
            else:
                text = f"Payment received against {order_code}."
            entries.append({
                "type": "payment", "subtype": p.payment_type, "date": p.date, "author": p.received_by,
                "text": text, "path": f"/orders/{p.order_id}",
            })

        docs = db.query(GenericDocument).filter(
            GenericDocument.parent_type == "order", GenericDocument.parent_id.in_(order_ids),
        ).all()
        for d in docs:
            entries.append({
                "type": "document", "subtype": "order_document", "date": d.created_at,
                "author": d.uploaded_by, "text": f"Document uploaded on order {d.parent_id}: {d.original_filename}",
                "path": f"/orders/{d.parent_id}",
            })

        # Business-level notifications about these orders (broadcasts -
        # recipient_user_id is None), not any one staff member's
        # personal inbox entry - a relationship timeline is about the
        # client relationship, not one employee's notifications.
        # Financial notification types stay master-only, matching
        # NotificationService.visible_to's own rule everywhere else.
        from app.modules.communications.services import FINANCIAL_NOTIFICATION_TYPES
        notif_query = db.query(Notification).filter(
            Notification.related_entity_type == "order", Notification.related_entity_id.in_(order_ids),
            Notification.recipient_user_id.is_(None),
        )
        if not is_privileged:
            notif_query = notif_query.filter(Notification.notification_type.notin_(FINANCIAL_NOTIFICATION_TYPES))
        for n in notif_query.all():
            entries.append({
                "type": "communication", "subtype": n.notification_type, "date": n.created_at, "author": None,
                "text": f"{n.title}: {n.message}", "path": n.action_path or "/",
            })

    entries.sort(key=lambda e: e["date"] or datetime.min, reverse=True)
    for e in entries:
        e["date"] = e["date"].isoformat() if e["date"] else None

    return {
        "client_id": client.id, "client_name": client.name,
        "total_entries": len(entries),
        "entries": entries[:limit],
        "generated_at": datetime.utcnow().isoformat(),
    }
