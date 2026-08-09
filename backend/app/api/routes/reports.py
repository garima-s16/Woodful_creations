from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.security import verify_token
from app.models.product import Product
from app.models.sales_order import SalesOrder, SalesOrderItem

router = APIRouter(prefix="/api/reports", tags=["reports"])
security = HTTPBearer()

HEADER_FILL = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def verify_auth(credentials: HTTPAuthorizationCredentials = Depends(security)):
    payload = verify_token(credentials.credentials)
    if not payload:
        raise HTTPException(status_code=401, detail="Invalid token")
    return payload


def _style_header_row(ws, row: int, col_count: int):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


@router.get("/inventory/excel")
def download_inventory_report(
    db: Session = Depends(get_db),
    auth=Depends(verify_auth),
):
    products = db.query(Product).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"

    headers = [
        "ID",
        "Material Type",
        "Category",
        "SKU",
        "Description",
        "Quantity",
        "Min Quantity",
        "Price Per Unit (Rs)",
        "Unit",
        "Supplier",
        "Status",
        "Last Restocked",
        "Created At",
    ]
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))

    for p in products:
        status = "Low Stock" if p.quantity <= p.min_quantity else "In Stock"
        if p.quantity == 0:
            status = "Out of Stock"
        ws.append(
            [
                p.id,
                p.material_type,
                p.category,
                p.sku or "",
                p.description or "",
                p.quantity,
                p.min_quantity,
                p.price_per_unit,
                p.unit,
                p.supplier or "",
                status,
                p.last_restocked.strftime("%Y-%m-%d") if p.last_restocked else "",
                p.created_at.strftime("%Y-%m-%d %H:%M") if p.created_at else "",
            ]
        )

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=inventory_report.xlsx"},
    )


@router.get("/sales/excel")
def download_sales_report(
    db: Session = Depends(get_db),
    auth=Depends(verify_auth),
):
    orders = db.query(SalesOrder).order_by(SalesOrder.created_at.desc()).all()

    wb = Workbook()

    ws_orders = wb.active
    ws_orders.title = "Orders"
    order_headers = [
        "Order Number",
        "Client Name",
        "Client Phone",
        "Client Email",
        "Status",
        "Total Amount (Rs)",
        "Notes",
        "Created At",
    ]
    ws_orders.append(order_headers)
    _style_header_row(ws_orders, 1, len(order_headers))

    for o in orders:
        ws_orders.append(
            [
                o.order_number,
                o.client_name,
                o.client_phone or "",
                o.client_email or "",
                o.status,
                o.total_amount,
                o.notes or "",
                o.created_at.strftime("%Y-%m-%d %H:%M") if o.created_at else "",
            ]
        )

    for col in ws_orders.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws_orders.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    ws_items = wb.create_sheet("Order Items")
    item_headers = [
        "Order Number",
        "Product Name",
        "Quantity",
        "Unit Price (Rs)",
        "Line Total (Rs)",
    ]
    ws_items.append(item_headers)
    _style_header_row(ws_items, 1, len(item_headers))

    for o in orders:
        for item in o.items:
            ws_items.append(
                [
                    o.order_number,
                    item.product_name,
                    item.quantity,
                    item.unit_price,
                    item.line_total,
                ]
            )

    for col in ws_items.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws_items.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=sales_report.xlsx"},
    )
