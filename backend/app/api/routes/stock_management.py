from datetime import datetime, timezone
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session, joinedload

from app.core.database import get_db
from app.models.stock_management import AppSetting, Material, StockIn, StockOut, Supplier
from app.schemas.stock_management import (
    DashboardSummary,
    MaterialCreate,
    MaterialRead,
    MaterialUpdate,
    SettingsRead,
    SettingsUpdate,
    StockInCreate,
    StockInRead,
    StockOutCreate,
    StockOutRead,
    SupplierCreate,
    SupplierRead,
)
from app.services.stock_management import dashboard_data, material_to_view

router = APIRouter(prefix="/api/stock-management", tags=["stock-management"])


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=14, bold=True)


def _get_settings(db: Session) -> AppSetting:
    settings = db.query(AppSetting).first()
    if not settings:
        settings = AppSetting()
        db.add(settings)
        db.commit()
        db.refresh(settings)
    return settings


def _materials_query(db: Session):
    return (
        db.query(Material)
        .options(joinedload(Material.stock_ins), joinedload(Material.stock_outs), joinedload(Material.supplier))
        .order_by(Material.name.asc())
    )


@router.get("/materials", response_model=list[MaterialRead])
def list_materials(db: Session = Depends(get_db)):
    settings = _get_settings(db)
    return [material_to_view(m, settings.reorder_buffer) for m in _materials_query(db).all()]


@router.post("/materials", response_model=MaterialRead)
def create_material(payload: MaterialCreate, db: Session = Depends(get_db)):
    if payload.supplier_id and not db.get(Supplier, payload.supplier_id):
        raise HTTPException(status_code=400, detail="Supplier not found")

    material = Material(**payload.model_dump())
    db.add(material)
    db.commit()
    db.refresh(material)
    settings = _get_settings(db)
    material = _materials_query(db).filter(Material.id == material.id).first()
    return material_to_view(material, settings.reorder_buffer)


@router.put("/materials/{material_id}", response_model=MaterialRead)
def update_material(material_id: int, payload: MaterialUpdate, db: Session = Depends(get_db)):
    material = db.get(Material, material_id)
    if not material:
        raise HTTPException(status_code=404, detail="Material not found")

    update_data = payload.model_dump(exclude_unset=True)
    if "supplier_id" in update_data and update_data["supplier_id"]:
        if not db.get(Supplier, update_data["supplier_id"]):
            raise HTTPException(status_code=400, detail="Supplier not found")

    for key, value in update_data.items():
        setattr(material, key, value)

    db.add(material)
    db.commit()
    db.refresh(material)
    settings = _get_settings(db)
    material = _materials_query(db).filter(Material.id == material_id).first()
    return material_to_view(material, settings.reorder_buffer)


@router.delete("/materials/{material_id}")
def delete_material(material_id: int, db: Session = Depends(get_db)):
    material = db.get(Material, material_id)
    if not material:
        raise HTTPException(status_code=404, detail="Material not found")
    db.delete(material)
    db.commit()
    return {"message": "Material deleted"}


@router.get("/stock-in", response_model=list[StockInRead])
def list_stock_in(db: Session = Depends(get_db)):
    return db.query(StockIn).order_by(StockIn.purchased_at.desc()).all()


@router.post("/stock-in", response_model=StockInRead)
def create_stock_in(payload: StockInCreate, db: Session = Depends(get_db)):
    material = db.get(Material, payload.material_id)
    if not material:
        raise HTTPException(status_code=404, detail="Material not found")
    stock_in = StockIn(**payload.model_dump(exclude_none=True))
    if stock_in.purchased_at is None:
        stock_in.purchased_at = datetime.now(timezone.utc)
    db.add(stock_in)
    db.commit()
    db.refresh(stock_in)
    return stock_in


@router.get("/stock-out", response_model=list[StockOutRead])
def list_stock_out(db: Session = Depends(get_db)):
    return db.query(StockOut).order_by(StockOut.issued_at.desc()).all()


@router.post("/stock-out", response_model=StockOutRead)
def create_stock_out(payload: StockOutCreate, db: Session = Depends(get_db)):
    material = _materials_query(db).filter(Material.id == payload.material_id).first()
    if not material:
        raise HTTPException(status_code=404, detail="Material not found")

    settings = _get_settings(db)
    current = material_to_view(material, settings.reorder_buffer)["current_stock"]
    if payload.quantity > current:
        raise HTTPException(status_code=400, detail="Issued quantity exceeds current stock")

    stock_out = StockOut(**payload.model_dump(exclude_none=True))
    if stock_out.issued_at is None:
        stock_out.issued_at = datetime.now(timezone.utc)
    db.add(stock_out)
    db.commit()
    db.refresh(stock_out)
    return stock_out


@router.get("/suppliers", response_model=list[SupplierRead])
def list_suppliers(db: Session = Depends(get_db)):
    return db.query(Supplier).order_by(Supplier.name.asc()).all()


@router.post("/suppliers", response_model=SupplierRead)
def create_supplier(payload: SupplierCreate, db: Session = Depends(get_db)):
    supplier = Supplier(**payload.model_dump())
    db.add(supplier)
    db.commit()
    db.refresh(supplier)
    return supplier


@router.get("/settings", response_model=SettingsRead)
def get_settings(db: Session = Depends(get_db)):
    return _get_settings(db)


@router.put("/settings", response_model=SettingsRead)
def update_settings(payload: SettingsUpdate, db: Session = Depends(get_db)):
    settings = _get_settings(db)
    for key, value in payload.model_dump(exclude_unset=True).items():
        setattr(settings, key, value)
    db.add(settings)
    db.commit()
    db.refresh(settings)
    return settings


@router.get("/dashboard", response_model=DashboardSummary)
def get_dashboard_summary(db: Session = Depends(get_db)):
    settings = _get_settings(db)
    materials = _materials_query(db).all()
    return dashboard_data(materials, settings.reorder_buffer)


def _style_header(ws, row: int, headers: list[str]):
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT


def _auto_width(ws):
    for index, col in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column), start=1):
        length = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(index)].width = min(length + 2, 40)


@router.get("/reports/stock/download")
def download_stock_report(db: Session = Depends(get_db)):
    settings = _get_settings(db)
    materials = _materials_query(db).all()
    stock_in_rows = db.query(StockIn).order_by(StockIn.purchased_at.desc()).all()
    stock_out_rows = db.query(StockOut).order_by(StockOut.issued_at.desc()).all()
    suppliers = db.query(Supplier).order_by(Supplier.name.asc()).all()
    summary = dashboard_data(materials, settings.reorder_buffer)
    material_views = [material_to_view(m, settings.reorder_buffer) for m in materials]
    supplier_map = {supplier.id: supplier.name for supplier in suppliers}
    material_map = {material.id: material.name for material in materials}

    wb = Workbook()

    ws_dashboard = wb.active
    ws_dashboard.title = "Dashboard"
    ws_dashboard.merge_cells("A1:D1")
    ws_dashboard["A1"] = f"{settings.company_name} - STOCK MANAGEMENT DASHBOARD"
    ws_dashboard["A1"].font = TITLE_FONT
    _style_header(ws_dashboard, 3, ["Metric", "Value", "", ""])
    metric_rows = []
    ws_dashboard.append(["Total Stock Value", summary["total_stock_value"]])
    metric_rows.append(ws_dashboard.max_row)
    ws_dashboard.append(["Low Stock Items", summary["low_stock_items"]])
    metric_rows.append(ws_dashboard.max_row)
    ws_dashboard.append(["Out of Stock", summary["out_of_stock_items"]])
    metric_rows.append(ws_dashboard.max_row)
    ws_dashboard.append(["Purchase Value", summary["purchase_value"]])
    metric_rows.append(ws_dashboard.max_row)

    ws_dashboard[f"B{metric_rows[0]}"].number_format = "#,##0.00"
    ws_dashboard[f"B{metric_rows[3]}"].number_format = "#,##0.00"

    ws_dashboard.append([])
    start_row = ws_dashboard.max_row + 1
    _style_header(ws_dashboard, start_row, ["Material", "Current", "Minimum", "Status", "Suggested Order"])
    for item in summary["low_stock_action_list"]:
        ws_dashboard.append([
            item["material"],
            item["current_stock"],
            item["minimum_stock"],
            item["status"],
            item["suggested_reorder_quantity"],
        ])

    ws_dashboard.append([])
    start_row = ws_dashboard.max_row + 1
    _style_header(ws_dashboard, start_row, ["Category", "Items", "Stock Qty", "Stock Value"])
    for item in summary["category_summary"]:
        ws_dashboard.append([item["category"], item["item_count"], item["stock_quantity"], item["stock_value"]])

    ws_materials = wb.create_sheet("Material Master")
    _style_header(ws_materials, 1, ["Material", "Category", "Supplier", "Unit", "Opening", "Current", "Minimum", "Status", "Unit Price", "Stock Value"])
    for row in material_views:
        ws_materials.append([
            row["name"],
            row["category"],
            supplier_map.get(row["supplier_id"], "") if row["supplier_id"] else "",
            row["unit"],
            row["opening_stock"],
            row["current_stock"],
            row["minimum_stock"],
            row["status"],
            row["unit_price"],
            round(row["current_stock"] * row["unit_price"], 2),
        ])

    ws_in = wb.create_sheet("Stock In - Purchases")
    _style_header(ws_in, 1, ["Date", "Material", "Qty", "Unit Price", "Invoice", "Notes"])
    for row in stock_in_rows:
        material_name = material_map.get(row.material_id, "")
        ws_in.append([row.purchased_at.isoformat(), material_name, row.quantity, row.unit_price, row.invoice_number, row.notes])

    ws_out = wb.create_sheet("Stock Out - Issues")
    _style_header(ws_out, 1, ["Date", "Material", "Qty", "Issued To", "Notes"])
    for row in stock_out_rows:
        material_name = material_map.get(row.material_id, "")
        ws_out.append([row.issued_at.isoformat(), material_name, row.quantity, row.issued_to, row.notes])

    ws_suppliers = wb.create_sheet("Suppliers")
    _style_header(ws_suppliers, 1, ["Name", "Contact Person", "Phone", "Email", "Address"])
    for row in suppliers:
        ws_suppliers.append([row.name, row.contact_person, row.phone, row.email, row.address])

    for sheet in wb.worksheets:
        _auto_width(sheet)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"stock-management-report-{datetime.now(timezone.utc).strftime('%Y%m%d-%H%M%S')}.xlsx"
    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )
